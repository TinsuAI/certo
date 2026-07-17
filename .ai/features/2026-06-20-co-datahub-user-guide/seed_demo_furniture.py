"""Seed a clean fictional demo company into the local Data Hub DB for the user guide.

Run from the data-hub repo so `app.*` resolves to Data Hub's package:

    cd /home/vp/workspace/client/data-hub
    uv run python /home/vp/workspace/client/barry-CO-main/.ai/features/2026-06-20-co-datahub-user-guide/seed_demo_furniture.py --reset

Company: "Demo Furniture Co." (demo-furniture) — a Vietnamese wooden-furniture
exporter. Identity code mode (customs_code == internal_code) so the flow is easy
to read in screenshots. Mirrors app/seed.py:_seed_growatt exactly (same parsers +
stores) so the BOM lands as a leaf-complete manual_flat artifact the CO picker accepts.
"""
from __future__ import annotations

import io
import sys

from openpyxl import Workbook

from app.database import connect
from app.parsers.bcct import parse_bcct_workbook
from app.parsers.bom import parse_bom_workbook
from app.parsers.materials import parse_materials_workbook
from app.routes.bcct import _insert_bcct
from app.routes.catalog import _insert_materials
from app.routes.clients import upsert_client
from app.stores.bom import create_artifact

CLIENT_ID = "demo-furniture"
CLIENT_NAME = "Demo Furniture Co."

# (Mã HQ, Mã NB, Tên, Loại, ĐVT, HS) — identity mode: Mã HQ == Mã NB.
NVL = [
    ("PLY18", "PLY18", "Van ep phu keo 18mm (plywood)", "nvl", "m2", "44123900"),
    ("OAKVNR", "OAKVNR", "Van lang soi 0.6mm (oak veneer)", "nvl", "m2", "44083190"),
    ("WSCREW", "WSCREW", "Bo vit go (wood screw set)", "nvl", "set", "73181500"),
    ("PUFOAM", "PUFOAM", "Mut PU boc nem (PU foam)", "nvl", "kg", "39211390"),
    ("PVCFAB", "PVCFAB", "Vai boc PVC (PVC upholstery)", "nvl", "m", "39219090"),
    ("PUVARN", "PUVARN", "Son phu PU (PU varnish)", "nvl", "L", "32089090"),
    ("WGLUE", "WGLUE", "Keo dan go (wood adhesive)", "nvl", "kg", "35069190"),
    ("SSHINGE", "SSHINGE", "Ban le inox (stainless hinge)", "nvl", "pcs", "83024110"),
]
TP = [
    ("CHAIR01", "CHAIR01", "Ghe an go boc nem (dining chair)", "tp", "pcs", "94016100"),
    ("TABLE01", "TABLE01", "Ban tra go (coffee table)", "tp", "pcs", "94036000"),
]

# (Số tờ khai, Dòng, Mã loại hình, Ngày, Mã NPL/SP, Tên hàng, HS, SL, ĐVT, Trị giá, Nguyên tệ)
BCCT = [
    # Imports (E11) — raw materials, non-origin, drive LVC + tồn.
    ("105100100001", 1, "E11", "2025-02-10", "PLY18", "Plywood sheet 18mm grade BB/CC", "44123900", 2000.0, "m2", 9000.0, "USD"),
    ("105100100001", 2, "E11", "2025-02-10", "OAKVNR", "Oak veneer 0.6mm rotary cut", "44083190", 1500.0, "m2", 6000.0, "USD"),
    ("105100100002", 1, "E11", "2025-02-28", "WSCREW", "Wood screw set zinc-plated", "73181500", 5000.0, "set", 2500.0, "USD"),
    ("105100100002", 2, "E11", "2025-02-28", "SSHINGE", "Stainless steel hinge 304", "83024110", 3000.0, "pcs", 1800.0, "USD"),
    ("105100100003", 1, "E11", "2025-03-15", "PUFOAM", "PU foam cushion density 30", "39211390", 800.0, "kg", 3200.0, "USD"),
    ("105100100003", 2, "E11", "2025-03-15", "PVCFAB", "PVC upholstery fabric brown", "39219090", 1200.0, "m", 4200.0, "USD"),
    ("105100100004", 1, "E11", "2025-04-05", "PUVARN", "PU varnish clear gloss", "32089090", 300.0, "L", 2700.0, "USD"),
    ("105100100004", 2, "E11", "2025-04-05", "WGLUE", "Wood adhesive PVA D3", "35069190", 400.0, "kg", 1600.0, "USD"),
    # Exports (E42) — finished goods.
    ("105100100100", 1, "E42", "2025-09-20", "CHAIR01", "Wooden dining chair upholstered", "94016100", 1000.0, "pcs", 45000.0, "USD"),
    ("105100100101", 1, "E42", "2025-10-18", "TABLE01", "Wooden coffee table oak top", "94036000", 400.0, "pcs", 32000.0, "USD"),
]

# (Mã SP, Mã NVL, Định mức, ĐVT)
BOM = [
    ("CHAIR01", "PLY18", 0.40, "m2"),
    ("CHAIR01", "OAKVNR", 0.50, "m2"),
    ("CHAIR01", "WSCREW", 1.0, "set"),
    ("CHAIR01", "PUFOAM", 0.60, "kg"),
    ("CHAIR01", "PVCFAB", 0.80, "m"),
    ("CHAIR01", "PUVARN", 0.05, "L"),
    ("CHAIR01", "WGLUE", 0.10, "kg"),
    ("TABLE01", "PLY18", 0.90, "m2"),
    ("TABLE01", "OAKVNR", 0.70, "m2"),
    ("TABLE01", "WSCREW", 1.0, "set"),
    ("TABLE01", "PUVARN", 0.08, "L"),
    ("TABLE01", "WGLUE", 0.15, "kg"),
    ("TABLE01", "SSHINGE", 2.0, "pcs"),
]


def _reset() -> None:
    """Idempotent wipe: delete every hub.* row for this client across all tables
    that carry a client_id column. Retry loop resolves FK ordering without
    hard-coding it."""
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            """
            select table_name from information_schema.columns
            where table_schema='hub' and column_name='client_id'
            """
        )
        tables = [r[0] for r in cur.fetchall()]
        remaining = set(tables)
        for _ in range(len(tables) + 2):
            if not remaining:
                break
            progressed = set()
            for t in list(remaining):
                try:
                    cur.execute("savepoint sp")
                    cur.execute(f"delete from hub.{t} where client_id=%s", (CLIENT_ID,))
                    cur.execute("release savepoint sp")
                    progressed.add(t)
                except Exception:
                    cur.execute("rollback to savepoint sp")
            remaining -= progressed
            if not progressed:
                break
        conn.commit()
        print(f"[reset] cleared client rows from {len(tables) - len(remaining)} tables"
              + (f" (still blocked: {sorted(remaining)})" if remaining else ""))


def _materials_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DM NVL"
    ws.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    for r in NVL:
        ws.append(r)
    ws_tp = wb.create_sheet("DM TP")
    ws_tp.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    for r in TP:
        ws_tp.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bcct_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BCCT"
    ws.append(["Số tờ khai", "Dòng", "Mã loại hình", "Ngày đăng ký", "Mã NPL/SP",
               "Tên hàng", "HS", "Tổng số lượng", "ĐVT", "Trị giá", "Nguyên tệ"])
    for r in BCCT:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bom_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["Mã SP", "Mã NVL", "Định mức", "ĐVT"])
    for r in BOM:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def seed() -> None:
    upsert_client(
        client_id=CLIENT_ID, name=CLIENT_NAME, tax_code="0312345678",
        code_resolution_mode="identity",
        bom_proposal_mode="auto", bom_proposal_qty_tolerance_pct=5.0,
        bom_approver_tier="edit",
        status="active",
        notes="Fictional demo company for the CO + Data Hub user guide. Identity code mode.",
    )

    rows, _ = parse_materials_workbook(_materials_workbook())
    _insert_materials(client_id=CLIENT_ID, rows=rows)

    parsed = parse_bcct_workbook(_bcct_workbook())
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            """
            insert into hub.file_uploads (upload_id, client_id, module,
              original_filename, stored_path, content_sha256, size_bytes,
              parse_status, row_count)
            values ('seed-bcct-demo-furniture', %s, 'bcct', 'seed-bcct.xlsx', 'seed/bcct',
              'seed', 0, 'done', %s)
            on conflict (upload_id) do nothing
            """,
            (CLIENT_ID, len(BCCT)),
        )
        conn.commit()
    _insert_bcct(client_id=CLIENT_ID, rows=parsed,
                 upload_id="seed-bcct-demo-furniture",
                 client={"client_id": CLIENT_ID, "code_resolution_mode": "identity"})

    products = parse_bom_workbook(_bom_workbook(), profile="manual_flat")
    for product_code, prows in products.items():
        create_artifact(
            client_id=CLIENT_ID, product_code=product_code, rows=prows,
            actor="agency_staff", intent="asserted_technical",
            parent_artifact_id=None,
            context={"channel": "agency_upload", "profile": "manual_flat", "seed": True},
            source_upload_id=None,
        )

    print(f"[seed] {CLIENT_NAME} ({CLIENT_ID}): "
          f"{len(NVL)} NVL + {len(TP)} TP, {len(BCCT)} BCCT lines, "
          f"{len(products)} BOM artifacts ({', '.join(products)}).")


if __name__ == "__main__":
    if "--reset" in sys.argv:
        _reset()
    seed()
