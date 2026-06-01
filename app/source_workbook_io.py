from __future__ import annotations

import re
import unicodedata
import xlrd

from app.co_stock_derivation import cell_text, normalize_decimal
from io import BytesIO
from openpyxl import load_workbook
from pathlib import Path
from zipfile import BadZipFile, ZipFile


UNIT_ALIASES = {
    "PC": "PCS",
    "PCS": "PCS",
    "PCE": "PCS",
    "PIECE": "PCS",
    "PIECES": "PCS",
    "CAI": "PCS",
    "CÁI": "PCS",
    "KG": "KG",
    "KGS": "KG",
    "EA": "EA",
    "SET": "SET",
}
DIRECTION_ALIASES = {
    "IMPORT": "import",
    "NHAP KHAU": "import",
    "NHẬP KHẨU": "import",
    "NK": "import",
    "EXPORT": "export",
    "XUAT KHAU": "export",
    "XUẤT KHẨU": "export",
    "XK": "export",
}
SOURCE_AUDIT_FIELDS = {
    "source_upload_id",
    "review_status",
    "import_row_id",
    "source_schema",
    "source_file",
    "source_sheet",
    "source_header_row",
    "source_row_number",
    "raw_fields",
}
CATALOG_CASE_RULE_FIELDS = {"rule"}
COMMON_HEADER_ALIASES = {
    "ten": "name",
    "ten_hang": "description",
    "mo_ta": "description",
    "ma_hs": "hs_code",
    "hs": "hs_code",
    "don_vi": "unit",
    "don_vi_tinh": "unit",
    "dvt": "unit",
    "uom": "unit",
    "quy_tac": "rule",
    "origin_rule": "rule",
    "trang_thai": "status",
    "status": "status",
    "muc_dich_sd": "purpose",
    "muc_dich_su_dung": "purpose",
    "canh_bao_khi_tao_to_khai": "declaration_warning",
    "noi_dung_canh_bao": "warning_content",
    "don_gia": "unit_price",
    "ma_bieu_thue_nk": "import_tariff_code",
    "ma_ap_dung_thue_ttdb": "excise_tax_code",
    "ma_ap_dung_thue_moi_truong": "environmental_tax_code",
    "ma_ap_dung_thue_vat": "vat_tax_code",
    "ghi_chu": "note",
    "ten_tieng_anh": "english_name",
}
MATERIAL_HEADER_ALIASES = {
    **COMMON_HEADER_ALIASES,
    "ma": "customs_code",
    "ma_hq": "customs_code",
    "ma_nvl": "customs_code",
    "ma_npl": "customs_code",
    "ma_noi_bo": "internal_code",
    "ten": "name",
}
PRODUCT_HEADER_ALIASES = {
    **COMMON_HEADER_ALIASES,
    "ma": "product_code",
    "ma_tp": "product_code",
    "ma_sp": "product_code",
    "customs_code": "product_code",
    "ma_dinh_danh_cua_lenh_sx": "production_order_identifier",
    "ten": "name",
}
BCCT_HEADER_ALIASES = {
    **COMMON_HEADER_ALIASES,
    "stt": "sequence_no",
    "ky": "coverage_period",
    "period": "coverage_period",
    "luong": "direction",
    "direction": "direction",
    "so_tk": "declaration_no",
    "so_to_khai": "declaration_no",
    "to_khai": "declaration_no",
    "declaration_no": "declaration_no",
    "ngay_tk": "declaration_date",
    "ngay_dk": "declaration_date",
    "declaration_date": "declaration_date",
    "hai_quan": "customs_office",
    "customs_office": "customs_office",
    "ma_loai_hinh": "declaration_type",
    "loai_hinh": "declaration_type",
    "declaration_type": "declaration_type",
    "ma_dia_diem_dich": "destination_location_code",
    "ten_dia_diem_dich_cho_van_chuyen_bao_thue": "destination_location_name",
    "dia_diem_do_hang": "unloading_location",
    "ma_hieu_ptvc": "transport_mode_code",
    "ngay_khoi_hanh_van_chuyen": "departure_date",
    "ky_hieu_va_so_hieu_bao_bi": "package_marks",
    "ty_gia_thanh_toan": "exchange_rate",
    "don_vi_tien_te": "currency",
    "so_luong_kien": "package_quantity",
    "ma_dvt_kien": "package_unit",
    "trong_luong": "gross_weight",
    "ma_dvt_trong_luong": "gross_weight_unit",
    "so_quan_ly_noi_bo": "internal_management_no",
    "dieu_kien_gia_hoa_don": "invoice_price_condition",
    "ghi_chu": "declaration_note",
    "stt_hang": "line_no",
    "line_no": "line_no",
    "ma_hang": "item_code",
    "ma_npl_sp": "item_code",
    "item_code": "item_code",
    "material_code": "item_code",
    "product_code": "item_code",
    "xuat_xu": "origin_country",
    "don_gia_tinh_thue": "taxable_unit_price",
    "tong_so_luong": "quantity",
    "so_luong": "quantity",
    "qty": "quantity",
    "tong_so_luong_2": "secondary_quantity",
    "don_vi_tinh_2": "secondary_unit",
    "tri_gia_nt": "foreign_currency_value",
    "tri_gia": "customs_value",
    "tong_tri_gia": "customs_value",
    "value": "customs_value",
    "ma_bieu_thue_xnk": "import_export_tariff_code",
    "thue_suat_xnk": "import_export_tax_rate",
    "tien_thue_xnk": "import_export_tax_amount",
    "so_tien_mien_thue_xnk": "import_export_tax_exempt_amount",
    "thue_suat_tv": "safeguard_tax_rate",
    "tien_thue_tv": "safeguard_tax_amount",
    "thue_suat_pb": "trade_remedy_tax_rate",
    "tien_thue_pb": "trade_remedy_tax_amount",
    "thue_suat_ttdb": "excise_tax_rate",
    "tien_thue_ttdb": "excise_tax_amount",
    "thue_suat_bvmt": "environmental_tax_rate",
    "tien_thue_mt": "environmental_tax_amount",
    "thue_suat_va": "vat_tax_rate",
    "tien_thue_vat": "vat_tax_amount",
    "tong_tien_thue": "total_tax_amount",
    "ma_doanh_nghiep": "company_tax_code",
    "ten_doanh_nghiep": "company_name",
    "ten_doi_tac": "partner_name",
    "so_hoa_don": "invoice_ref",
    "hoa_don": "invoice_ref",
    "ngay_hoa_don": "invoice_date",
    "so_hop_dong": "contract_no",
    "ngay_hop_dong": "contract_date",
}
DECLARATION_TYPE_DIRECTIONS = {
    "A11": "import",
    "A12": "import",
    "A21": "import",
    "A31": "import",
    "A41": "import",
    "E11": "import",
    "E13": "import",
    "E15": "import",
    "E21": "import",
    "E23": "import",
    "E31": "import",
    "G11": "import",
    "G12": "import",
    "B11": "export",
    "B12": "export",
    "B13": "export",
    "E42": "export",
    "E52": "export",
    "E54": "export",
    "E62": "export",
    "G21": "export",
    "G22": "export",
}
def parse_catalog_workbook(content: bytes, module: str) -> list[dict]:
    sheets = load_tabular_sheets(content, module)
    sheet = select_catalog_sheet(sheets, module)
    header_row_index, headers, raw_headers = find_header_row(
        sheet["rows"],
        module,
        {"customs_code", "name"} if module == "material_catalog" else {"product_code", "name"},
    )
    rows = []
    source_schema = "customs_material_catalog" if module == "material_catalog" else "customs_product_catalog"
    for values, raw_fields, source_row_number in iter_tabular_rows(
        sheet["rows"],
        headers,
        raw_headers,
        header_row_index + 1,
    ):
        if module == "material_catalog":
            customs_code = cell_text(values.get("customs_code"))
            if not customs_code:
                continue
            row = {
                "customs_code": customs_code,
                "internal_code": cell_text(values.get("internal_code")) or customs_code,
                "name": cell_text(values.get("name") or values.get("description")),
                "hs_code": cell_text(values.get("hs_code") or values.get("hs")),
                "unit": normalize_unit(values.get("unit")),
                "role": cell_text(values.get("role")) or "NVL",
                "origin_default": cell_text(values.get("origin_default") or values.get("origin")),
                "status": normalize_status(values.get("status")),
            }
        else:
            product_code = cell_text(values.get("product_code") or values.get("customs_code"))
            if not product_code:
                continue
            row = {
                "product_code": product_code,
                "name": cell_text(values.get("name") or values.get("description")),
                "hs_code": cell_text(values.get("hs_code") or values.get("hs")),
                "unit": normalize_unit(values.get("unit")),
                "status": normalize_status(values.get("status")),
            }
        attach_source_fields(row, values, raw_fields, sheet, header_row_index + 1, source_row_number, source_schema)
        rows.append(row)
    if not rows:
        raise SourceParseError("Workbook has no catalog rows.")
    return rows
def parse_bcct_workbook(content: bytes) -> list[dict]:
    sheets = load_tabular_sheets(content, "bcct")
    sheet = select_bcct_sheet(sheets)
    header_row_index, headers, raw_headers = find_header_row(
        sheet["rows"],
        "bcct",
        {"declaration_no", "declaration_type", "line_no", "item_code", "quantity", "unit"},
    )
    rows = []
    for values, raw_fields, source_row_number in iter_tabular_rows(
        sheet["rows"],
        headers,
        raw_headers,
        header_row_index + 1,
    ):
        declaration_type = normalize_declaration_type(values.get("declaration_type"))
        direction = normalize_direction(values.get("direction")) or infer_direction(declaration_type)
        declaration_no = cell_text(values.get("declaration_no"))
        line_no = cell_text(values.get("line_no") or values.get("stt") or values.get("stt_hang"))
        item_code = cell_text(values.get("item_code") or values.get("material_code") or values.get("product_code"))
        if not any([direction, declaration_no, line_no, item_code]):
            continue
        if not all([direction, declaration_no, line_no, item_code]):
            raise SourceParseError("BCCT rows require direction, declaration_no, line_no, and item_code.")
        quantity = normalize_decimal(values.get("quantity") or values.get("qty"))
        unit = normalize_unit(values.get("unit") or values.get("uom"))
        row = {
            "coverage_period": cell_text(values.get("coverage_period") or values.get("period")),
            "direction": direction,
            "declaration_no": declaration_no,
            "declaration_date": cell_text(values.get("declaration_date")),
            "customs_office": cell_text(values.get("customs_office")),
            "declaration_type": declaration_type,
            "line_no": line_no,
            "item_code": item_code,
            "description": cell_text(values.get("description") or values.get("name")),
            "hs_code": cell_text(values.get("hs_code") or values.get("hs")),
            "quantity": quantity,
            "unit": unit,
            "customs_value": normalize_decimal(values.get("customs_value") or values.get("value")),
            "invoice_ref": cell_text(values.get("invoice_ref") or values.get("invoice")),
        }
        attach_source_fields(row, values, raw_fields, sheet, header_row_index + 1, source_row_number, "customs_bcct")
        row["direction"] = direction
        row["declaration_no"] = declaration_no
        row["line_no"] = line_no
        row["item_code"] = item_code
        row["quantity"] = quantity
        row["unit"] = unit
        row["customs_value"] = normalize_decimal(row.get("customs_value"))
        row["value_currency"] = "VND" if row.get("customs_value") else row.get("currency", "") if row.get("foreign_currency_value") else ""
        row["transaction_key"] = transaction_key(row)
        rows.append(row)
    if not rows:
        raise SourceParseError("Workbook has no BCCT rows.")
    assert_unique_keys(rows, "transaction_key", allow_identical=True)
    return rows
def load_tabular_sheets(content: bytes, module: str) -> list[dict]:
    workbook_content, source_file = select_workbook_content(content, module)
    if workbook_content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return load_xls_sheets(workbook_content, source_file)
    return load_xlsx_sheets(workbook_content, source_file)
def select_workbook_content(content: bytes, module: str) -> tuple[bytes, str]:
    if not content.startswith(b"PK"):
        return content, "upload.xls"
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
            if "[Content_Types].xml" in names and any(name.startswith("xl/") for name in names):
                return content, "upload.xlsx"
            workbook_names = [
                name for name in names
                if not name.endswith("/") and Path(name).suffix.lower() in {".xls", ".xlsx", ".xlsm"}
            ]
            if not workbook_names:
                raise SourceParseError("ZIP upload has no Excel workbook.")
            selected_name = select_workbook_name(workbook_names, module)
            return archive.read(selected_name), selected_name
    except BadZipFile as exc:
        raise SourceParseError(f"Cannot read ZIP upload: {exc}") from exc
def select_workbook_name(names: list[str], module: str) -> str:
    scored = []
    for name in names:
        key = filename_key(Path(name).name)
        score = 0
        if module == "material_catalog" and any(token in key for token in ["npl", "nvl"]):
            score += 10
        if module == "product_catalog" and re.search(r"(^|_)sp($|_)", key):
            score += 10
        if module == "bcct" and any(token in key for token in ["baocaohangchitiet", "bao_cao_hang_chi_tiet", "bcct"]):
            score += 10
        if module == "bcct" and Path(name).suffix.lower() in {".xlsx", ".xlsm"}:
            score += 1
        if score:
            scored.append((score, name))
    if scored:
        return sorted(scored, reverse=True)[0][1]
    if len(names) == 1:
        return names[0]
    raise SourceParseError(f"ZIP upload has multiple workbooks; cannot choose one for {module}.")
def filename_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", strip_accents(value).lower()).strip("_")
def load_xls_sheets(content: bytes, source_file: str) -> list[dict]:
    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise SourceParseError(f"Cannot read .xls workbook: {exc}") from exc
    sheets = []
    for sheet in book.sheets():
        rows = [
            [sheet.cell_value(row_index, column_index) for column_index in range(sheet.ncols)]
            for row_index in range(sheet.nrows)
        ]
        sheets.append({"title": sheet.name, "rows": rows, "source_file": source_file})
    return sheets
def load_xlsx_sheets(content: bytes, source_file: str) -> list[dict]:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise SourceParseError(f"Cannot read .xlsx workbook: {exc}") from exc
    return [
        {
            "title": worksheet.title,
            "rows": [list(row) for row in worksheet.iter_rows(values_only=True)],
            "source_file": source_file,
        }
        for worksheet in workbook.worksheets
    ]
def select_catalog_sheet(sheets: list[dict], module: str) -> dict:
    preferred_titles = {"NVL"} if module == "material_catalog" else {"TP"}
    for sheet in sheets:
        if sheet["title"] in preferred_titles:
            return sheet
    return first_sheet_with_rows(sheets)
def select_bcct_sheet(sheets: list[dict]) -> dict:
    for sheet in sheets:
        try:
            find_header_row(
                sheet["rows"],
                "bcct",
                {"declaration_no", "declaration_type", "line_no", "item_code", "quantity", "unit"},
            )
            return sheet
        except SourceParseError:
            continue
    return first_sheet_with_rows(sheets)
def first_sheet_with_rows(sheets: list[dict]) -> dict:
    for sheet in sheets:
        if any(any(cell_text(value) for value in row) for row in sheet["rows"]):
            return sheet
    raise SourceParseError("Workbook has no non-empty sheets.")
def find_header_row(rows: list[list], module: str, required_headers: set[str]) -> tuple[int, list[str], list[str]]:
    for row_index, row in enumerate(rows[:30]):
        raw_headers = [cell_text(value) for value in row]
        headers = dedupe_headers([normalize_header(value, module) for value in raw_headers])
        if required_headers.issubset(set(headers)):
            return row_index, headers, raw_headers
    raise SourceParseError("Cannot find expected customs header row.")
def dedupe_headers(headers: list[str]) -> list[str]:
    counts: dict[str, int] = {}
    output = []
    for header in headers:
        if not header:
            output.append("")
            continue
        counts[header] = counts.get(header, 0) + 1
        output.append(header if counts[header] == 1 else f"{header}_{counts[header]}")
    return output
def iter_tabular_rows(rows: list[list], headers: list[str], raw_headers: list[str], first_data_row_index: int):
    for row_index, row in enumerate(rows[first_data_row_index:], start=first_data_row_index + 1):
        values = {}
        raw_fields = {}
        for index, header in enumerate(headers):
            raw_header = raw_headers[index] if index < len(raw_headers) else ""
            if not header and not raw_header:
                continue
            value = row[index] if index < len(row) else ""
            text_value = cell_text(value)
            if header:
                values[header] = text_value
            if raw_header:
                raw_fields[raw_header] = text_value
        if any(raw_fields.values()):
            yield values, raw_fields, row_index
def attach_source_fields(
    row: dict,
    values: dict,
    raw_fields: dict,
    sheet: dict,
    header_row_number: int,
    source_row_number: int,
    source_schema: str,
) -> None:
    for field, value in values.items():
        if source_schema in {"customs_material_catalog", "customs_product_catalog"} and field in CATALOG_CASE_RULE_FIELDS:
            continue
        if field in row or field in {"customs_code", "product_code", "name", "description"}:
            continue
        row[field] = normalize_source_value(field, value)
    row.update({
        "source_schema": source_schema,
        "source_file": sheet["source_file"],
        "source_sheet": sheet["title"],
        "source_header_row": header_row_number,
        "source_row_number": source_row_number,
        "raw_fields": raw_fields,
    })
def normalize_source_value(field: str, value) -> str:
    if field in {
        "quantity",
        "secondary_quantity",
        "unit_price",
        "taxable_unit_price",
        "foreign_currency_value",
        "customs_value",
        "exchange_rate",
        "package_quantity",
        "gross_weight",
        "import_export_tax_amount",
        "import_export_tax_exempt_amount",
        "safeguard_tax_amount",
        "trade_remedy_tax_amount",
        "excise_tax_amount",
        "environmental_tax_amount",
        "vat_tax_amount",
        "total_tax_amount",
    }:
        return normalize_decimal(value)
    if field in {"unit", "secondary_unit"}:
        return normalize_unit(value)
    return cell_text(value)
def infer_direction(declaration_type: str) -> str:
    code = cell_text(declaration_type).upper()
    if not code:
        return ""
    if code in DECLARATION_TYPE_DIRECTIONS:
        return DECLARATION_TYPE_DIRECTIONS[code]
    if code.startswith("A"):
        return "import"
    if code.startswith("B"):
        return "export"
    return ""
def assert_unique_keys(rows: list[dict], field: str, allow_identical: bool = False) -> None:
    seen = {}
    conflicts = []
    for row in rows:
        key = row[field]
        prior = seen.get(key)
        if prior is None:
            seen[key] = row
            continue
        if allow_identical and normalized_row(prior) == normalized_row(row):
            continue
        conflicts.append(key)
    if conflicts:
        raise SourceParseError(f"Duplicate row keys: {', '.join(conflicts[:5])}")
def transaction_key(row: dict) -> str:
    return "||".join([row["direction"], row["declaration_no"], row["line_no"], row["item_code"]])
def normalize_header(value: str, module: str = "") -> str:
    header = header_key(value)
    aliases = COMMON_HEADER_ALIASES
    if module == "material_catalog":
        aliases = MATERIAL_HEADER_ALIASES
    elif module == "product_catalog":
        aliases = PRODUCT_HEADER_ALIASES
    elif module == "bcct":
        aliases = BCCT_HEADER_ALIASES
    return aliases.get(header, header)
def header_key(value: str) -> str:
    text = strip_accents(cell_text(value)).lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")
def strip_accents(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.replace("đ", "d").replace("Đ", "D"))
    return "".join(char for char in normalized if not unicodedata.combining(char))
def normalize_status(value) -> str:
    status = cell_text(value)
    return status or "active"
def normalize_direction(value) -> str:
    direction = cell_text(value)
    return DIRECTION_ALIASES.get(direction.upper(), direction.lower())
def normalize_declaration_type(value) -> str:
    return cell_text(value).upper()
def normalize_unit(value) -> str:
    unit = cell_text(value).upper()
    return UNIT_ALIASES.get(unit, unit)
def normalized_row(row: dict) -> dict:
    normalized = {}
    for key, value in sorted(row.items()):
        if key in SOURCE_AUDIT_FIELDS:
            continue
        if isinstance(value, dict):
            normalized[key] = value
            continue
        text = cell_text(value)
        if text:
            normalized[key] = text
    return normalized
