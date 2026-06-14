"""system_template — the canonical year-end inventory snapshot workbook.

One data sheet "Tồn kho": header row 1, data row 2+. Snapshot date + client
come from the upload form. Single source of truth for the canonical snapshot
column set: renders AND parses it.
"""
from __future__ import annotations

import io

from app.parsers._excel import (
    cell_num, cell_str, header_row, index_headers, load_xlsx,
)
from app.parsers.inventory_adapters._common import ALIASES, NUMERIC_FIELDS

# (header text written into the template, logical field). "Chênh lệch" is shown
# for human convenience but derived on read, so it is NOT a parsed input column.
CANONICAL_COLUMNS: list[tuple[str, str]] = [
    ("Mã", "code"),
    ("Tên", "name"),
    ("ĐVT", "uom"),
    ("Kho", "warehouse"),
    ("Lô/Batch", "batch"),
    ("SL sổ sách", "qty_book"),
    ("SL thực đếm", "qty_physical"),
    ("Ghi chú", "note"),
]
DATA_SHEET = "Tồn kho"


class SystemTemplateInventoryAdapter:
    name = "system_template"
    label_key = "inventory.adapter.system_template.label"
    description_key = "inventory.adapter.system_template.desc"
    supports_mapping_override = False

    def detect(self, blob: bytes) -> float | None:
        try:
            wb = load_xlsx(blob)
        except Exception:
            return None
        for ws in wb.worksheets:
            hdr = header_row(ws, aliases=ALIASES)
            if not hdr:
                continue
            cols = index_headers(hdr[1], ALIASES)
            if "code" in cols and ("qty_book" in cols or "qty_physical" in cols):
                return 0.9
        return None

    def parse(self, blob: bytes, *,
              mapping_override: dict[str, str] | None = None) -> list[dict]:
        from app.parsers.inventory_adapters import InventoryParseError
        try:
            wb = load_xlsx(blob)
        except Exception as e:
            raise InventoryParseError(f"Cannot open workbook: {e}") from e

        lines: list[dict] = []
        for ws in wb.worksheets:
            hdr = header_row(ws, aliases=ALIASES)
            if not hdr:
                continue
            header_idx, headers = hdr
            cols = index_headers(headers, ALIASES)
            if "code" not in cols:
                continue
            for raw in _iter_data_rows(ws, header_idx):
                line = _row_to_line(raw, cols, ws.title)
                if line is not None:
                    lines.append(line)

        if not lines:
            raise InventoryParseError("No inventory rows recognized.")
        return lines


def _iter_data_rows(ws, header_row_idx: int):
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        yield row


def _row_to_line(raw, cols: dict[str, int], sheet_title: str) -> dict | None:
    code = cell_str(raw, cols.get("code"))
    if not code:
        return None
    line = {
        "code": code,
        "name": cell_str(raw, cols.get("name")),
        "uom": cell_str(raw, cols.get("uom")),
        # Sheet title is a useful warehouse fallback (DKE = one sheet per kho).
        "warehouse": cell_str(raw, cols.get("warehouse")) or sheet_title or None,
        "batch": cell_str(raw, cols.get("batch")),
        "note": cell_str(raw, cols.get("note")),
    }
    for f in NUMERIC_FIELDS:
        line[f] = cell_num(raw, cols.get(f))
    return line


def render_template_xlsx() -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    headers = [h for h, _ in CANONICAL_COLUMNS]
    # Show the derived chênh lệch column for human readers (ignored on parse).
    display_headers = headers[:-1] + ["Chênh lệch", "Ghi chú"]
    for ci, h in enumerate(display_headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    samples = [
        ["NVL-001", "Nhựa ABS", "KG", "Kho NVL", "LOT-2025-001", 1500, 1495, "", "Hụt 5"],
        ["TP-001", "Bộ biến tần 5kW", "PCE", "Kho TP", "", 20, 20, "", ""],
    ]
    for ri, sample in enumerate(samples, start=2):
        for ci, v in enumerate(sample, 1):
            ws.cell(row=ri, column=ci, value=v)
    for ci in range(1, len(display_headers) + 1):
        ws.column_dimensions[chr(64 + ci)].width = 18

    inst = wb.create_sheet("Hướng dẫn")
    inst["A1"] = "Mẫu chốt tồn kho cuối kỳ — chuẩn hệ thống"
    inst["A1"].font = Font(size=13, bold=True)
    lines = [
        "",
        "Một sheet 'Tồn kho': dòng 1 tiêu đề cột, dữ liệu từ dòng 2.",
        "Cột: " + " | ".join(h for h, _ in CANONICAL_COLUMNS),
        "",
        "Quy ước:",
        "• Mã là bắt buộc.",
        "• SL sổ sách = số theo sổ; SL thực đếm = số kiểm kê thực tế.",
        "• Chênh lệch (= thực đếm − sổ sách) do hệ thống tự tính — cột trong file chỉ để xem.",
        "• Ngày chốt chọn ở màn hình tải lên, không nhập trong file.",
        "• Nếu mỗi kho một sheet: để trống cột Kho, hệ thống lấy tên sheet làm kho.",
    ]
    for ri, line in enumerate(lines, start=2):
        inst.cell(row=ri, column=1, value=line).alignment = Alignment(
            wrap_text=True, vertical="top")
    inst.column_dimensions["A"].width = 100

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
