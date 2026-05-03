"""Auto-seed demo clients + sample data on first run."""
from __future__ import annotations

from app.database import connect
from app.parsers.bcct import parse_bcct_workbook
from app.parsers.bom import parse_bom_workbook
from app.parsers.code_mappings import parse_code_mappings_workbook
from app.parsers.goods_name import internal_code_parser_for
from app.parsers.materials import parse_materials_workbook
from app.routes.clients import slug, upsert_client
from app.routes.bcct import _insert_bcct
from app.routes.bqd import _insert_mappings
from app.routes.catalog import _insert_materials
from app.stores.bom import create_version


def auto_seed_demo_if_empty() -> str:
    """If no clients exist, seed two demo clients (Growatt + Johnson) with sample data.
    Returns a description of what was seeded, or empty string if skipped."""
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute("select count(*) from hub.clients")
            (n,) = cur.fetchone()
    if n > 0:
        return ""

    growatt_id = "growatt-vn"
    upsert_client(
        client_id=growatt_id, name="Growatt VN", tax_code="0307123456",
        code_resolution_mode="batch_aggregate_resolution",
        bom_proposal_mode="auto", bom_proposal_qty_tolerance_pct=5.0,
        bom_approver_tier="edit",
        status="active", notes="Reference Growatt — 1:n BQD with BCCT-aggregate disambiguation.",
    )

    johnson_id = "johnson-vn"
    upsert_client(
        client_id=johnson_id, name="Johnson VN", tax_code="0309876543",
        code_resolution_mode="identity",
        bom_proposal_mode="auto", bom_proposal_qty_tolerance_pct=5.0,
        bom_approver_tier="edit",
        status="active", notes="Identity mode — customs_code IS internal_code.",
    )

    # Seed Growatt with full data
    _seed_growatt(growatt_id)
    _seed_johnson(johnson_id)

    return f"2 clients ({growatt_id}, {johnson_id})"


def _seed_growatt(client_id: str) -> None:
    from openpyxl import Workbook
    import io

    # Materials
    wb = Workbook(); ws = wb.active; ws.title = "DM NVL"
    ws.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    materials = [
        ("PE-001", "PE-001", "Polyethylene resin grade A", "nvl", "kg", "39011010"),
        ("PE-002", "PE-002", "Polyethylene resin grade B", "nvl", "kg", "39011010"),
        ("AL-100", "AL-100", "Aluminum sheet 1mm", "nvl", "kg", "76061110"),
        ("CU-WIRE-2", "CU-WIRE-2", "Copper wire 2.5mm", "nvl", "m", "85447000"),
        ("PCB-12", "PCB-12", "Printed circuit board 12-layer", "nvl", "pcs", "85340010"),
        ("HEATSINK-A", "HEATSINK-A", "Aluminum heatsink type A", "nvl", "pcs", "76161000"),
        ("CASE-INV", "CASE-INV", "Inverter outer case ABS", "nvl", "pcs", "39269097"),
    ]
    for r in materials:
        ws.append(r)
    ws_tp = wb.create_sheet("DM TP")
    ws_tp.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    for r in [
        ("INV-3000", "INV-3000", "Solar inverter 3kW", "tp", "pcs", "85044090"),
        ("INV-5000", "INV-5000", "Solar inverter 5kW", "tp", "pcs", "85044090"),
        ("INV-10K", "INV-10K", "Solar inverter 10kW", "tp", "pcs", "85044090"),
    ]:
        ws_tp.append(r)
    buf = io.BytesIO(); wb.save(buf)
    _insert_materials(client_id=client_id, rows=parse_materials_workbook(buf.getvalue()))

    # BQD (with 1:n example for PE-001)
    wb = Workbook(); ws = wb.active; ws.title = "BQD NVL"
    ws.append(["Mã nội bộ", "Mã hải quan"])
    for ic, cc in [
        ("PE-001", "PE-001"),
        ("PE-001", "PE-ALT-FALLBACK"),
        ("PE-002", "PE-002"),
        ("AL-100", "AL-100"),
        ("CU-WIRE-2", "CU-WIRE-2"),
        ("PCB-12", "PCB-12"),
        ("HEATSINK-A", "HEATSINK-A"),
        ("CASE-INV", "CASE-INV"),
    ]:
        ws.append((ic, cc))
    buf = io.BytesIO(); wb.save(buf)
    _insert_mappings(client_id=client_id, rows=parse_code_mappings_workbook(buf.getvalue()))

    # BCCT (Growatt-style #&-prefix in goods name)
    wb = Workbook(); ws = wb.active; ws.title = "BCCT"
    ws.append(["Số tờ khai", "Dòng", "Mã loại hình", "Ngày đăng ký", "Mã NPL/SP",
               "Tên hàng", "HS", "Tổng số lượng", "ĐVT", "Trị giá", "Nguyên tệ"])
    bcct_rows = [
        ("104100100001", 1, "E11", "2025-02-12", "PE-001",
         "PE-001#&Polyethylene resin grade A premium quality", "39011010", 1500.0, "kg", 3750.0, "USD"),
        ("104100100001", 2, "E11", "2025-02-12", "PE-002",
         "PE-002#&Polyethylene resin grade B standard", "39011010", 800.0, "kg", 1760.0, "USD"),
        ("104100100002", 1, "E11", "2025-03-08", "AL-100",
         "AL-100#&Aluminum sheet 1mm rolled", "76061110", 600.0, "kg", 4800.0, "USD"),
        ("104100100003", 1, "E15", "2025-03-22", "CU-WIRE-2",
         "CU-WIRE-2#&Copper wire 2.5mm 100m roll", "85447000", 250.0, "m", 875.0, "USD"),
        ("104100100004", 1, "E11", "2025-04-15", "PCB-12",
         "PCB-12#&Printed circuit board 12-layer 200x150", "85340010", 500.0, "pcs", 12500.0, "USD"),
        ("104100100005", 1, "E11", "2025-05-20", "HEATSINK-A",
         "HEATSINK-A#&Aluminum heatsink type A 100x80", "76161000", 200.0, "pcs", 1800.0, "USD"),
        ("104100100006", 1, "E11", "2025-06-04", "CASE-INV",
         "CASE-INV#&Inverter outer case ABS molded", "39269097", 300.0, "pcs", 2400.0, "USD"),
        ("104100100100", 1, "E42", "2025-09-15", "INV-3000",
         "INV-3000#&Solar inverter 3kW model 2025", "85044090", 80.0, "pcs", 20000.0, "USD"),
        ("104100100101", 1, "E42", "2025-10-08", "INV-5000",
         "INV-5000#&Solar inverter 5kW model 2025", "85044090", 60.0, "pcs", 21000.0, "USD"),
        ("104100100102", 1, "E42", "2025-11-12", "INV-10K",
         "INV-10K#&Solar inverter 10kW commercial", "85044090", 30.0, "pcs", 18000.0, "USD"),
    ]
    for r in bcct_rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    parsed = parse_bcct_workbook(buf.getvalue())
    # Insert stub upload first so the FK is valid.
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                insert into hub.file_uploads (upload_id, client_id, module,
                  original_filename, stored_path, content_sha256, size_bytes,
                  parse_status, row_count)
                values ('seed-bcct-' || %s, %s, 'bcct', 'seed-bcct.xlsx', 'seed/bcct',
                  'seed', 0, 'done', %s)
                on conflict (upload_id) do nothing
                """,
                (client_id, client_id, len(bcct_rows)),
            )
    parser = internal_code_parser_for(client_id, "batch_aggregate_resolution")
    _insert_bcct(client_id=client_id, year=2025, rows=parsed,
                 upload_id=f"seed-bcct-{client_id}", parser=parser)

    # BOM — 3 finished products + 1 dual-source sub-assembly (HEATSINK-A: vừa import vừa tự sản xuất)
    wb = Workbook(); ws = wb.active; ws.title = "BOM"
    ws.append(["Mã SP", "Mã NVL", "Định mức", "ĐVT"])
    bom_rows = [
        ("INV-3000", "PE-001", 0.45, "kg"),
        ("INV-3000", "AL-100", 0.30, "kg"),
        ("INV-3000", "PCB-12", 1.0, "pcs"),
        ("INV-3000", "HEATSINK-A", 1.0, "pcs"),
        ("INV-3000", "CASE-INV", 1.0, "pcs"),
        ("INV-5000", "PE-001", 0.60, "kg"),
        ("INV-5000", "PE-002", 0.20, "kg"),
        ("INV-5000", "AL-100", 0.50, "kg"),
        ("INV-5000", "CU-WIRE-2", 1.20, "m"),
        ("INV-5000", "PCB-12", 1.0, "pcs"),
        ("INV-5000", "HEATSINK-A", 1.0, "pcs"),
        ("INV-5000", "CASE-INV", 1.0, "pcs"),
        ("INV-10K", "PE-001", 1.10, "kg"),
        ("INV-10K", "AL-100", 0.90, "kg"),
        ("INV-10K", "CU-WIRE-2", 2.50, "m"),
        ("INV-10K", "PCB-12", 2.0, "pcs"),
        ("INV-10K", "HEATSINK-A", 2.0, "pcs"),
        ("INV-10K", "CASE-INV", 1.0, "pcs"),
        # Dual-source: HEATSINK-A có thể vừa nhập (NVL trong BCCT) vừa tự đúc từ AL-100.
        # Khi BOM có HEATSINK-A là product_code, settlement nhận diện dual-source.
        ("HEATSINK-A", "AL-100", 0.18, "kg"),
    ]
    for r in bom_rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    products = parse_bom_workbook(buf.getvalue(), profile="manual_flat")
    for product_code, rows in products.items():
        create_version(
            client_id=client_id, product_code=product_code, rows=rows,
            actor="agency_staff", intent="asserted_technical",
            parent_version_id=None,
            context={"channel": "agency_upload", "profile": "manual_flat", "seed": True},
            source_upload_id=None,
        )

    # Submit a couple of BOM proposals to populate audit trail
    from app.stores.bom import submit_proposal
    parent_version = None
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "select version_id from hub.bom_versions where client_id=%s and product_code='INV-3000' order by version_no desc limit 1",
                (client_id,),
            )
            row = cur.fetchone()
            if row:
                parent_version = row[0]
    if parent_version:
        # auto-approve sample
        submit_proposal(
            client_id=client_id, product_code="INV-3000",
            actor="co_system", intent="modified_for_case",
            parent_version_id=parent_version,
            context={"case_id": "CO-2025-DEMO-001", "trigger": "rvc_threshold_pass"},
            rows=[
                {"material_code": "PE-001", "qty_per_unit": 0.46, "uom": "kg"},
                {"material_code": "AL-100", "qty_per_unit": 0.30, "uom": "kg"},
                {"material_code": "PCB-12", "qty_per_unit": 1.0, "uom": "pcs"},
                {"material_code": "HEATSINK-A", "qty_per_unit": 1.0, "uom": "pcs"},
                {"material_code": "CASE-INV", "qty_per_unit": 1.0, "uom": "pcs"},
            ],
        )
        # auto-reject sample (qty exceeds tolerance)
        submit_proposal(
            client_id=client_id, product_code="INV-3000",
            actor="co_system", intent="modified_for_case",
            parent_version_id=parent_version,
            context={"case_id": "CO-2025-DEMO-002"},
            rows=[
                {"material_code": "PE-001", "qty_per_unit": 0.80, "uom": "kg"},  # +77% — too big
                {"material_code": "AL-100", "qty_per_unit": 0.30, "uom": "kg"},
                {"material_code": "PCB-12", "qty_per_unit": 1.0, "uom": "pcs"},
                {"material_code": "HEATSINK-A", "qty_per_unit": 1.0, "uom": "pcs"},
                {"material_code": "CASE-INV", "qty_per_unit": 1.0, "uom": "pcs"},
            ],
        )


def _seed_johnson(client_id: str) -> None:
    """Identity-mode client with smaller dataset."""
    from openpyxl import Workbook
    import io

    wb = Workbook(); ws = wb.active; ws.title = "DM NVL"
    ws.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    for r in [
        ("3923301010", "3923301010", "Plastic bottle 500ml", "nvl", "pcs", "39233010"),
        ("4819100010", "4819100010", "Cardboard box 30x20x15", "nvl", "pcs", "48191000"),
        ("3402900090", "3402900090", "Cleaning agent surfactant", "nvl", "kg", "34029000"),
    ]:
        ws.append(r)
    ws_tp = wb.create_sheet("DM TP")
    ws_tp.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    ws_tp.append(("3401300090", "3401300090", "Liquid soap 500ml retail", "tp", "pcs", "34013000"))
    buf = io.BytesIO(); wb.save(buf)
    _insert_materials(client_id=client_id, rows=parse_materials_workbook(buf.getvalue()))

    # No BQD needed (identity mode)
    # Tiny BCCT
    wb = Workbook(); ws = wb.active; ws.title = "BCCT"
    ws.append(["Số tờ khai", "Dòng", "Mã loại hình", "Ngày đăng ký", "Mã NPL/SP",
               "Tên hàng", "HS", "Tổng số lượng", "ĐVT", "Trị giá", "Nguyên tệ"])
    for r in [
        ("104200200001", 1, "E11", "2025-04-10", "3923301010",
         "Plastic bottle 500ml HDPE white", "39233010", 10000.0, "pcs", 1500.0, "USD"),
        ("104200200002", 1, "E11", "2025-05-22", "3402900090",
         "Cleaning agent surfactant industrial grade", "34029000", 500.0, "kg", 2500.0, "USD"),
        ("104200200003", 1, "E42", "2025-07-15", "3401300090",
         "Liquid soap 500ml retail finished", "34013000", 5000.0, "pcs", 3500.0, "USD"),
    ]:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                insert into hub.file_uploads (upload_id, client_id, module,
                  original_filename, stored_path, content_sha256, size_bytes,
                  parse_status, row_count)
                values ('seed-johnson-bcct', %s, 'bcct', 'seed.xlsx', 'seed/bcct',
                  'seed', 0, 'done', 3)
                on conflict (upload_id) do nothing
                """,
                (client_id,),
            )
    parser = internal_code_parser_for(client_id, "identity")
    _insert_bcct(client_id=client_id, year=2025,
                 rows=parse_bcct_workbook(buf.getvalue()),
                 upload_id="seed-johnson-bcct", parser=parser)

    # 1 BOM
    wb = Workbook(); ws = wb.active; ws.title = "BOM"
    ws.append(["Mã SP", "Mã NVL", "Định mức", "ĐVT"])
    for r in [
        ("3401300090", "3923301010", 1.0, "pcs"),
        ("3401300090", "3402900090", 0.45, "kg"),
        ("3401300090", "4819100010", 0.04, "pcs"),
    ]:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    products = parse_bom_workbook(buf.getvalue(), profile="manual_flat")
    for product_code, rows in products.items():
        create_version(
            client_id=client_id, product_code=product_code, rows=rows,
            actor="agency_staff", intent="asserted_technical",
            parent_version_id=None,
            context={"channel": "agency_upload", "profile": "manual_flat", "seed": True},
            source_upload_id=None,
        )
