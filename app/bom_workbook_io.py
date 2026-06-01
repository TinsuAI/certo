from __future__ import annotations

import re
import unicodedata

from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import load_workbook
from pathlib import Path
from typing import Any


MANUAL_HEADER_ALIASES = {
    "product_code": {"product_code", "export_product_code", "finished_product_code", "ma_tp", "ma_sp", "ma_thanh_pham"},
    "bom_code": {"bom_code", "ma_bom"},
    "bom_variant_id": {"bom_variant_id", "variant", "version", "phien_ban"},
    "material_code": {"material_code", "component_code", "ma_nvl", "ma_vat_tu", "ma_nguyen_lieu"},
    "material_name": {"material_name", "component_name", "ten_nvl", "ten_vat_tu", "ten_nguyen_lieu"},
    "qty_per": {"qty_per", "quantity", "dinh_muc", "so_luong", "standard_usage"},
    "uom": {"uom", "unit", "dvt", "don_vi"},
    "scrap_rate": {"scrap_rate", "waste_rate", "hao_hut"},
}
GROWATT_REQUIRED_HEADERS = {"成品物料", "组件物料", "标准用量"}
JOHNSON_REQUIRED_HEADERS = {"Level", "Explosion level", "Component number"}
def parse_bom_workbook(content: bytes, filename: str, profile: str, upload_mode: str) -> list[dict]:
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise BomParseError(f"Không đọc được BOM workbook: {exc}") from exc

    if upload_mode == "technical_bom" and profile == "growatt_multi_workbook":
        rows = parse_growatt_workbook(wb)
        if rows:
            return rows

    if upload_mode == "technical_bom" and profile == "johnson_sap_exploded":
        rows = parse_johnson_workbook(wb, filename)
        if rows:
            return rows

    return parse_manual_flat_workbook(wb)
def parse_manual_flat_workbook(wb) -> list[dict]:
    ws = wb["BOM"] if "BOM" in wb.sheetnames else wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        raise BomParseError("Workbook BOM không có dữ liệu.")

    headers = [header_key(value) for value in raw_rows[0]]
    field_indexes = {}
    for field, aliases in MANUAL_HEADER_ALIASES.items():
        for index, header in enumerate(headers):
            if header in aliases:
                field_indexes[field] = index
                break

    for required in ["product_code", "material_code", "qty_per", "uom"]:
        if required not in field_indexes:
            raise BomParseError(f"Thiếu cột BOM bắt buộc: {required}.")

    rows = []
    for excel_index, raw in enumerate(raw_rows[1:], start=2):
        values = {field: cell_text(raw[index] if index < len(raw) else "") for field, index in field_indexes.items()}
        if not any(values.values()):
            continue
        if not values.get("product_code") or not values.get("material_code"):
            continue
        rows.append(normalized_bom_row(
            product_code=values["product_code"],
            bom_code=values.get("bom_code") or values["product_code"],
            bom_variant_id=values.get("bom_variant_id") or "default",
            material_code=values["material_code"],
            material_name=values.get("material_name", ""),
            qty_per=values.get("qty_per", "0"),
            uom=values.get("uom", ""),
            scrap_rate=values.get("scrap_rate", ""),
            source=f"{ws.title}:{excel_index}",
            row_class="material_candidate_leaf",
        ))

    if not rows:
        raise BomParseError("Workbook BOM không có dòng hợp lệ.")
    return rows
def parse_growatt_workbook(wb) -> list[dict]:
    rows = []
    for ws in wb.worksheets:
        raw_rows = list(ws.iter_rows(values_only=True))
        if not raw_rows:
            continue
        headers = [cell_text(value) for value in raw_rows[0]]
        if not GROWATT_REQUIRED_HEADERS.issubset(set(headers)):
            continue
        indexes = {header: headers.index(header) for header in headers}
        for excel_index, raw in enumerate(raw_rows[1:], start=2):
            product_code = cell_text(raw[indexes["成品物料"]])
            material_code = cell_text(raw[indexes["组件物料"]])
            if not product_code or not material_code:
                continue
            rows.append(normalized_bom_row(
                product_code=product_code,
                bom_code=product_code,
                bom_variant_id="technical",
                material_code=material_code,
                material_name=cell_text(raw[indexes.get("组件物料描述", -1)] if indexes.get("组件物料描述", -1) >= 0 else ""),
                qty_per=cell_text(raw[indexes["标准用量"]]),
                uom=cell_text(raw[indexes.get("单位", -1)] if indexes.get("单位", -1) >= 0 else ""),
                scrap_rate="",
                source=f"{ws.title}:{excel_index}",
                row_class="needs_graph_flatten_review",
            ))
    return rows
def parse_johnson_workbook(wb, filename: str) -> list[dict]:
    rows = []
    product_code = Path(filename).stem
    for ws in wb.worksheets:
        raw_rows = list(ws.iter_rows(values_only=True))
        if not raw_rows:
            continue
        headers = [cell_text(value) for value in raw_rows[0]]
        if not JOHNSON_REQUIRED_HEADERS.issubset(set(headers)):
            continue
        indexes = {header: headers.index(header) for header in headers}
        levels = [parse_int(cell_text(raw[indexes["Level"]])) for raw in raw_rows[1:]]
        for offset, raw in enumerate(raw_rows[1:]):
            level = levels[offset]
            next_level = levels[offset + 1] if offset + 1 < len(levels) else 0
            is_leaf = next_level <= level
            if not is_leaf:
                continue
            material_code = cell_text(raw[indexes["Component number"]])
            if not material_code:
                continue
            qty = cell_text(raw[indexes.get("Comp. Qty (CUn)", -1)] if indexes.get("Comp. Qty (CUn)", -1) >= 0 else "")
            if not qty:
                qty = cell_text(raw[indexes.get("Component quantity", -1)] if indexes.get("Component quantity", -1) >= 0 else "")
            rows.append(normalized_bom_row(
                product_code=product_code,
                bom_code=product_code,
                bom_variant_id="sap-exploded",
                material_code=material_code,
                material_name=cell_text(raw[indexes.get("Object description", -1)] if indexes.get("Object description", -1) >= 0 else ""),
                qty_per=qty,
                uom=cell_text(raw[indexes.get("Component unit", -1)] if indexes.get("Component unit", -1) >= 0 else ""),
                scrap_rate="",
                source=f"{ws.title}:{offset + 2}",
                row_class="material_candidate_leaf",
            ))
    return rows
def normalized_bom_row(**values) -> dict:
    product_code = cell_text(values["product_code"])
    bom_code = cell_text(values.get("bom_code") or product_code)
    material_code = cell_text(values["material_code"])
    uom = cell_text(values.get("uom", "")).upper()
    return {
        "product_code": product_code,
        "bom_code": bom_code,
        "bom_variant_id": cell_text(values.get("bom_variant_id") or "default"),
        "material_code": material_code,
        "material_name": cell_text(values.get("material_name", "")),
        "qty_per": quantity_text(values.get("qty_per", "0")),
        "uom": uom,
        "scrap_rate": cell_text(values.get("scrap_rate", "")),
        "source": cell_text(values.get("source", "")),
        "row_class": cell_text(values.get("row_class", "material_candidate_leaf")),
        "row_key": "||".join([product_code, bom_code, material_code, uom]),
    }
def cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
def quantity_text(value: Any) -> str:
    text = cell_text(value).replace(",", "")
    if not text:
        return "0"
    try:
        decimal = Decimal(text)
    except InvalidOperation:
        return text
    normalized = decimal.normalize()
    return format(normalized, "f")
def parse_int(value: str) -> int:
    try:
        return int(Decimal(value))
    except InvalidOperation:
        return 0
def header_key(value: Any) -> str:
    text = unicodedata.normalize("NFKD", cell_text(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
