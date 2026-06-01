from __future__ import annotations


from fastapi import APIRouter
from app.web.client_context import resolve_client
from app.web.templating import templates
from decimal import Decimal, InvalidOperation
from fastapi import File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, Response


router = APIRouter()

def _decimal_str(value: Decimal) -> str:
    """Render Decimal for the admin form: strip trailing zeros but keep at
    least one digit. Empty for zero (so the placeholder shows)."""
    if value is None or value == 0:
        return ""
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text
def _row_for_template(row) -> dict:
    return {
        "product_code": row.product_code,
        "coef_wages_str": _decimal_str(row.coef_wages),
        "coef_welfare_str": _decimal_str(row.coef_welfare),
        "coef_rent_str": _decimal_str(row.coef_rent),
        "coef_depreciation_str": _decimal_str(row.coef_depreciation),
        "coef_other_mfg_str": _decimal_str(row.coef_other_mfg),
        "coef_transport_storage_str": _decimal_str(row.coef_transport_storage),
        "note": row.note,
        "has_value": any([
            row.coef_wages, row.coef_welfare, row.coef_rent,
            row.coef_depreciation, row.coef_other_mfg, row.coef_transport_storage,
            row.note,
        ]),
    }
def _cost_allocation_context(client_id: str, **extra) -> dict:
    """Lightweight context for the cost-allocation admin page.

    Deliberately avoids `client_context()` because that helper pulls the full
    source workspace (BCCT scan ~2.6s on Growatt) which this page does not
    need. We render the nav with no counts; the rest of the template only
    needs client identity + the ratio rows.
    """
    from app import cost_allocation_store
    from app.cost_allocation_store import CostAllocationRow
    client = resolve_client(client_id)
    rows = cost_allocation_store.list_ratios(client_id)
    mode_b = cost_allocation_store.get_mode_b_default(client_id) or CostAllocationRow(product_code="")
    return {
        "client": client,
        "active": "cost-allocation",
        "rows": [_row_for_template(r) for r in sorted(rows, key=lambda r: r.product_code)],
        "mode_b": _row_for_template(mode_b),
        **extra,
    }
def _coef_from_form(form, key: str) -> Decimal:
    raw = str(form.get(key) or "").strip().replace(",", ".")
    if not raw:
        return Decimal(0)
    try:
        value = Decimal(raw)
    except (InvalidOperation, ValueError):
        return Decimal(0)
    return value if value >= 0 else Decimal(0)
@router.get("/clients/{client_id}/cost-allocation", response_class=HTMLResponse)
async def cost_allocation_page(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id),
    )
@router.post("/clients/{client_id}/cost-allocation/mode-b", response_class=HTMLResponse)
async def cost_allocation_save_mode_b(request: Request, client_id: str):
    resolve_client(client_id)  # validates client exists
    from app import cost_allocation_store
    form = await request.form()
    row = cost_allocation_store.CostAllocationRow(
        product_code="",
        coef_wages=_coef_from_form(form, "coef_wages"),
        coef_welfare=_coef_from_form(form, "coef_welfare"),
        coef_rent=_coef_from_form(form, "coef_rent"),
        coef_depreciation=_coef_from_form(form, "coef_depreciation"),
        coef_other_mfg=_coef_from_form(form, "coef_other_mfg"),
        coef_transport_storage=_coef_from_form(form, "coef_transport_storage"),
        note=str(form.get("note", "") or "").strip(),
    )
    cost_allocation_store.upsert_ratio(client_id, row)
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id, message="Đã lưu hệ số mặc định Mode B."),
    )
@router.post("/clients/{client_id}/cost-allocation/mode-b/delete", response_class=HTMLResponse)
async def cost_allocation_delete_mode_b(request: Request, client_id: str):
    resolve_client(client_id)
    from app import cost_allocation_store
    cost_allocation_store.delete_ratio(client_id, "")
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id, message="Đã xóa hệ số mặc định Mode B."),
    )
@router.post("/clients/{client_id}/cost-allocation/row/delete", response_class=HTMLResponse)
async def cost_allocation_delete_row(request: Request, client_id: str, product_code: str = Form(...)):
    resolve_client(client_id)
    from app import cost_allocation_store
    cost_allocation_store.delete_ratio(client_id, product_code.strip())
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id, message=f"Đã xóa hệ số cho {product_code}."),
    )
@router.post("/clients/{client_id}/cost-allocation/upload", response_class=HTMLResponse)
async def cost_allocation_upload(request: Request, client_id: str, file: UploadFile = File(...)):
    resolve_client(client_id)
    from app import cost_allocation_store, cost_allocation_importer
    try:
        parsed = cost_allocation_importer.parse_excel(await file.read())
    except Exception as exc:  # noqa: BLE001
        return templates.TemplateResponse(
            request=request,
            name="cost_allocation.html",
            status_code=400,
            context=_cost_allocation_context(client_id, error=f"Không đọc được file: {exc}"),
        )
    diff = cost_allocation_store.replace_all(client_id, parsed)
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(
            client_id,
            message=f"Đã import {len(parsed)} dòng từ {file.filename}.",
            upload_diff=diff,
        ),
    )
@router.get("/clients/{client_id}/cost-allocation/template.xlsx")
async def cost_allocation_template(client_id: str):
    resolve_client(client_id)
    from openpyxl import Workbook
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet3"
    ws["A1"] = "BẢNG PHÂN BỔ TỶ LỆ CHI PHÍ"
    ws["A2"] = "STT"
    ws["B2"] = "Mã SP"
    ws["C2"] = "Lương, thưởng"
    ws["D2"] = "Phúc lợi y tế"
    ws["E2"] = "Phí thuê nhà xưởng"
    ws["F2"] = "Phí khấu hao, BH, BD"
    ws["G2"] = "CP SX chung khác"
    ws["H2"] = "Lợi nhuận (bỏ qua khi import)"
    ws["I2"] = "Vận chuyển, lưu kho, dịch vụ"
    ws["J2"] = "Ghi chú"
    # Row 4 onward = data area.
    ws["A4"] = 1
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="cost-allocation-{client_id}-template.xlsx"'},
    )
@router.get("/clients/{client_id}/cost-allocation/resolve")
async def cost_allocation_resolve(client_id: str, product_code: str, fob: str = "0"):
    """JSON endpoint for the 'Áp hệ số' button on the origin product panel.

    Returns the resolved coefficient (Mode A → Mode B fallback) plus the
    multiplied detail values. `found` is false when neither mode matches.
    """
    resolve_client(client_id)
    from app import cost_allocation_store, cost_allocation_importer
    row = cost_allocation_store.get_ratio(client_id, product_code.strip())
    if row is None:
        return JSONResponse({"found": False, "product_code": product_code})
    try:
        fob_dec = Decimal(str(fob).replace(",", "."))
    except (InvalidOperation, ValueError):
        fob_dec = Decimal(0)
    detail = cost_allocation_importer.apply_to_fob(row, fob_dec)
    return JSONResponse({
        "found": True,
        "product_code": product_code,
        "matched_mode": "A" if row.product_code else "B",
        "matched_product_code": row.product_code,
        "fob": str(fob_dec),
        "coefficients": {
            "wages": str(row.coef_wages),
            "welfare": str(row.coef_welfare),
            "rent": str(row.coef_rent),
            "depreciation": str(row.coef_depreciation),
            "other_mfg": str(row.coef_other_mfg),
            "transport_storage": str(row.coef_transport_storage),
        },
        "details": {k: str(v) for k, v in detail.items()},
        "note": row.note,
    })
