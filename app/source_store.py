from __future__ import annotations

import fcntl
import hashlib
import json
import mimetypes
import os
import re
import tempfile
import unicodedata
import uuid
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
import xlrd

from app.client_config_store import get_client_config, resolve_allocation_code
from app.co_stock_derivation import (
    aggregate_co_stock_rows,
    cell_text,
    co_stock_fx_fields,
    co_stock_rows_from_bcct,
    co_stock_value_fields,
    decimal_text_value,
    decimal_to_text,
    display_bcct_row,
    first_normalized_decimal,
    import_row_id,
    normalize_decimal,
    normalize_numeric_text,
    resolve_stock_eligibility,
    review_allocation,
    sum_decimal_text,
    unit_value_from_total,
)
from app.source_workbook_io import (
    BCCT_HEADER_ALIASES,
    CATALOG_CASE_RULE_FIELDS,
    COMMON_HEADER_ALIASES,
    DECLARATION_TYPE_DIRECTIONS,
    DIRECTION_ALIASES,
    MATERIAL_HEADER_ALIASES,
    PRODUCT_HEADER_ALIASES,
    SOURCE_AUDIT_FIELDS,
    UNIT_ALIASES,
    assert_unique_keys,
    attach_source_fields,
    dedupe_headers,
    filename_key,
    find_header_row,
    first_sheet_with_rows,
    header_key,
    infer_direction,
    iter_tabular_rows,
    load_tabular_sheets,
    load_xls_sheets,
    load_xlsx_sheets,
    normalize_declaration_type,
    normalize_direction,
    normalize_header,
    normalize_source_value,
    normalize_status,
    normalize_unit,
    normalized_row,
    parse_bcct_workbook,
    parse_catalog_workbook,
    select_bcct_sheet,
    select_catalog_sheet,
    select_workbook_content,
    select_workbook_name,
    strip_accents,
    transaction_key,
)


CATALOG_MODULES = {"material": "material_catalog", "product": "product_catalog"}
SOURCE_MODULES = {"material_catalog", "product_catalog", "bcct"}
CUSTOMS_MATERIAL_HEADERS = [
    "STT",
    "Mã",
    "Tên",
    "Đơn vị tính",
    "Mã HS",
    "Mục đích SD",
    "Cảnh báo khi tạo tờ khai",
    "Nội dung cảnh báo",
    "Đơn giá",
    "Mã biểu thuế NK",
    "Mã áp dụng thuế TTĐB",
    "Mã áp dụng thuế môi trường",
    "Mã áp dụng thuế VAT",
    "Ghi chú",
    "Tên tiếng anh",
]
CUSTOMS_PRODUCT_HEADERS = [
    "STT",
    "Mã",
    "Tên",
    "Đơn vị tính",
    "Mã HS",
    "Mục đích sử dụng",
    "Mã định danh của lệnh SX",
    "Cảnh báo khi tạo tờ khai",
    "Nội dung cảnh báo",
    "Đơn giá",
    "Mã biểu thuế NK",
    "Mã áp dụng thuế TTĐB",
    "Mã áp dụng thuế môi trường",
    "Mã áp dụng thuế VAT",
    "Ghi chú",
    "Tên tiếng anh",
]
CUSTOMS_BCCT_HEADERS = [
    "STT",
    "Số To Khai",
    "Ngày ĐK",
    "Mã loại hình",
    "Mã địa điểm đích",
    "Tên địa điểm đích cho vận chuyển bảo thuế",
    "Địa điểm dỡ hàng",
    "Mã hiệu PTVC",
    "Ngày khởi hành vận chuyển",
    "Ký hiệu và số hiệu bao bì",
    "Tỷ giá thanh toán",
    "Đơn vị tiền tệ",
    "Số lượng kiện",
    "Mã ĐVT kiện",
    "Trọng lượng",
    "Mã ĐVT trọng lượng",
    "Số quản lý nội bộ",
    "Điều kiện giá hóa đơn",
    "Ghi chú",
    "STT hàng",
    "Mã NPL/SP",
    "Mã HS",
    "Tên hàng",
    "Xuất xứ",
    "Đơn giá",
    "Đơn giá tính thuế",
    "Tổng số lượng",
    "Đơn vị tính",
    "Tổng số lượng 2",
    "Đơn vị tính 2",
    "Trị giá NT",
    "Tổng trị giá",
    "Mã biểu thuế XNK",
    "Thuế suất XNK",
    "Tiền thuế XNK",
    "Số tiền miễn thuế XNK",
    "Thuế suất TV",
    "Tiền thuế TV",
    "Thuế suất PB",
    "Tiền thuế PB",
    "Thuế suất TTĐB",
    "Tiền thuế TTĐB",
    "Thuế suất BVMT",
    "Tiền thuế MT",
    "Thuế suất VA",
    "Tiền thuế VAT",
    "Tổng tiền thuế",
    "Mã doanh nghiệp",
    "Tên doanh nghiệp",
    "Tên đối tác",
    "Số hóa đơn",
    "Ngày hóa đơn",
    "Số hợp đồng",
    "Ngày hợp đồng",
]


class SourceParseError(ValueError):
    pass


def get_source_workspace(client: dict) -> dict:
    material = load_module_state(client, "material_catalog")
    product = load_module_state(client, "product_catalog")
    bcct = load_module_state(client, "bcct")
    client_config = get_client_config(client)
    return {
        "client_config": client_config,
        "material_catalog": module_workspace(material),
        "product_catalog": module_workspace(product),
        "bcct": module_workspace(bcct),
        "co_stock_rows": co_stock_rows_from_bcct(
            bcct["published_rows"],
            client_config,
            customs_fx_rows=_safe_customs_fx_rows(),
        ),
    }


def _safe_customs_fx_rows() -> list[dict]:
    """Load all customs FX rows once per derivation. Empty list on store failure
    so a missing customs FX backend never breaks stock materialization."""
    try:
        from app.customs_fx_store import CUSTOMS_FX_CLIENT_ID, get_customs_fx_store
        return get_customs_fx_store().rows(CUSTOMS_FX_CLIENT_ID)
    except Exception:  # noqa: BLE001 — best-effort; FX backfill is optional
        return []


def get_source_summary(client: dict) -> dict:
    material = load_module_state(client, "material_catalog")
    product = load_module_state(client, "product_catalog")
    bcct = load_module_state(client, "bcct")
    client_config = get_client_config(client)
    return source_summary_from_states(material, product, bcct, client_config)


def source_summary_from_states(material: dict, product: dict, bcct: dict, client_config: dict) -> dict:
    reviewed_bcct_rows = [
        row for row in bcct["published_rows"]
        if row.get("review_status") == "reviewed"
    ]
    return {
        "client_config": client_config,
        "material_catalog": source_module_summary(material),
        "product_catalog": source_module_summary(product),
        "bcct": {
            **source_module_summary(bcct),
            "reviewed_row_count": len(reviewed_bcct_rows),
        },
        "co_stock_row_count": len(co_stock_rows_from_bcct(bcct["published_rows"], client_config)),
    }


def source_module_summary(state: dict) -> dict:
    return {
        "module": state["module"],
        "published_row_count": len(state.get("published_rows", [])),
        "latest_version": dict(state.get("latest_version") or {}),
        "version_count": len(state.get("versions", [])),
        "upload_count": len(state.get("uploads", [])),
        "correction_candidate_count": len(state.get("correction_candidates", [])),
    }


def refresh_source_index_if_configured(client: dict) -> None:
    from app.source_index_store import rebuild_source_index_if_configured

    rebuild_source_index_if_configured(client)


def enrich_client_with_source_workspace(client: dict, workspace: dict) -> dict:
    client["material_catalog"] = [dict(row) for row in workspace["material_catalog"]["published_rows"]]
    client["product_catalog"] = [dict(row) for row in workspace["product_catalog"]["published_rows"]]
    client["bcct_rows"] = [display_bcct_row(row) for row in workspace["bcct"]["published_rows"]]
    client["co_stock"] = [dict(row) for row in workspace["co_stock_rows"]]
    client["counts"] = {
        **client.get("counts", {}),
        "materials": len(client["material_catalog"]),
        "products": len(client["product_catalog"]),
        "bcct": len(client["bcct_rows"]),
        "co_stock": len(client["co_stock"]),
    }
    return client


def enrich_client_with_source_modules(client: dict) -> dict:
    return enrich_client_with_source_workspace(client, get_source_workspace(client))


def attach_case_source_snapshot(case: dict, source_workspace: dict) -> dict:
    material = source_workspace["material_catalog"].get("latest_version") or {}
    product = source_workspace["product_catalog"].get("latest_version") or {}
    bcct = source_workspace["bcct"].get("latest_version") or {}
    reviewed_bcct_rows = [
        row for row in source_workspace["bcct"]["published_rows"]
        if row.get("review_status") == "reviewed"
    ]
    case["source_snapshot"] = {
        "material_catalog_version_id": material.get("version_id", ""),
        "material_catalog_version_no": material.get("version_no", ""),
        "product_catalog_version_id": product.get("version_id", ""),
        "product_catalog_version_no": product.get("version_no", ""),
        "bcct_version_id": bcct.get("version_id", ""),
        "bcct_version_no": bcct.get("version_no", ""),
        "bcct_reviewed_row_count": len(reviewed_bcct_rows),
        "correction_candidate_count": len(source_workspace["bcct"].get("correction_candidates", [])),
        "client_config_version": source_workspace["client_config"].get("config_version", ""),
        "client_config_hash": source_workspace["client_config"].get("config_hash", ""),
    }
    return case


def create_material_catalog_template_workbook(client: dict) -> bytes:
    return create_catalog_template_workbook(client, "material")


def create_product_catalog_template_workbook(client: dict) -> bytes:
    return create_catalog_template_workbook(client, "product")


def create_catalog_template_workbook(client: dict, catalog_type: str) -> bytes:
    module = catalog_module(catalog_type)
    workspace = get_source_workspace(client)[module]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    if module == "material_catalog":
        worksheet.append(CUSTOMS_MATERIAL_HEADERS)
        for index, row in enumerate(workspace["published_rows"], start=1):
            worksheet.append([
                index,
                row.get("customs_code", ""),
                row.get("name", ""),
                row.get("unit", ""),
                row.get("hs_code", ""),
                row.get("purpose", ""),
                row.get("declaration_warning", ""),
                row.get("warning_content", ""),
                row.get("unit_price", ""),
                row.get("import_tariff_code", ""),
                row.get("excise_tax_code", ""),
                row.get("environmental_tax_code", ""),
                row.get("vat_tax_code", ""),
                row.get("note", ""),
                row.get("english_name", ""),
            ])
    else:
        worksheet.append(CUSTOMS_PRODUCT_HEADERS)
        for index, row in enumerate(workspace["published_rows"], start=1):
            worksheet.append([
                index,
                row.get("product_code", ""),
                row.get("name", ""),
                row.get("unit", ""),
                row.get("hs_code", ""),
                row.get("purpose", ""),
                row.get("production_order_identifier", ""),
                row.get("declaration_warning", ""),
                row.get("warning_content", ""),
                row.get("unit_price", ""),
                row.get("import_tariff_code", ""),
                row.get("excise_tax_code", ""),
                row.get("environmental_tax_code", ""),
                row.get("vat_tax_code", ""),
                row.get("note", ""),
                row.get("english_name", ""),
            ])
    return workbook_to_bytes(workbook)


def create_bcct_template_workbook(client: dict) -> bytes:
    workspace = get_source_workspace(client)["bcct"]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append([])
    worksheet.append(["", "BÁO CÁO CHI TIẾT HÀNG HÓA XUẤT NHẬP KHẨU"])
    worksheet.append([])
    worksheet.append(["Đơn vị Hải quan:", "", ""])
    worksheet.append(["Mã doanh nghiệp:", "", ""])
    worksheet.append(["Tổng số dòng hàng:", "", len(workspace["published_rows"])])
    worksheet.append(["Tổng trị giá: ", "", ""])
    worksheet.append(["Tổng tiền thuế:", "", ""])
    worksheet.append([])
    worksheet.append(CUSTOMS_BCCT_HEADERS)
    for index, row in enumerate(workspace["published_rows"], start=1):
        worksheet.append([
            row.get("sequence_no", index),
            row.get("declaration_no", ""),
            row.get("declaration_date", ""),
            row.get("declaration_type", ""),
            row.get("destination_location_code", ""),
            row.get("destination_location_name", ""),
            row.get("unloading_location", ""),
            row.get("transport_mode_code", ""),
            row.get("departure_date", ""),
            row.get("package_marks", ""),
            row.get("exchange_rate", ""),
            row.get("currency", ""),
            row.get("package_quantity", ""),
            row.get("package_unit", ""),
            row.get("gross_weight", ""),
            row.get("gross_weight_unit", ""),
            row.get("internal_management_no", ""),
            row.get("invoice_price_condition", ""),
            row.get("declaration_note", ""),
            row.get("line_no", ""),
            row.get("item_code", ""),
            row.get("hs_code", ""),
            row.get("description", ""),
            row.get("origin_country", ""),
            row.get("unit_price", ""),
            row.get("taxable_unit_price", ""),
            row.get("quantity", ""),
            row.get("unit", ""),
            row.get("secondary_quantity", ""),
            row.get("secondary_unit", ""),
            row.get("foreign_currency_value", ""),
            row.get("customs_value", ""),
            row.get("import_export_tariff_code", ""),
            row.get("import_export_tax_rate", ""),
            row.get("import_export_tax_amount", ""),
            row.get("import_export_tax_exempt_amount", ""),
            row.get("safeguard_tax_rate", ""),
            row.get("safeguard_tax_amount", ""),
            row.get("trade_remedy_tax_rate", ""),
            row.get("trade_remedy_tax_amount", ""),
            row.get("excise_tax_rate", ""),
            row.get("excise_tax_amount", ""),
            row.get("environmental_tax_rate", ""),
            row.get("environmental_tax_amount", ""),
            row.get("vat_tax_rate", ""),
            row.get("vat_tax_amount", ""),
            row.get("total_tax_amount", ""),
            row.get("company_tax_code", ""),
            row.get("company_name", ""),
            row.get("partner_name", ""),
            row.get("invoice_ref", ""),
            row.get("invoice_date", ""),
            row.get("contract_no", ""),
            row.get("contract_date", ""),
        ])
    return workbook_to_bytes(workbook)


def process_catalog_upload(client: dict, catalog_type: str, content: bytes, filename: str, upload_scope: str = "full_catalog") -> dict:
    module = catalog_module(catalog_type)
    upload_scope = upload_scope if upload_scope in {"full_catalog", "partial_update"} else "full_catalog"
    with module_lock(client["id"], module):
        state = load_module_state(client, module)
        upload = add_upload_record(client["id"], module, state, content, filename, {"upload_scope": upload_scope})
        try:
            rows = parse_catalog_workbook(content, module)
            assert_unique_keys(rows, key_field(module))
        except SourceParseError as exc:
            upload["parse_status"] = "failed"
            upload["parse_error"] = str(exc)
            upload["result"] = "failed"
            append_audit(state, f"{module}.parse.failed", {"upload_id": upload["upload_id"], "error": str(exc)})
            save_module_state(client["id"], module, state)
            return {"status": "failed", "message": str(exc), "upload": upload}

        snapshot_id = write_snapshot(client["id"], module, upload["upload_id"], rows)
        upload["parse_status"] = "parsed"
        upload["snapshot_id"] = snapshot_id
        upload["snapshot_rows_hash"] = normalized_rows_hash(rows)
        upload["row_count"] = len(rows)

        result_rows, summary = merge_catalog_rows(state["published_rows"], rows, module, upload_scope)
        upload["diff_summary"] = summary
        if normalized_rows_hash(result_rows) == normalized_rows_hash(state["published_rows"]):
            upload["result"] = "no_change"
            append_audit(state, f"{module}.diff.no_change", {"upload_id": upload["upload_id"]})
            save_module_state(client["id"], module, state)
            refresh_source_index_if_configured(client)
            return {"status": "no_change", "message": "No catalog changes.", "summary": summary, "upload": upload}

        version = publish_version(client["id"], module, state, result_rows, upload["upload_id"], summary)
        version["snapshot_id"] = snapshot_id
        upload["result"] = "new_version"
        upload["created_version_id"] = version["version_id"]
        append_audit(state, f"{module}.version.published", {"version_id": version["version_id"], "upload_id": upload["upload_id"]})
        save_module_state(client["id"], module, state)
        refresh_source_index_if_configured(client)
        return {"status": "new_version", "message": f"Published {module} v{version['version_no']}.", "summary": summary, "upload": upload, "version": version}


def process_bcct_upload(client: dict, content: bytes, filename: str) -> dict:
    module = "bcct"
    with module_lock(client["id"], module):
        state = load_module_state(client, module)
        upload = add_upload_record(client["id"], module, state, content, filename, {"upload_scope": "append_or_review_by_transaction_key"})
        try:
            rows = parse_bcct_workbook(content)
        except SourceParseError as exc:
            upload["parse_status"] = "failed"
            upload["parse_error"] = str(exc)
            upload["result"] = "failed"
            append_audit(state, "bcct.parse.failed", {"upload_id": upload["upload_id"], "error": str(exc)})
            save_module_state(client["id"], module, state)
            return {"status": "failed", "message": str(exc), "upload": upload}

        snapshot_id = write_snapshot(client["id"], module, upload["upload_id"], rows)
        upload["parse_status"] = "parsed"
        upload["snapshot_id"] = snapshot_id
        upload["snapshot_rows_hash"] = normalized_rows_hash(rows)
        upload["row_count"] = len(rows)

        existing_by_key = {row["transaction_key"]: row for row in state["published_rows"]}
        added_rows = []
        unchanged_rows = 0
        new_candidates = []
        for row in rows:
            existing = existing_by_key.get(row["transaction_key"])
            if existing is None:
                published_row = {
                    **row,
                    "source_upload_id": upload["upload_id"],
                    "review_status": "reviewed",
                }
                if row["direction"] == "import":
                    published_row["import_row_id"] = import_row_id(row["transaction_key"])
                added_rows.append(published_row)
                existing_by_key[row["transaction_key"]] = published_row
                continue
            if bcct_rows_equal(existing, row):
                unchanged_rows += 1
                continue
            candidate = correction_candidate(existing, row, upload["upload_id"])
            if not correction_candidate_exists(state["correction_candidates"], candidate):
                new_candidates.append(candidate)

        if added_rows:
            state["published_rows"].extend(added_rows)
            state["published_rows"].sort(key=lambda row: (row["direction"], row["declaration_no"], row["line_no"], row["item_code"]))
        state["correction_candidates"].extend(new_candidates)

        summary = {
            "added_rows": len(added_rows),
            "unchanged_rows": unchanged_rows,
            "correction_candidates": len(new_candidates),
            "published_rows": len(state["published_rows"]),
        }
        upload["result"] = "review_required" if new_candidates else ("new_version" if added_rows else "no_change")
        upload["diff_summary"] = summary

        version = None
        if added_rows:
            version = publish_version(client["id"], module, state, state["published_rows"], upload["upload_id"], summary)
            version["snapshot_id"] = snapshot_id
            upload["created_version_id"] = version["version_id"]
        append_audit(state, "bcct.diff.completed", {"upload_id": upload["upload_id"], **summary})
        save_module_state(client["id"], module, state)
        refresh_source_index_if_configured(client)

        status = "review_required" if new_candidates else ("new_version" if added_rows else "no_change")
        message = "BCCT has correction candidates." if new_candidates else ("Published BCCT rows." if added_rows else "No BCCT changes.")
        return {"status": status, "message": message, "summary": summary, "upload": upload, "version": version, "correction_candidates": new_candidates}


def merge_catalog_rows(existing_rows: list[dict], uploaded_rows: list[dict], module: str, upload_scope: str) -> tuple[list[dict], dict]:
    key = key_field(module)
    existing = {row[key]: dict(row) for row in existing_rows}
    uploaded = {row[key]: dict(row) for row in uploaded_rows}
    result = {}
    added = changed = unchanged = inactive_pending_review = 0

    if upload_scope == "partial_update":
        result = {row_key: dict(row) for row_key, row in existing.items()}
        for row_key, row in uploaded.items():
            if row_key not in existing:
                added += 1
            elif normalized_row(row) != normalized_row(existing[row_key]):
                changed += 1
            else:
                unchanged += 1
            result[row_key] = row
    else:
        for row_key, row in uploaded.items():
            if row_key not in existing:
                added += 1
            elif normalized_row(row) != normalized_row(existing[row_key]):
                changed += 1
            else:
                unchanged += 1
            result[row_key] = row
        for row_key, row in existing.items():
            if row_key not in uploaded:
                retained = dict(row)
                if retained.get("status") != "inactive_pending_review":
                    inactive_pending_review += 1
                retained["status"] = "inactive_pending_review"
                result[row_key] = retained

    return sorted(result.values(), key=lambda row: row[key]), {
        "added": added,
        "changed": changed,
        "unchanged": unchanged,
        "inactive_pending_review": inactive_pending_review,
        "total_rows": len(result),
    }


def load_module_state(client: dict, module: str) -> dict:
    path = state_path(client["id"], module)
    if path.exists():
        return migrate_state(read_json(path), client, module)
    state = seed_module_state(client, module)
    save_module_state(client["id"], module, state)
    return state


def seed_module_state(client: dict, module: str) -> dict:
    rows = seed_rows(client, module)
    state = {
        "schema_version": 1,
        "client_id": client["id"],
        "module": module,
        "published_rows": rows,
        "versions": [],
        "uploads": [],
        "correction_candidates": [],
        "audit_events": [],
        "next_version_no": 1,
    }
    if rows:
        publish_version(client["id"], module, state, rows, "", {"seeded_rows": len(rows)})
        append_audit(state, f"{module}.version.seeded", {"rows": len(rows)})
    return state


def seed_rows(client: dict, module: str) -> list[dict]:
    if module == "material_catalog":
        return [normalize_seed_material(row) for row in client.get("material_catalog", [])]
    if module == "product_catalog":
        return [normalize_seed_product(row) for row in client.get("product_catalog", [])]
    rows = []
    for index, row in enumerate(client.get("bcct_rows", []), start=1):
        normalized = normalize_seed_bcct(row, index)
        if normalized:
            rows.append(normalized)
    return rows


def normalize_seed_material(row: dict) -> dict:
    customs_code = cell_text(row.get("customs_code"))
    return {
        "customs_code": customs_code,
        "internal_code": cell_text(row.get("internal_code")) or customs_code,
        "name": cell_text(row.get("name")),
        "hs_code": cell_text(row.get("hs_code")),
        "unit": normalize_unit(row.get("unit")),
        "role": cell_text(row.get("role")) or "NVL",
        "origin_default": cell_text(row.get("origin_default")),
        "status": normalize_status(row.get("status")),
    }


def normalize_seed_product(row: dict) -> dict:
    return {
        "product_code": cell_text(row.get("product_code")),
        "name": cell_text(row.get("name")),
        "hs_code": cell_text(row.get("hs_code")),
        "status": normalize_status(row.get("status")),
    }


def normalize_seed_bcct(row: dict, index: int) -> dict:
    direction = normalize_direction(row.get("direction"))
    declaration_no = cell_text(row.get("declaration_no"))
    item_code = cell_text(row.get("item_code"))
    if not direction or not declaration_no or not item_code:
        return {}
    normalized = {
        "coverage_period": cell_text(row.get("period") or row.get("coverage_period")),
        "direction": direction,
        "declaration_no": declaration_no,
        "declaration_date": cell_text(row.get("declaration_date")),
        "customs_office": cell_text(row.get("customs_office")),
        "declaration_type": normalize_declaration_type(row.get("declaration_type")),
        "line_no": cell_text(row.get("line_no") or index),
        "item_code": item_code,
        "description": cell_text(row.get("description")),
        "hs_code": cell_text(row.get("hs_code")),
        "quantity": normalize_decimal(row.get("qty") or row.get("quantity")),
        "unit": normalize_unit(row.get("unit")),
        "customs_value": normalize_decimal(row.get("customs_value")),
        "value_currency": "VND" if row.get("customs_value") else row.get("currency", "") if row.get("foreign_currency_value") else "",
        "invoice_ref": cell_text(row.get("invoice_ref")),
        "source_upload_id": "",
        "review_status": "reviewed",
    }
    normalized["transaction_key"] = transaction_key(normalized)
    if normalized["direction"] == "import":
        normalized["import_row_id"] = import_row_id(normalized["transaction_key"])
    return normalized


def publish_version(client_id: str, module: str, state: dict, rows: list[dict], upload_id: str, summary: dict, write_artifacts: bool = True) -> dict:
    version_no = state["next_version_no"]
    rows_hash = normalized_rows_hash(rows)
    version = {
        "version_id": f"{module}-v{version_no}-{rows_hash[:10]}",
        "version_no": version_no,
        "created_at": now_iso(),
        "source_upload_id": upload_id,
        "row_count": len(rows),
        "rows_hash": rows_hash,
        "summary": summary,
    }
    state["next_version_no"] += 1
    state["latest_version"] = version
    state["published_rows"] = [dict(row) for row in rows]
    state["versions"].insert(0, version)
    if write_artifacts:
        version_dir = module_root(client_id, module) / "versions" / f"v{version_no}"
        write_json(version_dir / "version.json", version)
        write_json(version_dir / "rows.json", rows)
        write_json(version_dir / "diff.json", summary)
    return version


def add_upload_record(client_id: str, module: str, state: dict, content: bytes, filename: str, extra: dict) -> dict:
    upload_id = make_id("upload")
    upload_dir = module_root(client_id, module) / "uploads" / upload_id
    raw_dir = upload_dir / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    original_filename = original_upload_filename(filename)
    stored_filename = f"{upload_id}-{storage_filename(original_filename)}"
    stored_path = raw_dir / stored_filename
    stored_path.write_bytes(content)
    relative_stored_path = stored_path.relative_to(source_root())
    content_sha256 = hashlib.sha256(content).hexdigest()
    mime_type = mimetypes.guess_type(original_filename)[0] or "application/octet-stream"
    upload = {
        "metadata_schema_version": 1,
        "upload_id": upload_id,
        "client_id": client_id,
        "module": module,
        "filename": original_filename,
        "original_filename": original_filename,
        "safe_filename": storage_filename(original_filename),
        "stored_filename": stored_filename,
        "stored_path": str(relative_stored_path),
        "storage_backend": "filesystem",
        "content_sha256": content_sha256,
        "size_bytes": len(content),
        "file_ext": Path(original_filename).suffix.lower(),
        "mime_type": mime_type,
        "content_type": mime_type,
        "created_at": now_iso(),
        "parse_status": "uploaded",
        "parse_error": "",
        "snapshot_id": "",
        "row_count": 0,
        "result": "uploaded",
        "created_version_id": "",
        "diff_summary": {},
        **extra,
    }
    state["uploads"].insert(0, upload)
    append_audit(state, f"{module}.uploaded", {"upload_id": upload_id, "filename": original_filename})
    return upload


def write_snapshot(client_id: str, module: str, upload_id: str, rows: list[dict], write_artifacts: bool = True) -> str:
    snapshot_id = make_id("snapshot")
    if write_artifacts:
        snapshot_dir = module_root(client_id, module) / "snapshots" / snapshot_id
        write_json(snapshot_dir / "snapshot.json", {
            "snapshot_id": snapshot_id,
            "client_id": client_id,
            "module": module,
            "upload_id": upload_id,
            "row_count": len(rows),
            "rows_hash": normalized_rows_hash(rows),
            "created_at": now_iso(),
        })
        write_json(snapshot_dir / "rows.json", rows)
    return snapshot_id


def module_workspace(state: dict) -> dict:
    return {
        "module": state["module"],
        "published_rows": [dict(row) for row in state["published_rows"]],
        "latest_version": state.get("latest_version"),
        "versions": [dict(version) for version in state["versions"]],
        "uploads": [dict(upload) for upload in state["uploads"]],
        "correction_candidates": [dict(candidate) for candidate in state.get("correction_candidates", [])],
        "audit_events": [dict(event) for event in state.get("audit_events", [])],
    }


def bcct_rows_equal(existing: dict, incoming: dict) -> bool:
    return normalized_row(existing) == normalized_row(incoming)


def correction_candidate(existing: dict, incoming: dict, upload_id: str) -> dict:
    return {
        "candidate_id": make_id("correction"),
        "status": "correction_candidate",
        "transaction_key": incoming["transaction_key"],
        "created_at": now_iso(),
        "source_upload_id": upload_id,
        "existing_row": dict(existing),
        "incoming_row": dict(incoming),
    }


def correction_candidate_exists(candidates: list[dict], candidate: dict) -> bool:
    incoming_hash = normalized_rows_hash([candidate["incoming_row"]])
    for existing in candidates:
        if existing.get("transaction_key") != candidate["transaction_key"]:
            continue
        if normalized_rows_hash([existing.get("incoming_row", {})]) == incoming_hash:
            return True
    return False


def iter_rows_as_dicts(worksheet, headers: list[str]):
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        values = {headers[index]: row[index] for index in range(min(len(headers), len(row))) if headers[index]}
        if any(cell_text(value) for value in values.values()):
            yield values


def key_field(module: str) -> str:
    return "customs_code" if module == "material_catalog" else "product_code"


def catalog_module(catalog_type: str) -> str:
    if catalog_type in CATALOG_MODULES:
        return CATALOG_MODULES[catalog_type]
    if catalog_type in CATALOG_MODULES.values():
        return catalog_type
    raise SourceParseError(f"Unknown catalog type: {catalog_type}")


def normalized_rows_hash(rows: list[dict]) -> str:
    payload = [normalized_row(row) for row in rows]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def append_audit(state: dict, event: str, details: dict) -> None:
    state.setdefault("audit_events", []).insert(0, {"event": event, "created_at": now_iso(), "details": details})


def migrate_state(state: dict, client: dict, module: str) -> dict:
    state.setdefault("schema_version", 1)
    state.setdefault("client_id", client["id"])
    state.setdefault("module", module)
    state.setdefault("published_rows", [])
    state.setdefault("versions", [])
    state.setdefault("uploads", [])
    state.setdefault("correction_candidates", [])
    state.setdefault("audit_events", [])
    state.setdefault("next_version_no", len(state["versions"]) + 1)
    state["next_version_no"] = max(state["next_version_no"], len(state["versions"]) + 1)
    if state["versions"] and "latest_version" not in state:
        state["latest_version"] = state["versions"][0]
    return state


def save_module_state(client_id: str, module: str, state: dict) -> None:
    write_json(state_path(client_id, module), state)


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    with os.fdopen(fd, "w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2, sort_keys=True)
        handle.write("\n")
    os.replace(tmp_name, path)


def workbook_to_bytes(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@contextmanager
def module_lock(client_id: str, module: str):
    root = module_root(client_id, module)
    root.mkdir(parents=True, exist_ok=True)
    lock_path = root / ".lock"
    with lock_path.open("w") as handle:
        fcntl.flock(handle, fcntl.LOCK_EX)
        try:
            yield
        finally:
            fcntl.flock(handle, fcntl.LOCK_UN)


def state_path(client_id: str, module: str) -> Path:
    return module_root(client_id, module) / "state.json"


def module_root(client_id: str, module: str) -> Path:
    if module not in SOURCE_MODULES:
        raise SourceParseError(f"Unknown source module: {module}")
    return source_root() / "clients" / client_id / module.replace("_", "-")


def source_root() -> Path:
    return Path(os.environ.get("SOURCE_STORE_ROOT", "data/local/source-modules"))


def original_upload_filename(filename: str) -> str:
    return Path(filename or "upload.xlsx").name.strip() or "upload.xlsx"


def storage_filename(filename: str) -> str:
    source_name = original_upload_filename(filename)
    suffix = Path(source_name).suffix.lower()
    stem = Path(source_name).stem or "upload"
    safe_stem = re.sub(r"[^A-Za-z0-9._-]+", "-", strip_accents(stem)).strip("-._")
    safe_stem = safe_stem or "upload"
    return f"{safe_stem}{suffix}"


def make_id(prefix: str) -> str:
    return f"{prefix}-{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S%f')}-{uuid.uuid4().hex[:8]}"


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()
