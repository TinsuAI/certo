"""Parser for agency cost-allocation Excel files.

Canonical layout (matched by column index, since headers are bilingual
and merged across rows 2-3):

    A=STT (ignored)         G=other_mfg
    B=Mã SP                  H=profit (IGNORED on import — residual)
    C=wages                  I=transport_storage
    D=welfare                J=note
    E=rent
    F=depreciation

Data starts at row 4. First row with empty Mã SP terminates parsing
(trailing blanks are common in agency-issued sheets).
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from app.cost_allocation_store import CostAllocationRow


_DATA_START_ROW = 4
_COL_PRODUCT_CODE = 2
_COL_WAGES = 3
_COL_WELFARE = 4
_COL_RENT = 5
_COL_DEPRECIATION = 6
_COL_OTHER_MFG = 7
# Column H (index 8) is the profit residual on the agency sheet — ignored.
_COL_TRANSPORT_STORAGE = 9
_COL_NOTE = 10


def parse_excel(data: bytes | Path) -> list[CostAllocationRow]:
    """Parse agency workbook bytes into CostAllocationRow records."""
    if isinstance(data, (bytes, bytearray)):
        wb = load_workbook(BytesIO(data), data_only=True)
    else:
        wb = load_workbook(Path(data), data_only=True)
    ws = wb.active  # GROWATT sample uses a single sheet ("Sheet3").
    rows: list[CostAllocationRow] = []
    for r in range(_DATA_START_ROW, ws.max_row + 1):
        product_code = _cell_text(ws.cell(r, _COL_PRODUCT_CODE).value)
        if not product_code:
            break
        rows.append(CostAllocationRow(
            product_code=product_code,
            coef_wages=_cell_decimal(ws.cell(r, _COL_WAGES).value),
            coef_welfare=_cell_decimal(ws.cell(r, _COL_WELFARE).value),
            coef_rent=_cell_decimal(ws.cell(r, _COL_RENT).value),
            coef_depreciation=_cell_decimal(ws.cell(r, _COL_DEPRECIATION).value),
            coef_other_mfg=_cell_decimal(ws.cell(r, _COL_OTHER_MFG).value),
            coef_transport_storage=_cell_decimal(ws.cell(r, _COL_TRANSPORT_STORAGE).value),
            note=_cell_text(ws.cell(r, _COL_NOTE).value),
        ))
    return rows


def apply_to_fob(row: CostAllocationRow, fob: Decimal) -> dict[str, Decimal]:
    """Multiply each coefficient by FOB → 6 detail values for cost_buildup.

    Profit is intentionally absent: per TT 05/2018 + the GROWATT sheet,
    profit is the residual (FOB − material − I+II+III − VII). The caller
    leaves the profit input blank; the bảng kê engine derives it via the
    formula chain (IV = I+II+III, VIII = FOB, V = inferred at print time).
    """
    fob = Decimal(fob)
    quant = Decimal("0.01")
    return {
        "wages":             (row.coef_wages * fob).quantize(quant),
        "welfare":           (row.coef_welfare * fob).quantize(quant),
        "rent":              (row.coef_rent * fob).quantize(quant),
        "depreciation":      (row.coef_depreciation * fob).quantize(quant),
        "other_mfg":         (row.coef_other_mfg * fob).quantize(quant),
        "transport_storage": (row.coef_transport_storage * fob).quantize(quant),
    }


_DETAIL_KEYS = ("wages", "welfare", "rent", "depreciation", "other_mfg", "transport_storage")
_RVC_LVC_TOKENS = ("RVC", "LVC")


def _has_filled_cost_buildup(cost_buildup) -> bool:
    """True when any of the 6 detail keys already holds a value (profit excluded —
    it is derived). Mirrors the per-product 'Áp hệ số' overwrite-confirm check."""
    if not isinstance(cost_buildup, dict):
        return False
    return any(str(cost_buildup.get(key) or "").strip() for key in _DETAIL_KEYS)


def _parse_fob(value) -> Decimal | None:
    text = str(value or "").strip().replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def bulk_apply_to_products(products, resolve_ratio, *, overwrite: bool = False) -> dict:
    """Bulk cost-buildup allocation over a case's products (Mục 4a).

    Pure: does NOT mutate inputs. The caller persists `updates` via the normal
    cost_buildup merge path. Only products whose effective criterion contains
    RVC/LVC are in scope (mirrors the template gate:
    `origin_sheet_effective_criteria_text or documented_result`).

    - `resolve_ratio(code)` → CostAllocationRow | None (Mode A → Mode B fallback;
      caller binds client_id, e.g. cost_allocation_store.get_ratio).
    - `overwrite=False` skips SP that already has any detail value (skipped_filled);
      `True` recomputes and overwrites.

    Returns {applied:[{code,mode}], skipped_no_ratio:[code], skipped_no_fob:[code],
    skipped_filled:[code], updates:{code:{detail:str}}}.
    """
    applied: list[dict] = []
    skipped_no_ratio: list[str] = []
    skipped_no_fob: list[str] = []
    skipped_filled: list[str] = []
    updates: dict[str, dict] = {}
    for product in products:
        criteria = str(
            product.get("origin_sheet_effective_criteria_text")
            or product.get("documented_result")
            or ""
        ).upper()
        if not any(token in criteria for token in _RVC_LVC_TOKENS):
            continue
        code = str(product.get("code") or "").strip()
        if not code:
            continue
        if not overwrite and _has_filled_cost_buildup(product.get("cost_buildup")):
            skipped_filled.append(code)
            continue
        fob = _parse_fob(product.get("fob"))
        if fob is None or fob <= 0:
            skipped_no_fob.append(code)
            continue
        row = resolve_ratio(code)
        if row is None:
            skipped_no_ratio.append(code)
            continue
        details = apply_to_fob(row, fob)
        updates[code] = {key: str(value) for key, value in details.items()}
        applied.append({"code": code, "mode": "A" if row.product_code else "B"})
    return {
        "applied": applied,
        "skipped_no_ratio": skipped_no_ratio,
        "skipped_no_fob": skipped_no_fob,
        "skipped_filled": skipped_filled,
        "updates": updates,
    }


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_decimal(value) -> Decimal:
    if value is None or value == "":
        return Decimal(0)
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", ".")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return Decimal(0)
