"""Standard CO stock template (xlsx) — schema + read/write helpers.

This is the canonical format the system import/export endpoints exchange.
The converter (scripts/convert_co_stock.py) turns the agency `Save` sheet
from `tru-lui-co-template.xlsm` into this layout; the import endpoint
expects exactly these columns.

Row identity: (declaration_no, line_no, customs_code) — re-import with the
same triplet overwrites prior values (snapshot model). See
app/co_stock_adjustments_store.py for the persistence layer.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

CO_STOCK_TEMPLATE_VERSION = "1.0"
CO_STOCK_SHEET_NAME = "co_stock"

CO_STOCK_COLUMNS: list[str] = [
    "declaration_no",
    "registration_date",
    "declaration_type",
    "line_no",
    "customs_code",
    "hs_code",
    "goods_name",
    "origin_country",
    "unit_price",
    "taxable_unit_price",
    "opening_qty",
    "unit",
    "partner",
    "invoice_no",
    "invoice_date",
    "exchange_rate",
    "used_qty",
    "source_co_no",
    "transaction_key",
]

CO_STOCK_REQUIRED_COLUMNS: set[str] = {
    "declaration_no",
    "line_no",
    "customs_code",
}

CO_STOCK_HEADER_LABELS: dict[str, str] = {
    "declaration_no": "Số TK",
    "registration_date": "Ngày ĐK",
    "declaration_type": "Mã loại hình",
    "line_no": "STT hàng",
    "customs_code": "Mã NPL/SP",
    "hs_code": "Mã HS",
    "goods_name": "Tên hàng",
    "origin_country": "Xuất xứ",
    "unit_price": "Đơn giá",
    "taxable_unit_price": "Đơn giá tính thuế",
    "opening_qty": "Tồn ban đầu",
    "unit": "Đơn vị tính",
    "partner": "Tên đối tác",
    "invoice_no": "Số hóa đơn",
    "invoice_date": "Ngày hóa đơn",
    "exchange_rate": "Tỷ giá thanh toán",
    "used_qty": "SL đã sử dụng",
    "source_co_no": "Số CO",
    "transaction_key": "Khóa giao dịch",
}

# Some columns store decimals; we coerce them at read/write time so the on-disk
# format stays human-readable (Excel-numeric) and the in-memory format stays
# explicit (Decimal, dates).
_DECIMAL_COLUMNS = {
    "unit_price",
    "taxable_unit_price",
    "exchange_rate",
    "opening_qty",
    "used_qty",
}
_DATE_COLUMNS = {"registration_date", "invoice_date"}


class CoStockTemplateError(ValueError):
    """Raised when the input workbook violates the standard template contract."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def _decimal_or_none(value: Any) -> Decimal | None:
    text = _text(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return None


def _date_or_none(value: Any):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    # Excel may surface a string; try common formats.
    text = _text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _coerce_row(raw: dict[str, Any]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for col in CO_STOCK_COLUMNS:
        value = raw.get(col)
        if col in _DECIMAL_COLUMNS:
            row[col] = _decimal_or_none(value)
        elif col in _DATE_COLUMNS:
            row[col] = _date_or_none(value)
        else:
            row[col] = _text(value)
    return row


def read_standard_co_stock(content: bytes) -> tuple[list[dict], list[str]]:
    """Read a standard template xlsx. Returns (rows, errors).

    Validates header row matches CO_STOCK_COLUMNS (extra trailing columns
    ignored; missing required columns surface as errors and abort parsing).
    Each data row is coerced into Python types per the column schema.
    """
    errors: list[str] = []
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise CoStockTemplateError(f"Không đọc được workbook: {exc}") from exc
    sheet_name = CO_STOCK_SHEET_NAME if CO_STOCK_SHEET_NAME in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    iterator = ws.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration:
        raise CoStockTemplateError("Workbook rỗng — không có header.")
    header_keys = [_text(cell).lower() for cell in header_row]
    # Accept either machine keys (declaration_no...) or Vietnamese labels (Số TK...).
    label_to_key = {v.lower(): k for k, v in CO_STOCK_HEADER_LABELS.items()}
    normalised_keys: list[str] = []
    for header in header_keys:
        if header in CO_STOCK_COLUMNS:
            normalised_keys.append(header)
        elif header in label_to_key:
            normalised_keys.append(label_to_key[header])
        else:
            normalised_keys.append("")
    missing_required = CO_STOCK_REQUIRED_COLUMNS - set(normalised_keys)
    if missing_required:
        raise CoStockTemplateError(
            "Header thiếu cột bắt buộc: " + ", ".join(sorted(missing_required))
        )
    rows: list[dict] = []
    for line_no, raw in enumerate(iterator, start=2):
        if raw is None:
            continue
        record_raw: dict[str, Any] = {}
        for key, value in zip(normalised_keys, raw):
            if not key:
                continue
            record_raw[key] = value
        if not any(record_raw.values()):
            continue
        record = _coerce_row(record_raw)
        # Required-field check per row.
        missing = [c for c in CO_STOCK_REQUIRED_COLUMNS if not _text(record.get(c))]
        if missing:
            errors.append(f"Dòng {line_no}: thiếu {', '.join(missing)}")
            continue
        rows.append(record)
    wb.close()
    return rows, errors


def write_standard_co_stock(rows: list[dict[str, Any]], *, use_labels: bool = False) -> bytes:
    """Render a list of dict rows into a standard template xlsx.

    `use_labels=True` writes Vietnamese header labels (for human export);
    default writes machine keys (for round-trip).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = CO_STOCK_SHEET_NAME
    headers = (
        [CO_STOCK_HEADER_LABELS[col] for col in CO_STOCK_COLUMNS]
        if use_labels
        else list(CO_STOCK_COLUMNS)
    )
    ws.append(headers)
    for raw in rows:
        record = _coerce_row(raw)
        row_values: list[Any] = []
        for col in CO_STOCK_COLUMNS:
            value = record.get(col)
            if isinstance(value, Decimal):
                row_values.append(float(value))
            else:
                row_values.append(value if value not in (None, "") else None)
        ws.append(row_values)
    # Set sensible column widths for human reading.
    widths = {
        "declaration_no": 16,
        "registration_date": 12,
        "declaration_type": 8,
        "line_no": 8,
        "customs_code": 16,
        "hs_code": 12,
        "goods_name": 40,
        "origin_country": 10,
        "unit_price": 14,
        "taxable_unit_price": 16,
        "opening_qty": 14,
        "unit": 10,
        "partner": 24,
        "invoice_no": 16,
        "invoice_date": 12,
        "exchange_rate": 12,
        "used_qty": 14,
        "source_co_no": 16,
        "transaction_key": 18,
    }
    for idx, col in enumerate(CO_STOCK_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 14)
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
