from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import re
import threading
from contextlib import asynccontextmanager
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import quote

import httpx
from fastapi import File, Form, HTTPException, Request, UploadFile
from fastapi import FastAPI
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, PlainTextResponse, RedirectResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles

from app import co_auth
from app.bom_store import attach_case_bom_snapshot
from app.bom_service import bom_service
from app.co_case_store import (
    MAX_SUPPORTING_FILE_BYTES,
    CaseClosedError,
    CaseHasActiveClaimsError,
    acquire_origin_calculation_lock,
    active_origin_calculation_lock,
    build_case_criteria_rows,
    case_from_record,
    co_case_delete_block_reason,
    co_case_is_completed,
    create_case_record,
    create_case_workbook,
    declaration_refs,
    delete_case_record,
    get_case_record,
    get_case_workspace,
    get_supporting_file,
    invoice_keys,
    json_safe,
    match_case_bcct_exports,
    safe_filename,
    save_supporting_file,
    release_origin_calculation_lock,
    update_case_record,
)
from app.co_forms import (
    COMMON_MARKET_PRESETS,
    common_market_guidance,
    criteria_preview_for_hs,
    form_candidates_for_market,
    prioritized_form_lanes,
    recommended_form_lane,
)
from app.co_form_config_store import (
    co_form_config_path,
    load_co_form_config,
    reset_co_form_config,
    sanitize_co_form_config,
    save_co_form_config,
    unique_text_list,
)
from app import co_stock_adjustments_store, co_stock_eligibility, co_stock_events_store, co_stock_ledger, co_stock_materializer
from app.co_stock_template import CoStockTemplateError, read_standard_co_stock, write_standard_co_stock
from app.co_market_hints import infer_market_from_invoice_matches
from app.client_registry import get_client as registry_get_client
from app.client_registry import get_client_case
from app.customs_fx_store import CUSTOMS_FX_CLIENT_ID, get_customs_fx_store
from app.database import apply_migrations, database_url
from app.data_hub_client import (
    DataHubClient,
    bom_product_code_from_material_identity,
    reset_current_data_hub_token,
    set_current_data_hub_token,
)
from app.data_hub_settings import (
    DATA_HUB_LINK_ENV_KEYS,
    DataHubLinkSettings,
    data_hub_config_path,
    data_hub_link_settings,
    load_data_hub_overrides,
    save_data_hub_overrides,
)
from app.demo_data import (
    DEMO_CASE,
    SOURCE_NOTES,
    attach_results,
    clone_case,
    update_products_from_form,
)
from app.origin import evaluate_tariff_shift
from app import material_search
from app.app_state_store import get_app_state_store
from app.portfolio import SourceBackendUnavailable, portfolio_app, portfolio_service
from app.routers import auth as auth_routes
from app.routers import catalog as catalog_routes
from app.routers import co_case as co_case_routes
from app.routers import customs_fx as customs_fx_routes
from app.routers import settings as settings_routes
from app.source_store import (
    attach_case_source_snapshot,
    co_stock_rows_from_bcct,
    enrich_client_with_source_workspace,
)
from app.table_view import build_table_view
from app.web.client_context import (
    _data_hub_overview_context,
    case_finished_hs_codes,
    client_case,
    client_context,
    default_client_case,
    effective_min_gap_days,
    resolve_client,
    source_workspace_for_client,
)
from app.web.deps import large_request_form, require_local_source_writes
from app.routers.co_case import (
    _hydrate_product_export_declaration_dates,
    _to_vietnamese_date,
    invoice_search_options,
    mark_origin_sheets_stale,
    merge_origin_action_payload,
    resolve_shipment_reference,
    set_origin_sheet_status,
)
from app.web.co_case_context import (
    CO_CASE_WORKFLOW_STEP_KEYS,
    _CO_CASE_SOURCE_CACHE,
    _co_stock_refresh_inflight,
    CALCULATE_SNAPSHOT_FRESHNESS_SECONDS,
    CO_CASE_STEP_STATUS_LABELS,
    CO_CASE_WORKFLOW_STEPS,
    ORIGIN_SHEET_STATUS_LABELS,
    SHEET_CURRENCY_MODES,
    SHEET_OPTIMIZATION_MODES,
    _CO_CASE_SOURCE_CACHE_TTL_SECONDS,
    _allocation_line_fx,
    _attach_fob_vnd,
    _calculate_stock_rows_from_snapshot,
    _co_stock_refresh_inflight_lock,
    _co_stock_snapshot_is_fresh,
    _full_refresh,
    _probe_server_time,
    _refresh_co_stock_delta_or_full,
    _schedule_background_co_stock_refresh,
    _try_delta_refresh,
    add_bom_code,
    allocate_material_stock,
    allocation_available_qty,
    allocation_currency_summary,
    allocation_document_ref,
    allocation_line_matches_stock,
    allocation_summary,
    allocation_unit_value_summary,
    allocation_valuation_source,
    apply_existing_origin_product_consumption,
    attach_case_source_summary_snapshot,
    attach_origin_bom_product_codes,
    attach_origin_demo,
    attach_origin_readiness,
    attach_origin_sheet_states,
    bom_code_candidates,
    bom_workspace_from_case_snapshot,
    cached_origin_source_context,
    calculate_lvc_result,
    case_allocation_pool,
    case_export_anchor_date,
    case_tkx_tkn_summary,
    co_case_bom_product_codes,
    co_case_context,
    co_case_hs_codes,
    co_case_light_context,
    co_case_source_context,
    co_case_source_context_cached,
    co_case_step_status,
    co_case_workflow_steps,
    co_stock_allocation_pool,
    co_stock_allocation_sort_key,
    co_stock_has_value,
    co_stock_is_usable,
    co_stock_key_candidates,
    compact_origin_signature_row,
    criterion_mode,
    decimal_text,
    decimal_value,
    declaration_file_count,
    durable_sheet_status,
    enrich_client_with_source_summary,
    enrich_origin_material,
    enrich_origin_product,
    first_decimal_source,
    first_non_empty,
    has_shipment_reference,
    invoice_form_lane_view,
    invoice_match_criteria_rows,
    invoice_match_preview_row,
    invoice_match_summary,
    invoice_preview_from_matches,
    latest_usable_product_version,
    lvc_threshold_from_criterion,
    market_inference_view,
    material_catalog_index,
    material_issue_summary,
    material_summary_row,
    minimal_bom_workspace,
    normalize_threshold,
    normalized_lvc_result,
    numeric_sequence,
    numeric_sort_text,
    order_invoice_matches_for_origin,
    origin_build_signature,
    origin_case_revision,
    origin_lock_actor,
    origin_match_from_existing_product,
    origin_material_count,
    origin_material_from_bom_row,
    origin_product_from_invoice_match,
    origin_product_order,
    origin_product_shell_from_invoice_match,
    origin_product_value,
    origin_sheet_action_error,
    origin_sheet_export_blockers,
    origin_source_context,
    origin_status_details_from_material,
    origin_status_from_material,
    origin_warning_summary,
    prepare_case_origin_product_shells,
    prepare_case_origin_products,
    prepare_case_origin_sheet,
    primary_shipment_reference,
    resolve_bom_product_code,
    selected_bom_rows_by_product,
    sheet_form_recommendation,
    shipment_reference_warnings,
    should_show_origin_demo,
    source_summary_from_case_snapshot,
    stock_allocation_consumption,
    stock_allocation_line,
    stock_allocation_remaining_qty,
    stock_available_qty,
    stock_candidates_for_material,
    stock_consumption_is_before,
    stock_consumption_label,
    stock_for_existing_allocation_line,
    stock_shortage_trace,
    tariff_shift_rule_from_criterion,
    text_list,
    unique_texts,
    usable_bom_product_version,
    valuation_source_label,
    valuation_status_label,
)
from app.web.templating import templates
from app.workbook_io import (
    WorkbookParseError,
    create_dossier_zip,
    create_evidence_workbook,
    create_hq_bang_ke_workbook,
    create_input_workbook,
    parse_input_workbook,
)


ROOT = Path(__file__).resolve().parent


@asynccontextmanager
async def lifespan(_app: FastAPI):
    if database_url():
        apply_migrations()
    yield


app = FastAPI(title="Barry CO Demo", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")
app.mount("/portfolio", portfolio_app, name="portfolio")
app.include_router(auth_routes.router)
app.include_router(catalog_routes.router)
app.include_router(co_case_routes.router)
app.include_router(customs_fx_routes.router)
app.include_router(settings_routes.router)


@app.exception_handler(CaseClosedError)
async def _case_closed_handler(request: Request, exc: CaseClosedError):
    """Mutating route hit a closed case → 409 with the friendly Vietnamese message."""
    return JSONResponse(status_code=409, content={"detail": str(exc)})


@app.exception_handler(SourceBackendUnavailable)
async def _source_backend_unavailable_handler(request: Request, exc: SourceBackendUnavailable):
    """Data Hub is the source of truth but unavailable, and local fallback is not
    allowed → 503 with a clear message instead of a silently-empty page."""
    return PlainTextResponse(str(exc), status_code=503)


@app.exception_handler(httpx.TransportError)
async def _data_hub_unreachable_handler(request: Request, exc: httpx.TransportError):
    """Data Hub is enabled but the API is unreachable (connection refused /
    timeout). Surface a clear 503 instead of a generic 500 so the operator knows
    it's a Data Hub outage, not a CO bug. Only fires for transport errors that
    propagate unhandled — local try/except (e.g. 404 fallbacks) still wins."""
    return PlainTextResponse(
        "Data Hub không phản hồi (kết nối thất bại/timeout). CO không dùng dữ liệu "
        "local backup; kiểm tra Data Hub rồi thử lại.",
        status_code=503,
    )


@app.exception_handler(httpx.HTTPStatusError)
async def _data_hub_error_status_handler(request: Request, exc: httpx.HTTPStatusError):
    """Data Hub returned an error status that no route handled → 502 (bad
    gateway): the upstream source failed, not CO. Routes that intentionally
    handle Data Hub statuses (e.g. 404 → fallback) catch the error themselves
    and never reach this handler."""
    upstream = exc.response.status_code if exc.response is not None else "?"
    return PlainTextResponse(
        f"Data Hub trả lỗi ({upstream}). CO không dùng dữ liệu local backup; "
        "kiểm tra Data Hub rồi thử lại.",
        status_code=502,
    )


def format_number_display(value, max_decimals: int = 2) -> str:
    text = "" if value is None else str(value).strip()
    if not text:
        return ""
    try:
        decimal = Decimal(text.replace(",", ""))
    except (InvalidOperation, ValueError):
        return text
    if decimal == decimal.to_integral():
        return f"{int(decimal):,}"
    max_decimals = max(0, min(int(max_decimals), 8))
    quant = Decimal("1").scaleb(-max_decimals)
    rounded = decimal.quantize(quant, rounding=ROUND_HALF_UP)
    return f"{rounded:,.{max_decimals}f}".rstrip("0").rstrip(".")


templates.env.filters["number"] = format_number_display


@app.get("/healthz")
async def healthz() -> dict[str, str]:
    return {"status": "ok"}


@app.middleware("http")
async def require_data_hub_auth(request: Request, call_next):
    redirect = co_auth.guard_response(request)
    if redirect:
        return redirect
    co_auth.load_optional_user(request)
    token_context = None
    user = co_auth.current_user(request)
    if user and user.access_token:
        token_context = set_current_data_hub_token(user.access_token)
    try:
        return await call_next(request)
    finally:
        if token_context:
            reset_current_data_hub_token(token_context)

BCCT_COLUMNS = [
    {"key": "coverage_period", "label": "Kỳ"},
    {"key": "declaration_no", "label": "Tờ khai", "class": "mono"},
    {"key": "line_no", "label": "STT", "class": "mono"},
    {"key": "declaration_type", "label": "LH", "class": "mono"},
    {"key": "direction_label", "label": "Luồng"},
    {"key": "item_code", "label": "Mã hàng", "class": "mono"},
    {"key": "hs_code", "label": "HS", "class": "mono"},
    {"key": "quantity", "label": "Số lượng", "class": "num"},
    {"key": "unit", "label": "ĐVT", "class": "mono"},
    {"key": "customs_value", "label": "Trị giá", "class": "num"},
    {"key": "invoice_ref", "label": "Hóa đơn", "class": "mono"},
]


CO_STOCK_COLUMNS = [
    {"key": "source_row", "label": "Dòng nguồn", "class": "mono"},
    {"key": "import_declaration_no", "label": "Tờ khai nhập", "class": "mono"},
    {"key": "line_no", "label": "STT", "class": "mono"},
    {"key": "declaration_type", "label": "LH", "class": "mono"},
    {"key": "customs_item_code", "label": "Mã HQ", "class": "mono"},
    {"key": "allocation_code", "label": "Mã phân bổ", "class": "mono"},
    {"key": "available_qty", "label": "Tồn CO", "class": "num"},
    {"key": "used_qty", "label": "Đã dùng", "class": "num"},
    {"key": "remaining_qty", "label": "Còn lại", "class": "num"},
    {"key": "status_label", "label": "Trạng thái"},
    {"key": "stock_reason_label", "label": "Lý do"},
    {"key": "history_action", "label": "Lịch sử", "kind": "history", "sortable": False},
]

BOM_PRODUCT_COLUMNS = [
    {"key": "product_code", "label": "Mã TP", "class": "mono", "link_key": "view_href"},
    {"key": "product_artifact_no", "label": "TP artifact", "class": "mono"},
    {"key": "row_count", "label": "Dòng BOM", "class": "num"},
    {"key": "status", "label": "Trạng thái"},
    {"key": "version_hash_short", "label": "Hash", "class": "mono"},
]

BOM_LINE_COLUMNS = [
    {"key": "material_code", "label": "Mã NVL", "class": "mono"},
    {"key": "material_name", "label": "Tên NVL"},
    {"key": "qty_per", "label": "Định mức", "class": "num"},
    {"key": "uom", "label": "ĐVT", "class": "mono"},
    {"key": "scrap_rate", "label": "Hao hụt", "class": "num"},
    {"key": "source", "label": "Nguồn"},
    {"key": "row_class", "label": "Trạng thái"},
]


def resolve_invoice_reference(client: dict, reference: str) -> dict:
    resolved = resolve_shipment_reference(client, reference)
    return {
        "invoice_no": resolved["invoice_no"],
        "source_reference": resolved["source_reference"],
        "source_reference_type": resolved["source_reference_type"],
    }


def co_stock_index(stock_rows: list[dict]) -> dict[str, dict]:
    output = {}
    for row in stock_rows:
        for key in co_stock_key_candidates(row):
            existing = output.get(key)
            if existing is None or co_stock_rank(row) > co_stock_rank(existing):
                output[key] = row
    return output


def co_stock_rank(row: dict) -> tuple[bool, bool, bool]:
    return (
        co_stock_is_usable(row),
        decimal_value(row.get("remaining_qty") or row.get("available_qty") or "0") > 0,
        co_stock_has_value(row),
    )


def first_decimal_value(*values) -> Decimal | None:
    for value in values:
        if value not in (None, ""):
            return decimal_value(value)
    return None


def bcct_table_context(request: Request, client_id: str, direction: str | None = None, **extra) -> dict:
    lean = _data_hub_overview_context(client_id, "bcct", dh_path="bcct")
    if lean is not None and not extra.get("bcct_result"):
        lean["bcct_view"] = direction or "all"
        title_by_direction = {"import": "BCCT nhập khẩu", "export": "BCCT xuất khẩu"}
        lean["bcct_title"] = title_by_direction.get(direction, "BCCT nhập khẩu / xuất khẩu")
        return lean
    context = client_context(client_id, "bcct", **extra)
    rows = [bcct_table_row(row) for row in context["source_workspace"]["bcct"]["published_rows"]]
    if direction:
        rows = [row for row in rows if row.get("direction") == direction]
    if direction == "export":
        relevant_types = set(context["client_config"]["bcct"].get("relevant_export_declaration_types", []))
        if relevant_types:
            rows = [row for row in rows if row.get("declaration_type") in relevant_types]
    context["bcct_view"] = direction or "all"
    title_by_direction = {"import": "BCCT nhập khẩu", "export": "BCCT xuất khẩu"}
    context["bcct_title"] = title_by_direction.get(direction, "BCCT nhập khẩu / xuất khẩu")
    context["source_table"] = build_table_view(
        rows,
        columns=BCCT_COLUMNS,
        query=request.query_params,
        filters=[
            {
                "name": "direction",
                "field": "direction",
                "label": "Luồng",
                "options": [
                    {"value": "import", "label": "Nhập khẩu"},
                    {"value": "export", "label": "Xuất khẩu"},
                ],
            },
            {"name": "type", "field": "declaration_type", "label": "Loại hình"},
            {"name": "hs", "field": "hs_code", "label": "HS"},
            {"name": "origin", "field": "origin_country", "label": "Xuất xứ"},
        ],
        summary_fields=[
            {"field": "direction_label", "label": "Luồng"},
            {"field": "declaration_type", "label": "Loại hình"},
        ],
        default_sort="declaration_no",
    )
    return context


def _co_stock_lean_client_context(client_id: str, co_stock_row_count: int | None = None) -> dict:
    """Lean context for /co-stock — skips full source_workspace pagination.

    Standard client_context calls source_workspace_for_client which paginates
    full BCCT from Data Hub (~10s on Johnson). /co-stock only needs client
    meta + client_config + counts + materialized co_stock_rows, so we build
    those directly.

    Also computes sync_status by comparing current BCCT row count against
    the snapshot's recorded count (see co_stock_materializer.compute_sync_status).
    """
    client = resolve_client(client_id)
    source_summary, source_backend = portfolio_service.source_summary(client)
    client_config = source_summary.get("client_config") or portfolio_service.get_client_config(client)
    bcct_row_count = source_summary.get("bcct", {}).get("published_row_count", 0)
    if co_stock_row_count is None:
        co_stock_row_count = co_stock_materializer.row_count(client["id"])
    counts = {
        **client.get("counts", {}),
        "materials": source_summary.get("material_catalog", {}).get("published_row_count", 0),
        "products": source_summary.get("product_catalog", {}).get("published_row_count", 0),
        "bcct": bcct_row_count,
        "co_stock": co_stock_row_count,
    }
    client = {**client, "counts": counts}
    sync_status = co_stock_materializer.compute_sync_status(client["id"], bcct_row_count)
    return {
        "client": client,
        "active": "co-stock",
        "source_backend": source_backend,
        "client_config": client_config,
        "source_workspace": {
            "client_config": client_config,
            "bcct": {"latest_version": source_summary.get("bcct", {}).get("latest_version") or {},
                     "correction_candidates": []},
            "material_catalog": {"latest_version": source_summary.get("material_catalog", {}).get("latest_version") or {}},
            "product_catalog": {"latest_version": source_summary.get("product_catalog", {}).get("latest_version") or {}},
        },
        "co_stock_last_refresh_at": co_stock_materializer.last_refresh_at(client["id"]),
        "co_stock_sync_status": sync_status,
    }


_CO_STOCK_STATUS_OPTIONS = [
    {"value": "available", "label": "Khả dụng"},
    {"value": "review_required", "label": "Cần review"},
    {"value": "inactive", "label": "Không dùng"},
    {"value": "depleted", "label": "Hết tồn"},
]


def co_stock_table_context(request: Request, client_id: str) -> dict:
    """SQL-paginated Tồn CO context. The DB does WHERE/ORDER/LIMIT so the
    page returns ~50 rows × few-ms even on 60k-row clients (Johnson).
    Ledger + adjustments apply only to the visible slice.

    Falls back to in-memory build_table_view for clients whose stock pool
    hasn't been materialized into `co_stock_rows` yet (file-mode demo
    fixtures + first-time-ever loads when no refresh has run).
    """
    from app.table_view import (
        DEFAULT_PAGE_SIZE,
        PAGE_SIZES,
        normalize_column,
        normalize_query,
        page_query,
        parse_int,
        table_field_names,
    )

    total_co_stock_rows = co_stock_materializer.row_count(client_id)
    # File-mode (no DB or no Data Hub configured) → legacy in-memory path so
    # Growatt-style fixtures keep working. Data Hub mode always uses the
    # materialized table — when empty, the UI shows a "Chưa có snapshot"
    # state that prompts the operator to click Refresh.
    use_legacy_path = total_co_stock_rows == 0 and not data_hub_link_settings().source_enabled
    if use_legacy_path:
        context = client_context(client_id, "co-stock")
        rows = [co_stock_table_row(row) for row in context["client"]["co_stock"]]
        context["co_stock_last_refresh_at"] = co_stock_materializer.last_refresh_at(client_id)
        context["co_stock_empty_needs_refresh"] = not context["client"]["co_stock"]
        context["co_stock_sync_status"] = {"status": "no_snapshot", "snapshot_rows": 0,
                                           "snapshot_bcct_rows": 0, "bcct_now_rows": 0,
                                           "refreshed_at": "", "delta": 0}
        context["source_table"] = build_table_view(
            rows,
            columns=CO_STOCK_COLUMNS,
            query=request.query_params,
            filters=[{"name": "status", "field": "status", "label": "Trạng thái",
                      "options": _CO_STOCK_STATUS_OPTIONS}],
            summary_fields=[
                {"field": "status_label", "label": "Trạng thái"},
                {"field": "declaration_type", "label": "Loại hình"},
            ],
            default_sort="import_declaration_no",
        )
        return context

    context = _co_stock_lean_client_context(client_id, co_stock_row_count=total_co_stock_rows)
    client = context["client"]
    context["co_stock_empty_needs_refresh"] = total_co_stock_rows == 0

    query_values = normalize_query(request.query_params)
    field_names = table_field_names("")
    q = query_values.get(field_names["q"], "").strip()
    status_value = query_values.get("status", "").strip()
    sort_key = query_values.get(field_names["sort"]) or "import_declaration_no"
    direction = "desc" if query_values.get(field_names["dir"]) == "desc" else "asc"
    per_page = min(max(1, parse_int(query_values.get(field_names["per_page"]), DEFAULT_PAGE_SIZE)), 500)
    page = max(1, parse_int(query_values.get(field_names["page"]), 1))

    page_rows, filtered_count = co_stock_materializer.read_co_stock_page(
        client["id"],
        q=q,
        status=status_value,
        sort=sort_key,
        direction=direction,
        offset=(page - 1) * per_page,
        limit=per_page,
    )
    # Apply ledger + adjustments only to the visible page.
    if page_rows:
        used_by_lot = co_stock_ledger.used_qty_by_lot(client["id"])
        if used_by_lot:
            page_rows = co_stock_ledger.apply_used_qty(page_rows, used_by_lot)
        adjustments = co_stock_adjustments_store.aggregate_by_lookup_key(client["id"])
        if adjustments:
            page_rows = co_stock_adjustments_store.apply_adjustments(page_rows, adjustments)
    rows = [co_stock_table_row(row) for row in page_rows]

    column_defs = [normalize_column(col) for col in CO_STOCK_COLUMNS]
    page_count = max(1, (filtered_count + per_page - 1) // per_page)
    page = min(page, page_count)

    prepared_query = {k: v for k, v in query_values.items() if k != field_names["page"]}
    owned_names = set(field_names.values()) | {"status"}
    reset_query = {k: v for k, v in query_values.items() if k not in owned_names}

    for col in column_defs:
        col["sort_active"] = col["key"] == sort_key
        col["sort_dir"] = direction if col["sort_active"] else ""
        col["sort_query"] = page_query(
            prepared_query,
            **{
                field_names["sort"]: col["key"],
                field_names["dir"]: "desc" if col["sort_active"] and direction == "asc" else "asc",
                field_names["page"]: 1,
            },
        )

    filter_defs = [
        {
            "name": "status",
            "query_name": "status",
            "label": "Trạng thái",
            "field": "status",
            "options": _CO_STOCK_STATUS_OPTIONS,
            "value": status_value,
        }
    ]

    context["source_table"] = {
        "rows": rows,
        "columns": column_defs,
        "filters": filter_defs,
        "summary_chips": [],
        "query": query_values,
        "field_names": field_names,
        "passthrough_params": reset_query,
        "param_prefix": "",
        "q": q,
        "sort": sort_key,
        "dir": direction,
        "page": page,
        "per_page": per_page,
        "page_sizes": PAGE_SIZES,
        "total_pages": page_count,
        "total_count": total_co_stock_rows,
        "filtered_count": filtered_count,
        "start_index": (page - 1) * per_page + 1 if rows else 0,
        "end_index": min(page * per_page, filtered_count),
        "has_previous": page > 1,
        "has_next": page < page_count,
        "previous_query": page_query(prepared_query, **{field_names["page"]: page - 1}),
        "next_query": page_query(prepared_query, **{field_names["page"]: page + 1}),
        "first_query": page_query(prepared_query, **{field_names["page"]: 1}),
        "last_query": page_query(prepared_query, **{field_names["page"]: page_count}),
        "reset_query": page_query(reset_query),
    }
    return context


def bom_context(request: Request, client_id: str, **extra) -> dict:
    lean = _data_hub_overview_context(client_id, "bom", dh_path="bom")
    if lean is not None and not extra.get("message") and not extra.get("error"):
        # DH mode: BOM is read-only; CO renders summary + link rather than
        # paginating the full bom_workspace (which fetches every product
        # version + line over HTTP from Data Hub).
        return lean
    context = client_context(client_id, "bom", **extra)
    workspace = context["bom_workspace"]
    selected_product = selected_bom_product(workspace, request.query_params.get("product", ""))
    product_rows = bom_product_table_rows(workspace, client_id, selected_product)
    line_rows = [
        bom_line_table_row(row)
        for row in workspace.get("latest_rows", [])
        if not selected_product or row.get("product_code") == selected_product
    ]
    product_table = build_table_view(
        product_rows,
        columns=BOM_PRODUCT_COLUMNS,
        query=request.query_params,
        filters=[{"name": "status", "field": "status", "label": "Trạng thái"}],
        summary_fields=[{"field": "status", "label": "Trạng thái"}],
        default_sort="product_code",
        default_per_page=25,
        param_prefix="tp_",
    )
    product_table["search_placeholder"] = "Mã thành phẩm, version, hash..."
    line_table = build_table_view(
        line_rows,
        columns=BOM_LINE_COLUMNS,
        query=request.query_params,
        filters=[
            {"name": "uom", "field": "uom", "label": "ĐVT"},
            {"name": "source", "field": "source", "label": "Nguồn"},
            {"name": "status", "field": "row_class", "label": "Trạng thái"},
        ],
        summary_fields=[
            {"field": "row_class", "label": "Trạng thái"},
            {"field": "uom", "label": "ĐVT"},
        ],
        default_sort="material_code",
        default_per_page=50,
        param_prefix="line_",
    )
    line_table["search_placeholder"] = "Mã NVL, tên NVL, trạng thái..."
    context["selected_bom_product"] = selected_product
    context["selected_bom_product_summary"] = next(
        (row for row in product_rows if row["product_code"] == selected_product),
        {},
    )
    context["bom_product_table"] = product_table
    context["bom_line_table"] = line_table
    return context


def selected_bom_product(workspace: dict, requested_product: str = "") -> str:
    codes = sorted({str(row.get("product_code", "")) for row in workspace.get("product_composition", []) if row.get("product_code")})
    if not codes:
        codes = sorted({str(row.get("product_code", "")) for row in workspace.get("latest_rows", []) if row.get("product_code")})
    requested = str(requested_product or "").strip()
    if requested in codes:
        return requested
    return codes[0] if codes else ""


def bom_product_table_rows(workspace: dict, client_id: str, selected_product: str) -> list[dict]:
    version_index = {
        version.get("product_version_id"): version
        for version in workspace.get("product_versions", [])
        if version.get("product_version_id")
    }
    rows = []
    composition = workspace.get("product_composition", [])
    if not composition:
        composition = [
            {
                "product_code": version.get("product_code", ""),
                "product_version_id": version.get("product_version_id", ""),
                "product_version_no": version.get("product_version_no", ""),
                "row_count": version.get("row_count", 0),
                "status": version.get("status", ""),
                "version_hash": version.get("version_hash", ""),
            }
            for version in workspace.get("product_versions", [])
        ]
    for row in composition:
        version = version_index.get(row.get("product_version_id"), {})
        product_code = str(row.get("product_code", ""))
        version_hash = str(row.get("version_hash") or version.get("version_hash") or "")
        rows.append({
            "product_code": product_code,
            "product_version_no": row.get("product_version_no", version.get("product_version_no", "")),
            "row_count": row.get("row_count", version.get("row_count", 0)),
            "status": row.get("status", version.get("status", "")),
            "version_hash": version_hash,
            "version_hash_short": version_hash[:10],
            "selected": "Đang xem" if product_code == selected_product else "",
            "view_href": f"/clients/{client_id}/bom?product={quote(product_code, safe='')}#bom-lines",
        })
    return rows


def bom_line_table_row(row: dict) -> dict:
    return {
        **row,
        "material_name": row.get("material_name", ""),
        "scrap_rate": row.get("scrap_rate", ""),
        "source": row.get("source", ""),
        "row_class": row.get("row_class", ""),
    }


def config_context(client_id: str, **extra) -> dict:
    # /config only renders client identity + client_config knobs. It does NOT
    # need source_workspace / bom_workspace, so skip the full pagination that
    # client_context triggers (Johnson: ~65k BCCT rows over HTTP per render).
    lean = _data_hub_overview_context(client_id, "config", dh_path="")
    if lean is not None:
        lean.update(extra)
        return lean
    return client_context(client_id, "config", **extra)


def bcct_table_row(row: dict) -> dict:
    direction = row.get("direction", "")
    return {
        **row,
        "direction_label": "Nhập khẩu" if direction == "import" else "Xuất khẩu",
    }


def co_stock_table_row(row: dict) -> dict:
    remaining_qty = str(row.get("remaining_qty", ""))
    if row.get("eligibility_status") == "inactive":
        status = "inactive"
    elif row.get("allocation_code_status") != "resolved":
        status = "review_required"
    else:
        status = "depleted" if remaining_qty in {"", "0", "0.0", "0.00"} else "available"
    return {
        **row,
        "status": status,
        "status_label": {
            "available": "Khả dụng",
            "depleted": "Hết tồn",
            "inactive": "Không dùng",
            "review_required": "Cần review",
        }[status],
        "stock_reason_label": stock_reason_label(row, status),
    }


def stock_reason_label(row: dict, status: str) -> str:
    if status == "inactive" and row.get("eligibility_reason") == "excluded_by_declaration_type_config":
        return "Loại hình không active trong config"
    if status == "review_required":
        return row.get("allocation_code_reason") or "Cần review mã phân bổ"
    return row.get("eligibility_reason", "")


@app.get("/", response_class=HTMLResponse)
@app.get("/clients", response_class=HTMLResponse)
async def clients(request: Request):
    clients = portfolio_service.clients()
    if co_auth.auth_required():
        clients = co_auth.filter_visible_clients(clients, co_auth.current_user(request))
    return templates.TemplateResponse(
        request=request,
        name="clients.html",
        context={"clients": clients},
    )


@app.get("/clients/{client_id}", response_class=HTMLResponse)
async def workspace(request: Request, client_id: str):
    # Workspace overview only renders client.counts tiles. Avoid the full
    # source_workspace + bom_service.workspace pagination here — those would
    # paginate every BCCT/material/BOM row from Data Hub on each render.
    return templates.TemplateResponse(
        request=request,
        name="workspace.html",
        context=client_overview_context(client_id),
    )


def client_overview_context(client_id: str) -> dict:
    client = resolve_client(client_id)
    try:
        source_summary, source_backend = portfolio_service.source_summary(client)
    except Exception:  # noqa: BLE001
        source_summary, source_backend = {
            "material_catalog": {"published_row_count": 0},
            "product_catalog": {"published_row_count": 0},
            "bcct": {"published_row_count": 0},
            "co_stock_row_count": 0,
        }, "n/a"
    client = enrich_client_with_source_summary(client, source_summary)
    # Workspace template references client.counts.bom_lines too; surface a
    # zero so the tile renders rather than crashes.
    client["counts"]["bom_lines"] = client["counts"].get("bom_lines", 0)
    return {
        "client": client,
        "case": client_case(client),
        "active": "overview",
        "source_backend": source_backend,
    }


@app.get("/clients/{client_id}/bom", response_class=HTMLResponse)
async def bom(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="bom.html",
        context=bom_context(request, client_id),
    )


@app.post("/clients/{client_id}/bom/config", response_class=HTMLResponse)
async def save_bom_config(request: Request, client_id: str):
    require_local_source_writes()
    client = resolve_client(client_id)
    form = await request.form()
    bom_service.update_config(client, {key: str(value) for key, value in form.items()})
    return templates.TemplateResponse(
        request=request,
        name="bom.html",
        context=bom_context(request, client_id, message="Đã lưu cấu hình BOM cho công ty này."),
    )


@app.get("/clients/{client_id}/config", response_class=HTMLResponse)
async def client_config(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="client_config.html",
        context=config_context(client_id),
    )


@app.post("/clients/{client_id}/config", response_class=HTMLResponse)
async def save_client_config_route(request: Request, client_id: str):
    client = resolve_client(client_id)
    form = await request.form()
    # Identity fields (legal_name / tax_code) are CO-side render metadata, not
    # source data — editable even when Data Hub source-mode is enabled. The
    # `co_stock_min_days_before_export` knob is also CO-side: it controls a
    # local CO eligibility predicate, not anything DH owns, so it persists to
    # the same local overlay without going through `require_local_source_writes`.
    if "legal_name" in form or "tax_code" in form or "co_stock_min_days_before_export" in form:
        client = dict(client)
        if "legal_name" in form:
            client["legal_name"] = str(form.get("legal_name") or "").strip()
        if "tax_code" in form:
            client["tax_code"] = str(form.get("tax_code") or "").strip()
        if "co_stock_min_days_before_export" in form:
            raw = str(form.get("co_stock_min_days_before_export") or "").strip()
            overrides = dict(client.get("co_stock_overrides") or {})
            if raw == "":
                overrides.pop("min_days_before_export", None)
            else:
                try:
                    n = int(raw)
                    if n < 0:
                        n = co_stock_eligibility.DEFAULT_MIN_GAP_DAYS
                except ValueError:
                    n = co_stock_eligibility.DEFAULT_MIN_GAP_DAYS
                overrides["min_days_before_export"] = n
            client["co_stock_overrides"] = overrides
        store = get_app_state_store()
        if store:
            store.upsert_client(client)
    require_local_source_writes()
    config = portfolio_service.get_client_config(client)
    config["co_stock"]["lot_policy"] = str(form.get("co_stock_lot_policy", "line_level"))
    config["allocation_code"]["strategy"] = str(form.get("allocation_code_strategy", "same_as_customs_code"))
    config["allocation_code"]["description_regex"] = str(form.get("description_regex", ""))
    config["allocation_code"]["fallback"] = str(form.get("allocation_code_fallback", "same_as_customs_code"))
    try:
        portfolio_service.save_client_config(client, config)
    except ValueError as exc:
        return templates.TemplateResponse(
            request=request,
            name="client_config.html",
            status_code=400,
            context=config_context(client_id, error=str(exc)),
        )
    portfolio_service.refresh_client_indexes(client)
    return templates.TemplateResponse(
        request=request,
        name="client_config.html",
        context=config_context(client_id, message="Đã lưu cấu hình công ty."),
    )


# ---------- Cost allocation ratios (LVC/RVC cost-buildup auto-fill) ----------


def _decimal_str(value: Decimal) -> str:
    """Render Decimal for the admin form: strip trailing zeros but keep at
    least one digit. Empty for zero (so the placeholder shows)."""
    if value is None or value == 0:
        return ""
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def _row_for_template(row) -> dict:
    return {
        "product_code": row.product_code,
        "coef_wages_str": _decimal_str(row.coef_wages),
        "coef_welfare_str": _decimal_str(row.coef_welfare),
        "coef_rent_str": _decimal_str(row.coef_rent),
        "coef_depreciation_str": _decimal_str(row.coef_depreciation),
        "coef_other_mfg_str": _decimal_str(row.coef_other_mfg),
        "coef_transport_storage_str": _decimal_str(row.coef_transport_storage),
        "note": row.note,
        "has_value": any([
            row.coef_wages, row.coef_welfare, row.coef_rent,
            row.coef_depreciation, row.coef_other_mfg, row.coef_transport_storage,
            row.note,
        ]),
    }


def _cost_allocation_context(client_id: str, **extra) -> dict:
    """Lightweight context for the cost-allocation admin page.

    Deliberately avoids `client_context()` because that helper pulls the full
    source workspace (BCCT scan ~2.6s on Growatt) which this page does not
    need. We render the nav with no counts; the rest of the template only
    needs client identity + the ratio rows.
    """
    from app import cost_allocation_store
    from app.cost_allocation_store import CostAllocationRow
    client = resolve_client(client_id)
    rows = cost_allocation_store.list_ratios(client_id)
    mode_b = cost_allocation_store.get_mode_b_default(client_id) or CostAllocationRow(product_code="")
    return {
        "client": client,
        "active": "cost-allocation",
        "rows": [_row_for_template(r) for r in sorted(rows, key=lambda r: r.product_code)],
        "mode_b": _row_for_template(mode_b),
        **extra,
    }


def _coef_from_form(form, key: str) -> Decimal:
    raw = str(form.get(key) or "").strip().replace(",", ".")
    if not raw:
        return Decimal(0)
    try:
        value = Decimal(raw)
    except (InvalidOperation, ValueError):
        return Decimal(0)
    return value if value >= 0 else Decimal(0)


@app.get("/clients/{client_id}/cost-allocation", response_class=HTMLResponse)
async def cost_allocation_page(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id),
    )


@app.post("/clients/{client_id}/cost-allocation/mode-b", response_class=HTMLResponse)
async def cost_allocation_save_mode_b(request: Request, client_id: str):
    resolve_client(client_id)  # validates client exists
    from app import cost_allocation_store
    form = await request.form()
    row = cost_allocation_store.CostAllocationRow(
        product_code="",
        coef_wages=_coef_from_form(form, "coef_wages"),
        coef_welfare=_coef_from_form(form, "coef_welfare"),
        coef_rent=_coef_from_form(form, "coef_rent"),
        coef_depreciation=_coef_from_form(form, "coef_depreciation"),
        coef_other_mfg=_coef_from_form(form, "coef_other_mfg"),
        coef_transport_storage=_coef_from_form(form, "coef_transport_storage"),
        note=str(form.get("note", "") or "").strip(),
    )
    cost_allocation_store.upsert_ratio(client_id, row)
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id, message="Đã lưu hệ số mặc định Mode B."),
    )


@app.post("/clients/{client_id}/cost-allocation/mode-b/delete", response_class=HTMLResponse)
async def cost_allocation_delete_mode_b(request: Request, client_id: str):
    resolve_client(client_id)
    from app import cost_allocation_store
    cost_allocation_store.delete_ratio(client_id, "")
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id, message="Đã xóa hệ số mặc định Mode B."),
    )


@app.post("/clients/{client_id}/cost-allocation/row/delete", response_class=HTMLResponse)
async def cost_allocation_delete_row(request: Request, client_id: str, product_code: str = Form(...)):
    resolve_client(client_id)
    from app import cost_allocation_store
    cost_allocation_store.delete_ratio(client_id, product_code.strip())
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(client_id, message=f"Đã xóa hệ số cho {product_code}."),
    )


@app.post("/clients/{client_id}/cost-allocation/upload", response_class=HTMLResponse)
async def cost_allocation_upload(request: Request, client_id: str, file: UploadFile = File(...)):
    resolve_client(client_id)
    from app import cost_allocation_store, cost_allocation_importer
    try:
        parsed = cost_allocation_importer.parse_excel(await file.read())
    except Exception as exc:  # noqa: BLE001
        return templates.TemplateResponse(
            request=request,
            name="cost_allocation.html",
            status_code=400,
            context=_cost_allocation_context(client_id, error=f"Không đọc được file: {exc}"),
        )
    diff = cost_allocation_store.replace_all(client_id, parsed)
    return templates.TemplateResponse(
        request=request,
        name="cost_allocation.html",
        context=_cost_allocation_context(
            client_id,
            message=f"Đã import {len(parsed)} dòng từ {file.filename}.",
            upload_diff=diff,
        ),
    )


@app.get("/clients/{client_id}/cost-allocation/template.xlsx")
async def cost_allocation_template(client_id: str):
    resolve_client(client_id)
    from openpyxl import Workbook
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet3"
    ws["A1"] = "BẢNG PHÂN BỔ TỶ LỆ CHI PHÍ"
    ws["A2"] = "STT"
    ws["B2"] = "Mã SP"
    ws["C2"] = "Lương, thưởng"
    ws["D2"] = "Phúc lợi y tế"
    ws["E2"] = "Phí thuê nhà xưởng"
    ws["F2"] = "Phí khấu hao, BH, BD"
    ws["G2"] = "CP SX chung khác"
    ws["H2"] = "Lợi nhuận (bỏ qua khi import)"
    ws["I2"] = "Vận chuyển, lưu kho, dịch vụ"
    ws["J2"] = "Ghi chú"
    # Row 4 onward = data area.
    ws["A4"] = 1
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="cost-allocation-{client_id}-template.xlsx"'},
    )


@app.get("/clients/{client_id}/cost-allocation/resolve")
async def cost_allocation_resolve(client_id: str, product_code: str, fob: str = "0"):
    """JSON endpoint for the 'Áp hệ số' button on the origin product panel.

    Returns the resolved coefficient (Mode A → Mode B fallback) plus the
    multiplied detail values. `found` is false when neither mode matches.
    """
    resolve_client(client_id)
    from app import cost_allocation_store, cost_allocation_importer
    row = cost_allocation_store.get_ratio(client_id, product_code.strip())
    if row is None:
        return JSONResponse({"found": False, "product_code": product_code})
    try:
        fob_dec = Decimal(str(fob).replace(",", "."))
    except (InvalidOperation, ValueError):
        fob_dec = Decimal(0)
    detail = cost_allocation_importer.apply_to_fob(row, fob_dec)
    return JSONResponse({
        "found": True,
        "product_code": product_code,
        "matched_mode": "A" if row.product_code else "B",
        "matched_product_code": row.product_code,
        "fob": str(fob_dec),
        "coefficients": {
            "wages": str(row.coef_wages),
            "welfare": str(row.coef_welfare),
            "rent": str(row.coef_rent),
            "depreciation": str(row.coef_depreciation),
            "other_mfg": str(row.coef_other_mfg),
            "transport_storage": str(row.coef_transport_storage),
        },
        "details": {k: str(v) for k, v in detail.items()},
        "note": row.note,
    })


@app.post("/clients/{client_id}/bom/upload", response_class=HTMLResponse)
async def upload_bom_workbook(
    request: Request,
    client_id: str,
    file: UploadFile = File(...),
    upload_mode: str = Form("direct_bom"),
    upload_scope: str = Form(""),
    accept_review_required: str = Form(""),
):
    require_local_source_writes()
    client = resolve_client(client_id)
    result = bom_service.process_upload(
        client,
        await file.read(),
        file.filename or "bom.xlsx",
        upload_mode,
        upload_scope or None,
        accept_review_required == "on",
    )
    status_code = 400 if result["status"] == "failed" else 200
    return templates.TemplateResponse(
        request=request,
        name="bom.html",
        status_code=status_code,
        context=bom_context(
            request,
            client_id,
            bom_result=result,
            message=result["message"] if status_code == 200 else "",
            error=result["message"] if status_code == 400 else "",
        ),
    )


@app.get("/clients/{client_id}/bom/template.xlsx")
async def download_bom_template(client_id: str):
    try:
        content = bom_service.template(resolve_client(client_id))
    except RuntimeError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{client_id}-bom-template.xlsx"'},
    )


@app.get("/clients/{client_id}/co-stock", response_class=HTMLResponse)
async def co_stock(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="co_stock.html",
        context=co_stock_table_context(request, client_id),
    )


@app.get("/clients/{client_id}/bcct", response_class=HTMLResponse)
async def bcct(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="bcct.html",
        context=bcct_table_context(request, client_id),
    )


@app.get("/clients/{client_id}/bcct/imports", response_class=HTMLResponse)
async def bcct_imports(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="bcct.html",
        context=bcct_table_context(request, client_id, "import"),
    )


@app.get("/clients/{client_id}/bcct/exports", response_class=HTMLResponse)
async def bcct_exports(request: Request, client_id: str):
    return templates.TemplateResponse(
        request=request,
        name="bcct.html",
        context=bcct_table_context(request, client_id, "export"),
    )


@app.get("/clients/{client_id}/bcct/template.xlsx")
async def download_bcct_template(client_id: str):
    require_local_source_writes()
    content = portfolio_service.bcct_template(resolve_client(client_id))
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{client_id}-bcct-template.xlsx"'},
    )


@app.post("/clients/{client_id}/bcct/upload", response_class=HTMLResponse)
async def upload_bcct_workbook(request: Request, client_id: str, file: UploadFile = File(...)):
    require_local_source_writes()
    client = resolve_client(client_id)
    result = portfolio_service.process_bcct_upload(client, await file.read(), file.filename or "bcct.xlsx")
    status_code = 400 if result["status"] == "failed" else 200
    return templates.TemplateResponse(
        request=request,
        name="bcct.html",
        status_code=status_code,
        context=bcct_table_context(
            request,
            client_id,
            bcct_result=result,
            message=result["message"] if status_code == 200 else "",
            error=result["message"] if status_code == 400 else "",
        ),
    )


@app.post("/clients/{client_id}/co-stock/import")
async def import_co_stock_workbook(client_id: str, file: UploadFile = File(...)):
    """Upload a standard CO stock template xlsx. Overwrites prior snapshot
    rows for the same (declaration_no, line_no, customs_code) keys; preserves
    rows untouched by this upload (so a partial upload only updates what it
    covers).

    Run scripts/convert_co_stock.py first if uploading from the agency
    `tru-lui-co-template.xlsm` workbook.
    """
    client = resolve_client(client_id)
    content = await file.read()
    try:
        rows, parse_errors = read_standard_co_stock(content)
    except CoStockTemplateError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    batch_id = "batch_" + hashlib.sha1(content).hexdigest()[:16]
    summary = co_stock_adjustments_store.upsert_batch(
        client["id"],
        rows,
        batch_id=batch_id,
        source_file_ref=file.filename or "co_stock.xlsx",
    )
    # Invalidate the cached source context so the next case page reflects the
    # new adjustments. The cache is process-local so this is cheap.
    _CO_CASE_SOURCE_CACHE.clear()
    return JSONResponse({
        "ok": not summary.get("errors"),
        "client_id": client["id"],
        "batch_id": batch_id,
        "filename": file.filename or "",
        "parsed_rows": len(rows),
        "parse_errors": parse_errors,
        "upsert": summary,
    })


@app.post("/clients/{client_id}/co-stock/refresh")
async def refresh_co_stock_endpoint(client_id: str):
    """Materialize the derived stock pool for one client into CO's
    co_stock_rows table.

    Tries the Data Hub delta path first (`list_bcct_with_envelope` with
    `since` + `include_tombstones`) when we have a stored
    `last_bcct_server_time` AND the response carries a fresh `server_time`.
    Falls back to full-pull when either is missing — the very first refresh
    of a client, or when Data Hub is on an older contract.
    """
    client = resolve_client(client_id)
    summary = _refresh_co_stock_delta_or_full(client)
    # Invalidate the case-source cache so the substitute modal / sheet calc
    # paths see the same fresh data.
    _CO_CASE_SOURCE_CACHE.clear()
    return JSONResponse({"ok": not summary.get("errors"), **summary})


@app.get("/clients/{client_id}/co-stock/lot-history")
async def co_stock_lot_history(
    client_id: str,
    declaration_no: str = "",
    line_no: str = "",
    customs_code: str = "",
    limit: int = 200,
):
    """Return chronological audit log for one lot: claim_lock/release events
    from sheet allocations + adjustment_import_insert/update from manual
    workbook uploads. Newest first.
    """
    client = resolve_client(client_id)
    if not (declaration_no.strip() and line_no.strip() and customs_code.strip()):
        raise HTTPException(status_code=400, detail="declaration_no, line_no, customs_code required")
    events = co_stock_events_store.events_for_lot(
        client["id"],
        declaration_no.strip(),
        line_no.strip(),
        customs_code.strip(),
        limit=max(1, min(int(limit), 500)),
    )
    return JSONResponse({
        "ok": True,
        "client_id": client["id"],
        "lot": {"declaration_no": declaration_no, "line_no": line_no, "customs_code": customs_code},
        "events": events,
        "count": len(events),
    })


@app.get("/clients/{client_id}/co-stock/export.xlsx")
async def export_co_stock_workbook(client_id: str):
    """Dump effective ton CO state (BCCT opening + ledger + adjustments) into
    a standard template xlsx. Heavy for big clients (full BCCT pagination
    on Data Hub mode); intended for on-demand download.
    """
    client = resolve_client(client_id)
    workspace, _backend = portfolio_service.source_workspace(client)
    stock_rows = [dict(row) for row in workspace.get("co_stock_rows") or []]
    client_id_value = client.get("id", "")
    used_by_lot = co_stock_ledger.used_qty_by_lot(client_id_value)
    if used_by_lot:
        stock_rows = co_stock_ledger.apply_used_qty(stock_rows, used_by_lot)
    adjustments = co_stock_adjustments_store.aggregate_by_lookup_key(client_id_value)
    if adjustments:
        stock_rows = co_stock_adjustments_store.apply_adjustments(stock_rows, adjustments)
    rows_for_template = []
    for row in stock_rows:
        rows_for_template.append({
            "declaration_no": row.get("import_declaration_no", ""),
            "registration_date": row.get("registration_date") or row.get("declaration_date") or row.get("import_declaration_date") or "",
            "declaration_type": row.get("declaration_type", ""),
            "line_no": row.get("line_no", ""),
            "customs_code": row.get("customs_item_code", ""),
            "hs_code": row.get("hs_code", ""),
            "goods_name": row.get("material_description") or row.get("goods_name", ""),
            "origin_country": row.get("origin_country", ""),
            "unit_price": row.get("unit_value") or row.get("unit_price", ""),
            "taxable_unit_price": row.get("taxable_unit_price", ""),
            "opening_qty": row.get("available_qty", ""),
            "unit": row.get("unit", ""),
            "partner": row.get("partner", ""),
            "invoice_no": row.get("invoice_no") or row.get("invoice_ref", ""),
            "invoice_date": row.get("invoice_date", ""),
            "exchange_rate": row.get("exchange_rate", ""),
            "used_qty": row.get("used_qty", ""),
            "source_co_no": "",  # Per-CO attribution requires reading co_stock_claims.case_id; out of scope here.
            "transaction_key": row.get("source_transaction_key", ""),
        })
    xlsx = write_standard_co_stock(rows_for_template, use_labels=True)
    filename = f"{client_id_value}-co-stock-{date.today().isoformat()}.xlsx"
    return Response(
        content=xlsx,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# Specific GET routes that share the /clients/{client_id}/co-case/{case_id}/...
# prefix MUST be defined before the catch-all {step} route below — FastAPI's
# router is order-sensitive, and {step} would otherwise swallow any single-
# segment GET (e.g. /export-bang-ke) and 404 it for not being a workflow key.


@app.post("/clients/{client_id}/evaluate", response_class=HTMLResponse)
async def evaluate(request: Request, client_id: str):
    form = await large_request_form(request)
    case = update_products_from_form({key: str(value) for key, value in form.items()})
    case_id = case.get("persisted_case_id", "")
    if case_id:
        client = resolve_client(client_id)
        lock_result = acquire_origin_calculation_lock(client, case_id, origin_lock_actor(request))
        if not lock_result["acquired"]:
            return templates.TemplateResponse(
                request=request,
                name="co_case.html",
                status_code=409,
                context=co_case_context(
                    client_id,
                    case=case,
                    current_step="origin",
                    error=f"Chưa thể tính lại: hồ sơ {lock_result['lock'].get('case_code') or lock_result['lock'].get('case_id')} đang giữ phiên tính tồn cho khách hàng này.",
                    origin_calculation_blocked=True,
                    preserve_origin_products=True,
                ),
            )
    context = co_case_context(
        client_id,
        case=case,
        current_step="origin",
        message="Đã tính lại theo dữ liệu đang sửa.",
        preserve_origin_products=False,
    )
    if context["case"].get("persisted_case_id"):
        try:
            update_case_record(
                resolve_client(client_id),
                case if context.get("origin_demo_active") else context["case"],
            )
        except KeyError:
            pass
    return templates.TemplateResponse(
        request=request,
        name="co_case.html",
        context=context,
    )


@app.post("/clients/{client_id}/upload", response_class=HTMLResponse)
async def upload_workbook(request: Request, client_id: str, file: UploadFile = File(...)):
    client = resolve_client(client_id)
    content = await file.read()
    try:
        case = parse_input_workbook(content, source_label=f"Upload: {file.filename}")
    except WorkbookParseError as exc:
        return templates.TemplateResponse(
            request=request,
            name="co_case.html",
            status_code=400,
            context=co_case_context(client_id, error=str(exc)),
        )
    case["customer"] = client.get("legal_name") or client["name"]
    case["customer_legal_name"] = client.get("legal_name", "")
    case["customer_tax_code"] = client.get("tax_code", "")
    return templates.TemplateResponse(
        request=request,
        name="co_case.html",
        context=co_case_context(client_id, case=case, current_step="origin", message=f"Đã parse {file.filename}."),
    )


@app.get("/clients/{client_id}/demo-input.xlsx")
async def download_demo_input(client_id: str):
    content = create_input_workbook(client_case(resolve_client(client_id)))
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{client_id}-demo-input.xlsx"'},
    )


@app.post("/clients/{client_id}/export")
async def export_evidence(request: Request, client_id: str):
    resolve_client(client_id)
    form = await large_request_form(request)
    case = update_products_from_form({key: str(value) for key, value in form.items()})
    content = create_evidence_workbook(case)
    filename = f"{case['case_code'] or 'co-case'}-evidence.xlsx"
    return StreamingResponse(
        iter([content]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
