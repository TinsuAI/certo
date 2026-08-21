"""Raw BOM edge parsers.

These parsers preserve factory/SAP direct structure as parent -> child
edges. They deliberately do not flatten to leaves; flattening is a later
materialization step.
"""
from __future__ import annotations

import io
from collections import defaultdict, deque
from decimal import Decimal

from hub.app.parsers._excel import header_row, index_headers, load_xlsx
from hub.app.parsers.bom_adapters import BomParseError
from hub.app.parsers.bom_adapters._common import (
    COMMON_ALIASES, cell_num, cell_str,
)
from hub.app.parsers.bom_adapters.sap_indented_walk import (
    _COMPONENT_ALIASES, _DESCRIPTION_ALIASES, _LEVEL_ALIASES,
    _PRODUCT_CODE_ALIASES, _QTY_ALIASES, _UNIT_ALIASES, _col_index,
)


RawEdge = dict


def parse_raw_edges_with_fallback(
    blob: bytes, *, root_code: str | None = None,
) -> tuple[list[RawEdge], str]:
    errors: list[str] = []
    for name, fn in (
        ("sap_indented_raw", parse_sap_indented_raw_edges),
        ("growatt_factory_technical", parse_growatt_factory_edges),
    ):
        try:
            return fn(blob, root_code=root_code), name
        except BomParseError as exc:
            errors.append(f"{name}: {exc}")
    raise BomParseError("; ".join(errors) or "No raw BOM edge parser matched")


def parse_growatt_factory_edges(
    blob: bytes, *, root_code: str | None = None,
) -> list[RawEdge]:
    try:
        wb = load_xlsx(blob)
    except Exception as exc:
        raise BomParseError(f"Cannot open workbook: {exc}") from exc

    edges: list[RawEdge] = []
    row_index = 0
    for ws in wb.worksheets:
        hdr = header_row(ws, aliases=COMMON_ALIASES)
        if not hdr:
            continue
        header_idx, headers = hdr
        cols = index_headers(headers, COMMON_ALIASES)
        if "product_code" not in cols or "material_code" not in cols:
            continue
        top_base_idx = _header_index(headers, ["顶层基本数量"])
        child_qty_idx = _header_index(headers, ["子项数量"])
        for source_row_no, raw in _iter_rows_with_index(ws, header_idx):
            parent = cell_str(raw, cols.get("product_code"))
            child = cell_str(raw, cols.get("material_code"))
            qty = cell_num(raw, cols.get("qty_per_unit"))
            if qty is None:
                qty = _qty_from_child_and_base(
                    raw, child_qty_idx=child_qty_idx, top_base_idx=top_base_idx,
                )
            if not parent or not child:
                continue
            row_index += 1
            edge_root = root_code or parent
            edges.append({
                "row_index": row_index - 1,
                "root_code": edge_root,
                "parent_code": parent,
                "child_code": child,
                "qty_per_parent": qty if qty is not None else 0.0,
                "uom": cell_str(raw, cols.get("uom")),
                "level": None,
                "node_path": None,
                "sheet_name": ws.title,
                "source_row_no": source_row_no,
                "payload": {
                    "adapter": "growatt_factory_technical",
                    "workbook_root_hint": root_code,
                },
            })
    if not edges:
        raise BomParseError("No raw parent/child BOM edges recognized")
    parents = {e["parent_code"] for e in edges}
    children = {e["child_code"] for e in edges}
    roots = parents - children
    if root_code and root_code not in roots and len(roots) > 1:
        raise BomParseError(
            "Workbook has multiple parent products but none matches the "
            "filename root hint; likely a combined/staff BOM, not one "
            "factory technical raw workbook",
        )
    _validate_positive_qty(edges)
    _assign_component_roots(edges, preferred_root=root_code)
    _annotate_paths(edges)
    return edges


def parse_sap_indented_raw_edges(
    blob: bytes, *, root_code: str | None = None,
) -> list[RawEdge]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    except Exception as exc:
        raise BomParseError(f"Cannot open workbook: {exc}") from exc

    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise BomParseError("Empty workbook")
    header = [str(c or "").strip() for c in rows[0]]
    if _col_index(header, _PRODUCT_CODE_ALIASES) is not None:
        raise BomParseError("Has product_code column - not SAP indented raw")

    c_level = _col_index(header, _LEVEL_ALIASES)
    c_comp = _col_index(header, _COMPONENT_ALIASES)
    c_qty = _col_index(header, _QTY_ALIASES)
    c_unit = _col_index(header, _UNIT_ALIASES)
    c_desc = _col_index(header, _DESCRIPTION_ALIASES)
    if c_level is None or c_comp is None or c_qty is None:
        raise BomParseError("Missing Level / Component / Qty columns")

    root = root_code or "__root__"
    parent_stack: list[str] = [root]
    edges: list[RawEdge] = []

    for source_row_no, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or c == "" for c in row):
            continue
        level = _to_int(row[c_level]) if c_level < len(row) else None
        if level is None or level < 1:
            continue
        comp = row[c_comp] if c_comp < len(row) else None
        if comp is None or str(comp).strip() == "":
            continue
        child = str(comp).strip()
        qty = _to_decimal(row[c_qty] if c_qty < len(row) else None)
        uom = (
            str(row[c_unit]).strip()
            if c_unit is not None and c_unit < len(row) and row[c_unit]
            else None
        )
        desc = (
            str(row[c_desc]).strip()
            if c_desc is not None and c_desc < len(row) and row[c_desc]
            else None
        )

        while len(parent_stack) > level:
            parent_stack.pop()
        parent = parent_stack[-1]
        node_path = " > ".join([*parent_stack, child])
        payload: dict = {"adapter": "sap_indented_raw"}
        if desc:
            payload["description"] = desc
        edges.append({
            "row_index": len(edges),
            "root_code": root,
            "parent_code": parent,
            "child_code": child,
            "qty_per_parent": float(qty),
            "uom": uom,
            "level": level,
            "node_path": node_path,
            "sheet_name": ws.title,
            "source_row_no": source_row_no,
            "payload": payload,
        })
        if len(parent_stack) == level:
            parent_stack.append(child)
        else:
            parent_stack[level:] = [child]

    if not edges:
        raise BomParseError("No raw SAP indented BOM edges recognized")
    _validate_positive_qty(edges)
    return edges


def _iter_rows_with_index(ws, header_row_idx: int):
    for offset, row in enumerate(
        ws.iter_rows(min_row=header_row_idx + 1, values_only=True), start=1,
    ):
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        yield header_row_idx + offset, row


def _header_index(headers: list[str], aliases: list[str]) -> int | None:
    norm = [str(h or "").strip().lower() for h in headers]
    for alias in aliases:
        a = alias.lower()
        for i, h in enumerate(norm):
            if h == a:
                return i
    return None


def _qty_from_child_and_base(
    row, *, child_qty_idx: int | None, top_base_idx: int | None,
) -> float | None:
    child_qty = _num(row, child_qty_idx)
    top_base = _num(row, top_base_idx)
    if child_qty is None or top_base in (None, 0):
        return None
    return child_qty / top_base


def _num(row, idx: int | None) -> float | None:
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _validate_positive_qty(edges: list[RawEdge]) -> None:
    bad = [
        e for e in edges
        if e.get("qty_per_parent") is None or float(e.get("qty_per_parent") or 0) <= 0
    ]
    if bad:
        sample = ", ".join(
            f"{e.get('parent_code')}->{e.get('child_code')}={e.get('qty_per_parent')!r}"
            for e in bad[:5]
        )
        raise BomParseError(
            f"Raw BOM has {len(bad)} edge(s) with qty_per_parent <= 0. "
            f"First few: {sample}",
        )


def _assign_component_roots(
    edges: list[RawEdge], *, preferred_root: str | None,
) -> None:
    children_by_parent: dict[str, list[RawEdge]] = defaultdict(list)
    parents: set[str] = set()
    children: set[str] = set()
    for edge in edges:
        parents.add(edge["parent_code"])
        children.add(edge["child_code"])
        children_by_parent[edge["parent_code"]].append(edge)

    roots = sorted(parents - children)
    if not roots:
        roots = [preferred_root or edges[0]["root_code"]]
    if preferred_root in roots:
        roots = [preferred_root, *[r for r in roots if r != preferred_root]]

    assigned: set[int] = set()
    for root in roots:
        q: deque[str] = deque([root])
        seen_nodes = {root}
        while q:
            parent = q.popleft()
            for edge in children_by_parent.get(parent, []):
                assigned.add(id(edge))
                edge["root_code"] = root
                child = edge["child_code"]
                if child not in seen_nodes:
                    seen_nodes.add(child)
                    q.append(child)

    fallback = preferred_root or roots[0]
    for edge in edges:
        if id(edge) not in assigned:
            edge["root_code"] = fallback


def _annotate_paths(edges: list[RawEdge]) -> None:
    by_root: dict[str, list[RawEdge]] = defaultdict(list)
    for edge in edges:
        by_root[edge["root_code"]].append(edge)

    for root, root_edges in by_root.items():
        children_by_parent: dict[str, list[RawEdge]] = defaultdict(list)
        for edge in root_edges:
            children_by_parent[edge["parent_code"]].append(edge)

        path_by_parent: dict[str, tuple[int, str]] = {root: (0, root)}
        q: deque[str] = deque([root])
        visited_edges: set[int] = set()
        while q:
            parent = q.popleft()
            parent_level, parent_path = path_by_parent[parent]
            for edge in children_by_parent.get(parent, []):
                visited_edges.add(id(edge))
                level = parent_level + 1
                child = edge["child_code"]
                path = f"{parent_path} > {child}"
                edge["level"] = level
                edge["node_path"] = path
                if child not in path_by_parent:
                    path_by_parent[child] = (level, path)
                    q.append(child)

        for edge in root_edges:
            if id(edge) in visited_edges:
                continue
            edge["level"] = edge.get("level") or 1
            edge["node_path"] = edge.get("node_path") or (
                f"{edge['parent_code']} > {edge['child_code']}"
            )


def _to_int(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return None


def _to_decimal(value) -> Decimal:
    if value is None or value == "":
        return Decimal(0)
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value))
    except Exception:  # noqa: BLE001
        return Decimal(0)
