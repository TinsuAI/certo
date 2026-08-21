"""system_template — the canonical NXT workbook Data Hub itself ships.

Three data sheets named exactly NVL / TP / BTP (+ a Hướng dẫn sheet). Each
sheet: header row 1, data row 2+. The sheet name sets `reported_role`
(provenance). Period + client are supplied by the upload form, not the file,
so the sheets stay pure data tables and parsing is deterministic.

This module is the single source of truth for the canonical column set: it
both RENDERS the downloadable template and PARSES it back.
"""
from __future__ import annotations

import io

from app.parsers._excel import (
    cell_num, cell_str, header_row, index_headers, load_xlsx,
)
from app.parsers.nxt_adapters._common import ALIASES, NUMERIC_FIELDS, OUT_BUCKETS

# Canonical column order: (header text written into the template, logical field).
CANONICAL_COLUMNS: list[tuple[str, str]] = [
    ("Mã nội bộ", "internal_code"),
    ("Mã hải quan", "customs_code"),
    ("Tên", "name"),
    ("ĐVT", "uom"),
    ("Tồn đầu kỳ", "opening"),
    ("Nhập trong kỳ", "inbound_total"),
    ("Tái xuất", "out_tai_xuat"),
    ("Chuyển MĐSD/TTNĐ/tiêu hủy", "out_chuyen_mdsd"),
    ("Xuất kho SX", "out_xuat_sx"),
    ("Xuất kho khác", "out_xuat_khac"),
    ("Tồn cuối kỳ", "closing_reported"),
    ("Ghi chú", "note"),
]

# Sheet name → reported_role provenance.
SHEET_ROLES = {"NVL": "nvl", "TP": "tp", "BTP": "btp"}


class SystemTemplateNxtAdapter:
    name = "system_template"
    label_key = "nxt.adapter.system_template.label"
    description_key = "nxt.adapter.system_template.desc"
    supports_mapping_override = False

    def detect(self, blob: bytes) -> float | None:
        """High precision: a NVL/TP/BTP-named sheet that actually carries our
        canonical headers (opening + closing + a code column) is unambiguously
        ours. Checking headers — not just the sheet name — avoids claiming an
        agency's own bilingual NVL/TP/BTP sheets (handled by ezsoft_3tsoft)."""
        try:
            wb = load_xlsx(blob)
        except Exception:
            return None
        for ws in wb.worksheets:
            if ws.title.strip().upper() not in SHEET_ROLES:
                continue
            hdr = header_row(ws, aliases=ALIASES)
            if not hdr:
                continue
            cols = index_headers(hdr[1], ALIASES)
            if ("opening" in cols and "closing_reported" in cols
                    and ("internal_code" in cols or "customs_code" in cols)):
                return 0.95
        return None

    def parse(self, blob: bytes, *,
              mapping_override: dict[str, str] | None = None) -> list[dict]:
        from app.parsers.nxt_adapters import NxtParseError
        try:
            wb = load_xlsx(blob)
        except Exception as e:
            raise NxtParseError(f"Cannot open workbook: {e}") from e

        lines: list[dict] = []
        for ws in wb.worksheets:
            role = SHEET_ROLES.get(ws.title.strip().upper())
            if role is None:
                continue
            hdr = header_row(ws, aliases=ALIASES)
            if not hdr:
                continue
            header_idx, headers = hdr
            cols = index_headers(headers, ALIASES)
            if "internal_code" not in cols and "customs_code" not in cols:
                continue
            for raw in _iter_data_rows(ws, header_idx):
                line = _row_to_line(raw, cols, role)
                if line is not None:
                    lines.append(line)

        if not lines:
            raise NxtParseError(
                "No NXT rows recognized in NVL/TP/BTP sheets.")
        return lines


def _iter_data_rows(ws, header_row_idx: int):
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        yield row


def _row_to_line(raw, cols: dict[str, int], role: str) -> dict | None:
    internal = cell_str(raw, cols.get("internal_code"))
    customs = cell_str(raw, cols.get("customs_code"))
    if not internal and not customs:
        return None
    line = {
        "internal_code": internal,
        "customs_code": customs,
        "name": cell_str(raw, cols.get("name")),
        "uom": cell_str(raw, cols.get("uom")),
        "reported_role": role,
        "note": cell_str(raw, cols.get("note")),
    }
    for f in NUMERIC_FIELDS:
        line[f] = cell_num(raw, cols.get(f))
    # Canonical total = sum of the 4 regulatory buckets the template carries.
    buckets = [line.get(f) for f in OUT_BUCKETS]
    line["outbound_total"] = (
        sum(b or 0.0 for b in buckets) if any(b is not None for b in buckets)
        else None
    )
    return line


def render_template_xlsx() -> bytes:
    """Build the downloadable canonical NXT template: NVL/TP/BTP data sheets
    (header row + 1 sample row each) + a Hướng dẫn sheet."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    headers = [h for h, _ in CANONICAL_COLUMNS]
    sample = {
        "NVL": ["NVL-001", "A0123", "Nhựa ABS", "KG", 1000, 5000, 0, 0, 4500, 0, 1500, ""],
        "TP": ["TP-001", "SP0123", "Bộ biến tần 5kW", "PCE", 50, 0, 0, 0, 0, 30, 20, ""],
        "BTP": ["BTP-001", "", "Bo mạch điều khiển", "PCE", 100, 200, 0, 0, 250, 0, 50, ""],
    }
    for sheet in ("NVL", "TP", "BTP"):
        ws = wb.create_sheet(sheet)
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9E1F2")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for ci, v in enumerate(sample[sheet], 1):
            ws.cell(row=2, column=ci, value=v)
        for ci in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + ci)].width = 16

    inst = wb.create_sheet("Hướng dẫn")
    inst["A1"] = "Mẫu nhập Nhập-Xuất-Tồn (NXT) — chuẩn hệ thống"
    inst["A1"].font = Font(size=13, bold=True)
    lines = [
        "",
        "Mỗi sheet = một loại: NVL (nguyên vật liệu), TP (thành phẩm), BTP (bán thành phẩm).",
        "Tên sheet quyết định loại được ghi nhận (reported_role) — không cần cột phân loại.",
        "Phân loại chính thức của mã vẫn lấy từ Danh Mục; cột này chỉ là dấu vết nguồn.",
        "",
        "Cấu trúc mỗi sheet: dòng 1 là tiêu đề cột, dữ liệu từ dòng 2.",
        "Cột: " + " | ".join(headers),
        "",
        "Quy ước:",
        "• Một trong hai mã (Mã nội bộ / Mã hải quan) là bắt buộc.",
        "• Các cột số để trống = 0.",
        "• Tồn cuối kỳ là giá trị khai báo. Hệ thống tự tính 'tồn cuối ngụ ý' = "
        "Tồn đầu + Nhập − (Tái xuất + Chuyển MĐSD + Xuất SX + Xuất khác) để đối chiếu.",
        "• Kỳ báo cáo (từ ngày / đến ngày) chọn ở màn hình tải lên, không nhập trong file.",
    ]
    for ri, line in enumerate(lines, start=2):
        inst.cell(row=ri, column=1, value=line).alignment = Alignment(
            wrap_text=True, vertical="top")
    inst.column_dimensions["A"].width = 100

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
