"""Phase 2 step 7 — admin UI for hub.client_uom_overrides.

Tests the store helpers + HTTP routes."""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from hub.app import auth
from hub.app.auth.session import create_session, hash_password, SESSION_COOKIE
from hub.app.database import connect
from hub.app.main import app
from hub.app.stores import client_uom_overrides as factors


CLIENT = "_admin_factors_test"
USER_ID = "test_admin_factors_user"


@pytest.fixture(autouse=True)
def setup():
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "insert into hub.clients (client_id, name) values (%s, %s) "
            "on conflict do nothing", (CLIENT, "factors test"))
        cur.execute(
            "insert into hub.users (user_id, email, display_name, "
            "password_hash, role, status) values (%s, %s, 'tester', %s, "
            "'admin', 'active') on conflict (user_id) do update set "
            "role='admin', status='active'",
            (USER_ID, "admin_factors@test.local", hash_password("test")))
    yield
    with connect() as conn, conn.cursor() as cur:
        cur.execute("delete from hub.client_uom_overrides where client_id=%s",
                     (CLIENT,))
        cur.execute("delete from hub.sessions where user_id=%s", (USER_ID,))
        cur.execute("delete from hub.users where user_id=%s", (USER_ID,))
        cur.execute("delete from hub.clients where client_id=%s", (CLIENT,))


@pytest.fixture
def http():
    sid = create_session(USER_ID)
    c = TestClient(app, follow_redirects=False)
    c.cookies.set(SESSION_COOKIE, sid)
    return c


# ── Store helpers ────────────────────────────────────────────────────


def test_create_factor_basic():
    factors.create_factor(
        client_id=CLIENT, material_code="M1",
        from_uom="EA", to_uom="KG", factor="0.5",
        source="supplier_data", notes="from spec sheet")
    rows = factors.list_factors(CLIENT)
    assert len(rows) == 1
    assert rows[0]["material_code"] == "M1"
    assert rows[0]["factor"] == "0.500000000"
    assert rows[0]["is_cross_family"] is True
    assert rows[0]["notes"] == "from spec sheet"


def test_create_factor_client_wide():
    factors.create_factor(
        client_id=CLIENT, material_code=None,
        from_uom="EA", to_uom="kg", factor="0.5", source="staff_form")
    rows = factors.list_factors(CLIENT)
    assert rows[0]["material_code"] is None


def test_validation_rejects_zero_factor():
    with pytest.raises(factors.FactorError):
        factors.create_factor(
            client_id=CLIENT, material_code="M1",
            from_uom="EA", to_uom="KG", factor="0", source="staff_form")


def test_validation_rejects_negative_factor():
    with pytest.raises(factors.FactorError):
        factors.create_factor(
            client_id=CLIENT, material_code="M1",
            from_uom="EA", to_uom="KG", factor="-1", source="staff_form")


def test_validation_rejects_same_uom():
    with pytest.raises(factors.FactorError):
        factors.create_factor(
            client_id=CLIENT, material_code="M1",
            from_uom="kg", to_uom="kg", factor="1", source="staff_form")


def test_validation_rejects_unknown_source():
    with pytest.raises(factors.FactorError):
        factors.create_factor(
            client_id=CLIENT, material_code="M1",
            from_uom="EA", to_uom="KG", factor="1", source="weird")


def test_update_existing_factor():
    factors.create_factor(
        client_id=CLIENT, material_code="M_U", from_uom="EA",
        to_uom="KG", factor="1.0", source="staff_form")
    factors.update_factor(
        client_id=CLIENT, material_code="M_U", from_uom="EA",
        to_uom="KG", factor="2.5", source="supplier_data",
        notes="updated")
    rows = factors.list_factors(CLIENT)
    assert rows[0]["factor"] == "2.500000000"
    assert rows[0]["source"] == "supplier_data"
    assert rows[0]["notes"] == "updated"


def test_update_nonexistent_raises():
    with pytest.raises(factors.FactorError):
        factors.update_factor(
            client_id=CLIENT, material_code="NOPE",
            from_uom="EA", to_uom="KG", factor="1", source="staff_form")


def test_delete_factor():
    factors.create_factor(
        client_id=CLIENT, material_code="M_D",
        from_uom="EA", to_uom="KG", factor="1.0", source="staff_form")
    factors.delete_factor(
        client_id=CLIENT, material_code="M_D",
        from_uom="EA", to_uom="KG")
    assert factors.list_factors(CLIENT) == []


def test_stats_counts_cross_family():
    factors.create_factor(
        client_id=CLIENT, material_code="M_SAME",
        from_uom="g", to_uom="kg", factor="0.001", source="staff_form")
    factors.create_factor(
        client_id=CLIENT, material_code="M_XFAM",
        from_uom="EA", to_uom="KG", factor="0.5", source="staff_form")
    s = factors.stats(CLIENT)
    assert s["total"] == 2
    assert s["cross_family"] == 1


# ── CSV import ────────────────────────────────────────────────────────


def test_csv_import_minimal():
    csv = "material_code,from_uom,to_uom,factor\nM1,EA,KG,0.5\nM2,EA,SETS,4\n"
    result = factors.import_csv(client_id=CLIENT, csv_text=csv)
    assert result["inserted"] == 2
    assert result["failed"] == 0
    rows = factors.list_factors(CLIENT)
    assert {r["material_code"] for r in rows} == {"M1", "M2"}
    # Imported defaults to source='imported'
    assert all(r["source"] == "imported" for r in rows)


def test_csv_import_with_notes_and_blank_material():
    csv = ("material_code,from_uom,to_uom,factor,notes\n"
           ",EA,KG,0.5,client-wide test\n"
           "M1,EA,SETS,4,per supplier\n")
    result = factors.import_csv(client_id=CLIENT, csv_text=csv)
    assert result["inserted"] == 2
    rows = factors.list_factors(CLIENT)
    by_mat = {r["material_code"]: r for r in rows}
    assert None in by_mat
    assert by_mat[None]["notes"] == "client-wide test"


def test_csv_import_skips_empty_lines():
    csv = "material_code,from_uom,to_uom,factor\n\nM1,EA,KG,0.5\n\n"
    result = factors.import_csv(client_id=CLIENT, csv_text=csv)
    assert result["inserted"] == 1


def test_csv_import_records_failures():
    csv = ("material_code,from_uom,to_uom,factor\n"
           "M1,EA,KG,not_a_number\n"
           "M2,EA,EA,1.0\n"  # same uom
           "M3,EA,KG,0.5\n")
    result = factors.import_csv(client_id=CLIENT, csv_text=csv)
    assert result["inserted"] == 1
    assert result["failed"] == 2
    assert len(result["errors"]) == 2


def test_csv_import_rejects_missing_columns():
    csv = "material_code,from_uom\nM1,EA\n"
    with pytest.raises(factors.FactorError) as exc_info:
        factors.import_csv(client_id=CLIENT, csv_text=csv)
    assert "missing required columns" in str(exc_info.value)


# ── HTTP routes ──────────────────────────────────────────────────────


def test_list_view_renders(http):
    factors.create_factor(
        client_id=CLIENT, material_code="M_VIEW",
        from_uom="EA", to_uom="KG", factor="0.5", source="supplier_data")
    r = http.get(f"/clients/{CLIENT}/uom-factors")
    assert r.status_code == 200
    body = r.text
    assert "Hệ số quy đổi" in body
    assert "M_VIEW" in body
    assert "0.500000000" in body or "0.5" in body
    assert "khác họ" in body  # cross-family badge


def test_create_via_http(http):
    r = http.post(
        f"/clients/{CLIENT}/uom-factors/new",
        data={"material_code": "M_NEW", "from_uom": "EA",
              "to_uom": "KG", "factor": "0.7",
              "source": "staff_form", "notes": ""})
    assert r.status_code == 303
    assert "saved=created" in r.headers["location"]
    assert factors.list_factors(CLIENT)[0]["factor"] == "0.700000000"


def test_create_invalid_redirects_with_error(http):
    r = http.post(
        f"/clients/{CLIENT}/uom-factors/new",
        data={"material_code": "M_BAD", "from_uom": "EA",
              "to_uom": "KG", "factor": "0",  # invalid
              "source": "staff_form", "notes": ""})
    assert r.status_code == 303
    assert "error=" in r.headers["location"]
    assert factors.list_factors(CLIENT) == []


def test_update_via_http(http):
    factors.create_factor(
        client_id=CLIENT, material_code="M_UP",
        from_uom="EA", to_uom="KG", factor="1.0", source="staff_form")
    r = http.post(
        f"/clients/{CLIENT}/uom-factors/update",
        data={"material_code": "M_UP", "from_uom": "EA",
              "to_uom": "KG", "factor": "2.5",
              "source": "supplier_data", "notes": "updated"})
    assert r.status_code == 303
    assert "saved=updated" in r.headers["location"]
    assert factors.list_factors(CLIENT)[0]["factor"] == "2.500000000"


def test_delete_via_http(http):
    factors.create_factor(
        client_id=CLIENT, material_code="M_DEL",
        from_uom="EA", to_uom="KG", factor="1.0", source="staff_form")
    r = http.post(
        f"/clients/{CLIENT}/uom-factors/delete",
        data={"material_code": "M_DEL", "from_uom": "EA", "to_uom": "KG"})
    assert r.status_code == 303
    assert factors.list_factors(CLIENT) == []


def test_csv_import_via_http(http):
    csv_bytes = (b"material_code,from_uom,to_uom,factor,notes\n"
                  b"M_CSV1,EA,KG,0.5,first\n"
                  b"M_CSV2,EA,SETS,4,second\n")
    r = http.post(
        f"/clients/{CLIENT}/uom-factors/import",
        files={"upload": ("test.csv", io.BytesIO(csv_bytes), "text/csv")},
        data={"source": "imported"})
    assert r.status_code == 303
    loc = r.headers["location"]
    assert "import%20ok" in loc or "import ok" in loc
    assert "2%20inserted" in loc or "2 inserted" in loc
    rows = factors.list_factors(CLIENT)
    assert {r["material_code"] for r in rows} == {"M_CSV1", "M_CSV2"}


def test_xlsx_import_via_http(http):
    """Phase 2 step 2 follow-up: XLSX import alongside CSV."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["material_code", "from_uom", "to_uom", "factor", "notes"])
    ws.append(["M_XLSX1", "EA", "KG", 0.5, "from xlsx 1"])
    ws.append(["M_XLSX2", "EA", "SETS", 4, "from xlsx 2"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = http.post(
        f"/clients/{CLIENT}/uom-factors/import",
        files={"upload": ("test.xlsx", buf,
                          "application/vnd.openxmlformats-officedocument."
                          "spreadsheetml.sheet")},
        data={"source": "imported"})
    assert r.status_code == 303
    rows = factors.list_factors(CLIENT)
    assert {r["material_code"] for r in rows} == {"M_XLSX1", "M_XLSX2"}


def test_template_download(http):
    r = http.get(f"/clients/{CLIENT}/uom-factors/template.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml.sheet" in r.headers["content-type"]
    assert "uom-factors-template" in r.headers["content-disposition"]
    # Verify the bytes are a valid xlsx
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert "uom-factors" in wb.sheetnames
    # Header in first sheet row 1.
    ws = wb["uom-factors"]
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "material_code" in headers
    assert "from_uom" in headers
    assert "factor" in headers
    wb.close()


def test_xlsx_import_helper_directly():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["material_code", "from_uom", "to_uom", "factor"])
    ws.append(["M_HELPER", "EA", "KG", 0.7])
    buf = io.BytesIO()
    wb.save(buf)
    result = factors.import_xlsx(
        client_id=CLIENT, xlsx_bytes=buf.getvalue(), source="imported")
    assert result["inserted"] == 1
    assert result["failed"] == 0


def test_xlsx_import_rejects_missing_columns():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["material_code", "from_uom"])  # missing to_uom + factor
    ws.append(["M1", "EA"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(factors.FactorError) as exc_info:
        factors.import_xlsx(client_id=CLIENT, xlsx_bytes=buf.getvalue())
    assert "missing required columns" in str(exc_info.value)
