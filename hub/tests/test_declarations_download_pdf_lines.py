"""Provider tests for page-selective declaration export
`POST /v1/hub/clients/{c}/declarations/download.pdf` (issue #50).

A real Growatt import declaration is 52-54 pages for 50 goods lines
(VNACCS caps 50 lines/TK) — roughly ONE goods line per page. A CO dossier
uses a few lines per declaration, so ~90% of the printed pages are
irrelevant. The rendered ECUS page text carries an angle-bracket line
marker (`<01>` on line 1's page); `<IMP>`/`<EXP>` is the watermark, so the
marker regex is numeric-only.

The rule under test (derived per page, never a hardcoded header count):
a page with NO numeric marker is framing (header/trailer) → ALWAYS kept;
a page WITH a marker is kept only if its line_no was requested. Real data
shows both a 2-page header (108077837340) and a 3-page header + trailing
page (108234677720) — a fixed page count would drop evidence.

Whole pages are dropped, never edited: the stored `.xls` is never
re-numbered, so a kept page still carries the exact `X/N` marker it was
filed with and stays a faithful copy of the filed declaration.

These build synthetic `.xlsx` fixtures whose rendered pages carry the
same marker shape as the real form, and invoke real soffice.
"""
from __future__ import annotations

import secrets
import shutil
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.pagebreak import Break
from pypdf import PdfReader, PdfWriter

from hub.app import jwt_issuer, settings_store
from hub.app.database import connect
from hub.app.main import app
from hub.app.storage import get_backend, sha256_bytes
from hub.app.stores import service_accounts as sa_store

pytestmark = pytest.mark.skipif(
    shutil.which("soffice") is None,
    reason="LibreOffice (soffice) not installed",
)

URL_TMPL = "/v1/hub/clients/{cid}/declarations/download.pdf"


def _url(client_id: str) -> str:
    return URL_TMPL.format(cid=client_id)


def _client() -> TestClient:
    return TestClient(app)


def _decl_xlsx(*, nlines: int = 4, header_pages: int = 2,
               trailer: bool = True, tag: str = "TAGX") -> bytes:
    """A declaration workbook whose render mirrors the real ECUS form:
    `header_pages` framing pages (page 1 carrying the `<IMP>` watermark and
    the per-declaration `tag`, but NO numeric marker), then one goods page
    per line carrying `<NN>`, then an optional trailing framing page.
    Explicit row breaks pin the one-line-per-page layout the real form has.

    `tag` is unique per declaration so merge order can be asserted from the
    extracted text (the header text alone repeats across declarations)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "TKN"
    row = 1
    ws.cell(row=row, column=2, value=f"To khai hang hoa nhap khau <IMP> {tag}")
    row += 1
    for h in range(header_pages - 1):
        ws.row_breaks.append(Break(id=row - 1))
        ws.cell(row=row, column=2, value=f"HEADERCONT{h}")
        row += 1
    for i in range(1, nlines + 1):
        ws.row_breaks.append(Break(id=row - 1))
        ws.cell(row=row, column=2, value=f"<{i:02d}> goods line MARKLINE{i:02d}")
        row += 1
    if trailer:
        ws.row_breaks.append(Break(id=row - 1))
        ws.cell(row=row, column=2, value="TRAILERPAGE")
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = 9  # A4
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf_pages_text(content: bytes) -> list[str]:
    return [(p.extract_text() or "").strip()
            for p in PdfReader(BytesIO(content)).pages]


def _pdf_text(content: bytes) -> str:
    return "\n".join(_pdf_pages_text(content))


def _page_count(content: bytes) -> int:
    return len(PdfReader(BytesIO(content)).pages)


# ─── Fixtures ────────────────────────────────────────────────────────


@pytest.fixture
def files_root(tmp_path, monkeypatch):
    root = tmp_path / "files"
    root.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("DATA_HUB_FILES_ROOT", str(root))
    import hub.app.storage as storage_mod
    storage_mod._BACKEND = None
    yield root
    storage_mod._BACKEND = None


@pytest.fixture
def auth_disabled(monkeypatch):
    monkeypatch.setenv("DATA_HUB_API_AUTH_DISABLED", "1")


def _seed_file(cur, *, client_id, decl_no, direction, filename, blob,
               file_kind="xls"):
    # `file_kind` is constrained to xls|pdf|scan|other (mig 059); the render
    # path feeds soffice by content, not extension, so an .xlsx blob rides
    # the `xls` kind fine — same as the real ECUS forms.
    backend = get_backend()
    key = f"customs_declarations/{client_id}/{secrets.token_hex(4)}_{filename}"
    backend.put(blob, key=key)
    cur.execute(
        """insert into hub.customs_declaration_files
           (client_id, declaration_no, direction, file_kind,
            backend_key, original_filename, sha256, size_bytes)
           values (%s, %s, %s, %s, %s, %s, %s, %s)""",
        (client_id, decl_no, direction, file_kind, key, filename,
         sha256_bytes(blob), len(blob)),
    )


def _teardown(cid: str) -> None:
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "delete from hub.customs_declaration_files where client_id=%s",
            (cid,),
        )
        cur.execute("delete from hub.bcct_rows where client_id=%s", (cid,))
        cur.execute("delete from hub.clients where client_id=%s", (cid,))


@pytest.fixture
def seeded(files_root, auth_disabled):
    """DEC001: 2 header pages + 4 goods lines + trailer = 7 pages.
    DEC002: 3 header pages + 2 goods lines, no trailer = 5 pages."""
    cid = "pdfln-" + secrets.token_hex(4)
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "insert into hub.clients (client_id, name) values (%s, %s) "
            "on conflict do nothing", (cid, "pdf line selection test"),
        )
        _seed_file(cur, client_id=cid, decl_no="DEC001", direction="import",
                   filename="dec001.xls",
                   blob=_decl_xlsx(nlines=4, tag="TAGONE"))
        _seed_file(cur, client_id=cid, decl_no="DEC002", direction="import",
                   filename="dec002.xls",
                   blob=_decl_xlsx(nlines=2, header_pages=3, trailer=False,
                                   tag="TAGTWO"))
    yield cid
    _teardown(cid)


# ─── Seam 1: build_merged_pdf page selection ─────────────────────────


def _files_for(cid: str, decl_no: str, direction: str = "import"):
    from hub.app.stores.customs_declaration_files import list_files_for_declarations
    return list_files_for_declarations(cid, [decl_no], direction=direction)


def _build(cid, decl_nos, lines_by_decl, tmp_path, **kw):
    from hub.app.declarations_pdf import build_merged_pdf
    from hub.app.stores.customs_declaration_files import list_files_for_declarations

    files = list_files_for_declarations(cid, decl_nos, direction="import")
    files_by_decl: dict[str, list] = {d: [] for d in decl_nos}
    for f in files:
        files_by_decl.setdefault(f.declaration_no, []).append(f)
    dest = tmp_path / f"out-{secrets.token_hex(3)}"
    dest.mkdir()
    return build_merged_pdf(
        requested=decl_nos, decl_order=sorted(decl_nos),
        files_by_decl=files_by_decl, backend=get_backend(),
        dest_dir=dest, lines_by_decl=lines_by_decl, **kw,
    )


def test_selection_keeps_framing_pages_and_requested_lines_only(
    seeded, tmp_path,
):
    """DEC001 renders 7 pages: [header, header, <01>, <02>, <03>, <04>,
    trailer]. Requesting lines 1,3 keeps the 3 framing pages (2 header + 1
    trailer — derived, not a fixed count) + exactly lines 1 and 3."""
    res = _build(seeded, ["DEC001"], {"DEC001": {1, 3}}, tmp_path)
    body = res.path.read_bytes()
    pages = _pdf_pages_text(body)
    assert len(pages) == 5
    assert "<IMP>" in pages[0]          # framing page 1 kept (watermark only)
    assert "HEADERCONT0" in pages[1]    # framing page 2 kept
    assert "MARKLINE01" in pages[2]
    assert "MARKLINE03" in pages[3]
    assert "TRAILERPAGE" in pages[4]    # trailing framing page kept
    # Non-requested goods lines dropped.
    assert "MARKLINE02" not in "\n".join(pages)
    assert "MARKLINE04" not in "\n".join(pages)
    assert res.included == 1
    assert res.missing_line_nos == []


def test_selection_derives_framing_per_declaration(seeded, tmp_path):
    """DEC002 has a 3-page header and NO trailer (the real 108234677720 vs
    108077837340 divergence). Requesting line 2 keeps all 3 header pages —
    "keep page 1" or "keep pages 1-2" would drop evidence."""
    res = _build(seeded, ["DEC002"], {"DEC002": {2}}, tmp_path)
    pages = _pdf_pages_text(res.path.read_bytes())
    assert len(pages) == 4
    assert "<IMP>" in pages[0]
    assert "HEADERCONT0" in pages[1]
    assert "HEADERCONT1" in pages[2]
    assert "MARKLINE02" in pages[3]
    assert "MARKLINE01" not in "\n".join(pages)


def test_kept_pages_are_preserved_verbatim(seeded, tmp_path):
    """The evidence guarantee: a kept page is the SAME page, not a rebuilt
    one. Selection drops whole pages and never edits content, so each kept
    page's extracted text is byte-for-byte what the unselected render
    produced — including whatever `X/N` marker the form was filed with.
    (Real forms carry their own numbering: 108234677720's line-1 goods page
    prints "3/52", matching neither the PDF page index nor its 54-page
    render. Preserving that verbatim is exactly the point.)"""
    full = _pdf_pages_text(_build(seeded, ["DEC001"], None, tmp_path)
                           .path.read_bytes())
    sel = _pdf_pages_text(_build(seeded, ["DEC001"], {"DEC001": {2}}, tmp_path)
                          .path.read_bytes())
    # DEC001 renders [hdr, hdr, <01>, <02>, <03>, <04>, trailer]; line 2 keeps
    # the framing pages + page index 3, untouched.
    assert sel == [full[0], full[1], full[3], full[6]]


def test_line_with_no_marker_page_is_reported_not_silently_ignored(
    seeded, tmp_path,
):
    """DEC001 has lines 1-4; requesting 2 and 9 keeps line 2 and reports 9."""
    res = _build(seeded, ["DEC001"], {"DEC001": {2, 9}}, tmp_path)
    assert "MARKLINE02" in _pdf_text(res.path.read_bytes())
    assert res.missing_line_nos == ["DEC001:9"]


def test_selection_is_per_declaration(seeded, tmp_path):
    """Each declaration's own line set applies to its own pages only."""
    res = _build(
        seeded, ["DEC001", "DEC002"],
        {"DEC001": {4}, "DEC002": {1}}, tmp_path,
    )
    text = _pdf_text(res.path.read_bytes())
    assert "MARKLINE04" in text and "MARKLINE01" in text
    assert "MARKLINE03" not in text
    assert "MARKLINE02" not in text
    assert res.included == 2


def test_declaration_absent_from_selection_keeps_all_pages(seeded, tmp_path):
    """A declaration with no entry in the line map is untouched — the
    never-drop-evidence fallback."""
    res = _build(seeded, ["DEC001", "DEC002"], {"DEC001": {1}}, tmp_path)
    text = _pdf_text(res.path.read_bytes())
    # DEC002 keeps every goods line.
    assert "MARKLINE01" in text and "MARKLINE02" in text
    # DEC001 kept only line 1 → its lines 2-4 are gone; DEC002's line 2 is
    # present, so count pages instead: DEC001 3 framing + 1 = 4, DEC002 all 5.
    assert _page_count(res.path.read_bytes()) == 9


# ─── The byte-identity gate ──────────────────────────────────────────
#
# docs/API_CONTRACT.md promises that omitting `quality` + `max_part_bytes`
# leaves the body byte-identical to the original contract. Page selection
# threads through build_merged_pdf's page-add loop — the exact code that
# promise depends on — so these pin it. No such test existed before #50.


def _original_merge(pdfs: list[bytes]) -> bytes:
    """The original contract's merge algorithm, reimplemented here: one
    writer, every page of every source in order. Pinning against an
    independent implementation (rather than a stored golden) keeps the gate
    honest if the merge is ever refactored."""
    writer = PdfWriter()
    for b in pdfs:
        for page in PdfReader(BytesIO(b)).pages:
            writer.add_page(page)
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


@pytest.mark.parametrize("selection", [None, {}], ids=["none", "empty"])
def test_no_selection_body_is_byte_identical_to_original_merge(
    seeded, tmp_path, selection,
):
    """No selection (None or empty map) → byte-for-byte the original merge."""
    from hub.app.declarations_pdf import ensure_pdfs

    decls = ["DEC001", "DEC002"]
    res = _build(seeded, decls, selection, tmp_path)
    body = res.path.read_bytes()

    # Independently render + merge the same sources in the same order.
    srcs: list[bytes] = []
    for d in sorted(decls):
        files = sorted(_files_for(seeded, d),
                       key=lambda x: ((x.original_filename or ""), x.id))
        rendered = ensure_pdfs(files, get_backend())
        srcs.extend(rendered[f.id] for f in files if rendered.get(f.id))
    assert body == _original_merge(srcs)
    assert res.missing_line_nos == []


def test_empty_line_list_for_a_declaration_keeps_all_pages(seeded, tmp_path):
    """`lines: []` on an entry means ALL pages (the never-drop-evidence
    fallback), so it must produce the same bytes as no selection at all."""
    a = _build(seeded, ["DEC001"], None, tmp_path).path.read_bytes()
    b = _build(seeded, ["DEC001"], {"DEC001": set()}, tmp_path).path.read_bytes()
    assert a == b
    assert _page_count(a) == 7


# ─── Seam 2: the POST route ──────────────────────────────────────────


def _post(cid: str, body: dict, headers: dict | None = None):
    return _client().post(_url(cid), json=body, headers=headers or {})


def test_post_selects_lines_end_to_end(seeded):
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001", "lines": [1, 3]}],
    })
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    assert "attachment" in r.headers["content-disposition"]
    assert r.headers["X-Declarations-Requested"] == "1"
    assert r.headers["X-Declarations-Included"] == "1"
    assert _page_count(r.content) == 5  # 3 framing + 2 goods
    text = _pdf_text(r.content)
    assert "MARKLINE01" in text and "MARKLINE03" in text
    assert "MARKLINE02" not in text and "MARKLINE04" not in text
    assert "X-Lines-Missing-Nos" not in r.headers


def test_post_without_lines_returns_all_pages(seeded):
    """`lines` omitted → every page. POST is a strict superset of GET."""
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001"}],
    })
    assert r.status_code == 200
    assert _page_count(r.content) == 7


def test_post_no_lines_is_byte_identical_to_get(seeded):
    """The GET is unchanged and the POST's no-selection path is the same
    code path — so the two bodies must match byte-for-byte."""
    g = _client().get(_url(seeded), params={
        "direction": "import", "declaration_nos": "DEC001,DEC002",
    })
    p = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001"},
                         {"declaration_no": "DEC002", "lines": []}],
    })
    assert g.status_code == 200 and p.status_code == 200
    assert p.content == g.content
    assert p.headers["X-Pdf-Quality"] == g.headers["X-Pdf-Quality"] == "print"


def test_post_reports_missing_lines_header(seeded):
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001", "lines": [2, 9]},
                         {"declaration_no": "DEC002", "lines": [7]}],
    })
    assert r.status_code == 200
    assert r.headers["X-Lines-Missing-Nos"] == "DEC001:9,DEC002:7"
    assert "MARKLINE02" in _pdf_text(r.content)


def test_post_missing_declaration_still_reported(seeded):
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001", "lines": [1]},
                         {"declaration_no": "NOPE", "lines": [1]}],
    })
    assert r.status_code == 200
    assert r.headers["X-Declarations-Requested"] == "2"
    assert r.headers["X-Declarations-Included"] == "1"
    assert r.headers["X-Declarations-Missing-Nos"] == "NOPE"
    # A wholly-missing declaration is reported once, at declaration level —
    # not duplicated per requested line.
    assert "X-Lines-Missing-Nos" not in r.headers


def test_post_sort_registration_date(seeded):
    with connect() as conn, conn.cursor() as cur:
        for decl, d in (("DEC001", "2026-04-19"), ("DEC002", "2026-04-01")):
            cur.execute(
                """insert into hub.bcct_rows
                   (client_id, transaction_key, line_no, declaration_no,
                    declaration_type, direction, registration_date,
                    customs_code, goods_name, payload)
                   values (%s,%s,'1',%s,'E11','import',%s,'MAT-X','d','{}'::jsonb)""",
                (seeded, f"TX-{decl}", decl, d),
            )
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001", "lines": [1]},
                         {"declaration_no": "DEC002", "lines": [1]}],
        "sort": "registration_date",
    })
    assert r.status_code == 200
    text = _pdf_text(r.content)
    # DEC002 (04-01) body precedes DEC001 (04-19) — opposite of decl_no order.
    assert text.index("TAGTWO") < text.index("TAGONE")

    # And the default sort puts them the other way round, proving the
    # assertion above tracks the sort rather than the seed order.
    d = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001", "lines": [1]},
                         {"declaration_no": "DEC002", "lines": [1]}],
    })
    assert d.status_code == 200
    dtext = _pdf_text(d.content)
    assert dtext.index("TAGONE") < dtext.index("TAGTWO")


def test_post_custom_filename_sanitized(seeded):
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001"}],
        "filename": "6. TKN GHEP",
    })
    assert r.status_code == 200
    assert 'filename="6. TKN GHEP.pdf"' in r.headers["content-disposition"]


def test_post_selection_composes_with_compact(seeded):
    """Selection runs BEFORE compact dedup: fewer pages AND smaller bytes,
    with the selected content intact."""
    base = {"direction": "import",
            "declarations": [{"declaration_no": "DEC001", "lines": [1, 3]}]}
    rp = _post(seeded, {**base, "quality": "print"})
    rc = _post(seeded, {**base, "quality": "compact"})
    assert rp.status_code == 200 and rc.status_code == 200
    assert rc.headers["X-Pdf-Quality"] == "pypdf-dedup-1"
    assert int(rc.headers["X-Pdf-Bytes"]) <= int(rp.headers["X-Pdf-Bytes"])
    assert _page_count(rc.content) == _page_count(rp.content) == 5
    assert "MARKLINE01" in _pdf_text(rc.content)


def test_post_selection_composes_with_max_part_bytes(seeded):
    """Selection runs BEFORE splitting: the selected pages fit one part."""
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001", "lines": [1]}],
        "max_part_bytes": 2_000_000,
    })
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["X-Pdf-Parts"] == "1"
    assert _page_count(r.content) == 4  # 3 framing + 1 goods


def test_post_zero_match_returns_info_page(seeded):
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DOES-NOT-EXIST", "lines": [1]}],
    })
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["X-Declarations-Included"] == "0"
    assert _page_count(r.content) == 1
    assert "Không có tờ khai" in _pdf_text(r.content)


# ─── POST validation ─────────────────────────────────────────────────


def test_post_get_still_works_unchanged(seeded):
    """The GET keeps its exact contract; `lines` is POST-only."""
    r = _client().get(_url(seeded), params={
        "direction": "import", "declaration_nos": "DEC001", "lines": "1",
    })
    assert r.status_code == 200
    assert _page_count(r.content) == 7  # `lines` ignored on the GET


@pytest.mark.parametrize("params,detail", [
    # No declaration_nos AND a bad sort: the declaration_nos check runs first.
    ({"direction": "import", "sort": "banana"}, "declaration_nos_required"),
    # Too many nos AND a bad quality: the too-many check runs first.
    ({"direction": "import", "declaration_nos": ",".join(f"D{i}" for i in range(501)),
      "quality": "ultra"}, "too_many_declaration_nos"),
    # Bad direction AND a bad sort: direction runs first.
    ({"direction": "sideways", "declaration_nos": "DEC001", "sort": "banana"},
     "invalid_direction"),
])
def test_get_error_precedence_unchanged(seeded, params, detail):
    """The GET's validation ORDER is contract: direction → declaration_nos →
    render options, first bad check wins. The POST shares the render-option
    validator, and hoisting it above the declaration_nos checks silently
    changes which `detail` CO sees when two params are bad at once. The
    pre-existing tests miss this because each supplies exactly one bad
    param."""
    r = _client().get(_url(seeded), params=params)
    assert r.status_code == 400
    assert r.json()["detail"] == detail


@pytest.mark.parametrize("body,detail", [
    ({"declarations": [{"declaration_no": "DEC001"}]}, "invalid_direction"),
    # The POST applies the same order as the GET.
    ({"direction": "import", "sort": "banana"}, "declarations_required"),
    ({"direction": "sideways", "declarations": [], "sort": "banana"},
     "invalid_direction"),
    ({"direction": "sideways",
      "declarations": [{"declaration_no": "DEC001"}]}, "invalid_direction"),
    ({"direction": "import"}, "declarations_required"),
    ({"direction": "import", "declarations": []}, "declarations_required"),
    ({"direction": "import", "declarations": "DEC001"}, "declarations_required"),
    ({"direction": "import", "declarations": [{"lines": [1]}]},
     "declarations_required"),
    ({"direction": "import", "declarations": [{"declaration_no": "DEC001"}],
      "sort": "banana"}, "invalid_sort"),
    ({"direction": "import", "declarations": [{"declaration_no": "DEC001"}],
      "quality": "ultra"}, "invalid_quality"),
    ({"direction": "import", "declarations": [{"declaration_no": "DEC001"}],
      "max_part_bytes": 0}, "invalid_max_part_bytes"),
    ({"direction": "import", "declarations": [{"declaration_no": "DEC001"}],
      "max_part_bytes": "abc"}, "invalid_max_part_bytes"),
    ({"direction": "import",
      "declarations": [{"declaration_no": "DEC001", "lines": "1,2"}]},
     "invalid_lines"),
    ({"direction": "import",
      "declarations": [{"declaration_no": "DEC001", "lines": [0]}]},
     "invalid_lines"),
    ({"direction": "import",
      "declarations": [{"declaration_no": "DEC001", "lines": [1000]}]},
     "invalid_lines"),
    ({"direction": "import",
      "declarations": [{"declaration_no": "DEC001", "lines": ["x"]}]},
     "invalid_lines"),
    ({"direction": "import",
      "declarations": [{"declaration_no": "DEC001", "lines": [1.5]}]},
     "invalid_lines"),
])
def test_post_validation_errors(seeded, body, detail):
    r = _post(seeded, body)
    assert r.status_code == 400, r.text
    assert r.json()["detail"] == detail


def test_post_too_many_declarations(seeded):
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": f"D{i:04d}"} for i in range(501)],
    })
    assert r.status_code == 400
    assert r.json()["detail"] == "too_many_declaration_nos"


def test_post_invalid_body_returns_400(seeded):
    r = _client().post(
        _url(seeded), content=b"not json",
        headers={"content-type": "application/json"},
    )
    assert r.status_code == 400
    assert r.json()["detail"] == "invalid_body"


def test_post_unknown_client_returns_404(auth_disabled, files_root):
    r = _post("does-not-exist", {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001"}],
    })
    assert r.status_code == 404


def test_post_duplicate_declaration_merges_lines(seeded):
    """Repeating a declaration_no unions its line sets rather than losing
    one — declaration order/dedupe matches the GET's contract."""
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001", "lines": [1]},
                         {"declaration_no": "DEC001", "lines": [3]}],
    })
    assert r.status_code == 200
    assert r.headers["X-Declarations-Requested"] == "1"
    text = _pdf_text(r.content)
    assert "MARKLINE01" in text and "MARKLINE03" in text
    assert "MARKLINE02" not in text


def test_post_duplicate_declaration_bare_entry_wins_all_pages(seeded):
    """An entry with no `lines` means ALL pages; a duplicate that also names
    lines cannot narrow it — never drop evidence the caller asked for."""
    r = _post(seeded, {
        "direction": "import",
        "declarations": [{"declaration_no": "DEC001"},
                         {"declaration_no": "DEC001", "lines": [3]}],
    })
    assert r.status_code == 200
    assert _page_count(r.content) == 7


# ─── Strict-mode auth (parity with the GET mirror) ───────────────────


@pytest.fixture
def strict_mode_on(monkeypatch, tmp_path):
    monkeypatch.setenv("DATA_HUB_KEYS_DIR", str(tmp_path / "keys"))
    settings_store.set_many({"api_auth_strict": "true"})
    yield
    settings_store.set_many({"api_auth_strict": "false"})


@pytest.fixture
def auth_seeded(files_root, strict_mode_on):
    cid = "pdfln-auth-" + secrets.token_hex(4)
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "insert into hub.clients (client_id, name) values (%s, %s) "
            "on conflict do nothing", (cid, "pdf line auth test"),
        )
        _seed_file(cur, client_id=cid, decl_no="DEC001", direction="import",
                   filename="dec001.xls", blob=_decl_xlsx(nlines=4))
    yield cid
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "delete from hub.service_accounts where name like 'sa_test_%'"
        )
        cur.execute(
            "delete from hub.revoked_service_tokens "
            "where revoked_by='sa_test_runner'"
        )
    _teardown(cid)


_BODY = {"direction": "import",
         "declarations": [{"declaration_no": "DEC001", "lines": [1]}]}


def _token(scopes: list[str], client_ids=None) -> str:
    sa_store.create_account(
        name="sa_test_co", description="", scopes=scopes,
        client_ids=client_ids, created_by="sa_test_runner",
    )
    return jwt_issuer.make_service_token(
        name="sa_test_co", scopes=scopes, client_ids=client_ids,
    )["access_token"]


def test_post_service_token_with_hub_read_returns_pdf(auth_seeded):
    r = _post(auth_seeded, _BODY,
              {"authorization": f"Bearer {_token(['hub:read'])}"})
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"


def test_post_service_token_without_hub_read_rejected(auth_seeded):
    r = _post(auth_seeded, _BODY,
              {"authorization": f"Bearer {_token(['bom:propose'])}"})
    assert r.status_code == 403


def test_post_client_whitelist_blocks_non_listed(auth_seeded):
    tok = _token(["hub:read"], client_ids=["a-different-client"])
    r = _post(auth_seeded, _BODY, {"authorization": f"Bearer {tok}"})
    assert r.status_code == 403


def test_post_no_bearer_in_strict_mode_returns_401(auth_seeded):
    r = _post(auth_seeded, _BODY)
    assert r.status_code == 401


def test_post_response_is_no_store(seeded):
    """The /v1/hub no-store guard covers the POST too (2026-06-06 leak)."""
    r = _post(seeded, _BODY)
    assert r.status_code == 200
    assert r.headers["Cache-Control"] == "no-store, private"
