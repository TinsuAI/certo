"""Declarability filter for the bảng kê / C/O export.

Data Hub (mig 078+) classifies each material with `customs_relevance`; CO consumes
it instead of guessing. See
`.ai/sister-app-notes/2026-06-09-co-consumer-spec-declarability.md`.

  excluded_non_material  -> rác (drawing/label/doc): excluded from export.
  declarable             -> physical material with import match: emitted.
  declarable_unmatched   -> real material, no import match: excluded from export
                            BUT surfaced as a review queue (NOT plain noise).
  review / null          -> needs attention; KEPT, never silently dropped.

No client-specific heuristic: an unclassified material (null/review/absent) is
kept, so an unmapped client (e.g. growatt) has nothing dropped.
"""
from __future__ import annotations

import openpyxl

from app.origin_material_filters import is_bom_technical_noise, is_declarable_unmatched
from app.web.co_case_context import origin_warning_summary
from app import workbook_io


def _mat(code, *, customs_relevance="", name="", hs="", alloc=False, bom_source="data-hub"):
    return {
        "material_code": code,
        "customs_relevance": customs_relevance,
        "material_description": name,
        "hs_code": hs,
        "bom_source": bom_source,
        "allocation_lines": [{"import_declaration_no": "X"}] if alloc else [],
    }


def test_customs_relevance_drives_export_exclusion():
    drawing = _mat("DRW", customs_relevance="excluded_non_material", name="Blueprint;Semi-Assy")
    steel = _mat("STEEL", customs_relevance="declarable_unmatched", name="Tube;Round;45#")
    real = _mat("AL1", customs_relevance="declarable", name="Nhôm", hs="76061190", alloc=True)
    review = _mat("REV", customs_relevance="review", name="Unknown MG")

    # rác AND real-but-unmatched are both kept OUT of the export
    assert is_bom_technical_noise(drawing) is True
    assert is_bom_technical_noise(steel) is True
    # a declarable material is emitted; a review/unclassified row is NOT auto-dropped
    assert is_bom_technical_noise(real) is False
    assert is_bom_technical_noise(review) is False


def test_declarable_unmatched_is_a_review_flag_not_plain_noise():
    steel = _mat("STEEL", customs_relevance="declarable_unmatched", name="Tube;Round;45#")
    drawing = _mat("DRW", customs_relevance="excluded_non_material", name="Checklist;;A")
    # the steel cut-piece is export-excluded but flagged for reconciliation...
    assert is_bom_technical_noise(steel) is True
    assert is_declarable_unmatched(steel) is True
    # ...whereas true rác is NOT in the review bucket
    assert is_declarable_unmatched(drawing) is False
    assert is_declarable_unmatched(_mat("AL1", customs_relevance="declarable", hs="76061190", alloc=True)) is False


def test_unclassified_material_is_never_dropped():
    # NO client-specific heuristic. Anything DH doesn't classify is kept.
    # field absent (legacy/pre-078 DH)
    assert is_bom_technical_noise({"material_code": "X", "bom_source": "technical_flattened",
                                   "hs_code": "", "allocation_lines": []}) is False
    # DH classified it null (mapped client, code only in BCCT)
    assert is_bom_technical_noise({"material_code": "X", "customs_relevance": None}) is False
    # review (has material group but no map row) -> kept (needs attention, not dropped)
    assert is_bom_technical_noise({"material_code": "X", "customs_relevance": "review"}) is False


def test_no_overfit_unmapped_client_excludes_nothing():
    # Growatt shape: technical_flattened BOM, no HS, no allocation — but DH returns
    # customs_relevance=null because the client's material groups aren't mapped yet.
    # The removed heuristic dropped 94% of these; the DH-driven filter drops ZERO.
    growatt_rows = [
        {"material_code": f"G{i}", "bom_source": "technical_flattened",
         "hs_code": "", "allocation_lines": [], "customs_relevance": None}
        for i in range(50)
    ]
    assert sum(is_bom_technical_noise(m) for m in growatt_rows) == 0


def test_summary_splits_non_material_from_unmatched_review():
    materials = [
        {**_mat("AL1", customs_relevance="declarable", hs="76061190", alloc=True)},
        {**_mat("DRW", customs_relevance="excluded_non_material", name="Blueprint"),
         "bom_technical_noise": True, "declarable_unmatched": False},
        {**_mat("STEEL", customs_relevance="declarable_unmatched", name="Tube;Round;45#"),
         "bom_technical_noise": True, "declarable_unmatched": True},
        {**_mat("STEEL2", customs_relevance="declarable_unmatched", name="welding rod"),
         "bom_technical_noise": True, "declarable_unmatched": True},
    ]
    summary = {row["kind"]: row for row in origin_warning_summary({"code": "P1", "fob": "1000000"}, materials, [])}
    assert summary["excluded_non_material"]["count"] == 1
    assert summary["declarable_unmatched"]["count"] == 2
    # steel rows are NOT counted as rác
    assert "STEEL" not in summary["excluded_non_material"]["examples"]


def test_hq_export_skips_excluded_and_unmatched_keeps_declarable():
    wb = openpyxl.Workbook()
    ws = wb.active
    product = {
        "code": "P1", "fob": "0", "source_declaration_no": "308352475340",
        "materials": [
            _mat("AL1", customs_relevance="declarable", name="Nhôm", hs="76061190", alloc=True),
            _mat("DRW", customs_relevance="excluded_non_material", name="Blueprint"),
            _mat("STEEL", customs_relevance="declarable_unmatched", name="Tube;Round;45#"),
            _mat("CU1", customs_relevance="declarable", name="Đồng", hs="74091100", alloc=True),
        ],
        "origin_sheet_material_overrides": {},
    }
    next_row = workbook_io.write_hq_sheet_materials(ws, product, 16, sheet_code="CTH")
    codes = [ws.cell(r, 16).value for r in range(16, next_row) if ws.cell(r, 1).value]
    assert codes == ["AL1", "CU1"], "only declarable+matched rows are emitted"
    assert "DRW" not in codes and "STEEL" not in codes
