"""Tests for the Excel importer that ingests agency cost-allocation files.

The canonical sample is .ai/samples/BANG-PHAN-BO-TY-LE-CHI-PHI.xlsx (GROWATT).
Column layout (matched by index, not header text — headers are bilingual
and merged across rows 2-3):

    A=STT, B=Mã SP,
    C=wages, D=welfare,
    E=rent, F=depreciation, G=other_mfg,
    H=profit (IGNORED on import — residual),
    I=transport_storage,
    J=note

Data rows start at row 4; stop at first row with no Mã SP.
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook


SAMPLE_PATH = Path(".ai/samples/BANG-PHAN-BO-TY-LE-CHI-PHI.xlsx")


def _make_workbook(rows: list[list]) -> bytes:
    """Build a workbook that mimics the GROWATT layout: title row 1, header
    rows 2-3, data from row 4. `rows` are the data cells starting at row 4.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet3"
    ws["A1"] = "BẢNG PHÂN BỔ TỶ LỆ CHI PHÍ"
    # Header rows are essentially noise to the parser (it matches by index),
    # but we keep representative labels so the test workbook looks like the real one.
    ws["A2"] = "STT"; ws["B2"] = "Mã SP"
    ws["A3"] = None
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            ws.cell(row=4 + offset, column=col, value=value)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_growatt_sample_file():
    """The real agency sample yields 24 per-product rows with the documented mapping."""
    if not SAMPLE_PATH.exists():
        pytest.skip(f"{SAMPLE_PATH} missing")
    from app.cost_allocation_importer import parse_excel
    rows = parse_excel(SAMPLE_PATH.read_bytes())
    assert len(rows) == 24
    first = next(r for r in rows if r.product_code == "PV00.0048400")
    # Match the file values (see .ai/samples raw dump).
    assert first.coef_wages == Decimal("0.008987374915832628")
    assert first.coef_welfare == Decimal("0.0008226583189315068")
    assert first.coef_rent == Decimal("0.025810078727409962")
    assert first.coef_depreciation == Decimal("0.017336729849674937")
    assert first.coef_other_mfg == Decimal("0.009490219070589674")
    assert first.coef_transport_storage == Decimal("0.0034125978093402214")
    # Profit column (H) is ignored — even though the sample has a formula string there.
    # CostAllocationRow does not expose profit.
    assert not hasattr(first, "coef_profit")


def test_skips_blank_trailing_rows():
    from app.cost_allocation_importer import parse_excel
    raw = _make_workbook([
        [1, "A.1", 0.001, 0.002, 0.003, 0.004, 0.005, "=...", 0.006, "n1"],
        [2, "A.2", 0.01,  0.02,  0.03,  0.04,  0.05,  None,    0.06,  None],
        [None] * 10,  # blank → stop
        [3, "A.3", 0.1, 0.2, 0.3, 0.4, 0.5, None, 0.6, None],  # should NOT be parsed
    ])
    rows = parse_excel(raw)
    assert [r.product_code for r in rows] == ["A.1", "A.2"]


def test_ignores_profit_column_h():
    from app.cost_allocation_importer import parse_excel
    raw = _make_workbook([
        [1, "X.1", 0.001, 0, 0, 0, 0, 0.999, 0, "ignore profit"],
    ])
    rows = parse_excel(raw)
    assert len(rows) == 1
    assert rows[0].coef_wages == Decimal("0.001")
    # Profit not surfaced on the row at all.
    assert not hasattr(rows[0], "coef_profit")


def test_blank_coefficients_become_zero():
    from app.cost_allocation_importer import parse_excel
    raw = _make_workbook([
        [1, "Y.1", 0.001, None, "", 0.004, None, None, None, None],
    ])
    rows = parse_excel(raw)
    r = rows[0]
    assert r.coef_wages == Decimal("0.001")
    assert r.coef_welfare == Decimal(0)
    assert r.coef_rent == Decimal(0)
    assert r.coef_depreciation == Decimal("0.004")
    assert r.coef_other_mfg == Decimal(0)
    assert r.coef_transport_storage == Decimal(0)


def test_note_column_captured():
    from app.cost_allocation_importer import parse_excel
    raw = _make_workbook([
        [1, "Z.1", 0.001, 0, 0, 0, 0, None, 0, "hệ số mới 2026"],
    ])
    rows = parse_excel(raw)
    assert rows[0].note == "hệ số mới 2026"


def test_string_coefficients_parsed_with_comma_separator():
    """Some agency files use comma decimal separators on locale-VN exports."""
    from app.cost_allocation_importer import parse_excel
    raw = _make_workbook([
        [1, "C.1", "0,001", "0,002", 0, 0, 0, None, 0, None],
    ])
    rows = parse_excel(raw)
    assert rows[0].coef_wages == Decimal("0.001")
    assert rows[0].coef_welfare == Decimal("0.002")


def test_malformed_coef_treated_as_zero():
    from app.cost_allocation_importer import parse_excel
    raw = _make_workbook([
        [1, "D.1", "not a number", 0, 0, 0, 0, None, 0, None],
    ])
    rows = parse_excel(raw)
    assert rows[0].coef_wages == Decimal(0)


def test_apply_ratio_to_fob():
    """Helper that resolves a ratio and multiplies by FOB to produce detail values."""
    from app.cost_allocation_importer import apply_to_fob
    from app.cost_allocation_store import CostAllocationRow
    row = CostAllocationRow(
        product_code="X",
        coef_wages=Decimal("0.01"),
        coef_welfare=Decimal("0.02"),
        coef_rent=Decimal("0.03"),
        coef_depreciation=Decimal("0.04"),
        coef_other_mfg=Decimal("0.05"),
        coef_transport_storage=Decimal("0.06"),
    )
    out = apply_to_fob(row, Decimal("1000"))
    assert out == {
        "wages": Decimal("10.00"),
        "welfare": Decimal("20.00"),
        "rent": Decimal("30.00"),
        "depreciation": Decimal("40.00"),
        "other_mfg": Decimal("50.00"),
        "transport_storage": Decimal("60.00"),
    }


# ---------- Bulk apply across a case (Mục 4a) ----------


def _row(code, w="0.01", note=""):
    from app.cost_allocation_store import CostAllocationRow
    return CostAllocationRow(
        product_code=code,
        coef_wages=Decimal(w),
        coef_welfare=Decimal("0.02"),
        coef_rent=Decimal("0.03"),
        coef_depreciation=Decimal("0.04"),
        coef_other_mfg=Decimal("0.05"),
        coef_transport_storage=Decimal("0.06"),
        note=note,
    )


def _product(code, *, criteria="RVC 35%", fob="1000", cost_buildup=None, documented=""):
    return {
        "code": code,
        "fob": fob,
        "origin_sheet_effective_criteria_text": criteria,
        "documented_result": documented,
        "cost_buildup": cost_buildup or {},
    }


def _resolver(table):
    """Fake Mode-A→Mode-B resolver. `table` maps product_code → CostAllocationRow.
    Key '' is the Mode B default. Mirrors cost_allocation_store.get_ratio."""
    def resolve(code):
        if code in table:
            return table[code]
        return table.get("")  # Mode B default, or None
    return resolve


def test_bulk_applies_mode_a_then_mode_b_fallback():
    from app.cost_allocation_importer import bulk_apply_to_products
    products = [_product("P-A"), _product("P-B")]
    resolver = _resolver({"P-A": _row("P-A"), "": _row("", w="0.005")})  # P-B → Mode B
    result = bulk_apply_to_products(products, resolver)
    assert result["applied"] == [
        {"code": "P-A", "mode": "A"},
        {"code": "P-B", "mode": "B"},
    ]
    assert result["updates"]["P-A"]["wages"] == "10.00"   # 0.01 × 1000
    assert result["updates"]["P-B"]["wages"] == "5.00"    # 0.005 × 1000 (Mode B)
    assert result["skipped_no_ratio"] == []


def test_bulk_skips_non_rvc_lvc_products():
    from app.cost_allocation_importer import bulk_apply_to_products
    products = [_product("CTH-1", criteria="CTH"), _product("RVC-1", criteria="RVC 40%")]
    resolver = _resolver({"CTH-1": _row("CTH-1"), "RVC-1": _row("RVC-1")})
    result = bulk_apply_to_products(products, resolver)
    assert [a["code"] for a in result["applied"]] == ["RVC-1"]
    # CTH product is out of scope entirely — not applied, not in any skipped bucket.
    assert "CTH-1" not in result["updates"]
    assert "CTH-1" not in result["skipped_no_ratio"]


def test_bulk_skips_product_without_ratio():
    from app.cost_allocation_importer import bulk_apply_to_products
    products = [_product("P-X")]
    resolver = _resolver({})  # no Mode A, no Mode B
    result = bulk_apply_to_products(products, resolver)
    assert result["applied"] == []
    assert result["skipped_no_ratio"] == ["P-X"]
    assert result["updates"] == {}


def test_bulk_skips_product_missing_fob():
    from app.cost_allocation_importer import bulk_apply_to_products
    products = [_product("P-NOFOB", fob=""), _product("P-ZERO", fob="0")]
    resolver = _resolver({"": _row("")})
    result = bulk_apply_to_products(products, resolver)
    assert result["applied"] == []
    assert set(result["skipped_no_fob"]) == {"P-NOFOB", "P-ZERO"}


def test_bulk_skips_already_filled_unless_overwrite():
    from app.cost_allocation_importer import bulk_apply_to_products
    filled = {"wages": "999"}
    products = [_product("P-FILLED", cost_buildup=filled)]
    resolver = _resolver({"P-FILLED": _row("P-FILLED")})
    # Default: skip filled, do not touch.
    result = bulk_apply_to_products(products, resolver)
    assert result["applied"] == []
    assert result["skipped_filled"] == ["P-FILLED"]
    # overwrite=True: recompute and overwrite.
    result2 = bulk_apply_to_products(products, resolver, overwrite=True)
    assert result2["applied"] == [{"code": "P-FILLED", "mode": "A"}]
    assert result2["updates"]["P-FILLED"]["wages"] == "10.00"


def test_bulk_criterion_falls_back_to_documented_result():
    """Mirror the template: criterion = effective_criteria_text or documented_result."""
    from app.cost_allocation_importer import bulk_apply_to_products
    products = [_product("P-DOC", criteria="", documented="40.19% RVC")]
    resolver = _resolver({"P-DOC": _row("P-DOC")})
    result = bulk_apply_to_products(products, resolver)
    assert [a["code"] for a in result["applied"]] == ["P-DOC"]


def test_bulk_single_product_matches_per_product_apply_to_fob():
    """Regression: bulk on one SP must equal the per-product 'Áp hệ số' path."""
    from app.cost_allocation_importer import bulk_apply_to_products, apply_to_fob
    row = _row("P-1")
    products = [_product("P-1", fob="100000")]
    result = bulk_apply_to_products(products, _resolver({"P-1": row}))
    expected = {k: str(v) for k, v in apply_to_fob(row, Decimal("100000")).items()}
    assert result["updates"]["P-1"] == expected
