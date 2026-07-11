"""Column (9) materialization at Tính (VN-origin ticket #9).

The mode (country | qualification_label) is resolved once per Tính (case
override > client default > code default), threaded down the funnel, and
materialized per allocation line and per material as `bang_ke_origin_text`
(+ the (12)/(13) slots). Renderers and the web cell are pure readers, so
export == web holds by construction and money never depends on the mode.
"""
from __future__ import annotations

from decimal import Decimal

from app.web.co_case_context import (
    bang_ke_settings,
    enrich_origin_product,
    materialize_bang_ke_origin_fields,
    origin_material_from_bom_row,
    resolve_case_column9_mode,
)


def _stock(**extra):
    base = {
        "source_row": "ROW-1",
        "material_code": "NVL-1",
        "allocation_code": "NVL-1",
        "customs_item_code": "NVL-1",
        "import_declaration_no": "D100",
        "registration_date": "2026-04-21",
        "line_no": "1",
        "value_currency": "VND",
        "currency": "VND",
        "unit_value": "1000",
        "remaining_qty": "100",
        "available_qty": "100",
        "exchange_rate_to_vnd": "1",
        "exchange_rate_source": "vnd_native",
    }
    base.update(extra)
    return base


def _case_with_lots(*origins, consumed="4"):
    pool = {
        "NVL-1": [
            _stock(
                source_row=f"ROW-{index}",
                line_no=str(index),
                origin_country=origin,
                remaining_qty="2",
                available_qty="2",
            )
            for index, origin in enumerate(origins, start=1)
        ]
    }
    material = origin_material_from_bom_row(
        {"material_code": "NVL-1", "qty_per": "1"}, Decimal(consumed), {}, pool,
        product_code="TP1", product_name="SP",
    )
    return {"products": [{"code": "TP1", "materials": [material]}]}


# --- resolution precedence ---

def test_mode_precedence_case_override_beats_client_default():
    client = {"bang_ke_overrides": {"column9_mode": "qualification_label"}}
    assert resolve_case_column9_mode({}, client) == "qualification_label"
    assert resolve_case_column9_mode({"bang_ke_column9_mode_override": "country"}, client) == "country"
    assert resolve_case_column9_mode({}, {}) == "country"  # code default


def test_bang_ke_settings_defaults_and_overlay():
    assert bang_ke_settings({}) == {"column9_mode": "country", "unknown_origin_label": "Không xác định"}
    client = {"bang_ke_overrides": {"column9_mode": "qualification_label", "unknown_origin_label": "Không xuất xứ"}}
    assert bang_ke_settings(client) == {
        "column9_mode": "qualification_label",
        "unknown_origin_label": "Không xuất xứ",
    }


# --- materialization at Tính ---

def test_country_mode_materializes_per_line_and_material_text():
    case = materialize_bang_ke_origin_fields(_case_with_lots("VIETNAM", "CHINA"), {})
    product = case["products"][0]
    material = product["materials"][0]
    assert product["bang_ke_column9_mode"] == "country"
    assert [line["bang_ke_origin_text"] for line in material["allocation_lines"]] == ["Việt Nam", "Trung Quốc"]
    assert material["bang_ke_origin_text"] == "Việt Nam, Trung Quốc"
    assert material["bang_ke_co_doc_no"] == ""
    assert material["bang_ke_co_doc_date"] == ""


def test_qualification_mode_yields_single_label():
    client = {"bang_ke_overrides": {"column9_mode": "qualification_label"}}
    case = materialize_bang_ke_origin_fields(_case_with_lots("VIETNAM", "CHINA"), client)
    material = case["products"][0]["materials"][0]
    assert [line["bang_ke_origin_text"] for line in material["allocation_lines"]] == [
        "Không xuất xứ", "Không xuất xứ",
    ]
    assert material["bang_ke_origin_text"] == "Không xuất xứ"


def test_country_mode_split_activates_fan_out_but_qualification_does_not():
    from app.bang_ke_rows import material_render_parts

    country = materialize_bang_ke_origin_fields(_case_with_lots("CHINA", "TAIWAN"), {})
    country_parts = material_render_parts(country["products"][0]["materials"][0])
    assert [part["bang_ke_origin_text"] for part in country_parts] == ["Trung Quốc", "Đài Loan"]

    qualification = materialize_bang_ke_origin_fields(
        _case_with_lots("CHINA", "TAIWAN"),
        {"bang_ke_overrides": {"column9_mode": "qualification_label"}},
    )
    assert len(material_render_parts(qualification["products"][0]["materials"][0])) == 1


def test_modes_never_change_money():
    country = materialize_bang_ke_origin_fields(_case_with_lots("VIETNAM", "CHINA"), {})
    qualification = materialize_bang_ke_origin_fields(
        _case_with_lots("VIETNAM", "CHINA"),
        {"bang_ke_overrides": {"column9_mode": "qualification_label"}},
    )
    money_fields = ["material_value", "non_origin_cif_value", "consumed_qty", "unit_value", "origin_status"]
    for field in money_fields:
        assert (
            country["products"][0]["materials"][0][field]
            == qualification["products"][0]["materials"][0][field]
        ), field
    enriched_country = enrich_origin_product({"code": "P", "fob": "100", **country["products"][0]})
    enriched_qualification = enrich_origin_product({"code": "P", "fob": "100", **qualification["products"][0]})
    assert enriched_country["lvc_percentage"] == enriched_qualification["lvc_percentage"]
    assert enriched_country["lvc_status"] == enriched_qualification["lvc_status"]


def test_unknown_bucket_gets_label_but_unmapped_string_renders_raw_with_warning():
    case = materialize_bang_ke_origin_fields(_case_with_lots("UNKNOWN", "WAKANDA"), {})
    material = case["products"][0]["materials"][0]
    assert [line["bang_ke_origin_text"] for line in material["allocation_lines"]] == [
        "Không xác định", "WAKANDA",
    ]
    warnings = material["material_warnings"]
    assert any("WAKANDA" in warning and "bảng quy đổi" in warning for warning in warnings)
    assert not any("Không xác định" in warning for warning in warnings)


def test_custom_unknown_label_from_client_overlay():
    client = {"bang_ke_overrides": {"unknown_origin_label": "Không xuất xứ"}}
    case = materialize_bang_ke_origin_fields(_case_with_lots("UNKNOWN"), client)
    assert case["products"][0]["materials"][0]["allocation_lines"][0]["bang_ke_origin_text"] == "Không xuất xứ"


def test_material_without_lot_data_stays_blank():
    material = origin_material_from_bom_row(
        {"material_code": "GHOST", "qty_per": "1"}, Decimal("1"), {}, {},
        product_code="TP1", product_name="SP",
    )
    case = materialize_bang_ke_origin_fields({"products": [{"code": "TP1", "materials": [material]}]}, {})
    assert case["products"][0]["materials"][0]["bang_ke_origin_text"] == ""


# --- renderers are pure readers ---

def test_renderer_prefers_materialized_text_over_raw_country():
    from openpyxl import Workbook
    from app.bang_ke_renderer import load_form_config, render_into_sheet

    cfg = load_form_config("LVC")
    wb = Workbook()
    ws = wb.active
    product = {"code": "P1", "materials": [{
        "material_code": "NVL-1", "material_description": "Thép",
        "origin_status": "non_origin", "hs_code": "73182200", "uom": "PCE",
        "bom_qty_per": "1", "consumed_qty": "1", "unit_value": "1", "material_value": "1",
        "origin_country": "VIETNAM", "bang_ke_origin_text": "Việt Nam",
        "bang_ke_co_doc_no": "Phụ lục X/NCC A", "bang_ke_co_doc_date": "01/07/2026",
    }]}
    render_into_sheet(ws, cfg, case={"case_code": "C", "products": [product]}, product=product, sheet_title="P1")
    cols, start = cfg["body"]["columns"], cfg["body"]["start_row"]
    assert ws[f"{cols['country']}{start}"].value == "Việt Nam"
    assert ws[f"{cols['co_doc_no']}{start}"].value == "Phụ lục X/NCC A"
    assert ws[f"{cols['co_doc_date']}{start}"].value == "01/07/2026"


def test_renderer_falls_back_to_legacy_raw_then_blank():
    from openpyxl import Workbook
    from app.bang_ke_renderer import load_form_config, render_into_sheet

    cfg = load_form_config("LVC")
    base = {
        "material_code": "NVL-1", "material_description": "Thép",
        "origin_status": "non_origin", "hs_code": "73182200", "uom": "PCE",
        "bom_qty_per": "1", "consumed_qty": "1", "unit_value": "1", "material_value": "1",
    }
    for material, expected in [
        ({**base, "origin_country": "VIETNAM"}, ["VIETNAM"]),  # ticket-#6-era sheet
        (dict(base), [None, ""]),  # pre-feature sheet stays blank
    ]:
        wb = Workbook()
        ws = wb.active
        product = {"code": "P1", "materials": [material]}
        render_into_sheet(ws, cfg, case={"case_code": "C", "products": [product]}, product=product, sheet_title="P1")
        cols, start = cfg["body"]["columns"], cfg["body"]["start_row"]
        assert ws[f"{cols['country']}{start}"].value in expected


def test_xml_generator_and_legacy_workbook_read_materialized_fields():
    from openpyxl import Workbook
    from app.bang_ke_xml_generator import _build_material_row
    from app.workbook_io import _hq_layout_for, write_hq_sheet_materials

    material = {
        "material_code": "NVL-1", "material_description": "Thép",
        "origin_status": "non_origin", "hs_code": "73182200", "uom": "PCE",
        "bom_qty_per": "1", "consumed_qty": "1", "unit_value": "1", "material_value": "1",
        "origin_country": "VIETNAM", "bang_ke_origin_text": "Việt Nam",
        "bang_ke_co_doc_no": "Phụ lục X/NCC A", "bang_ke_co_doc_date": "01/07/2026",
    }
    values, _o, _n = _build_material_row(material, {}, {"code": "P1"}, 1)
    assert values["country"] == "Việt Nam"
    assert values["co_doc_no"] == "Phụ lục X/NCC A"
    assert values["co_doc_date"] == "01/07/2026"

    layout = _hq_layout_for("LVC")["cols"]
    wb = Workbook()
    ws = wb.active
    write_hq_sheet_materials(ws, {"materials": [material]}, 5, sheet_code="LVC")
    assert ws.cell(row=5, column=layout["country"]).value == "Việt Nam"
    assert ws.cell(row=5, column=layout["co_no"]).value == "Phụ lục X/NCC A"
    assert ws.cell(row=5, column=layout["co_date"]).value == "01/07/2026"


def test_column9_parity_across_all_three_export_formats():
    # One calculated material rendered through every export format shows EXACTLY
    # the materialized web-grid text — export == web by construction.
    from openpyxl import Workbook
    from app.bang_ke_renderer import load_form_config, render_into_sheet
    from app.bang_ke_xml_generator import _build_material_row
    from app.workbook_io import _hq_layout_for, write_hq_sheet_materials

    material = {
        "material_code": "NVL-1", "material_description": "Thép",
        "origin_status": "non_origin", "hs_code": "73182200", "uom": "PCE",
        "bom_qty_per": "1", "consumed_qty": "1", "unit_value": "1", "material_value": "1",
        "origin_country": "CHINA", "bang_ke_origin_text": "Trung Quốc",
    }
    web_text = material["bang_ke_origin_text"]

    cfg = load_form_config("LVC")
    wb = Workbook()
    ws = wb.active
    product = {"code": "P1", "materials": [dict(material)]}
    render_into_sheet(ws, cfg, case={"case_code": "C", "products": [product]}, product=product, sheet_title="P1")
    xlsx_text = ws[f"{cfg['body']['columns']['country']}{cfg['body']['start_row']}"].value

    xml_text = _build_material_row(dict(material), {}, {"code": "P1"}, 1)[0]["country"]

    wb2 = Workbook()
    ws2 = wb2.active
    write_hq_sheet_materials(ws2, {"materials": [dict(material)]}, 5, sheet_code="LVC")
    legacy_text = ws2.cell(row=5, column=_hq_layout_for("LVC")["cols"]["country"]).value

    assert xlsx_text == xml_text == legacy_text == web_text


# --- config form surface ---

def test_config_page_renders_bang_ke_settings_and_readonly_mapping(monkeypatch):
    from fastapi.testclient import TestClient
    import app.main as main_module
    from app import co_stock_materializer

    monkeypatch.setattr(
        co_stock_materializer,
        "read_co_stock_rows_cached",
        lambda client_id: [{"declaration_type": "E11", "origin_country": "SWITZLD"},
                           {"declaration_type": "E11", "origin_country": "WAKANDA"}],
    )
    client = TestClient(main_module.app)
    page = client.get("/clients/growatt/config")
    assert page.status_code == 200
    assert 'name="bang_ke_column9_mode"' in page.text
    assert 'name="bang_ke_unknown_origin_label"' in page.text
    assert "data-origin-country-mapping" in page.text
    assert "VIETNAM → Việt Nam" in page.text
    # SWITZLD is mapped (→ Thụy Sĩ); WAKANDA is the client's own unmapped string
    assert "data-origin-country-unmapped" in page.text
    assert "WAKANDA" in page.text


# --- form round-trip ---

def test_form_roundtrips_materialized_fields():
    from app.demo_data import update_products_from_form

    form = {
        "product_count": "1",
        "product_0_code": "TP1",
        "product_0_bang_ke_column9_mode": "country",
        "product_0_material_count": "1",
        "product_0_material_0_material_code": "NVL-1",
        "product_0_material_0_bang_ke_origin_text": "Việt Nam, Trung Quốc",
        "product_0_material_0_bang_ke_co_doc_no": "Phụ lục X/NCC A",
        "product_0_material_0_bang_ke_co_doc_date": "01/07/2026",
        "product_0_material_0_allocation_line_count": "1",
        "product_0_material_0_allocation_0_bang_ke_origin_text": "Việt Nam",
    }
    case = update_products_from_form(form)
    product = case["products"][0]
    material = product["materials"][0]
    assert product["bang_ke_column9_mode"] == "country"
    assert material["bang_ke_origin_text"] == "Việt Nam, Trung Quốc"
    assert material["bang_ke_co_doc_no"] == "Phụ lục X/NCC A"
    assert material["bang_ke_co_doc_date"] == "01/07/2026"
    assert material["allocation_lines"][0]["bang_ke_origin_text"] == "Việt Nam"
