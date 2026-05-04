import hashlib
import html
import json
import os
import re
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.bom_store import get_bom_workspace
from app.co_case_store import MAX_SUPPORTING_FILE_BYTES, match_case_bcct_exports, update_case_record
from app.client_config_store import (
    get_client_config,
    resolve_allocation_code,
    save_client_config,
)
from app.demo_data import get_client, get_client_case, get_clients, get_demo_case, update_products_from_form
from app.main import app
from app.origin import calculate_rvc, evaluate_tariff_shift
from app.source_store import (
    create_bcct_template_workbook,
    create_material_catalog_template_workbook,
    create_product_catalog_template_workbook,
    get_source_summary,
    get_source_workspace,
    parse_bcct_workbook,
    parse_catalog_workbook,
    process_bcct_upload,
    process_catalog_upload,
)
from app.source_index_store import (
    build_bcct_index_records,
    build_catalog_index_records,
    build_source_state_records,
    PostgresSourceIndexStore,
    source_state_from_workspace,
)
from app.table_view import build_table_view
from app.workbook_io import create_evidence_workbook, create_input_workbook, parse_input_workbook


@pytest.fixture(autouse=True)
def isolate_bom_store(tmp_path, monkeypatch):
    monkeypatch.setenv("BOM_STORE_ROOT", str(tmp_path / "bom-store"))
    monkeypatch.setenv("SOURCE_STORE_ROOT", str(tmp_path / "source-store"))
    monkeypatch.setenv("CLIENT_CONFIG_ROOT", str(tmp_path / "client-config"))
    monkeypatch.setenv("CO_CASE_STORE_ROOT", str(tmp_path / "co-case-store"))
    monkeypatch.setenv("CUSTOMS_FX_STORE_ROOT", str(tmp_path / "customs-fx-store"))
    monkeypatch.setenv("CO_FORM_CONFIG_PATH", str(tmp_path / "co-form-index.json"))


def workbook_bytes(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def hidden_form_data(markup: str) -> dict[str, str]:
    return {
        name: html.unescape(value)
        for name, value in re.findall(r'<input type="hidden" name="([^"]+)" value="([^"]*)"', markup)
    }


def co_form_settings_payload(config: dict) -> dict[str, str]:
    data = {
        "source_note": config["source_note"],
        "form_priority": ", ".join(config["form_priority"]),
        "form_count": str(len(config["forms"])),
        "market_count": str(len(config["market_presets"])),
        "psr_count": str(len(config.get("psr_rules", []))),
    }
    for index, form in enumerate(config["forms"]):
        prefix = f"form_{index}"
        data[f"{prefix}_enabled"] = "1" if form.get("enabled") else ""
        for key in [
            "form_code",
            "display_name",
            "agreement",
            "instrument",
            "instrument_note",
            "source_label",
            "source_url",
            "verification_status",
        ]:
            data[f"{prefix}_{key}"] = form.get(key, "")
    for index, market in enumerate(config["market_presets"]):
        prefix = f"market_{index}"
        data[f"{prefix}_enabled"] = "1" if market.get("enabled") else ""
        data[f"{prefix}_show_in_picker"] = "1" if market.get("show_in_picker") else ""
        data[f"{prefix}_market"] = market.get("market", "")
        data[f"{prefix}_label"] = market.get("label", "")
        data[f"{prefix}_form_code"] = market.get("form_code", "")
        data[f"{prefix}_aliases"] = ", ".join(market.get("aliases", []))
        data[f"{prefix}_selection_reason"] = market.get("selection_reason", "")
        data[f"{prefix}_source_label"] = market.get("source_label", "")
    for index, rule in enumerate(config.get("psr_rules", [])):
        prefix = f"psr_{index}"
        data[f"{prefix}_enabled"] = "1" if rule.get("enabled") else ""
        data[f"{prefix}_form_code"] = rule.get("form_code", "")
        data[f"{prefix}_hs_scope"] = rule.get("hs_scope", "")
        data[f"{prefix}_criteria"] = rule.get("criteria", "")
        data[f"{prefix}_source_reference"] = rule.get("source_reference", "")
        data[f"{prefix}_note"] = rule.get("note", "")
        data[f"{prefix}_status"] = rule.get("status", "")
    return data


def bcct_workbook(rows: list[dict]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "BCCT"
    worksheet.append([
        "coverage_period",
        "direction",
        "declaration_no",
        "declaration_date",
        "customs_office",
        "declaration_type",
        "line_no",
        "item_code",
        "description",
        "hs_code",
        "quantity",
        "unit",
        "customs_value",
        "currency",
        "invoice_ref",
    ])
    for row in rows:
        worksheet.append([
            row.get("coverage_period", ""),
            row["direction"],
            row["declaration_no"],
            row.get("declaration_date", ""),
            row.get("customs_office", ""),
            row.get("declaration_type", ""),
            row["line_no"],
            row["item_code"],
            row.get("description", ""),
            row.get("hs_code", ""),
            row["quantity"],
            row["unit"],
            row.get("customs_value", ""),
            row.get("currency", ""),
            row.get("invoice_ref", ""),
        ])
    return workbook_bytes(workbook)


def material_catalog_value_workbook(rows: list[dict]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "NVL"
    worksheet.append(["customs_code", "name", "unit", "hs_code", "unit_price", "origin_default", "status"])
    for row in rows:
        worksheet.append([
            row["customs_code"],
            row.get("name", ""),
            row.get("unit", ""),
            row.get("hs_code", ""),
            row.get("unit_price", ""),
            row.get("origin_default", ""),
            row.get("status", "active"),
        ])
    return workbook_bytes(workbook)


def customs_zip_entry(filename: str) -> bytes:
    archive_path = Path("temp/drive-download-20260428T145419Z-3-001.zip")
    if not archive_path.exists():
        pytest.skip("Customs sample ZIP is local-only.")
    with ZipFile(archive_path) as archive:
        return archive.read(filename)


def test_calculates_growatt_documented_rvc_seed_from_material_rows():
    case = get_demo_case()

    assert case["products"][0]["result"].rvc.percentage == Decimal("35.88")
    assert case["products"][1]["result"].rvc.percentage == Decimal("40.19")
    assert case["summary"]["passed_products"] == 2


def test_evaluates_ctsh_by_six_digit_subheading():
    result = evaluate_tariff_shift("8504.40", ["8542.39", "8536.90"], "CTSH")

    assert result.finished_key == "850440"
    assert result.input_keys == ("854239", "853690")
    assert result.passed is True


def test_form_update_recomputes_rvc_from_edited_material_value():
    case = update_products_from_form(
        {
            "customer": "Growatt",
            "case_code": "TEST",
            "product_count": "1",
            "product_0_code": "PV00.0048500",
            "product_0_name": "Model inverter PV00.0048500",
            "product_0_finished_hs": "8504.40",
            "product_0_fob": "100000",
            "product_0_rvc_threshold": "35",
            "product_0_material_count": "1",
            "product_0_material_0_source_row": "IMP-001/1",
            "product_0_material_0_hs_code": "8542.39",
            "product_0_material_0_origin_status": "non_origin",
            "product_0_material_0_available_qty": "520",
            "product_0_material_0_consumed_qty": "120",
            "product_0_material_0_non_origin_cif_value": "70000",
        }
    )

    assert case["products"][0]["result"].rvc.percentage == Decimal("30.00")
    assert case["products"][0]["result"].passed is False


def test_seed_workbook_round_trips_through_parser():
    parsed = parse_input_workbook(create_input_workbook(get_demo_case()), "roundtrip")

    assert parsed["source_label"] == "roundtrip"
    assert parsed["products"][0]["code"] == "PV00.0048500"
    assert parsed["products"][0]["materials"][0]["customs_material_code"] == "DEMO-NPL-001"
    assert parsed["products"][0]["materials"][0]["internal_material_code"] == "DEMO-NPL-001"
    assert parsed["products"][0]["result"].rvc.percentage == Decimal("35.88")


def test_evidence_workbook_is_generated():
    content = create_evidence_workbook(get_demo_case())

    assert content.startswith(b"PK")


def test_workspace_renders_interactive_controls():
    client = TestClient(app)

    response = client.get("/clients/growatt")

    assert response.status_code == 200
    assert "Growatt" in response.text
    assert "/clients/growatt/catalog" in response.text
    assert "/clients/growatt/bom" in response.text
    assert "/clients/growatt/co-stock" in response.text
    assert "/clients/growatt/bcct" in response.text
    assert "/clients/growatt/co-case" in response.text
    assert "Workspace theo công ty" in response.text


def test_theme_toggle_persists_dark_theme_cookie():
    client = TestClient(app)

    initial = client.get("/clients")
    assert initial.status_code == 200
    assert 'data-theme="light"' in initial.text

    response = client.post(
        "/settings/theme",
        data={"theme": "dark", "next_url": "/clients/growatt"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"] == "/clients/growatt"
    assert "co_theme=dark" in response.headers["set-cookie"]

    themed = client.get("/clients/growatt")
    assert themed.status_code == 200
    assert 'data-theme="dark"' in themed.text
    assert 'name="theme" value="light"' in themed.text


def test_client_navigation_hides_data_modules_inside_co_workflow():
    client = TestClient(app)

    response = client.get("/clients/growatt/co-case")

    assert response.status_code == 200
    assert 'aria-label="Luồng làm C/O"' in response.text
    assert 'aria-label="Dữ liệu nền công ty"' not in response.text
    assert "Làm hồ sơ C/O" in response.text
    assert 'href="/clients/growatt/co-case">Hồ sơ C/O</a>' not in response.text
    assert "Overview" not in response.text
    assert "Danh mục mã hàng" not in response.text
    assert "/clients/growatt/catalog" not in response.text
    assert "/clients/growatt/bom" not in response.text
    assert "/clients/growatt/co-stock" not in response.text
    assert "/clients/growatt/bcct" not in response.text

    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Hidden data nav case", "case_code": "NAV-HIDE", "destination_market": "EU"},
        follow_redirects=False,
    )
    detail = client.get(created.headers["location"])

    assert detail.status_code == 200
    assert 'aria-label="Dữ liệu nền công ty"' not in detail.text
    assert "Overview" not in detail.text
    assert "Danh mục mã hàng" not in detail.text

    catalog = client.get("/clients/growatt/catalog")

    assert catalog.status_code == 200
    assert 'aria-label="Dữ liệu nền công ty"' in catalog.text
    assert "Danh mục mã hàng" in catalog.text


def test_table_view_filters_searches_sorts_and_paginates():
    rows = [
        {"code": "MAT-001", "name": "Capacitor", "status": "inactive"},
        {"code": "MAT-002", "name": "Diode bridge", "status": "active"},
        {"code": "MAT-003", "name": "Small diode", "status": "active"},
    ]

    table = build_table_view(
        rows,
        columns=[
            {"key": "code", "label": "Code"},
            {"key": "name", "label": "Name"},
            {"key": "status", "label": "Status"},
        ],
        query={"q": "diode", "status": "active", "sort": "code", "dir": "desc", "page": "2", "per_page": "1"},
        filters=[{"name": "status", "field": "status", "label": "Status"}],
    )

    assert table["total_count"] == 3
    assert table["filtered_count"] == 2
    assert table["page"] == 2
    assert table["total_pages"] == 2
    assert [row["code"] for row in table["rows"]] == ["MAT-002"]


def test_table_view_clamps_page_to_available_results():
    table = build_table_view(
        [{"code": "A"}, {"code": "B"}, {"code": "C"}],
        columns=[{"key": "code", "label": "Code"}],
        query={"page": "99", "per_page": "2"},
    )

    assert table["page"] == 2
    assert table["total_pages"] == 2
    assert [row["code"] for row in table["rows"]] == ["C"]


def test_catalog_bom_stock_bcct_are_data_views_and_co_case_is_workflow_entry():
    client = TestClient(app)

    catalog_response = client.get("/clients/growatt/catalog")
    material_catalog_response = client.get("/clients/growatt/catalog/materials")
    product_catalog_response = client.get("/clients/growatt/catalog/products")
    bom_response = client.get("/clients/growatt/bom")
    stock_response = client.get("/clients/growatt/co-stock")
    bcct_response = client.get("/clients/growatt/bcct")
    customs_fx_response = client.get("/customs-exchange-rates")
    co_case_response = client.get("/clients/growatt/co-case")

    assert catalog_response.status_code == 200
    assert material_catalog_response.status_code == 200
    assert product_catalog_response.status_code == 200
    assert bom_response.status_code == 200
    assert stock_response.status_code == 200
    assert bcct_response.status_code == 200
    assert customs_fx_response.status_code == 200
    assert co_case_response.status_code == 200
    assert "Upload danh mục" in catalog_response.text
    assert "/clients/growatt/catalog/materials" in catalog_response.text
    assert "/clients/growatt/catalog/products" in catalog_response.text
    assert "DS NVL DK HQ" in material_catalog_response.text
    assert "DEMO-NPL-001" in material_catalog_response.text
    assert "Mã nội bộ" not in material_catalog_response.text
    assert "DS SP DK HQ" in product_catalog_response.text
    assert "PV00.0048500" in product_catalog_response.text
    assert "DEMO-NPL-001" in bom_response.text
    assert "Tồn CO khác tồn kho vật lý" in stock_response.text
    assert "BCCT nhập khẩu / xuất khẩu" in bcct_response.text
    assert "107101950210" in bcct_response.text
    assert "Tỷ giá hải quan" in customs_fx_response.text
    assert "app-level" in customs_fx_response.text
    assert "Refresh tỷ giá" in customs_fx_response.text
    assert "Quy trình làm C/O" in co_case_response.text
    assert "Tạo hoặc mở hồ sơ" in co_case_response.text
    assert "Danh sách hồ sơ C/O" in co_case_response.text
    assert "Các bước xử lý" not in co_case_response.text
    assert "Dữ liệu nền đang sẵn sàng" not in co_case_response.text
    assert "Đánh giá RVC + CTSH" not in co_case_response.text
    assert "BCCT xuất khẩu theo invoice" not in co_case_response.text
    assert "BTP" not in catalog_response.text


def test_catalog_child_route_search_limits_material_rows():
    client = TestClient(app)

    response = client.get("/clients/growatt/catalog/materials?q=DEMO-NPL-002")

    assert response.status_code == 200
    assert "DEMO-NPL-002" in response.text
    assert "DEMO-NPL-001" not in response.text
    assert "1 / 3 dòng" in response.text


def test_catalog_product_route_search_limits_product_rows():
    client = TestClient(app)

    response = client.get("/clients/growatt/catalog/products?q=PV01.0117600")

    assert response.status_code == 200
    assert "PV01.0117600" in response.text
    assert "PV00.0048500" not in response.text
    assert "1 / 2 dòng" in response.text


def test_product_catalog_does_not_render_origin_rule():
    client = TestClient(app)

    response = client.get("/clients/growatt/catalog/products")

    assert response.status_code == 200
    assert "Quy tắc" not in response.text
    assert "RVC 35% + CTSH" not in response.text


def test_product_catalog_template_does_not_include_origin_rule_column():
    workbook = load_workbook(BytesIO(create_product_catalog_template_workbook(get_client("growatt"))))
    headers = [cell.value for cell in workbook.active[1]]

    assert "Quy tắc" not in headers
    assert "origin_rule" not in headers
    assert "rule" not in headers


def test_product_catalog_upload_ignores_legacy_origin_rule_column():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Mã", "Tên", "Đơn vị tính", "Mã HS", "Quy tắc"])
    worksheet.append(["TP-001", "Finished product", "PCS", "85044090", "RVC 35% + CTSH"])

    result = process_catalog_upload(
        get_client("do-thanh"),
        "product",
        workbook_bytes(workbook),
        "legacy-ds-sp.xlsx",
        "full_catalog",
    )

    rows = get_source_workspace(get_client("do-thanh"))["product_catalog"]["published_rows"]
    assert result["status"] == "new_version"
    assert rows[0]["product_code"] == "TP-001"
    assert "rule" not in rows[0]
    assert rows[0]["raw_fields"]["Quy tắc"] == "RVC 35% + CTSH"


def test_bcct_route_paginates_rows_server_side():
    client = TestClient(app)

    response = client.get("/clients/growatt/bcct?per_page=1&page=2")

    assert response.status_code == 200
    assert "GIN01425L031" in response.text
    assert "107101950210" not in response.text
    assert "Trang 2 / 3" in response.text


def customs_fx_payloads() -> tuple[dict, dict]:
    return (
        {
            "d": [
                {"DONG_TIEN": "USD", "TEN_DONG_TIEN": "Đô-la Mỹ"},
                {"DONG_TIEN": "JPY", "TEN_DONG_TIEN": "Yên Nhật"},
            ]
        },
        {
            "d": [
                {
                    "LOAI_NGOAI_TE": "USD",
                    "TEN_NGOAI_TE": "Đô-la Mỹ",
                    "HIEU_LUC_TU_NGAY": "27/04/2026",
                    "TY_GIA": "26.130 VNĐ",
                },
                {
                    "LOAI_NGOAI_TE": "JPY",
                    "TEN_NGOAI_TE": "Yên Nhật",
                    "HIEU_LUC_TU_NGAY": "27/04/2026",
                    "TY_GIA": "177 VNĐ",
                },
                {
                    "LOAI_NGOAI_TE": "JPY",
                    "TEN_NGOAI_TE": "Yên Nhật",
                    "HIEU_LUC_TU_NGAY": "20/04/2026",
                    "TY_GIA": "178 VNĐ",
                },
            ]
        },
    )


def test_customs_fx_parser_preserves_vietnamese_rate_format():
    from app.customs_fx_store import lookup_exchange_rate, parse_customs_exchange_rate_payloads, parse_vnd_rate_text

    rows = parse_customs_exchange_rate_payloads(*customs_fx_payloads())

    usd = [row for row in rows if row["currency_code"] == "USD"][0]
    assert parse_vnd_rate_text("26.130 VNĐ") == Decimal("26130")
    assert usd["effective_date"] == "2026-04-27"
    assert usd["rate_vnd_per_unit"] == "26130"
    assert usd["rate_display"] == "26.130 VNĐ"
    assert lookup_exchange_rate(rows, "USD", "2026-05-01")["rate_vnd_per_unit"] == "26130"


def test_customs_fx_fetch_uses_customs_public_json_endpoints():
    from app.customs_fx_store import fetch_customs_exchange_rates

    currency_payload, history_payload = customs_fx_payloads()

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "GetListDongTienTyGia" in url:
            return httpx.Response(200, json=currency_payload)
        if "GetListRateByNameOrDate" in url:
            payload = json.loads(request.content.decode())
            assert payload["ten_ngoai_te"] == ""
            assert payload["hieu_luc_tu_ngay"] == "01-04-2026"
            assert payload["hieu_luc_den_ngay"] == "01-05-2026"
            assert payload["captcha"] == ""
            return httpx.Response(200, json=history_payload)
        return httpx.Response(404, json={})

    rows = fetch_customs_exchange_rates(
        history_start_date="2026-04-01",
        history_end_date="2026-05-01",
        transport=httpx.MockTransport(handler),
    )

    assert {row["currency_code"] for row in rows} == {"USD", "JPY"}
    assert len([row for row in rows if row["currency_code"] == "JPY"]) == 2
    assert rows[0]["source"] == "customs.gov.vn"
    assert rows[0]["source_endpoint"] == "GetListRateByNameOrDate"


def test_customs_fx_file_store_upserts_global_rate_rows():
    from app.customs_fx_store import CUSTOMS_FX_CLIENT_ID, FileCustomsFxStore, parse_customs_exchange_rate_payloads

    rows = parse_customs_exchange_rate_payloads(*customs_fx_payloads())
    store = FileCustomsFxStore()

    first = store.save_refresh(CUSTOMS_FX_CLIENT_ID, rows)
    second = store.save_refresh(CUSTOMS_FX_CLIENT_ID, rows)

    assert first["fetched_row_count"] == 3
    assert first["saved_row_count"] == 3
    assert second["upserted_row_count"] == 0
    assert store.summary()["currency_count"] == 2
    assert store.rows()[0]["effective_date"] == "2026-04-27"


def test_customs_fx_route_refreshes_and_filters_rows(monkeypatch):
    from app import customs_fx_store as customs_fx_module
    from app.customs_fx_store import parse_customs_exchange_rate_payloads

    rows = parse_customs_exchange_rate_payloads(*customs_fx_payloads())
    monkeypatch.setattr(customs_fx_module, "fetch_customs_exchange_rates", lambda **_kwargs: rows)

    client = TestClient(app)
    refresh = client.post("/customs-exchange-rates/refresh")
    filtered = client.get("/customs-exchange-rates?currency=JPY")

    assert refresh.status_code == 200
    assert "Đã cập nhật 3 dòng tỷ giá hải quan" in refresh.text
    assert "26.130 VNĐ" in refresh.text
    assert filtered.status_code == 200
    assert "JPY" in filtered.text
    assert "26.130 VNĐ" not in filtered.text


def test_customs_fx_client_route_redirects_to_app_level_surface():
    client = TestClient(app)

    response = client.get("/clients/growatt/customs-exchange-rates", follow_redirects=False)

    assert response.status_code == 303
    assert response.headers["location"] == "/customs-exchange-rates"


def test_co_stock_route_search_and_status_filter():
    client = TestClient(app)

    response = client.get("/clients/growatt/co-stock?q=DEMO-NPL-001&status=available")

    assert response.status_code == 200
    assert "DEMO-NPL-001" in response.text
    assert "GIN01425L031" not in response.text
    assert "Hiển thị 1-1 / 1 dòng" in response.text


def test_clients_page_is_entry_point():
    client = TestClient(app)

    response = client.get("/clients")

    assert response.status_code == 200
    assert "Danh sách công ty" in response.text
    assert "Growatt" in response.text
    assert "Johnson" in response.text
    assert "Danh mục TP" in response.text
    assert "Danh mục NVL" in response.text
    assert "BOM" in response.text
    assert "Tồn CO" in response.text
    assert "BTP" not in response.text


def test_upload_seed_workbook_runs_parser():
    client = TestClient(app)
    content = create_input_workbook(get_demo_case())

    response = client.post(
        "/clients/growatt/upload",
        files={"file": ("growatt.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "Đã parse growatt.xlsx" in response.text
    assert "PV01.0117600" in response.text


def test_download_seed_input_and_export_routes_return_xlsx():
    client = TestClient(app)

    seed_response = client.get("/clients/growatt/demo-input.xlsx")
    assert seed_response.status_code == 200
    assert seed_response.content.startswith(b"PK")

    export_response = client.post(
        "/clients/growatt/export",
        data={
            "customer": "Growatt",
            "case_code": "TEST",
            "product_count": "1",
            "product_0_code": "PV00.0048500",
            "product_0_name": "Model inverter PV00.0048500",
            "product_0_finished_hs": "8504.40",
            "product_0_fob": "100000",
            "product_0_rvc_threshold": "35",
            "product_0_material_count": "1",
            "product_0_material_0_hs_code": "8542.39",
            "product_0_material_0_origin_status": "non_origin",
            "product_0_material_0_available_qty": "520",
            "product_0_material_0_consumed_qty": "120",
            "product_0_material_0_non_origin_cif_value": "64120",
        },
    )
    assert export_response.status_code == 200
    assert export_response.content.startswith(b"PK")


def test_low_level_rvc_calculation_still_matches_formula():
    result = calculate_rvc("100000", "64120", "35")

    assert result.percentage == Decimal("35.88")
    assert result.passed is True


def test_client_seed_has_required_company_modules():
    growatt = [client for client in get_clients() if client["id"] == "growatt"][0]
    johnson = [client for client in get_clients() if client["id"] == "johnson"][0]
    case = get_client_case("growatt")

    assert growatt["counts"]["materials"] > 0
    assert growatt["counts"]["products"] > 0
    assert growatt["counts"]["bom_lines"] > 0
    assert growatt["counts"]["co_stock"] > 0
    assert growatt["counts"]["bcct"] > 0
    assert "boms" not in growatt["counts"]
    assert johnson["counts"]["products"] > 0
    assert johnson["counts"]["materials"] > 0
    assert johnson["counts"]["bom_lines"] > 0
    assert case["customer"] == "Growatt"


def test_bom_view_surfaces_company_source_profiles():
    client = TestClient(app)

    growatt_response = client.get("/clients/growatt/bom")
    johnson_response = client.get("/clients/johnson/bom")

    assert growatt_response.status_code == 200
    assert johnson_response.status_code == 200
    assert "Growatt multi-workbook technical BOM graph" in growatt_response.text
    assert "Johnson SAP exploded BOM export" in johnson_response.text
    assert "SAP WebAS" in johnson_response.text


def test_bom_template_download_and_no_change_upload_keeps_current_version():
    client = TestClient(app)

    template = client.get("/clients/growatt/bom/template.xlsx")
    assert template.status_code == 200
    assert template.content.startswith(b"PK")

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom.xlsx", template.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "Không có thay đổi BOM" in response.text
    workspace = get_bom_workspace(get_client("growatt"))
    assert workspace["latest_version"]["version_no"] == 1
    assert {row["product_version_no"] for row in workspace["product_composition"]} == {1}


def test_bom_changed_upload_creates_next_version():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx")
    workbook = load_workbook(BytesIO(template.content))
    worksheet = workbook["BOM"]
    worksheet["F2"] = "2.00"
    stream = BytesIO()
    workbook.save(stream)

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom-changed.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "Đã tạo BOM tổng hợp v2" in response.text
    assert "BOM tổng hợp v2" in response.text
    assert "Đổi <strong>1</strong>" in response.text
    workspace = get_bom_workspace(get_client("growatt"))
    composition = {row["product_code"]: row["product_version_no"] for row in workspace["product_composition"]}
    assert workspace["latest_version"]["version_no"] == 2
    assert composition["PV00.0048500"] == 2
    assert composition["PV01.0117600"] == 1
    assert [version["product_version_no"] for version in workspace["product_versions"] if version["product_code"] == "PV00.0048500"] == [2, 1]


def test_direct_bom_full_aggregate_retires_missing_products():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx")
    workbook = load_workbook(BytesIO(template.content))
    worksheet = workbook["BOM"]
    for row_index in range(worksheet.max_row, 1, -1):
        if worksheet.cell(row_index, 1).value == "PV01.0117600":
            worksheet.delete_rows(row_index)
    worksheet["F2"] = "2.00"
    stream = BytesIO()
    workbook.save(stream)

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom-full.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "PV01.0117600 v1 -&gt; retired" in response.text
    workspace = get_bom_workspace(get_client("growatt"))
    composition = {row["product_code"]: row["product_version_no"] for row in workspace["product_composition"]}
    assert composition == {"PV00.0048500": 2}


def test_direct_bom_partial_upload_preserves_missing_products():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx")
    workbook = load_workbook(BytesIO(template.content))
    worksheet = workbook["BOM"]
    for row_index in range(worksheet.max_row, 1, -1):
        if worksheet.cell(row_index, 1).value == "PV01.0117600":
            worksheet.delete_rows(row_index)
    worksheet["F2"] = "2.00"
    stream = BytesIO()
    workbook.save(stream)

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom", "upload_scope": "partial_product"},
        files={"file": ("growatt-bom-partial.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    workspace = get_bom_workspace(get_client("growatt"))
    composition = {row["product_code"]: row["product_version_no"] for row in workspace["product_composition"]}
    assert composition["PV00.0048500"] == 2
    assert composition["PV01.0117600"] == 1


def test_duplicate_bom_rows_are_rejected_before_versioning():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx")
    workbook = load_workbook(BytesIO(template.content))
    worksheet = workbook["BOM"]
    duplicate_values = [worksheet.cell(2, column_index).value for column_index in range(1, worksheet.max_column + 1)]
    worksheet.append(duplicate_values)
    stream = BytesIO()
    workbook.save(stream)

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom-duplicate.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 400
    assert "BOM có dòng trùng khóa" in response.text
    assert get_bom_workspace(get_client("growatt"))["latest_version"]["version_no"] == 1


def test_reupload_historical_bom_can_create_new_product_version_against_current():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx").content
    client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom.xlsx", template, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    workbook = load_workbook(BytesIO(template))
    workbook["BOM"]["F2"] = "2.00"
    changed = BytesIO()
    workbook.save(changed)
    client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom-changed.xlsx", changed.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom.xlsx", template, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "PV00.0048500 v2 -&gt; v3" in response.text
    workspace = get_bom_workspace(get_client("growatt"))
    composition = {row["product_code"]: row["product_version_no"] for row in workspace["product_composition"]}
    assert workspace["latest_version"]["version_no"] == 3
    assert composition["PV00.0048500"] == 3
    assert composition["PV01.0117600"] == 1
    assert [version["product_version_no"] for version in workspace["product_versions"] if version["product_code"] == "PV00.0048500"] == [3, 2, 1]


def test_bom_company_config_can_be_updated():
    client = TestClient(app)

    response = client.post(
        "/clients/growatt/bom/config",
        data={
            "bom_profile": "johnson_sap_exploded",
            "default_import_mode": "technical_bom",
            "code_system_mode": "single_code",
        },
    )

    assert response.status_code == 200
    assert "Đã lưu cấu hình BOM" in response.text
    assert '<option value="johnson_sap_exploded" selected>' in response.text
    assert '<option value="technical_bom" selected>' in response.text
    assert '<option value="single_code" selected>' in response.text


def test_growatt_technical_bom_upload_requires_flatten_review_before_publish():
    client = TestClient(app)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Growatt"
    worksheet.append(["成品物料", "组件物料", "组件物料描述", "标准用量", "单位"])
    worksheet.append(["PV00.0048500", "GW-NVL-001", "Growatt board", "1.5", "PCE"])
    stream = BytesIO()
    workbook.save(stream)

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "technical_bom"},
        files={"file": ("growatt-technical.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "cần review flatten" in response.text
    assert get_bom_workspace(get_client("growatt"))["latest_version"]["version_no"] == 1


def test_growatt_technical_bom_upload_can_be_accepted_as_flat_for_demo():
    client = TestClient(app)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Growatt"
    worksheet.append(["成品物料", "组件物料", "组件物料描述", "标准用量", "单位"])
    worksheet.append(["PV00.0048500", "GW-NVL-001", "Growatt board", "1.5", "PCE"])
    stream = BytesIO()
    workbook.save(stream)

    response = client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "technical_bom", "accept_review_required": "on"},
        files={"file": ("growatt-technical.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "Đã tạo BOM tổng hợp v2" in response.text
    assert "GW-NVL-001" in response.text
    assert "needs_graph_flatten_review" in response.text
    workspace = get_bom_workspace(get_client("growatt"))
    composition = {row["product_code"]: row["product_version_no"] for row in workspace["product_composition"]}
    assert composition["PV00.0048500"] == 2
    assert composition["PV01.0117600"] == 1


def test_repeated_uploads_get_distinct_upload_ids():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx").content

    for _ in range(2):
        response = client.post(
            "/clients/growatt/bom/upload",
            data={"upload_mode": "direct_bom"},
            files={"file": ("growatt-bom.xlsx", template, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 200

    uploads = get_bom_workspace(get_client("growatt"))["uploads"]
    upload_ids = [upload["upload_id"] for upload in uploads]
    assert len(upload_ids) == len(set(upload_ids))


def test_bom_state_prefers_postgres_store_when_available(monkeypatch):
    from app import bom_store

    saved_states = []

    class FakeBomStateStore:
        def __init__(self):
            self.state = None

        def get_state(self, client_id: str) -> dict | None:
            assert client_id == "growatt"
            return self.state

        def save_state(self, client_id: str, state: dict) -> None:
            assert client_id == "growatt"
            self.state = dict(state)
            saved_states.append(dict(state))

    fake_store = FakeBomStateStore()
    monkeypatch.setattr(bom_store, "get_bom_state_store", lambda: fake_store)

    state = bom_store.load_state(get_client("growatt"))
    bom_store.update_bom_config(get_client("growatt"), {"bom_profile": "manual_flat"})

    assert state["client_id"] == "growatt"
    assert saved_states
    assert fake_store.state["config"]["bom_profile"] == "manual_flat"
    assert not bom_store.state_path("growatt").exists()


def test_co_case_can_select_aggregate_bom_version_snapshot():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx")
    workbook = load_workbook(BytesIO(template.content))
    workbook["BOM"]["F2"] = "2.00"
    stream = BytesIO()
    workbook.save(stream)
    client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom-changed.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    workspace = get_bom_workspace(get_client("growatt"))
    v1 = [version for version in workspace["versions"] if version["version_no"] == 1][0]

    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "BOM snapshot case", "case_code": "CO-BOM", "destination_market": "Ấn Độ"},
        follow_redirects=False,
    )
    get_response = client.get(f"{created.headers['location']}/origin")
    assert get_response.status_code == 200
    assert "BOM snapshot" in get_response.text
    assert "BOM tổng hợp v2" in get_response.text

    post_response = client.post(
        "/clients/growatt/evaluate",
        data={
            "case_id": "TEST",
            "customer": "Growatt",
            "case_code": "TEST",
            "document_count": "0",
            "product_count": "1",
            "bom_version_id": v1["version_id"],
            "product_0_code": "PV00.0048500",
            "product_0_name": "Model inverter PV00.0048500",
            "product_0_finished_hs": "8504.40",
            "product_0_quantity": "demo",
            "product_0_fob": "100000",
            "product_0_rvc_threshold": "35",
            "product_0_documented_result": "demo",
            "product_0_material_count": "0",
        },
    )

    assert post_response.status_code == 200
    assert f'value="{v1["version_id"]}" selected' in post_response.text
    assert "TP BOM v1" in post_response.text


def test_johnson_technical_bom_upload_keeps_sap_leaf_rows_only():
    client = TestClient(app)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.append(["Level", "Explosion level", "Component number", "Object description", "Comp. Qty (CUn)", "Component unit"])
    worksheet.append([1, 1, "ASM-001", "Assembly parent", 1, "EA"])
    worksheet.append([2, 2, "004426-00", "Screw; Flat Head; M6x1.0Px12L", 2, "EA"])
    worksheet.append([2, 2, "1000461274", "Tube; Round; 20#", 1.306, "EA"])
    stream = BytesIO()
    workbook.save(stream)

    response = client.post(
        "/clients/johnson/bom/upload",
        data={"upload_mode": "technical_bom"},
        files={"file": ("MFW0502-39.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "Đã tạo BOM tổng hợp v2" in response.text
    assert "004426-00" in response.text
    assert "1000461274" in response.text
    assert "ASM-001" not in response.text
    workspace = get_bom_workspace(get_client("johnson"))
    product_versions = [version for version in workspace["product_versions"] if version["product_code"] == "MFW0502-39"]
    assert [version["product_version_no"] for version in product_versions] == [2, 1]


def test_material_catalog_full_upload_marks_omitted_code_inactive_pending_review():
    client = get_client("growatt")
    template = create_material_catalog_template_workbook(client)
    workbook = load_workbook(BytesIO(template))
    worksheet = workbook.active
    for row_index in range(worksheet.max_row, 1, -1):
        if worksheet.cell(row_index, 2).value == "DEMO-NPL-003":
            worksheet.delete_rows(row_index)

    result = process_catalog_upload(client, "material", workbook_bytes(workbook), "ds-nvl.xlsx", "full_catalog")

    assert result["status"] == "new_version"
    assert result["summary"]["inactive_pending_review"] == 1
    workspace = get_source_workspace(client)
    omitted_row = [row for row in workspace["material_catalog"]["published_rows"] if row["customs_code"] == "DEMO-NPL-003"][0]
    assert omitted_row["status"] == "inactive_pending_review"


def test_source_upload_metadata_uses_standard_file_contract():
    client = get_client("growatt")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "NVL"
    worksheet.append(["customs_code", "internal_code", "name", "hs_code", "unit", "role", "origin_default", "status"])
    worksheet.append(["STD-MAT-001", "STD-MAT-001", "Standard material", "8542.39", "PCS", "NVL", "Không xuất xứ", "active"])

    content = workbook_bytes(workbook)
    result = process_catalog_upload(client, "material", content, "../DS NVL chuẩn.xlsx", "partial_update")

    assert result["status"] == "new_version"
    upload = result["upload"]
    assert upload["metadata_schema_version"] == 1
    assert upload["client_id"] == "growatt"
    assert upload["module"] == "material_catalog"
    assert upload["original_filename"] == "DS NVL chuẩn.xlsx"
    assert upload["stored_filename"].endswith(".xlsx")
    assert upload["stored_filename"] != upload["original_filename"]
    assert upload["stored_path"].startswith("clients/growatt/material-catalog/uploads/")
    assert upload["storage_backend"] == "filesystem"
    assert upload["content_sha256"] == hashlib.sha256(content).hexdigest()
    assert upload["size_bytes"] == len(content)
    assert upload["file_ext"] == ".xlsx"
    assert upload["upload_scope"] == "partial_update"
    assert upload["parse_status"] == "parsed"
    assert upload["parse_error"] == ""
    assert upload["row_count"] == 1
    assert upload["snapshot_id"].startswith("snapshot-")
    assert upload["snapshot_rows_hash"]
    assert upload["created_version_id"] == result["version"]["version_id"]
    assert upload["result"] == "new_version"
    assert upload["diff_summary"] == result["summary"]
    assert result["version"]["snapshot_id"] == upload["snapshot_id"]
    assert Path(os.environ["SOURCE_STORE_ROOT"], upload["stored_path"]).exists()


def test_postgres_source_state_records_include_standard_upload_versions_and_audit():
    client = get_client("growatt")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "NVL"
    worksheet.append(["customs_code", "internal_code", "name", "hs_code", "unit", "role", "origin_default", "status"])
    worksheet.append(["PG-MAT-001", "PG-MAT-001", "Postgres metadata material", "8542.39", "PCS", "NVL", "Không xuất xứ", "active"])

    result = process_catalog_upload(client, "material", workbook_bytes(workbook), "metadata.xlsx", "partial_update")
    state = get_source_workspace(client)["material_catalog"]
    records = build_source_state_records(client["id"], "material_catalog", state)

    upload = records["uploads"][0]
    assert upload["upload_id"] == result["upload"]["upload_id"]
    assert upload["client_id"] == "growatt"
    assert upload["module"] == "material_catalog"
    assert upload["original_filename"] == "metadata.xlsx"
    assert upload["stored_path"] == result["upload"]["stored_path"]
    assert upload["content_sha256"] == result["upload"]["content_sha256"]
    assert upload["size_bytes"] == result["upload"]["size_bytes"]
    assert upload["parse_status"] == "parsed"
    assert upload["snapshot_id"] == result["upload"]["snapshot_id"]
    assert upload["created_version_id"] == result["version"]["version_id"]
    assert upload["diff_summary"] == result["summary"]
    assert records["raw_files"][0]["upload_id"] == upload["upload_id"]
    assert records["raw_files"][0]["storage_key"] == upload["stored_path"]
    assert records["snapshots"][0]["snapshot_id"] == result["upload"]["snapshot_id"]
    assert records["snapshots"][0]["rows_hash"] == result["upload"]["snapshot_rows_hash"]
    assert records["versions"][0]["version_id"] == result["version"]["version_id"]
    assert records["versions"][0]["snapshot_id"] == result["upload"]["snapshot_id"]
    snapshot_rows = [
        row for row in records["snapshot_rows"]
        if row["snapshot_id"] == result["upload"]["snapshot_id"]
    ]
    assert snapshot_rows[0]["row_key"] == "PG-MAT-001"
    assert snapshot_rows[0]["row_index"] == 1
    assert snapshot_rows[0]["customs_code"] == "PG-MAT-001"
    assert snapshot_rows[0]["payload"]["name"] == "Postgres metadata material"
    version_rows = [
        row for row in records["version_rows"]
        if row["version_id"] == result["version"]["version_id"] and row["row_key"] == "PG-MAT-001"
    ]
    assert version_rows[0]["customs_code"] == "PG-MAT-001"
    assert version_rows[0]["payload"]["name"] == "Postgres metadata material"
    assert records["audit_events"][0]["event"].startswith("material_catalog.")
    assert records["module_state"]["latest_version_id"] == result["version"]["version_id"]


def test_postgres_source_state_records_keep_direct_upload_history_rows_without_artifacts():
    workspace = {
        "module": "material_catalog",
        "published_rows": [{"customs_code": "PG-DIRECT-001", "name": "Direct row", "unit": "PCS"}],
        "latest_version": {"version_id": "material_catalog-v1-direct", "version_no": 1},
        "versions": [
            {
                "version_id": "material_catalog-v1-direct",
                "version_no": 1,
                "source_upload_id": "upload-direct",
                "snapshot_id": "snapshot-direct",
                "row_count": 1,
                "rows_hash": "hash-direct",
                "summary": {"added": 1},
            }
        ],
        "uploads": [
            {
                "upload_id": "upload-direct",
                "original_filename": "direct.xlsx",
                "stored_filename": "upload-direct-direct.xlsx",
                "stored_path": "clients/growatt/material-catalog/uploads/upload-direct/raw/upload-direct-direct.xlsx",
                "parse_status": "parsed",
                "snapshot_id": "snapshot-direct",
                "row_count": 1,
                "result": "new_version",
            }
        ],
        "snapshot_rows": {
            "snapshot-direct": [{"customs_code": "PG-DIRECT-001", "name": "Direct row", "unit": "PCS"}],
        },
        "version_rows": {
            "material_catalog-v1-direct": [{"customs_code": "PG-DIRECT-001", "name": "Direct row", "unit": "PCS"}],
        },
        "correction_candidates": [],
        "audit_events": [],
    }

    state = source_state_from_workspace("growatt", "material_catalog", workspace)
    records = build_source_state_records("growatt", "material_catalog", state)

    assert records["snapshot_rows"][0]["snapshot_id"] == "snapshot-direct"
    assert records["snapshot_rows"][0]["row_key"] == "PG-DIRECT-001"
    assert records["version_rows"][0]["version_id"] == "material_catalog-v1-direct"
    assert records["version_rows"][0]["payload"]["name"] == "Direct row"


def test_material_catalog_partial_update_does_not_deactivate_missing_codes():
    client = get_client("growatt")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "NVL"
    worksheet.append(["customs_code", "internal_code", "name", "hs_code", "unit", "role", "origin_default", "status"])
    worksheet.append(["DEMO-NPL-001", "DEMO-NPL-001", "Main control board - edited", "8542.39", "PCS", "NVL", "Không xuất xứ", "active"])

    result = process_catalog_upload(client, "material", workbook_bytes(workbook), "ds-nvl-partial.xlsx", "partial_update")

    assert result["status"] == "new_version"
    workspace = get_source_workspace(client)
    retained_row = [row for row in workspace["material_catalog"]["published_rows"] if row["customs_code"] == "DEMO-NPL-003"][0]
    edited_row = [row for row in workspace["material_catalog"]["published_rows"] if row["customs_code"] == "DEMO-NPL-001"][0]
    assert retained_row["status"] != "inactive_pending_review"
    assert edited_row["name"] == "Main control board - edited"


def test_bcct_arbitrary_overlap_upload_adds_new_rows_without_deleting_missing_rows():
    client = get_client("do-thanh")
    first_upload = bcct_workbook([
        {
            "coverage_period": "2026-01",
            "direction": "import",
            "declaration_no": "TK-001",
            "line_no": "1",
            "item_code": "MAT-001",
            "description": "Material 1",
            "hs_code": "8542.39",
            "quantity": "100",
            "unit": "PCS",
            "customs_value": "1000",
        }
    ])
    second_upload = bcct_workbook([
        {
            "coverage_period": "2026-02",
            "direction": "import",
            "declaration_no": "TK-002",
            "line_no": "1",
            "item_code": "MAT-002",
            "description": "Material 2",
            "hs_code": "8536.90",
            "quantity": "50",
            "unit": "PCS",
            "customs_value": "500",
        }
    ])

    process_bcct_upload(client, first_upload, "bcct-jan.xlsx")
    result = process_bcct_upload(client, second_upload, "bcct-feb.xlsx")

    assert result["status"] == "new_version"
    assert result["summary"]["added_rows"] == 1
    workspace = get_source_workspace(client)
    assert {row["transaction_key"] for row in workspace["bcct"]["published_rows"]} == {
        "import||TK-001||1||MAT-001",
        "import||TK-002||1||MAT-002",
    }


def test_bcct_reupload_identical_row_with_unit_alias_is_noop():
    client = get_client("do-thanh")
    first_upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"}
    ])
    alias_upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCE"}
    ])

    process_bcct_upload(client, first_upload, "bcct-a.xlsx")
    result = process_bcct_upload(client, alias_upload, "bcct-alias.xlsx")

    assert result["status"] == "no_change"
    assert result["summary"]["unchanged_rows"] == 1
    workspace = get_source_workspace(client)
    assert workspace["bcct"]["published_rows"][0]["unit"] == "PCS"
    assert workspace["bcct"]["correction_candidates"] == []


def test_bcct_changed_unit_same_key_creates_correction_candidate():
    client = get_client("do-thanh")
    first_upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"}
    ])
    changed_upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "KG"}
    ])

    process_bcct_upload(client, first_upload, "bcct-a.xlsx")
    result = process_bcct_upload(client, changed_upload, "bcct-conflict.xlsx")

    assert result["status"] == "review_required"
    assert result["summary"]["correction_candidates"] == 1
    workspace = get_source_workspace(client)
    assert workspace["bcct"]["published_rows"][0]["unit"] == "PCS"
    assert workspace["bcct"]["correction_candidates"][0]["incoming_row"]["unit"] == "KG"


def test_bcct_parser_infers_configured_import_types_when_direction_is_blank():
    rows = parse_bcct_workbook(bcct_workbook([
        {"direction": "", "declaration_type": "E21", "declaration_no": "GC-001", "line_no": "1", "item_code": "MAT-GC-1", "quantity": "10", "unit": "PCS"},
        {"direction": "", "declaration_type": "E23", "declaration_no": "GC-002", "line_no": "1", "item_code": "MAT-GC-2", "quantity": "20", "unit": "PCS"},
        {"direction": "", "declaration_type": "E31", "declaration_no": "SXXK-001", "line_no": "1", "item_code": "MAT-SX-1", "quantity": "30", "unit": "PCS"},
        {"direction": "", "declaration_type": "E62", "declaration_no": "XK-001", "line_no": "1", "item_code": "TP-1", "quantity": "5", "unit": "PCS"},
    ]))

    assert {row["declaration_type"]: row["direction"] for row in rows} == {
        "E21": "import",
        "E23": "import",
        "E31": "import",
        "E62": "export",
    }


def test_growatt_client_config_defaults_to_dncx_and_description_allocation():
    config = get_client_config(get_client("growatt"))

    assert config["bcct"]["eligible_import_declaration_types"] == ["E11", "E15"]
    assert config["bcct"]["relevant_export_declaration_types"] == ["E42"]
    assert config["co_stock"]["lot_policy"] == "line_level"
    assert config["allocation_code"]["strategy"] == "description_regex"
    assert config["allocation_code"]["fallback"] == "same_as_customs_code"
    assert config["config_hash"]


def test_allocation_code_resolver_handles_regex_fallback_and_ambiguity():
    config = get_client_config(get_client("growatt"))

    resolved = resolve_allocation_code(
        {"item_code": "DIENTRO", "description": "DIENTRO#&Điện trở. Hàng mới 100% (001.0001400)"},
        config,
    )
    fallback = resolve_allocation_code(
        {"item_code": "DIENTRO", "description": "DIENTRO#&Điện trở không có mã trong ngoặc"},
        config,
    )
    ambiguous = resolve_allocation_code(
        {"item_code": "DIENTRO", "description": "DIENTRO#&Điện trở (001.0001400) hoặc (001.0001500)"},
        config,
    )

    assert resolved["allocation_code"] == "001.0001400"
    assert resolved["status"] == "resolved"
    assert resolved["source"] == "description_regex"
    assert fallback["allocation_code"] == "DIENTRO"
    assert fallback["source"] == "same_as_customs_code"
    assert ambiguous["allocation_code"] == ""
    assert ambiguous["status"] == "requires_review"
    assert ambiguous["reason"] == "multiple_regex_matches"


def test_client_config_rejects_invalid_regex():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["allocation_code"]["strategy"] = "description_regex"
    config["allocation_code"]["description_regex"] = "("

    with pytest.raises(ValueError, match="Invalid allocation code regex"):
        save_client_config(client, config)


def test_bcct_import_rows_create_immutable_co_stock_source_rows():
    client = get_client("do-thanh")
    upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS", "customs_value": "1000"},
        {"direction": "export", "declaration_no": "XK-001", "line_no": "1", "item_code": "TP-001", "quantity": "10", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    workspace = get_source_workspace(client)
    import_rows = [row for row in workspace["bcct"]["published_rows"] if row["direction"] == "import"]
    assert len(import_rows) == 1
    assert import_rows[0]["import_row_id"].startswith("import-row-")
    assert len(workspace["co_stock_rows"]) == 1
    stock_row = workspace["co_stock_rows"][0]
    assert stock_row["source_row"] == import_rows[0]["import_row_id"]
    assert stock_row["source_line_ids"] == [import_rows[0]["import_row_id"]]
    assert stock_row["import_declaration_no"] == "TK-001"
    assert stock_row["line_no"] == "1"
    assert stock_row["customs_item_code"] == "MAT-001"
    assert stock_row["allocation_code"] == "MAT-001"
    assert stock_row["material_code"] == "MAT-001"
    assert stock_row["allocation_code_source"] == "same_as_customs_code"
    assert stock_row["allocation_code_status"] == "resolved"
    assert stock_row["available_qty"] == "100"
    assert stock_row["used_qty"] == "0"
    assert stock_row["remaining_qty"] == "100"
    assert stock_row["customs_value"] == "1000"
    assert stock_row["unit_value"] == "10"
    assert stock_row["unit_value_source"] == "bcct_customs_value_per_qty"


def test_co_stock_line_level_keeps_duplicate_codes_as_separate_lots():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    config["allocation_code"] = {
        "strategy": "description_regex",
        "description_regex": r"\(([A-Z0-9][A-Z0-9._/-]{3,})\)",
        "fallback": "same_as_customs_code",
    }
    config["co_stock"]["lot_policy"] = "line_level"
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "DIENTRO", "description": "DIENTRO#&Điện trở (001.0001400)", "quantity": "100", "unit": "PCS"},
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "2", "item_code": "DIENTRO", "description": "DIENTRO#&Điện trở (001.0001400)", "quantity": "50", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    stock_rows = get_source_workspace(client)["co_stock_rows"]
    assert len(stock_rows) == 2
    assert [row["line_no"] for row in stock_rows] == ["1", "2"]
    assert {row["customs_item_code"] for row in stock_rows} == {"DIENTRO"}
    assert {row["allocation_code"] for row in stock_rows} == {"001.0001400"}


def test_co_stock_can_aggregate_within_one_declaration_without_losing_source_lines():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    config["allocation_code"] = {
        "strategy": "description_regex",
        "description_regex": r"\(([A-Z0-9][A-Z0-9._/-]{3,})\)",
        "fallback": "same_as_customs_code",
    }
    config["co_stock"]["lot_policy"] = "aggregate_by_declaration_and_allocation_code"
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "DIENTRO", "description": "DIENTRO#&Điện trở (001.0001400)", "quantity": "100", "unit": "PCS"},
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "2", "item_code": "DIENTRO", "description": "DIENTRO#&Điện trở (001.0001400)", "quantity": "50", "unit": "PCS"},
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-002", "line_no": "1", "item_code": "DIENTRO", "description": "DIENTRO#&Điện trở (001.0001400)", "quantity": "25", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    stock_rows = get_source_workspace(client)["co_stock_rows"]
    assert len(stock_rows) == 2
    grouped = [row for row in stock_rows if row["import_declaration_no"] == "TK-001"][0]
    assert grouped["line_no"] == "1,2"
    assert grouped["allocation_code"] == "001.0001400"
    assert grouped["available_qty"] == "150"
    assert len(grouped["source_line_ids"]) == 2


def test_co_stock_aggregation_handles_thousands_separators():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    config["co_stock"]["lot_policy"] = "aggregate_by_declaration_and_allocation_code"
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "1,000", "unit": "PCS"},
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "2", "item_code": "MAT-001", "quantity": "250.5", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    stock_rows = get_source_workspace(client)["co_stock_rows"]
    assert len(stock_rows) == 1
    assert stock_rows[0]["available_qty"] == "1250.5"


def test_client_config_declaration_types_marks_excluded_import_stock_inactive():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11", "E15"]
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"},
        {"direction": "import", "declaration_type": "E13", "declaration_no": "TK-002", "line_no": "1", "item_code": "TOOL-001", "quantity": "1", "unit": "PCS"},
        {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-001", "line_no": "1", "item_code": "TP-001", "quantity": "10", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    stock_rows = get_source_workspace(client)["co_stock_rows"]
    assert [row["customs_item_code"] for row in stock_rows] == ["MAT-001", "TOOL-001"]
    active_row = [row for row in stock_rows if row["customs_item_code"] == "MAT-001"][0]
    inactive_row = [row for row in stock_rows if row["customs_item_code"] == "TOOL-001"][0]
    assert active_row["eligibility_status"] == "active"
    assert active_row["remaining_qty"] == "100"
    assert inactive_row["eligibility_status"] == "inactive"
    assert inactive_row["eligibility_reason"] == "excluded_by_declaration_type_config"
    assert inactive_row["available_qty"] == "1"
    assert inactive_row["remaining_qty"] == "0"


def test_config_change_deactivates_and_reactivates_existing_stock_candidate():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11", "E15"]
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E15", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"},
    ])
    process_bcct_upload(client, upload, "bcct.xlsx")

    initial_stock = get_source_workspace(client)["co_stock_rows"][0]
    source_row = initial_stock["source_row"]
    assert initial_stock["eligibility_status"] == "active"

    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    save_client_config(client, config)
    deactivated_stock = get_source_workspace(client)["co_stock_rows"][0]
    assert deactivated_stock["source_row"] == source_row
    assert deactivated_stock["eligibility_status"] == "inactive"
    assert deactivated_stock["remaining_qty"] == "0"

    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11", "E15"]
    save_client_config(client, config)
    reactivated_stock = get_source_workspace(client)["co_stock_rows"][0]
    assert reactivated_stock["source_row"] == source_row
    assert reactivated_stock["eligibility_status"] == "active"
    assert reactivated_stock["remaining_qty"] == "100"


def test_co_stock_aggregation_keeps_active_and_inactive_candidates_separate():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    config["co_stock"]["lot_policy"] = "aggregate_by_declaration_and_allocation_code"
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"},
        {"direction": "import", "declaration_type": "E13", "declaration_no": "TK-001", "line_no": "2", "item_code": "MAT-001", "quantity": "50", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    stock_rows = get_source_workspace(client)["co_stock_rows"]
    assert len(stock_rows) == 2
    assert {row["eligibility_status"] for row in stock_rows} == {"active", "inactive"}
    assert sorted(row["remaining_qty"] for row in stock_rows) == ["0", "100"]


def test_co_stock_declaration_type_filter_normalizes_uploaded_case():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "", "declaration_type": "e11", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    workspace = get_source_workspace(client)
    assert workspace["bcct"]["published_rows"][0]["declaration_type"] == "E11"
    assert [row["customs_item_code"] for row in workspace["co_stock_rows"]] == ["MAT-001"]
    assert workspace["co_stock_rows"][0]["eligibility_status"] == "active"


def test_unresolved_allocation_code_is_review_only_stock():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    config["allocation_code"] = {
        "strategy": "description_regex",
        "description_regex": r"\(([A-Z0-9][A-Z0-9._/-]{3,})\)",
        "fallback": "requires_review",
    }
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "DIENTRO", "description": "DIENTRO không có mã nội bộ", "quantity": "100", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    stock_row = get_source_workspace(client)["co_stock_rows"][0]
    assert stock_row["allocation_code_status"] == "requires_review"
    assert stock_row["allocation_code"] == ""
    assert stock_row["material_code"] == ""
    response = TestClient(app).get("/clients/do-thanh/co-stock")
    assert response.status_code == 200
    assert "Cần review" in response.text


def test_manual_review_lot_policy_keeps_stock_rows_unusable_until_review():
    client = get_client("do-thanh")
    config = get_client_config(client)
    config["bcct"]["eligible_import_declaration_types"] = ["E11"]
    config["co_stock"]["lot_policy"] = "manual_review"
    save_client_config(client, config)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"},
    ])

    process_bcct_upload(client, upload, "bcct.xlsx")

    stock_row = get_source_workspace(client)["co_stock_rows"][0]
    assert stock_row["allocation_code_status"] == "requires_review"
    assert stock_row["allocation_code_reason"] == "manual_stock_review"
    assert stock_row["material_code"] == ""


def test_export_bcct_view_filters_to_relevant_configured_types():
    client = TestClient(app)
    config = get_client_config(get_client("do-thanh"))
    config["bcct"]["relevant_export_declaration_types"] = ["E42"]
    save_client_config(get_client("do-thanh"), config)
    upload = bcct_workbook([
        {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-001", "line_no": "1", "item_code": "TP-CO", "quantity": "10", "unit": "PCS"},
        {"direction": "export", "declaration_type": "B11", "declaration_no": "XK-002", "line_no": "1", "item_code": "TP-NORMAL", "quantity": "20", "unit": "PCS"},
    ])
    client.post(
        "/clients/do-thanh/bcct/upload",
        files={"file": ("bcct.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    response = client.get("/clients/do-thanh/bcct/exports")

    assert response.status_code == 200
    assert "TP-CO" in response.text
    assert "TP-NORMAL" not in response.text


def test_config_and_bcct_child_routes_render():
    client = TestClient(app)

    config_response = client.get("/clients/growatt/config")
    imports_response = client.get("/clients/growatt/bcct/imports")
    exports_response = client.get("/clients/growatt/bcct/exports")

    assert config_response.status_code == 200
    assert imports_response.status_code == 200
    assert exports_response.status_code == 200
    assert "Cấu hình công ty" in config_response.text
    assert "E11" in config_response.text
    assert "BCCT nhập khẩu" in imports_response.text
    assert "BCCT xuất khẩu" in exports_response.text


def test_co_case_snapshots_client_config_hash():
    client = TestClient(app)
    upload = bcct_workbook([
        {"direction": "import", "declaration_type": "E11", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"}
    ])
    client.post(
        "/clients/do-thanh/bcct/upload",
        files={"file": ("bcct.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    created = client.post(
        "/clients/do-thanh/co-case/create",
        data={"title": "Config snapshot case", "case_code": "CO-CONFIG", "destination_market": "Ấn Độ"},
        follow_redirects=False,
    )
    response = client.get(f"{created.headers['location']}/review")

    assert response.status_code == 200
    assert "Config snapshot" in response.text
    assert get_client_config(get_client("do-thanh"))["config_hash"] in response.text


def test_catalog_and_bcct_routes_offer_templates_and_upload_forms():
    client = TestClient(app)

    material_template = client.get("/clients/growatt/catalog/material-template.xlsx")
    bcct_template = client.get("/clients/growatt/bcct/template.xlsx")
    catalog_page = client.get("/clients/growatt/catalog")
    bcct_page = client.get("/clients/growatt/bcct")

    assert material_template.status_code == 200
    assert material_template.content.startswith(b"PK")
    assert bcct_template.status_code == 200
    assert bcct_template.content.startswith(b"PK")
    assert "Upload danh mục" in catalog_page.text
    assert "Upload BCCT" in bcct_page.text
    assert "append_or_review_by_transaction_key" in bcct_page.text


def test_catalog_upload_renders_selected_child_table():
    client = TestClient(app)
    content = create_material_catalog_template_workbook(get_client("growatt"))

    response = client.post(
        "/clients/do-thanh/catalog/upload",
        data={"catalog_type": "material", "upload_scope": "full_catalog"},
        files={"file": ("ds-nvl.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "DS NVL DK HQ" in response.text
    assert "DEMO-NPL-001" in response.text
    assert "Mã nội bộ" not in response.text


def test_real_customs_material_catalog_xls_preserves_hq_schema_fields():
    rows = parse_catalog_workbook(customs_zip_entry("DANH MUC NPL DK HQ MOI.xls"), "material_catalog")

    assert rows[0]["customs_code"] == "DIOT"
    assert rows[0]["name"] == "Đi ốt"
    assert rows[0]["unit"] == "PCS"
    assert rows[0]["source_schema"] == "customs_material_catalog"
    assert rows[0]["source_row_number"] == 2
    assert rows[0]["raw_fields"]["Mã"] == "DIOT"
    assert "Mã biểu thuế NK" in rows[0]["raw_fields"]
    assert "import_tariff_code" in rows[0]


def test_real_customs_product_catalog_xls_preserves_hq_schema_fields():
    rows = parse_catalog_workbook(customs_zip_entry("DANH MUC SP DK HQ MOI.xls"), "product_catalog")

    assert rows[0]["product_code"] == "BIENTAN.01"
    assert rows[0]["hs_code"] == "85044090"
    assert rows[0]["unit"] == "PCS"
    assert rows[0]["source_schema"] == "customs_product_catalog"
    assert rows[0]["raw_fields"]["Mã định danh của lệnh SX"] == ""
    assert "production_order_identifier" in rows[0]


def test_real_customs_bcct_xlsx_reads_header_row_10_and_preserves_hq_fields():
    rows = parse_bcct_workbook(customs_zip_entry("BaoCaoHangChiTiet 01.01.2025 - 31.12.2025 08.01 or.xlsx"))

    assert len(rows) == 19898
    assert rows[0]["declaration_no"] == "106865355330"
    assert rows[0]["declaration_date"] == "2025-01-07"
    assert rows[0]["declaration_type"] == "E15"
    assert rows[0]["direction"] == "import"
    assert rows[0]["line_no"] == "1"
    assert rows[0]["item_code"] == "LKN-VO"
    assert rows[0]["quantity"] == "2000"
    assert rows[0]["unit"] == "PCS"
    assert rows[0]["origin_country"] == "VIETNAM"
    assert rows[0]["customs_value"] == "1288209000"
    assert rows[0]["invoice_ref"] == "00000020"
    assert rows[0]["raw_fields"]["Số To Khai"] == "106865355330"
    assert rows[0]["source_header_row"] == 10
    assert any(row["direction"] == "export" and row["declaration_type"] == "E42" for row in rows)


def test_bcct_upload_route_surfaces_correction_candidates():
    client = TestClient(app)
    first_upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"}
    ])
    changed_upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "KG"}
    ])
    client.post(
        "/clients/do-thanh/bcct/upload",
        files={"file": ("bcct.xlsx", first_upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    response = client.post(
        "/clients/do-thanh/bcct/upload",
        files={"file": ("bcct-conflict.xlsx", changed_upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 200
    assert "correction_candidate" in response.text
    assert "TK-001" in response.text


def test_co_case_snapshots_reviewed_source_versions_without_correction_candidates():
    client = TestClient(app)
    upload = bcct_workbook([
        {"direction": "import", "declaration_no": "TK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "100", "unit": "PCS"}
    ])
    client.post(
        "/clients/do-thanh/bcct/upload",
        files={"file": ("bcct.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    created = client.post(
        "/clients/do-thanh/co-case/create",
        data={"title": "Source snapshot case", "case_code": "CO-SOURCE", "destination_market": "Ấn Độ"},
        follow_redirects=False,
    )
    response = client.get(f"{created.headers['location']}/review")

    assert response.status_code == 200
    assert "Source evidence snapshot" in response.text
    assert "BCCT reviewed rows: 1" in response.text
    assert "correction_candidate" not in response.text


def test_co_case_can_create_persisted_dossier_and_select_it():
    client = TestClient(app)

    response = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "C/O GROWATT INV-77",
            "case_code": "CO-INV-77",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-77",
            "bill_of_lading_no": "BL-77",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    location = response.headers["location"]
    assert location.startswith("/clients/growatt/co-case/")
    detail = client.get(location)
    assert detail.status_code == 200
    assert "C/O GROWATT INV-77" in detail.text
    assert "INV-77" in detail.text
    assert "BL-77" in detail.text
    index = client.get("/clients/growatt/co-case")
    assert 'aria-label="Danh sách hồ sơ C/O"' in index.text
    assert "Danh sách hồ sơ C/O" in index.text
    assert "Tên hồ sơ" in index.text
    assert "Invoice" in index.text
    assert "B/L" in index.text
    assert "Mở hồ sơ" in index.text
    assert "C/O GROWATT INV-77" in index.text
    assert "CO-INV-77" in index.text
    assert "INV-77" in index.text
    assert "BL-77" in index.text
    assert "Hồ sơ lưu local: CO-INV-77" not in index.text


def test_co_case_auto_generates_case_code_and_step_status_labels():
    client = TestClient(app)

    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Auto code", "destination_market": "Canada", "invoice_no": "INV-AUTO"},
        follow_redirects=False,
    )

    assert created.status_code == 303
    detail = client.get(created.headers["location"])
    index = client.get("/clients/growatt/co-case")
    assert "CO-GROWATT-INV-AUTO-" in detail.text
    assert "CO-GROWATT-INV-AUTO-" in index.text
    assert "Chưa nhập · Auto code" not in index.text
    assert "Đủ" in detail.text
    assert "Cần soát" in detail.text


def test_co_case_detail_is_split_into_workflow_step_views():
    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "Workflow dossier",
            "case_code": "CO-WORKFLOW",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-WORKFLOW",
            "bill_of_lading_no": "BL-WORKFLOW",
        },
        follow_redirects=False,
    )
    case_url = created.headers["location"]

    shipment = client.get(case_url)
    documents = client.get(f"{case_url}/documents")
    exports = client.get(f"{case_url}/exports")
    guidance = client.get(f"{case_url}/guidance")
    origin = client.get(f"{case_url}/origin")
    review = client.get(f"{case_url}/review")

    assert shipment.status_code == 200
    assert documents.status_code == 200
    assert exports.status_code == 200
    assert guidance.status_code == 200
    assert origin.status_code == 200
    assert review.status_code == 200
    assert "Thông tin lô hàng" in shipment.text
    assert "Supporting files" not in shipment.text
    assert "Supporting files" in documents.text
    assert "BCCT xuất khẩu theo invoice" in exports.text
    assert "Form và thông tư" in guidance.text
    assert "Bảng kê LVC" in origin.text
    assert "Tính lại snapshot" in origin.text
    assert 'class="table-input"' not in origin.text
    assert "Upload và parse" not in origin.text
    assert "Tải seed XLSX" not in origin.text
    assert "Xuất evidence XLSX" not in origin.text
    assert "Xuất dossier XLSX" in review.text
    assert f"{case_url}/documents" in shipment.text
    assert f"{case_url}/origin" in shipment.text


def test_co_case_shipment_step_updates_metadata_without_dropping_origin_view():
    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "Shipment edit",
            "case_code": "CO-SHIP",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-OLD",
            "bill_of_lading_no": "BL-OLD",
        },
        follow_redirects=False,
    )
    case_url = created.headers["location"]

    response = client.post(
        f"{case_url}/shipment",
        data={
            "title": "Shipment edit",
            "case_code": "CO-SHIP-NEW",
            "destination_market": "Canada",
            "invoice_no": "INV-NEW",
            "bill_of_lading_no": "BL-NEW",
            "agreement": "CPTPP",
            "co_form_type": "Form CPTPP",
            "rule": "Cần tra cứu PSR theo HS",
        },
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"] == case_url
    shipment = client.get(case_url)
    origin = client.get(f"{case_url}/origin")
    assert "CO-SHIP-NEW" in shipment.text
    assert "INV-NEW" in shipment.text
    assert "BL-NEW" in shipment.text
    assert "PV00.0048500" in origin.text
    assert "Bảng kê LVC" in origin.text


def test_co_case_origin_preloads_demo_when_case_has_no_invoice_source_data():
    client = TestClient(app)
    created = client.post(
        "/clients/do-thanh/co-case/create",
        data={"title": "Origin demo", "case_code": "CO-DEMO", "destination_market": "Canada"},
        follow_redirects=False,
    )
    assert created.status_code == 303

    origin = client.get(f"{created.headers['location']}/origin")
    review = client.get(f"{created.headers['location']}/review")

    assert origin.status_code == 200
    assert "Demo tự nạp" in origin.text
    assert "2 TP mẫu" in origin.text
    assert "4 NVL mẫu" in origin.text
    assert "PV00.0048500" in origin.text
    assert "Upload và parse" not in origin.text
    assert "Demo tự nạp" not in review.text
    assert "2 TP mẫu" not in review.text


def test_co_case_origin_builds_and_persists_invoice_bom_snapshot():
    client = TestClient(app)
    client.post(
        "/clients/growatt/bcct/upload",
        files={
            "file": (
                "bcct.xlsx",
                bcct_workbook([
                    {
                        "direction": "import",
                        "declaration_type": "E11",
                        "declaration_no": "NK-BOM-1",
                        "line_no": "1",
                        "item_code": "DEMO-NPL-001",
                        "description": "Main control board",
                        "hs_code": "8542.39",
                        "quantity": "100",
                        "unit": "PCE",
                        "customs_value": "1000",
                        "currency": "VND",
                    },
                    {
                        "direction": "import",
                        "declaration_type": "E11",
                        "declaration_no": "NK-BOM-2",
                        "line_no": "1",
                        "item_code": "DEMO-NPL-002",
                        "description": "Connector set",
                        "hs_code": "8536.90",
                        "quantity": "100",
                        "unit": "PCE",
                        "customs_value": "2000",
                        "currency": "VND",
                    },
                    {
                        "direction": "export",
                        "declaration_type": "E42",
                        "declaration_no": "XK-BOM",
                        "line_no": "1",
                        "item_code": "PV00.0048500",
                        "description": "Growatt inverter",
                        "hs_code": "850440",
                        "quantity": "3",
                        "unit": "PCS",
                        "customs_value": "1000",
                        "currency": "VND",
                        "invoice_ref": "INV-BOM",
                    }
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Invoice BOM", "case_code": "CO-BOM-INV", "destination_market": "Ấn Độ", "invoice_no": "INV-BOM"},
        follow_redirects=False,
    )
    location = created.headers["location"]
    case_id = location.rstrip("/").split("/")[-1]
    update_case_record(
        get_client("growatt"),
        {
            "persisted_case_id": case_id,
            "shipment": {"invoice_no": "INV-BOM", "bill_of_lading_no": ""},
            "products": [
                {
                    "code": "STALE-TP",
                    "materials": [{"material_code": "STALE-MAT", "unit_value": ""}],
                }
            ],
            "origin_snapshot": {"source": "invoice_bcct_bom", "invoice_no": "INV-BOM"},
        },
    )

    origin = client.get(f"{location}/origin")

    assert origin.status_code == 200
    assert "Demo tự nạp" not in origin.text
    assert "STALE-MAT" not in origin.text
    assert "PV00.0048500" in origin.text
    assert "DEMO-NPL-001" in origin.text
    assert 'name="product_0_material_0_material_code" value="DEMO-NPL-001"' in origin.text
    form_data = hidden_form_data(origin.text)
    assert form_data["product_0_fob"] == "1000"
    assert form_data["product_0_currency"] == "VND"
    assert form_data["product_0_material_0_consumed_qty"] == "3"
    assert form_data["product_0_material_0_unit_value"] == "10"
    assert form_data["product_0_material_0_currency"] == "VND"
    assert form_data["product_0_material_0_material_value"] == "30"
    assert form_data["product_0_material_1_material_value"] == "105"
    assert form_data["product_0_vnm_value"] == "135"
    assert form_data["product_0_lvc_percentage"] == "86.50"
    assert "1,000" in origin.text
    assert "VND" in origin.text
    assert 'name="product_0_bom_product_version_id"' in origin.text
    assert "86.50%" in origin.text
    assert "Đạt LVC" in origin.text
    assert "<th>Tờ khai nhập</th>" not in origin.text
    assert "<th>Tồn CO</th>" not in origin.text
    assert "<th>Còn lại</th>" not in origin.text
    assert "<th>Định mức</th>" in origin.text
    assert "<th>Lượng dùng</th>" in origin.text
    assert "<th>Đơn giá</th>" in origin.text
    assert "<th>Trị giá NVL</th>" in origin.text
    assert "<th>Trị giá KXX/VNM</th>" in origin.text

    recalculated = client.post("/clients/growatt/evaluate", data=form_data)
    persisted = client.get(f"{location}/origin")

    assert recalculated.status_code == 200
    assert "Đã tính lại theo dữ liệu đang sửa." in recalculated.text
    assert "DEMO-NPL-001" in persisted.text
    assert "TP BOM v1" in persisted.text


def test_co_case_origin_switches_product_bom_version_from_dropdown():
    client = TestClient(app)
    template = client.get("/clients/growatt/bom/template.xlsx")
    workbook = load_workbook(BytesIO(template.content))
    workbook["BOM"]["F2"] = "2.00"
    stream = BytesIO()
    workbook.save(stream)
    client.post(
        "/clients/growatt/bom/upload",
        data={"upload_mode": "direct_bom"},
        files={"file": ("growatt-bom-v2.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    versions = get_bom_workspace(get_client("growatt"))["product_version_options_by_code"]["PV00.0048500"]
    v1 = [version for version in versions if version["product_version_no"] == 1][0]
    v2 = [version for version in versions if version["product_version_no"] == 2][0]
    client.post(
        "/clients/growatt/bcct/upload",
        files={
            "file": (
                "bcct.xlsx",
                bcct_workbook([
                    {"direction": "import", "declaration_type": "E11", "declaration_no": "NK-BOM-SWITCH-1", "line_no": "1", "item_code": "DEMO-NPL-001", "description": "Main control board", "hs_code": "8542.39", "quantity": "100", "unit": "PCE", "customs_value": "1000", "currency": "VND"},
                    {"direction": "import", "declaration_type": "E11", "declaration_no": "NK-BOM-SWITCH-2", "line_no": "1", "item_code": "DEMO-NPL-002", "description": "Connector set", "hs_code": "8536.90", "quantity": "100", "unit": "PCE", "customs_value": "2000", "currency": "VND"},
                    {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-BOM-SWITCH", "line_no": "1", "item_code": "PV00.0048500", "description": "Growatt inverter", "hs_code": "850440", "quantity": "3", "unit": "PCS", "customs_value": "1000", "currency": "VND", "invoice_ref": "INV-BOM-SWITCH"},
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "BOM switch", "case_code": "CO-BOM-SWITCH", "destination_market": "Ấn Độ", "invoice_no": "INV-BOM-SWITCH"},
        follow_redirects=False,
    )

    origin = client.get(f"{created.headers['location']}/origin")

    assert origin.status_code == 200
    assert f'value="{v2["product_version_id"]}" selected' in origin.text
    assert "v1 · 2 dòng" in origin.text
    assert "v2 · 2 dòng" in origin.text

    form_data = hidden_form_data(origin.text)
    form_data["product_0_bom_product_version_id"] = v1["product_version_id"]
    switched = client.post("/clients/growatt/evaluate", data=form_data)
    switched_data = hidden_form_data(switched.text)

    assert switched.status_code == 200
    assert f'value="{v1["product_version_id"]}" selected' in switched.text
    assert switched_data["product_0_material_0_consumed_qty"] == "3"
    assert switched_data["product_0_material_0_material_value"] == "30"


def test_co_case_origin_page_surfaces_method_readiness_and_evidence_gaps():
    client = TestClient(app)
    client.post(
        "/clients/growatt/bcct/upload",
        files={
            "file": (
                "bcct.xlsx",
                bcct_workbook([
                    {
                        "direction": "import",
                        "declaration_type": "E11",
                        "declaration_no": "NK-ORIGIN-READY",
                        "line_no": "1",
                        "item_code": "DEMO-NPL-001",
                        "description": "Main control board",
                        "hs_code": "8542.39",
                        "quantity": "100",
                        "unit": "PCE",
                        "customs_value": "1000",
                        "currency": "VND",
                    },
                    {
                        "direction": "export",
                        "declaration_type": "E42",
                        "declaration_no": "XK-ORIGIN-READY",
                        "line_no": "1",
                        "item_code": "PV00.0048500",
                        "description": "Growatt inverter",
                        "hs_code": "850440",
                        "quantity": "3",
                        "unit": "PCS",
                        "customs_value": "1000",
                        "currency": "VND",
                        "invoice_ref": "INV-ORIGIN-READY",
                    },
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    created = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "Origin readiness",
            "case_code": "CO-ORIGIN-READY",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-ORIGIN-READY",
        },
        follow_redirects=False,
    )

    origin = client.get(f"{created.headers['location']}/origin")

    assert origin.status_code == 200
    assert "Bảng tính Xuất xứ" in origin.text
    assert "Build-down LVC/RVC" in origin.text
    assert "(FOB - VNM) / FOB x 100" in origin.text
    assert "Cần bổ sung evidence" in origin.text
    assert "Thiếu đơn giá NVL" in origin.text
    assert "DEMO-NPL-002: thiếu đơn giá để tính trị giá NVL/VNM." in origin.text
    assert "CTSH preview" in origin.text
    assert "chưa thay thế PSR engine/legal review" in origin.text
    assert "Nguồn giá" in origin.text
    assert "Trạng thái dữ liệu" in origin.text

    form_data = hidden_form_data(origin.text)
    assert form_data["product_0_origin_method"] == "build_down_lvc"
    assert form_data["product_0_origin_readiness_status"] == "blocked"
    assert form_data["product_0_material_1_valuation_status"] == "missing_unit_value"


def test_co_case_export_workbook_contains_bom_snapshot_rows_from_origin_form():
    client = TestClient(app)
    client.post(
        "/clients/growatt/bcct/upload",
        files={
            "file": (
                "bcct.xlsx",
                bcct_workbook([
                    {
                        "direction": "import",
                        "declaration_type": "E11",
                        "declaration_no": "NK-BOM-XLSX-1",
                        "line_no": "1",
                        "item_code": "DEMO-NPL-001",
                        "description": "Main control board",
                        "hs_code": "8542.39",
                        "quantity": "100",
                        "unit": "PCE",
                        "customs_value": "1000",
                    },
                    {
                        "direction": "import",
                        "declaration_type": "E11",
                        "declaration_no": "NK-BOM-XLSX-2",
                        "line_no": "1",
                        "item_code": "DEMO-NPL-002",
                        "description": "Connector set",
                        "hs_code": "8536.90",
                        "quantity": "100",
                        "unit": "PCE",
                        "customs_value": "2000",
                    },
                    {
                        "direction": "export",
                        "declaration_type": "E42",
                        "declaration_no": "XK-BOM-XLSX",
                        "line_no": "1",
                        "item_code": "PV00.0048500",
                        "description": "Growatt inverter",
                        "hs_code": "850440",
                        "quantity": "2",
                        "unit": "PCS",
                        "customs_value": "1000",
                        "invoice_ref": "INV-BOM-XLSX",
                    }
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "BOM export", "case_code": "CO-BOM-XLSX", "destination_market": "Ấn Độ", "invoice_no": "INV-BOM-XLSX"},
        follow_redirects=False,
    )
    origin = client.get(f"{created.headers['location']}/origin")

    response = client.post(f"{created.headers['location']}/export", data=hidden_form_data(origin.text))

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert "LVC Statement" in workbook.sheetnames
    values = [cell.value for row in workbook["LVC Statement"].iter_rows(values_only=False) for cell in row]
    assert "PV00.0048500" in values
    assert "DEMO-NPL-001" in values
    assert "91.00" in values
    assert "90" in values


def test_co_case_export_workbook_contains_origin_snapshot_metadata_from_web():
    client = TestClient(app)
    client.post(
        "/clients/growatt/bcct/upload",
        files={
            "file": (
                "bcct.xlsx",
                bcct_workbook([
                    {"direction": "import", "declaration_type": "E11", "declaration_no": "NK-ORIGIN-XLSX-1", "line_no": "1", "item_code": "DEMO-NPL-001", "description": "Main control board", "hs_code": "8542.39", "quantity": "100", "unit": "PCE", "customs_value": "1000", "currency": "VND"},
                    {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-ORIGIN-XLSX", "line_no": "1", "item_code": "PV00.0048500", "description": "Growatt inverter", "hs_code": "850440", "quantity": "3", "unit": "PCS", "customs_value": "1000", "currency": "VND", "invoice_ref": "INV-ORIGIN-XLSX"},
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Origin export", "case_code": "CO-ORIGIN-XLSX", "destination_market": "Ấn Độ", "invoice_no": "INV-ORIGIN-XLSX"},
        follow_redirects=False,
    )
    origin = client.get(f"{created.headers['location']}/origin")

    response = client.post(f"{created.headers['location']}/export", data=hidden_form_data(origin.text))

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert "Origin Snapshot" in workbook.sheetnames
    values = [cell.value for row in workbook["Origin Snapshot"].iter_rows(values_only=False) for cell in row]
    assert "build_down_lvc" in values
    assert "blocked" in values
    assert "DEMO-NPL-002: thiếu đơn giá để tính trị giá NVL/VNM." in values
    assert "CTSH preview" in values


def test_co_case_supporting_upload_saves_invoice_metadata_and_matches_bcct_exports():
    client = TestClient(app)
    upload = bcct_workbook([
        {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-001", "line_no": "1", "item_code": "TP-001", "description": "Finished good", "hs_code": "850440", "quantity": "10", "unit": "PCS", "invoice_ref": "INV-42"},
        {"direction": "import", "declaration_type": "E11", "declaration_no": "NK-001", "line_no": "1", "item_code": "MAT-001", "description": "Material", "hs_code": "853690", "quantity": "100", "unit": "PCS", "invoice_ref": "INV-42"},
        {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-002", "line_no": "1", "item_code": "TP-OTHER", "description": "Other finished good", "hs_code": "850440", "quantity": "5", "unit": "PCS", "invoice_ref": "INV-99"},
    ])
    client.post(
        "/clients/growatt/bcct/upload",
        files={"file": ("bcct.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Invoice lookup", "case_code": "CO-INV-42", "destination_market": "Ấn Độ"},
        follow_redirects=False,
    )
    location = created.headers["location"]

    response = client.post(
        f"{location}/supporting-files",
        data={"document_slot": "invoice", "invoice_no": "INV-42", "bill_of_lading_no": "BL-42"},
        files={"file": ("invoice-INV-42.pdf", b"%PDF-1.4 invoice", "application/pdf")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"] == f"{location}/documents"
    documents = client.get(f"{location}/documents")
    exports = client.get(f"{location}/exports")
    assert "invoice-INV-42.pdf" in documents.text
    download_href = re.search(r'href="([^"]+/supporting-files/[^"]+)"', documents.text)
    assert download_href is not None
    downloaded = client.get(download_href.group(1))
    assert downloaded.status_code == 200
    assert downloaded.content == b"%PDF-1.4 invoice"
    assert "INV-42" in documents.text
    assert "BL-42" in documents.text
    assert "TP-001" in exports.text
    assert "XK-001" in exports.text
    assert "MAT-001" not in exports.text
    assert "TP-OTHER" not in exports.text


def test_co_case_guidance_maps_invoice_bcct_products_to_form_instrument_and_hs_criteria():
    client = TestClient(app)
    upload = bcct_workbook([
        {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-PSR", "line_no": "1", "item_code": "PV00.0048500", "description": "Growatt inverter", "hs_code": "850440", "quantity": "12", "unit": "PCS", "invoice_ref": "INV-PSR"},
    ])
    client.post(
        "/clients/growatt/bcct/upload",
        files={"file": ("bcct.xlsx", upload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Criteria lookup", "case_code": "CO-PSR", "destination_market": "Canada", "invoice_no": "INV-PSR"},
        follow_redirects=False,
    )
    location = created.headers["location"]

    index = client.get("/clients/growatt/co-case")
    shipment = client.get(location)
    guidance = client.get(f"{location}/guidance")

    assert 'role="combobox"' in index.text
    assert "Form CPTPP" in shipment.text
    assert "03/2019/TT-BCT" in shipment.text
    assert "PV00.0048500" in shipment.text
    assert "850440" in shipment.text
    assert "CTH; hoặc RVC không thấp hơn" in guidance.text
    assert "03/2019/TT-BCT, Phụ lục I" in guidance.text


def test_co_case_state_prefers_postgres_store_and_keeps_supporting_file_metadata(monkeypatch):
    from app import co_case_store

    saved_states = []

    class FakeCoCaseStateStore:
        def __init__(self):
            self.state = None

        def get_state(self, client_id: str) -> dict | None:
            assert client_id == "growatt"
            return self.state

        def save_state(self, client_id: str, state: dict) -> None:
            assert client_id == "growatt"
            self.state = dict(state)
            saved_states.append(dict(state))

    fake_store = FakeCoCaseStateStore()
    monkeypatch.setattr(co_case_store, "get_co_case_state_store", lambda: fake_store)
    client = get_client("growatt")
    record = co_case_store.create_case_record(
        client,
        {"title": "Postgres C/O", "case_code": "CO-PG", "destination_market": "Ấn Độ"},
    )

    file_row = co_case_store.save_supporting_file(
        client,
        record["case_id"],
        b"%PDF-1.4 invoice",
        "invoice-CO-PG.pdf",
        "invoice",
        "INV-PG",
        "BL-PG",
    )

    assert saved_states
    assert fake_store.state["cases"][0]["case_id"] == record["case_id"]
    assert file_row["original_filename"] == "invoice-CO-PG.pdf"
    assert file_row["stored_filename"].startswith(file_row["upload_id"])
    assert file_row["storage_backend"] == "filesystem"
    assert file_row["content_sha256"] == hashlib.sha256(b"%PDF-1.4 invoice").hexdigest()
    assert fake_store.state["cases"][0]["supporting_files"][0]["content_sha256"] == file_row["content_sha256"]
    assert not co_case_store.state_path("growatt").exists()


def test_co_case_evaluate_keeps_persisted_dossier_supporting_metadata():
    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Evaluate dossier", "case_code": "CO-EVAL", "destination_market": "Ấn Độ", "invoice_no": "INV-OLD"},
        follow_redirects=False,
    )
    location = created.headers["location"]
    case_id = location.rsplit("/", 1)[-1]
    client.post(
        f"{location}/supporting-files",
        data={"document_slot": "invoice", "invoice_no": "INV-OLD", "bill_of_lading_no": "BL-OLD"},
        files={"file": ("invoice-INV-OLD.pdf", b"%PDF-1.4 invoice", "application/pdf")},
        follow_redirects=False,
    )

    response = client.post(
        "/clients/growatt/evaluate",
        data={
            "case_id": case_id,
            "persisted_case_id": case_id,
            "customer": "Growatt",
            "case_code": "CO-EVAL",
            "title": "Evaluate dossier",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-NEW",
            "bill_of_lading_no": "BL-NEW",
            "agreement": "AIFTA",
            "co_form_type": "Form AI",
            "rule": "Cần tra cứu PSR theo HS",
            "document_count": "0",
            "product_count": "0",
        },
    )

    assert response.status_code == 200
    assert "INV-NEW" in response.text
    assert "BL-NEW" in response.text
    detail = client.get(f"{location}/documents")
    assert "invoice-INV-OLD.pdf" in detail.text
    assert "INV-NEW" in detail.text
    assert "BL-NEW" in detail.text


def test_invoice_matching_uses_only_reviewed_export_rows():
    matches = match_case_bcct_exports(
        {"shipment": {"invoice_no": "INV-REVIEW"}},
        {
            "bcct": {
                "published_rows": [
                    {"direction": "export", "review_status": "correction_candidate", "declaration_no": "XK-DRAFT", "line_no": "1", "declaration_type": "E42", "item_code": "TP-DRAFT", "quantity": "1", "unit": "PCS", "invoice_ref": "INV-REVIEW"},
                    {"direction": "export", "review_status": "reviewed", "declaration_no": "XK-OK", "line_no": "1", "declaration_type": "E42", "item_code": "TP-OK", "quantity": "1", "unit": "PCS", "invoice_ref": "INV-REVIEW"},
                    {"direction": "import", "review_status": "reviewed", "declaration_no": "NK-OK", "line_no": "1", "declaration_type": "E11", "item_code": "MAT-OK", "quantity": "1", "unit": "PCS", "invoice_ref": "INV-REVIEW"},
                ]
            }
        },
        {"bcct": {"relevant_export_declaration_types": ["E42"]}},
    )

    assert [row["item_code"] for row in matches] == ["TP-OK"]


def test_co_case_page_uses_lightweight_source_summary(monkeypatch):
    from app import portfolio as portfolio_module
    from app import source_store as source_store_module

    def fail_full_workspace_load(*_args, **_kwargs):
        raise AssertionError("C/O pages should not load the full source workspace.")

    monkeypatch.setattr(portfolio_module, "get_source_workspace", fail_full_workspace_load)
    monkeypatch.setattr(source_store_module, "get_source_workspace", fail_full_workspace_load)

    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "Indexed source case",
            "case_code": "CO-INDEX",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-INDEX",
            "bill_of_lading_no": "BL-INDEX",
        },
        follow_redirects=False,
    )

    response = client.get(created.headers["location"])

    assert response.status_code == 200
    assert "CO-INDEX" in response.text


def test_source_summary_exposes_counts_and_snapshot_metadata():
    client = get_client("do-thanh")
    process_bcct_upload(
        client,
        bcct_workbook([
            {"direction": "import", "declaration_no": "NK-001", "line_no": "1", "item_code": "MAT-001", "quantity": "10", "unit": "PCS"},
            {"direction": "export", "declaration_type": "E42", "declaration_no": "XK-001", "line_no": "1", "item_code": "TP-001", "quantity": "2", "unit": "PCS", "invoice_ref": "INV-001"},
        ]),
        "bcct.xlsx",
    )

    summary = get_source_summary(client)

    assert summary["bcct"]["published_row_count"] == 2
    assert summary["bcct"]["reviewed_row_count"] == 2
    assert summary["co_stock_row_count"] == 1
    assert summary["bcct"]["latest_version"]["version_no"] == 1


def test_co_case_page_uses_postgres_source_index_when_available(monkeypatch):
    from app import portfolio as portfolio_module

    class FakeSourceIndexStore:
        def has_client(self, client_id: str) -> bool:
            return client_id == "growatt"

        def source_summary(self, _client_id: str, client_config: dict) -> dict:
            return {
                "client_config": client_config,
                "material_catalog": {"published_row_count": 10, "latest_version": {"version_no": 2}},
                "product_catalog": {"published_row_count": 3, "latest_version": {"version_no": 4}},
                "bcct": {
                    "published_row_count": 20,
                    "reviewed_row_count": 9,
                    "correction_candidate_count": 0,
                    "latest_version": {"version_no": 5},
                },
                "co_stock_row_count": 8,
            }

        def match_bcct_exports(self, _client_id: str, invoice_no: str, _relevant_types: list[str]) -> list[dict]:
            assert invoice_no == "INV-PG"
            return [
                {
                    "declaration_no": "XK-PG",
                    "line_no": "1",
                    "declaration_type": "E42",
                    "item_code": "TP-PG",
                    "hs_code": "8504.40",
                    "quantity": "2",
                    "unit": "PCS",
                    "invoice_ref": "INV-PG",
                }
            ]

    def fail_file_source_load(*_args, **_kwargs):
        raise AssertionError("Postgres-indexed C/O pages should not load source JSON.")

    monkeypatch.setattr(portfolio_module, "get_source_index_store", lambda: FakeSourceIndexStore())
    monkeypatch.setattr(portfolio_module, "load_module_state", fail_file_source_load)

    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "Postgres indexed case",
            "case_code": "CO-PG",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-PG",
        },
        follow_redirects=False,
    )

    response = client.get(f"{created.headers['location']}/exports")

    assert response.status_code == 200
    assert "XK-PG" in response.text
    assert "TP-PG" in response.text
    assert "20 dòng BCCT" in response.text


def test_postgres_catalog_index_records_keep_payload_and_keys():
    rows = [
        {
            "customs_code": "MAT-PG-001",
            "name": "Postgres material",
            "hs_code": "8504.40",
            "unit": "PCS",
            "status": "active",
        }
    ]

    records = build_catalog_index_records("growatt", "material_catalog", rows)

    assert records == [
        {
            "client_id": "growatt",
            "module": "material_catalog",
            "row_key": "MAT-PG-001",
            "customs_code": "MAT-PG-001",
            "product_code": "",
            "hs_code": "8504.40",
            "unit": "PCS",
            "status": "active",
            "payload": rows[0],
        }
    ]


def test_source_tables_use_postgres_workspace_when_available(monkeypatch):
    from app import portfolio as portfolio_module

    class FakeSourceIndexStore:
        def has_client(self, client_id: str) -> bool:
            return client_id == "growatt"

        def source_workspace(self, _client_id: str, client_config: dict) -> dict:
            return {
                "client_config": client_config,
                "material_catalog": {
                    "module": "material_catalog",
                    "published_rows": [
                        {
                            "customs_code": "MAT-PG-001",
                            "name": "Postgres material",
                            "hs_code": "8504.40",
                            "unit": "PCS",
                            "status": "active",
                        }
                    ],
                    "latest_version": {"version_no": 7},
                    "versions": [],
                    "uploads": [],
                    "correction_candidates": [],
                    "audit_events": [],
                },
                "product_catalog": {
                    "module": "product_catalog",
                    "published_rows": [
                        {
                            "product_code": "TP-PG-001",
                            "name": "Postgres product",
                            "hs_code": "8504.40",
                            "unit": "PCS",
                            "status": "active",
                        }
                    ],
                    "latest_version": {"version_no": 8},
                    "versions": [],
                    "uploads": [],
                    "correction_candidates": [],
                    "audit_events": [],
                },
                "bcct": {
                    "module": "bcct",
                    "published_rows": [
                        {
                            "transaction_key": "export||XK-PG||1||TP-PG-001",
                            "direction": "export",
                            "review_status": "reviewed",
                            "declaration_no": "XK-PG",
                            "line_no": "1",
                            "declaration_type": "E42",
                            "item_code": "TP-PG-001",
                            "hs_code": "8504.40",
                            "quantity": "2",
                            "unit": "PCS",
                            "invoice_ref": "INV-PG",
                        }
                    ],
                    "latest_version": {"version_no": 9},
                    "versions": [],
                    "uploads": [],
                    "correction_candidates": [],
                    "audit_events": [],
                },
                "co_stock_rows": [
                    {
                        "source_row": "PG-STOCK-001",
                        "import_declaration_no": "NK-PG",
                        "line_no": "1",
                        "declaration_type": "E11",
                        "customs_item_code": "MAT-PG-001",
                        "allocation_code": "MAT-PG-001",
                        "allocation_code_status": "resolved",
                        "eligibility_status": "eligible",
                        "remaining_qty": "5",
                        "unit": "PCS",
                    }
                ],
            }

    monkeypatch.setattr(portfolio_module, "get_source_index_store", lambda: FakeSourceIndexStore())

    client = TestClient(app)

    materials = client.get("/clients/growatt/catalog/materials")
    bcct = client.get("/clients/growatt/bcct")
    stock = client.get("/clients/growatt/co-stock")

    assert materials.status_code == 200
    assert "MAT-PG-001" in materials.text
    assert "MAT-001" not in materials.text
    assert "XK-PG" in bcct.text
    assert "TP-PG-001" in bcct.text
    assert "NK-PG" in stock.text


def test_postgres_source_index_initializes_empty_client():
    class FakePostgresSourceIndexStore(PostgresSourceIndexStore):
        def __init__(self):
            super().__init__("postgresql:///unused")
            self.exists = False
            self.replaced = None

        def ensure_schema(self) -> None:
            return None

        def _client_index_exists(self, client_id: str) -> bool:
            assert client_id == "new-client"
            return self.exists

        def replace_client_indexes(
            self,
            client_id: str,
            catalog_records: list[dict],
            bcct_records: list[dict],
            invoice_records: list[dict],
            stock_records: list[dict],
            correction_records: list[dict],
            metadata_records: list[dict],
            source_state_records: list[dict] | None = None,
        ) -> None:
            self.exists = True
            self.replaced = {
                "client_id": client_id,
                "catalog_records": catalog_records,
                "bcct_records": bcct_records,
                "invoice_records": invoice_records,
                "stock_records": stock_records,
                "correction_records": correction_records,
                "metadata_records": metadata_records,
                "source_state_records": source_state_records or [],
            }

    store = FakePostgresSourceIndexStore()

    assert store.has_client("new-client") is True
    assert store.replaced["client_id"] == "new-client"
    assert {row["module"] for row in store.replaced["metadata_records"]} == {
        "material_catalog",
        "product_catalog",
        "bcct",
        "co_stock",
    }
    assert [records["module_state"]["module"] for records in store.replaced["source_state_records"]] == [
        "material_catalog",
        "product_catalog",
        "bcct",
    ]
    assert all(records["module_state"]["published_row_count"] == 0 for records in store.replaced["source_state_records"])


def test_portfolio_app_exposes_source_dashboard_and_summary_api():
    client = TestClient(app)

    dashboard = client.get("/portfolio")
    summary = client.get("/portfolio/api/clients/growatt/source-summary")

    assert dashboard.status_code == 200
    assert "Source Portfolio" in dashboard.text
    assert "Growatt" in dashboard.text
    assert summary.status_code == 200
    payload = summary.json()
    assert payload["client"]["id"] == "growatt"
    assert payload["source_backend"] in {"files", "postgres"}
    assert payload["source_summary"]["material_catalog"]["published_row_count"] >= 3
    assert payload["source_summary"]["product_catalog"]["published_row_count"] >= 2


def test_portfolio_service_prefers_postgres_clients_and_config(monkeypatch):
    from app import portfolio as portfolio_module

    class FakeAppStateStore:
        def has_clients(self) -> bool:
            return True

        def clients(self) -> list[dict]:
            return [
                {
                    "id": "pg-client",
                    "name": "Postgres Client",
                    "code": "PG",
                    "status": "active",
                    "tax_code": "010-PG",
                    "contact": "db",
                    "module_status": {},
                    "material_catalog": [],
                    "product_catalog": [],
                    "bom_rows": [],
                    "co_stock": [],
                    "bcct_rows": [],
                    "counts": {"materials": 0, "products": 0, "bom_lines": 0, "co_stock": 0, "bcct": 0},
                }
            ]

        def client(self, client_id: str) -> dict:
            assert client_id == "pg-client"
            return self.clients()[0]

        def get_client_config(self, client: dict) -> dict:
            return {
                "schema_version": 1,
                "client_id": client["id"],
                "config_version": 3,
                "config_hash": "pg-hash",
                "bcct": {"eligible_import_declaration_types": ["E11"], "relevant_export_declaration_types": ["E42"]},
                "co_stock": {"lot_policy": "line_level"},
                "allocation_code": {"strategy": "same_as_customs_code", "description_regex": "", "fallback": "same_as_customs_code"},
            }

    monkeypatch.setattr(portfolio_module, "get_app_state_store", lambda: FakeAppStateStore(), raising=False)

    service = portfolio_module.PortfolioService()

    assert service.client("pg-client")["name"] == "Postgres Client"
    assert service.clients()[0]["id"] == "pg-client"
    assert service.get_client_config({"id": "pg-client"})["config_hash"] == "pg-hash"


def test_portfolio_service_saves_client_config_to_postgres(monkeypatch):
    from app import portfolio as portfolio_module

    saved_configs = []

    class FakeAppStateStore:
        def save_client_config(self, client: dict, config: dict) -> dict:
            saved_configs.append((client["id"], config["config_hash"]))
            return {**config, "config_version": 9, "config_hash": "saved-pg-hash"}

    monkeypatch.setattr(portfolio_module, "get_app_state_store", lambda: FakeAppStateStore(), raising=False)

    service = portfolio_module.PortfolioService()
    saved = service.save_client_config({"id": "growatt"}, {"config_hash": "draft"})

    assert saved["config_version"] == 9
    assert saved["config_hash"] == "saved-pg-hash"
    assert saved_configs == [("growatt", "draft")]


def test_portfolio_service_uses_postgres_source_writer_for_uploads(monkeypatch):
    from app import portfolio as portfolio_module

    calls = []

    class FakeSourceWriteStore:
        def has_client(self, client_id: str) -> bool:
            return client_id == "growatt"

        def process_catalog_upload(self, client: dict, catalog_type: str, content: bytes, filename: str, upload_scope: str, client_config: dict) -> dict:
            calls.append(("catalog", client["id"], catalog_type, filename, upload_scope, client_config["config_hash"]))
            return {"status": "postgres_catalog", "upload": {"upload_id": "pg-catalog-upload"}}

        def process_bcct_upload(self, client: dict, content: bytes, filename: str, client_config: dict) -> dict:
            calls.append(("bcct", client["id"], filename, client_config["config_hash"]))
            return {"status": "postgres_bcct", "upload": {"upload_id": "pg-bcct-upload"}}

    def fail_file_catalog(*_args, **_kwargs):
        raise AssertionError("Postgres source uploads should not use JSON catalog writer.")

    def fail_file_bcct(*_args, **_kwargs):
        raise AssertionError("Postgres source uploads should not use JSON BCCT writer.")

    monkeypatch.setattr(portfolio_module, "get_source_write_store", lambda: FakeSourceWriteStore(), raising=False)
    monkeypatch.setattr(portfolio_module, "process_catalog_upload", fail_file_catalog)
    monkeypatch.setattr(portfolio_module, "process_bcct_upload", fail_file_bcct)

    service = portfolio_module.PortfolioService()
    monkeypatch.setattr(service, "get_client_config", lambda client: {"config_hash": "cfg-pg"})

    catalog = service.process_catalog_upload({"id": "growatt"}, "material", b"catalog", "catalog.xlsx", "full_catalog")
    bcct = service.process_bcct_upload({"id": "growatt"}, b"bcct", "bcct.xlsx")

    assert catalog["status"] == "postgres_catalog"
    assert bcct["status"] == "postgres_bcct"
    assert calls == [
        ("catalog", "growatt", "material", "catalog.xlsx", "full_catalog", "cfg-pg"),
        ("bcct", "growatt", "bcct.xlsx", "cfg-pg"),
    ]


def test_co_routes_use_portfolio_service_adapter(monkeypatch):
    from app import main as main_module

    class FakePortfolioService:
        def source_workspace(self, client: dict) -> tuple[dict, str]:
            return {
                "client_config": {"config_version": 1, "config_hash": "fake", "bcct": {"eligible_import_declaration_types": [], "relevant_export_declaration_types": []}},
                "material_catalog": {
                    "module": "material_catalog",
                    "published_rows": [
                        {
                            "customs_code": "PF-MAT-001",
                            "name": "Portfolio material",
                            "hs_code": "8504.40",
                            "unit": "PCS",
                            "status": "active",
                        }
                    ],
                    "latest_version": {"version_no": 1},
                    "versions": [],
                    "uploads": [],
                    "correction_candidates": [],
                    "audit_events": [],
                },
                "product_catalog": {
                    "module": "product_catalog",
                    "published_rows": [],
                    "latest_version": {},
                    "versions": [],
                    "uploads": [],
                    "correction_candidates": [],
                    "audit_events": [],
                },
                "bcct": {
                    "module": "bcct",
                    "published_rows": [],
                    "latest_version": {},
                    "versions": [],
                    "uploads": [],
                    "correction_candidates": [],
                    "audit_events": [],
                },
                "co_stock_rows": [],
            }, "portfolio-fake"

        def co_case_source_context(self, client: dict, case: dict) -> dict:
            return {
                "source_backend": "portfolio-fake",
                "source_summary": {
                    "client_config": {"config_version": 1, "config_hash": "fake"},
                    "material_catalog": {"published_row_count": 1, "latest_version": {"version_no": 1}},
                    "product_catalog": {"published_row_count": 0, "latest_version": {}},
                    "bcct": {
                        "published_row_count": 1,
                        "reviewed_row_count": 1,
                        "correction_candidate_count": 0,
                        "latest_version": {"version_no": 1},
                    },
                    "co_stock_row_count": 0,
                },
                "invoice_matches": [
                    {
                        "declaration_no": "PF-XK-001",
                        "line_no": "1",
                        "declaration_type": "E42",
                        "item_code": "PF-TP-001",
                        "hs_code": "8504.40",
                        "quantity": "1",
                        "unit": "PCS",
                        "invoice_ref": case.get("shipment", {}).get("invoice_no", ""),
                    }
                ],
            }

    monkeypatch.setattr(main_module, "portfolio_service", FakePortfolioService())

    client = TestClient(app)
    materials = client.get("/clients/growatt/catalog/materials")
    created = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "Portfolio case",
            "case_code": "CO-PF",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-PF",
        },
        follow_redirects=False,
    )
    exports = client.get(f"{created.headers['location']}/exports")

    assert materials.status_code == 200
    assert "PF-MAT-001" in materials.text
    assert exports.status_code == 200
    assert "PF-XK-001" in exports.text
    assert "PF-TP-001" in exports.text


def test_postgres_bcct_index_records_tokenize_invoice_refs():
    rows = [
        {
            "transaction_key": "export||XK-001||1||TP-001",
            "direction": "export",
            "review_status": "reviewed",
            "declaration_no": "XK-001",
            "line_no": "1",
            "declaration_type": "E42",
            "item_code": "TP-001",
            "hs_code": "8504.40",
            "quantity": "2",
            "unit": "PCS",
            "invoice_ref": "INV-001 / INV 002",
        }
    ]

    bcct_records, invoice_records = build_bcct_index_records("growatt", rows)

    assert bcct_records[0]["transaction_key"] == "export||XK-001||1||TP-001"
    assert bcct_records[0]["payload"]["item_code"] == "TP-001"
    assert {(row["invoice_key"], row["transaction_key"]) for row in invoice_records} == {
        ("INV001INV002", "export||XK-001||1||TP-001"),
        ("INV001", "export||XK-001||1||TP-001"),
        ("INV", "export||XK-001||1||TP-001"),
        ("002", "export||XK-001||1||TP-001"),
    }


def test_co_case_supporting_upload_rejects_unsupported_or_oversized_files():
    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Upload validation", "case_code": "CO-UP", "destination_market": "Ấn Độ"},
        follow_redirects=False,
    )
    location = created.headers["location"]

    unsupported = client.post(
        f"{location}/supporting-files",
        data={"document_slot": "invoice"},
        files={"file": ("invoice.exe", b"bad", "application/octet-stream")},
        follow_redirects=False,
    )
    oversized = client.post(
        f"{location}/supporting-files",
        data={"document_slot": "invoice"},
        files={"file": ("invoice.pdf", b"x" * (MAX_SUPPORTING_FILE_BYTES + 1), "application/pdf")},
        follow_redirects=False,
    )

    assert unsupported.status_code == 400
    assert "Không hỗ trợ định dạng file" in unsupported.text
    assert oversized.status_code == 400
    assert "File supporting vượt quá giới hạn" in oversized.text


def test_co_case_destination_market_shows_verified_form_candidates():
    client = TestClient(app)

    india = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "India shipment", "case_code": "CO-IN", "destination_market": "Ấn Độ"},
        follow_redirects=False,
    )
    france = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "France shipment", "case_code": "CO-FR", "destination_market": "Pháp"},
        follow_redirects=False,
    )
    canada = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Canada shipment", "case_code": "CO-CA", "destination_market": "Canada"},
        follow_redirects=False,
    )

    india_page = client.get(f"{india.headers['location']}/guidance")
    france_page = client.get(f"{france.headers['location']}/guidance")
    canada_page = client.get(f"{canada.headers['location']}/guidance")

    assert "Form AI" in india_page.text
    assert "15/2010/TT-BCT" in india_page.text
    assert "Form EUR.1" in france_page.text
    assert "11/2020/TT-BCT" in france_page.text
    assert "Form CPTPP" in canada_page.text
    assert "03/2019/TT-BCT" in canada_page.text
    assert "Cần tra cứu PSR theo HS" in canada_page.text


def test_co_form_index_defaults_cover_initial_priority_forms():
    from app.co_forms import criteria_preview_for_hs, form_candidates_for_market, hs_scope_is_ex, prioritized_form_lanes, recommended_form_lane
    from app.co_form_config_store import default_co_form_config

    config = default_co_form_config()
    assert form_candidates_for_market("United States")[0]["form_code"] == "B"
    assert recommended_form_lane(prioritized_form_lanes("India"))["form_code"] == "AI"
    assert recommended_form_lane(prioritized_form_lanes("Canada"))["form_code"] == "CPTPP"
    assert recommended_form_lane(prioritized_form_lanes("Pháp"))["form_code"] == "EUR.1"
    assert len(config["psr_rules"]) > 6000
    assert all("Chương" not in rule["hs_scope"] for rule in config["psr_rules"])
    assert any(rule["hs_scope"] == "01" for rule in config["psr_rules"])
    assert criteria_preview_for_hs("AI", "850440")["criteria"] == "AIFTA 35% FOB + CTSH"
    assert "Product Specific Rules" in criteria_preview_for_hs("AI", "850440")["note"]
    assert criteria_preview_for_hs("CPTPP", "850440")["criteria"].startswith("CTH; hoặc RVC không thấp hơn")
    assert criteria_preview_for_hs("B", "999999")["criteria"] == "Tra PSR Form B theo Phụ lục I"
    assert hs_scope_is_ex("ex 0307")


def test_co_form_ex_hs_scope_requires_product_description_confirmation():
    from app.co_forms import criteria_preview_for_hs

    preview = criteria_preview_for_hs("EUR.1", "030600")

    assert preview["status"] == "requires_manual_lookup"
    assert preview["criteria"].startswith("Cần đối chiếu mô tả hàng hóa trước khi áp dụng")
    assert "Match by HS code alone is not enough" in preview["note"]


def test_co_form_settings_page_saves_market_alias_config():
    from app.co_form_config_store import default_co_form_config
    from app.co_forms import form_candidates_for_market

    config = default_co_form_config()
    data = co_form_settings_payload(config)
    new_index = len(config["market_presets"])
    data[f"market_{new_index}_enabled"] = "1"
    data[f"market_{new_index}_show_in_picker"] = "1"
    data[f"market_{new_index}_market"] = "Bharat"
    data[f"market_{new_index}_label"] = "Bharat / Form AI"
    data[f"market_{new_index}_form_code"] = "AI"
    data[f"market_{new_index}_aliases"] = "Bharat"
    data[f"market_{new_index}_selection_reason"] = "Test alias maps to Form AI."
    data[f"market_{new_index}_source_label"] = "test"

    client = TestClient(app)
    settings = client.get("/settings")
    saved = client.post("/settings/co-forms", data=data, follow_redirects=False)
    page = client.get("/settings/co-forms?saved=1")

    assert 'href="/settings/co-forms"' in settings.text
    assert saved.status_code == 303
    assert saved.headers["location"] == "/settings/co-forms?saved=1"
    assert "Đã lưu cấu hình form" in page.text
    assert form_candidates_for_market("Bharat")[0]["form_code"] == "AI"


def test_co_form_settings_page_saves_psr_rule_config():
    from app.co_form_config_store import default_co_form_config
    from app.co_forms import criteria_preview_for_hs

    config = default_co_form_config()
    data = co_form_settings_payload(config)
    new_index = len(config["psr_rules"])
    data[f"psr_{new_index}_enabled"] = "1"
    data[f"psr_{new_index}_form_code"] = "AI"
    data[f"psr_{new_index}_hs_scope"] = "850760"
    data[f"psr_{new_index}_criteria"] = "AIFTA custom battery rule"
    data[f"psr_{new_index}_source_reference"] = "Test source"
    data[f"psr_{new_index}_note"] = "Test note"
    data[f"psr_{new_index}_status"] = "confirmed_by_trong_tin"

    client = TestClient(app)
    saved = client.post("/settings/co-forms", data=data, follow_redirects=False)
    page = client.get("/settings/co-forms?saved=1")
    preview = criteria_preview_for_hs("AI", "85076039")

    assert saved.status_code == 303
    assert "HS Criteria" in page.text
    assert preview["criteria"] == "AIFTA custom battery rule"
    assert preview["status"] == "confirmed_by_trong_tin"


def test_co_form_settings_page_shows_readable_status_labels():
    client = TestClient(app)
    page = client.get("/settings/co-forms?tab=psr&psr_form=AI")
    forms_page = client.get("/settings/co-forms?tab=forms")

    assert page.status_code == 200
    assert forms_page.status_code == 200
    assert "Chờ Trọng Tín xác nhận" in page.text
    assert "Chờ Trọng Tín xác nhận" in forms_page.text
    assert 'placeholder="pending_trong_tin_confirmation"' not in page.text
    assert 'name="psr_0_status"' in page.text
    assert '<select name="psr_0_status">' in page.text


def test_co_form_settings_page_updates_filtered_psr_rule_without_reposting_all_rules():
    from app.co_form_config_store import default_co_form_config
    from app.co_forms import criteria_preview_for_hs

    config = default_co_form_config()
    original_index, rule = next(
        (index, row)
        for index, row in enumerate(config["psr_rules"])
        if row["form_code"] == "CPTPP" and row["hs_scope"] == "85.04"
    )
    data = {
        "active_tab": "psr",
        "psr_visible_count": "1",
        "psr_0_original_index": str(original_index),
        "psr_0_enabled": "1",
        "psr_0_form_code": rule["form_code"],
        "psr_0_hs_scope": rule["hs_scope"],
        "psr_0_criteria": "CPTPP updated filtered 8504 rule",
        "psr_0_source_reference": rule["source_reference"],
        "psr_0_note": rule["note"],
        "psr_0_status": "confirmed_by_trong_tin",
    }

    client = TestClient(app)
    filtered_page = client.get("/settings/co-forms?tab=psr&psr_form=CPTPP&psr_query=8504")
    saved = client.post("/settings/co-forms", data=data, follow_redirects=False)

    assert "Đang hiện" in filtered_page.text
    assert 'name="psr_0_original_index"' in filtered_page.text
    assert saved.status_code == 303
    assert saved.headers["location"] == "/settings/co-forms?saved=1&tab=psr"
    assert criteria_preview_for_hs("CPTPP", "850440")["criteria"] == "CPTPP updated filtered 8504 rule"
    assert criteria_preview_for_hs("B", "850440")["criteria"] == "LVC 30% hoặc CTH"


def test_co_case_create_explains_invoice_market_hint_without_auto_selecting(monkeypatch):
    from app import main as main_module

    class FakePortfolioService:
        def client(self, client_id: str) -> dict:
            return {"id": client_id, "name": "Growatt VN", "code": client_id, "counts": {}}

        def co_case_source_context(self, client: dict, case: dict) -> dict:
            return {
                "source_backend": "data-hub",
                "source_summary": {
                    "client_config": {"config_version": 1, "config_hash": "fake"},
                    "material_catalog": {"published_row_count": 0, "latest_version": {}},
                    "product_catalog": {"published_row_count": 0, "latest_version": {}},
                    "bcct": {"published_row_count": 1, "reviewed_row_count": 1, "latest_version": {}},
                    "co_stock_row_count": 0,
                },
                "invoice_matches": [
                    {
                        "declaration_no": "XK1",
                        "line_no": "1",
                        "item_code": "TP-US",
                        "invoice_ref": case.get("shipment", {}).get("invoice_no", ""),
                        "market_hint": {
                            "country_code": "US",
                            "country_name": "United States",
                            "source_field": "unloading_location",
                            "source_value": "USLAX - LOS ANGELES - CA",
                            "confidence": "high",
                        },
                    }
                ],
            }

    monkeypatch.setattr(main_module, "portfolio_service", FakePortfolioService())
    client = TestClient(app)

    created = client.post(
        "/clients/growatt-vn/co-case/create",
        data={"title": "Invoice hinted", "case_code": "CO-HINT", "invoice_no": "GUS28826A131-3F"},
        follow_redirects=False,
    )
    page = client.get(created.headers["location"])
    preview = client.get(
        "/clients/growatt-vn/co-case/invoice-preview",
        params={"invoice_no": "GUS28826A131-3F"},
    ).json()

    assert created.status_code == 303
    assert "Thị trường Chưa nhập" in page.text
    assert "Gợi ý thị trường" in page.text
    assert "United States" in page.text
    assert "Form B" in page.text
    assert 'data-market-value="United States"' in page.text
    assert preview["market_inference"]["destination_market"] == "United States"
    assert "unloading_location" in preview["market_inference"]["explanation"]
    assert preview["suggested_forms"][0]["form_code"] == "B"


def test_invoice_market_hint_requires_single_high_confidence_country():
    from app.co_market_hints import infer_market_from_invoice_matches

    conflict = infer_market_from_invoice_matches([
        {"market_hint": {"country_code": "US", "country_name": "United States", "confidence": "high"}},
        {"market_hint": {"country_code": "IN", "country_name": "India", "confidence": "high"}},
    ])
    low_confidence = infer_market_from_invoice_matches([
        {"market_hint": {"country_code": "US", "country_name": "United States", "confidence": "low"}}
    ])

    assert conflict["status"] == "conflict"
    assert conflict["destination_market"] == ""
    assert low_confidence["status"] == "missing"


def test_co_case_export_workbook_contains_dossier_sheets_and_criteria_rows():
    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={
            "title": "Export dossier",
            "case_code": "CO-XLSX",
            "destination_market": "Ấn Độ",
            "invoice_no": "INV-XLSX",
        },
        follow_redirects=False,
    )

    response = client.post(f"{created.headers['location']}/export")

    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    workbook = load_workbook(BytesIO(response.content))
    assert set(["Case", "Supporting Files", "BCCT Invoice Matches", "Form Guidance", "Criteria"]).issubset(workbook.sheetnames)
    assert workbook["Case"]["B2"].value == "CO-XLSX"
    assert workbook["Form Guidance"]["A2"].value == "Form AI"
    criteria_values = [cell.value for row in workbook["Criteria"].iter_rows(values_only=False) for cell in row]
    assert "PV00.0048500" in criteria_values


def test_co_case_export_filename_is_sanitized():
    client = TestClient(app)
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "Unsafe filename", "case_code": "CO/../../bad", "destination_market": "Ấn Độ"},
        follow_redirects=False,
    )

    response = client.post(f"{created.headers['location']}/export")

    assert response.status_code == 200
    assert 'filename="bad-dossier.xlsx"' in response.headers["content-disposition"]
