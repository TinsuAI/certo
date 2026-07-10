"""VN-origin ticket #6: plumb lot origin fields into sheet materials at Tính.

Lots (co_stock rows) already carry origin_country from BCCT but drop consignee_name;
allocation lines and sheet materials carried neither, so bảng kê column (9) rendered
blank in prod. After this change every allocation line copies origin_country +
consignee_name + supplier_key from its lot, and the material folds the distinct
per-line values (same unique_texts pattern as import_declaration_no) so a line whose
lots disagree on country surfaces ALL countries — never silently one.
"""
from __future__ import annotations

from decimal import Decimal

from app.co_stock_derivation import aggregate_co_stock_rows, co_stock_rows_from_bcct
from app.demo_data import update_products_from_form
from app.web.co_case_context import origin_material_from_bom_row, stock_allocation_line


CLIENT_CONFIG = {
    "co_stock": {"lot_policy": "line_level"},
    "bcct": {"eligible_import_declaration_types": []},
    "allocation_code": {"strategy": "same_as_customs_code", "fallback": "same_as_customs_code"},
}


def _bcct_row(**extra):
    base = {
        "direction": "import",
        "transaction_key": "tk-1",
        "declaration_no": "NK-1",
        "declaration_type": "E15",
        "line_no": "1",
        "item_code": "NVL-1",
        "description": "Vòng đệm",
        "hs_code": "73182200",
        "quantity": "100",
        "unit": "PCE",
        "customs_value": "1000",
        "currency": "VND",
        "origin_country": "VIETNAM",
    }
    base.update(extra)
    return base


def _stock(**extra):
    base = {
        "source_row": "ROW-1",
        "material_code": "NVL-1",
        "allocation_code": "NVL-1",
        "customs_item_code": "NVL-1",
        "import_declaration_no": "D100",
        "registration_date": "2026-04-21",
        "line_no": "1",
        "value_currency": "VND",
        "currency": "VND",
        "unit_value": "1000",
        "remaining_qty": "100",
        "available_qty": "100",
        "exchange_rate_to_vnd": "1",
        "exchange_rate_source": "vnd_native",
    }
    base.update(extra)
    return base


# --- lot derivation: consignee_name must survive BCCT → co_stock row ---

def test_co_stock_row_carries_consignee_name():
    rows = co_stock_rows_from_bcct([_bcct_row(consignee_name="CONG TY TNHH MINGJIE VIET NAM")], CLIENT_CONFIG)
    assert rows[0]["consignee_name"] == "CONG TY TNHH MINGJIE VIET NAM"
    assert rows[0]["origin_country"] == "VIETNAM"


def test_co_stock_row_falls_back_to_partner_name():
    # File-mode BCCT workbooks map "Tên đối tác" → partner_name, not consignee_name.
    rows = co_stock_rows_from_bcct([_bcct_row(partner_name="NCC NOI DIA")], CLIENT_CONFIG)
    assert rows[0]["consignee_name"] == "NCC NOI DIA"


def test_co_stock_row_consignee_defaults_empty():
    rows = co_stock_rows_from_bcct([_bcct_row()], CLIENT_CONFIG)
    assert rows[0]["consignee_name"] == ""


def test_aggregate_lots_do_not_merge_across_consignees():
    rows = co_stock_rows_from_bcct(
        [
            _bcct_row(consignee_name="NCC A", line_no="1", transaction_key="tk-1"),
            _bcct_row(consignee_name="NCC B", line_no="2", transaction_key="tk-2"),
        ],
        CLIENT_CONFIG,
    )
    merged = aggregate_co_stock_rows(rows)
    assert len(merged) == 2
    assert {row["consignee_name"] for row in merged} == {"NCC A", "NCC B"}


def test_aggregate_lots_still_merge_same_consignee():
    rows = co_stock_rows_from_bcct(
        [
            _bcct_row(consignee_name="NCC A", line_no="1", transaction_key="tk-1", quantity="10"),
            _bcct_row(consignee_name="NCC A", line_no="2", transaction_key="tk-2", quantity="5"),
        ],
        CLIENT_CONFIG,
    )
    merged = aggregate_co_stock_rows(rows)
    assert len(merged) == 1
    assert merged[0]["available_qty"] == "15"


# --- allocation line: copies the lot's origin fields ---

def test_allocation_line_carries_lot_origin_and_supplier():
    line = stock_allocation_line(
        _stock(origin_country="VIETNAM", consignee_name="CONG TY TNHH  MINGJIE VIET NAM"),
        allocated_qty=Decimal("1"),
        available_qty=Decimal("10"),
        bom_row={},
        material={},
    )
    assert line["origin_country"] == "VIETNAM"
    assert line["consignee_name"] == "CONG TY TNHH  MINGJIE VIET NAM"  # raw, unmodified
    assert line["supplier_key"] == "CONG TY TNHH MINGJIE VIET NAM"  # normalized


def test_allocation_line_defaults_empty_for_legacy_lots():
    # Lots materialized before this feature have neither field — no KeyError, no None.
    line = stock_allocation_line(
        _stock(),
        allocated_qty=Decimal("1"),
        available_qty=Decimal("10"),
        bom_row={},
        material={},
    )
    assert line["origin_country"] == ""
    assert line["consignee_name"] == ""
    assert line["supplier_key"] == ""


# --- material: folds distinct per-line values up from allocation lines ---

def test_material_folds_lot_origin_up_from_allocation_lines():
    row = {"material_code": "NVL-1", "qty_per": "1"}
    stock_pool = {"NVL-1": [_stock(origin_country="VIETNAM", consignee_name="NCC A")]}
    material = origin_material_from_bom_row(
        row, Decimal("5"), {}, stock_pool, product_code="TP1", product_name="SP",
    )
    assert material["origin_country"] == "VIETNAM"
    assert material["consignee_name"] == "NCC A"
    assert material["supplier_key"] == "NCC A"


def test_material_surfaces_all_distinct_countries():
    # One BOM line drawn from a VN lot AND a CN lot: column (9) must surface BOTH
    # (interim single-row rendering until the split machinery of tickets #7/#9).
    row = {"material_code": "NVL-1", "qty_per": "1"}
    stock_pool = {
        "NVL-1": [
            _stock(origin_country="VIETNAM", consignee_name="NCC A", remaining_qty="2", available_qty="2"),
            _stock(source_row="ROW-2", line_no="2", origin_country="CHINA", consignee_name="NCC B", remaining_qty="100", available_qty="100"),
        ]
    }
    material = origin_material_from_bom_row(
        row, Decimal("5"), {}, stock_pool, product_code="TP1", product_name="SP",
    )
    assert material["origin_country"] == "VIETNAM, CHINA"
    assert material["consignee_name"] == "NCC A, NCC B"
    assert material["supplier_key"] == "NCC A, NCC B"
    assert [line["origin_country"] for line in material["allocation_lines"]] == ["VIETNAM", "CHINA"]


def test_material_supplier_key_dedups_by_normalized_name():
    # Two spellings of ONE company (whitespace-only difference) → one supplier_key;
    # raw consignee_name keeps both spellings.
    row = {"material_code": "NVL-1", "qty_per": "1"}
    stock_pool = {
        "NVL-1": [
            _stock(origin_country="VIETNAM", consignee_name="MYS GROUP(VIET NAM)", remaining_qty="2", available_qty="2"),
            _stock(source_row="ROW-2", line_no="2", origin_country="VIETNAM", consignee_name="MYS GROUP (VIET NAM)", remaining_qty="100", available_qty="100"),
        ]
    }
    material = origin_material_from_bom_row(
        row, Decimal("5"), {}, stock_pool, product_code="TP1", product_name="SP",
    )
    assert material["origin_country"] == "VIETNAM"
    assert material["supplier_key"] == "MYS GROUP (VIET NAM)"
    assert material["consignee_name"] == "MYS GROUP(VIET NAM), MYS GROUP (VIET NAM)"


def test_material_without_lots_has_empty_origin_fields():
    row = {"material_code": "GHOST", "qty_per": "1"}
    material = origin_material_from_bom_row(
        row, Decimal("1"), {}, {}, product_code="TP1", product_name="SP",
    )
    assert material["origin_country"] == ""
    assert material["consignee_name"] == ""
    assert material["supplier_key"] == ""


def test_structure_only_material_has_empty_origin_fields():
    row = {"material_code": "NVL-1", "qty_per": "1"}
    material = origin_material_from_bom_row(
        row, Decimal("1"), {}, {}, product_code="TP1", product_name="SP", allocate=False,
    )
    assert material["origin_country"] == ""
    assert material["consignee_name"] == ""
    assert material["supplier_key"] == ""


# --- form round-trip: export rebuilds the case from the posted form ---

def test_form_roundtrips_material_and_allocation_origin_fields():
    form = {
        "product_count": "1",
        "product_0_code": "TP1",
        "product_0_material_count": "1",
        "product_0_material_0_material_code": "NVL-1",
        "product_0_material_0_origin_country": "VIETNAM, CHINA",
        "product_0_material_0_consignee_name": "NCC A, NCC B",
        "product_0_material_0_supplier_key": "NCC A, NCC B",
        "product_0_material_0_allocation_line_count": "1",
        "product_0_material_0_allocation_0_import_declaration_no": "NK-1",
        "product_0_material_0_allocation_0_origin_country": "VIETNAM",
        "product_0_material_0_allocation_0_consignee_name": "NCC A",
        "product_0_material_0_allocation_0_supplier_key": "NCC A",
    }
    case = update_products_from_form(form)
    material = case["products"][0]["materials"][0]
    assert material["origin_country"] == "VIETNAM, CHINA"
    assert material["consignee_name"] == "NCC A, NCC B"
    assert material["supplier_key"] == "NCC A, NCC B"
    line = material["allocation_lines"][0]
    assert line["origin_country"] == "VIETNAM"
    assert line["consignee_name"] == "NCC A"
    assert line["supplier_key"] == "NCC A"


# --- column (9): a filled material now lands in the exported country cell ---

def _country_material(**extra):
    base = {
        "material_code": "NVL-1", "material_description": "Vòng đệm",
        "origin_status": "non_origin", "hs_code": "73182200", "uom": "PCE",
        "bom_qty_per": "1", "consumed_qty": "1", "unit_value": "1", "material_value": "1",
        "origin_country": "VIETNAM",
    }
    base.update(extra)
    return base


def test_config_renderer_writes_origin_country_into_country_column():
    from openpyxl import Workbook
    from app.bang_ke_renderer import load_form_config, render_into_sheet

    cfg = load_form_config("LVC")
    wb = Workbook()
    ws = wb.active
    product = {"code": "P1", "materials": [_country_material()]}
    render_into_sheet(ws, cfg, case={"case_code": "C", "products": [product]}, product=product, sheet_title="P1")
    col, start = cfg["body"]["columns"]["country"], cfg["body"]["start_row"]
    assert ws[f"{col}{start}"].value == "VIETNAM"


def test_xml_generator_material_row_carries_origin_country():
    from app.bang_ke_xml_generator import _build_material_row

    values, _origin, _non_origin = _build_material_row(_country_material(), {}, {"code": "P1"}, 1)
    assert values["country"] == "VIETNAM"


def test_legacy_hq_sheet_writes_origin_country():
    from openpyxl import Workbook
    from app.workbook_io import _hq_layout_for, write_hq_sheet_materials

    col = _hq_layout_for("LVC")["cols"].get("country")
    assert col, "LVC layout must map a country column"
    wb = Workbook()
    ws = wb.active
    write_hq_sheet_materials(ws, {"materials": [_country_material()]}, 5, sheet_code="LVC")
    assert ws.cell(row=5, column=col).value == "VIETNAM"
