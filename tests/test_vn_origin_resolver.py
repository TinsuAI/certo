"""Per-row VN-origin resolver (VN-origin ticket #12 — the feature).

A line is originating iff its lot's origin_country normalizes to VN AND its
supplier holds a current evidence flag. The resolver yields an originating
AMOUNT per row (never a boolean on the material), qualifying value leaves VNM
so RVC/LVC rises, column (12) composes "Phụ lục X/<NCC>", and a zero-flag
client (Johnson) is byte-identical to pre-feature output.
"""
from __future__ import annotations

from decimal import Decimal

from app.web.co_case_context import (
    enrich_origin_product,
    materialize_bang_ke_origin_fields,
    origin_material_from_bom_row,
    origin_product_from_invoice_match,
)

MINGJIE = "CONG TY TNHH MINGJIE VIET NAM"
FLAGS = {MINGJIE: {"supplier_key": MINGJIE, "supplier_name": MINGJIE, "action": "on", "evidence_kind": "phu_luc_x"}}


def _stock(**extra):
    base = {
        "source_row": "ROW-1",
        "material_code": "NVL-1",
        "allocation_code": "NVL-1",
        "customs_item_code": "NVL-1",
        "import_declaration_no": "D100",
        "registration_date": "2026-04-21",
        "line_no": "1",
        "value_currency": "VND",
        "currency": "VND",
        "unit_value": "10",
        "remaining_qty": "100",
        "available_qty": "100",
        "exchange_rate_to_vnd": "1",
        "exchange_rate_source": "vnd_native",
        "consignee_name": MINGJIE,
        "origin_country": "VIETNAM",
    }
    base.update(extra)
    return base


def _material(pool, consumed="4", flags=FLAGS):
    return origin_material_from_bom_row(
        {"material_code": "NVL-1", "qty_per": "1"}, Decimal(consumed), {}, pool,
        product_code="TP1", product_name="SP", supplier_flags=flags,
    )


# --- the AND rule ---

def test_vn_lot_of_flagged_supplier_is_originating():
    material = _material({"NVL-1": [_stock()]})
    line = material["allocation_lines"][0]
    assert line["origin_status"] == "origin"
    assert line["bang_ke_co_doc_no"] == f"Phụ lục X/{MINGJIE}"
    assert material["origin_amount"] == "40"  # 4 × 10
    assert material["non_origin_cif_value"] == "0"


def test_flagged_supplier_non_vn_lot_stays_non_originating():
    # 24 of 64 VN-capable suppliers sell mixed-origin goods — the flag alone
    # must never promote a Chinese-origin lot.
    material = _material({"NVL-1": [_stock(origin_country="CHINA")]})
    assert material["allocation_lines"][0]["origin_status"] == "non_origin"
    assert material["origin_amount"] == ""
    assert material["non_origin_cif_value"] == "40"


def test_vn_lot_of_unflagged_supplier_stays_non_originating():
    material = _material({"NVL-1": [_stock(consignee_name="NCC KHONG PHU LUC X")]})
    assert material["allocation_lines"][0]["origin_status"] == "non_origin"
    assert material["origin_amount"] == ""
    assert material["non_origin_cif_value"] == "40"


def test_unknown_origin_lot_never_qualifies():
    material = _material({"NVL-1": [_stock(origin_country="UNKNOWN")]})
    assert material["allocation_lines"][0]["origin_status"] == "non_origin"


def test_supplier_key_normalization_matches_flag():
    # the lot's consignee has a double space; the flag key is normalized — the
    # ONE shared normalization must bridge them.
    material = _material({"NVL-1": [_stock(consignee_name="CONG TY TNHH  MINGJIE VIET NAM")]})
    assert material["allocation_lines"][0]["origin_status"] == "origin"


# --- amount, not boolean: mixed line credits exactly its VN portion ---

def test_mixed_line_yields_partial_originating_amount():
    pool = {"NVL-1": [
        _stock(remaining_qty="2", available_qty="2"),
        _stock(source_row="ROW-2", line_no="2", origin_country="CHINA",
               consignee_name="NCC CN", remaining_qty="100", available_qty="100"),
    ]}
    material = _material(pool)
    assert material["origin_amount"] == "20"  # 2 × 10 VN portion only
    assert material["non_origin_cif_value"] == "20"  # 40 total − 20 originating
    statuses = [line["origin_status"] for line in material["allocation_lines"]]
    assert statuses == ["origin", "non_origin"]
    # (13) blank + non-blocking warning
    assert any("(13)" in warning for warning in material["material_warnings"])


def test_rvc_rises_by_exactly_the_qualifying_amount():
    pool = {"NVL-1": [
        _stock(remaining_qty="2", available_qty="2"),
        _stock(source_row="ROW-2", line_no="2", origin_country="CHINA",
               consignee_name="NCC CN", remaining_qty="100", available_qty="100"),
    ]}
    match = {"item_code": "TP1", "description": "SP", "quantity": "4",
             "fob_value": "100", "fob_currency": "VND", "currency": "VND", "hs_code": "850440"}
    with_flags = origin_product_from_invoice_match(
        match, [{"material_code": "NVL-1", "qty_per": "1"}], {}, {}, pool, supplier_flags=FLAGS,
    )
    pool2 = {"NVL-1": [
        _stock(remaining_qty="2", available_qty="2"),
        _stock(source_row="ROW-2", line_no="2", origin_country="CHINA",
               consignee_name="NCC CN", remaining_qty="100", available_qty="100"),
    ]}
    without_flags = origin_product_from_invoice_match(
        match, [{"material_code": "NVL-1", "qty_per": "1"}], {}, {}, pool2,
    )
    assert without_flags["non_origin_value"] == "40"
    assert with_flags["non_origin_value"] == "20"  # VNM shrank by the VN portion
    lvc_with = enrich_origin_product({**with_flags, "fob": "100"})["lvc_percentage"]
    lvc_without = enrich_origin_product({**without_flags, "fob": "100"})["lvc_percentage"]
    assert float(lvc_with) > float(lvc_without)
    assert lvc_with == "80.00"  # (100−20)/100
    assert lvc_without == "60.00"  # (100−40)/100


# --- zero-flag client byte-identical ---

def test_zero_flags_yield_byte_identical_material():
    pool_a = {"NVL-1": [_stock()]}
    pool_b = {"NVL-1": [_stock()]}
    without_param = origin_material_from_bom_row(
        {"material_code": "NVL-1", "qty_per": "1"}, Decimal("4"), {}, pool_a,
        product_code="TP1", product_name="SP",
    )
    with_empty_flags = origin_material_from_bom_row(
        {"material_code": "NVL-1", "qty_per": "1"}, Decimal("4"), {}, pool_b,
        product_code="TP1", product_name="SP", supplier_flags={},
    )
    assert without_param == with_empty_flags
    assert "origin_status" not in without_param["allocation_lines"][0]


# --- split + column (9)/(12) integration ---

def test_qualifying_split_renders_origin_part_with_evidence_text():
    from app.bang_ke_rows import material_render_parts

    pool = {"NVL-1": [
        _stock(remaining_qty="2", available_qty="2"),
        _stock(source_row="ROW-2", line_no="2", origin_country="CHINA",
               consignee_name="NCC CN", remaining_qty="100", available_qty="100"),
    ]}
    material = _material(pool)
    case = materialize_bang_ke_origin_fields(
        {"products": [{"code": "TP1", "materials": [material]}]}, {},
    )
    parts = material_render_parts(case["products"][0]["materials"][0])
    assert len(parts) == 2
    by_status = {part["origin_status"]: part for part in parts}
    origin_part = by_status["origin"]
    assert origin_part["bang_ke_origin_text"] == "Việt Nam"
    assert origin_part["bang_ke_co_doc_no"] == f"Phụ lục X/{MINGJIE}"
    assert origin_part["material_value"] == "20"
    assert by_status["non_origin"]["bang_ke_origin_text"] == "Trung Quốc"
    assert by_status["non_origin"]["bang_ke_co_doc_no"] == ""


def test_fully_qualifying_material_renders_single_origin_part():
    # all lots qualify → ONE render row, but it must carry origin status +
    # evidence text (not the material's conservative default).
    material = _material({"NVL-1": [_stock()]})
    case = materialize_bang_ke_origin_fields(
        {"products": [{"code": "TP1", "materials": [material]}]}, {},
    )
    from app.bang_ke_rows import material_render_parts
    parts = material_render_parts(case["products"][0]["materials"][0])
    assert len(parts) == 1
    assert parts[0]["origin_status"] == "origin"
    assert parts[0]["bang_ke_co_doc_no"] == f"Phụ lục X/{MINGJIE}"
    assert parts[0]["material_value"] == "40"
    assert parts[0]["non_origin_cif_value"] == ""


def test_renderer_puts_qualifying_value_in_column_7():
    from openpyxl import Workbook
    from app.bang_ke_renderer import load_form_config, render_into_sheet

    material = _material({"NVL-1": [_stock()]})
    case = materialize_bang_ke_origin_fields(
        {"products": [{"code": "TP1", "materials": [material]}]}, {},
    )
    product = {"code": "TP1", "materials": case["products"][0]["materials"],
               "fob": "100", "quantity": "4"}
    cfg = load_form_config("LVC")
    wb = Workbook()
    ws = wb.active
    render_into_sheet(ws, cfg, case={"case_code": "C", "products": [product]}, product=product, sheet_title="P1")
    cols, start = cfg["body"]["columns"], cfg["body"]["start_row"]
    assert ws[f"{cols['origin_value']}{start}"].value == "40"
    assert ws[f"{cols['non_origin_value']}{start}"].value == "0"
    assert ws[f"{cols['co_doc_no']}{start}"].value == f"Phụ lục X/{MINGJIE}"
    assert ws[f"{cols['country']}{start}"].value == "Việt Nam"


# --- end-to-end through the real calculate route (flags patched in) ---

def test_calculate_route_resolves_flagged_vn_lots_end_to_end(monkeypatch):
    from fastapi.testclient import TestClient
    import app.main as main_module
    from app import supplier_evidence_store

    monkeypatch.setattr(supplier_evidence_store, "flagged_suppliers", lambda client_id: dict(FLAGS))
    from tests.test_co_demo import bcct_workbook, hidden_form_data

    client = TestClient(main_module.app)
    client.post(
        "/clients/growatt/bcct/upload",
        files={
            "file": (
                "bcct.xlsx",
                bcct_workbook([
                    {
                        "direction": "import", "declaration_type": "E15",
                        "declaration_no": "NK-AAA-VN", "line_no": "1",
                        "item_code": "DEMO-NPL-001", "description": "Board VN",
                        "hs_code": "8542.39", "quantity": "3", "unit": "PCE",
                        "customs_value": "30", "currency": "VND",
                        "origin_country": "VIETNAM", "partner_name": MINGJIE,
                    },
                    {
                        "direction": "export", "declaration_type": "E42",
                        "declaration_no": "XK-VN", "line_no": "1",
                        "item_code": "PV00.0048500", "description": "Growatt inverter",
                        "hs_code": "850440", "quantity": "3", "unit": "PCS",
                        "customs_value": "1000", "currency": "VND",
                        "invoice_ref": "INV-VN",
                    },
                ]),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    created = client.post(
        "/clients/growatt/co-case/create",
        data={"title": "VN origin", "case_code": "CO-VN-E2E", "destination_market": "Ấn Độ", "invoice_no": "INV-VN"},
        follow_redirects=False,
    )
    origin = client.get(f"{created.headers['location']}/origin")
    client.post(
        f"{created.headers['location']}/origin/sheet/PV00.0048500/calculate",
        data=hidden_form_data(origin.text),
    )
    page = client.get(f"{created.headers['location']}/origin")
    form_data = hidden_form_data(page.text)
    assert form_data["product_0_material_0_allocation_0_origin_status"] == "origin"
    assert form_data["product_0_material_0_allocation_0_bang_ke_co_doc_no"] == f"Phụ lục X/{MINGJIE}"
    assert form_data["product_0_material_0_bang_ke_co_doc_no"] == f"Phụ lục X/{MINGJIE}"
    assert form_data["product_0_material_0_non_origin_cif_value"] == "0"
    assert form_data["product_0_material_0_bang_ke_origin_text"] == "Việt Nam"
