"""Tests for the standard CO stock template + converter
(no Postgres required for the pure-logic paths)."""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.co_stock_template import (
    CO_STOCK_COLUMNS,
    CO_STOCK_HEADER_LABELS,
    CO_STOCK_SHEET_NAME,
    CoStockTemplateError,
    read_standard_co_stock,
    write_standard_co_stock,
)
from scripts.convert_co_stock import convert as convert_co_stock


def _make_agency_save_workbook(rows: list[dict]) -> bytes:
    """Build a tiny agency-shaped workbook with a `Save` sheet for the converter."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Save"
    headers = [
        "Số TK", "Ngày ĐK", "Mã loại hình", "STT hàng", "Mã NPL/SP",
        "Mã HS", "Tên hàng", "Xuất xứ", "Đơn giá", "Đơn giá tính thuế",
        "Tồn", "Đơn vị tính", "Tên đối tác", "Số hóa đơn", "Ngày hóa đơn",
        "Tỷ giá thanh toán", "Đã xuất", "TKX", "Mã ", "Số lượng XUẤT",
        "Số CO", "SL sử dụng", "Khóa giao dịch",
    ]
    ws.append(headers)
    for row in rows:
        ws.append([
            row.get("declaration_no"),
            row.get("registration_date"),
            row.get("declaration_type"),
            row.get("line_no"),
            row.get("customs_code"),
            row.get("hs_code"),
            row.get("goods_name"),
            row.get("origin_country"),
            row.get("unit_price"),
            row.get("taxable_unit_price"),
            row.get("opening_qty"),
            row.get("unit"),
            row.get("partner"),
            row.get("invoice_no"),
            row.get("invoice_date"),
            row.get("exchange_rate"),
            0,  # Đã xuất (Q) — unused
            row.get("source_tkx", ""),
            row.get("customs_code"),  # Mã (S) — repeated
            row.get("used_qty"),
            row.get("source_co_no"),
            row.get("used_qty"),  # SL sử dụng (V) — typically mirrors T
            row.get("transaction_key"),
        ])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _make_agency_nk2_workbook(rows: list[dict]) -> bytes:
    """Build a tiny agency-shaped workbook with the `NK2` sheet.

    NK2 has 3 preamble rows + header at row 4 + data from row 5.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "NK2"
    ws.append([None] * 22)  # row 1 — preamble
    ws.append([None] * 22)  # row 2 — preamble
    ws.append([None] * 22)  # row 3 — preamble
    ws.append([
        "Số TK", "Ngày ĐK", "Mã loại hình", "STT hàng", "Mã NPL/SP",
        "Mã HS", "Tên hàng", "Xuất xứ", "Đơn giá", "Đơn giá tính thuế",
        "Tổng số lượng", "Đơn vị tính", "Tên đối tác", "Số hóa đơn", "Ngày hóa đơn",
        "Tỷ giá thanh toán", "Đã xuất", "Tồn", "Check", "TKX",
        "", "SL sử dụng", "",
    ])  # row 4 — header
    for row in rows:
        ws.append([
            row.get("declaration_no"),
            row.get("registration_date"),
            row.get("declaration_type"),
            row.get("line_no"),
            row.get("customs_code"),
            row.get("hs_code"),
            row.get("goods_name"),
            row.get("origin_country"),
            row.get("unit_price"),
            row.get("taxable_unit_price"),
            row.get("opening_qty"),
            row.get("unit"),
            row.get("partner"),
            row.get("invoice_no"),
            row.get("invoice_date"),
            row.get("exchange_rate"),
            row.get("used_qty"),         # Q: Đã xuất
            None,                        # R: Tồn (formula in real file; left blank in tests)
            None,                        # S: Check
            row.get("transaction_key"),  # T: TKX cộng dồn
        ])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_standard_template_round_trip_preserves_decimal_and_date_types():
    rows = [
        {
            "declaration_no": "107700688100",
            "registration_date": date(2025, 11, 12),
            "declaration_type": "E11",
            "line_no": "1",
            "customs_code": "0000096074",
            "hs_code": "39079940",
            "goods_name": "Sơn bột tĩnh",
            "origin_country": "CHINA",
            "opening_qty": Decimal("500"),
            "unit": "KILO-GRAMMES",
            "used_qty": Decimal("11.489"),
            "source_co_no": "VNG26000001",
        },
        {
            "declaration_no": "107700957930",
            "line_no": "4",
            "customs_code": "000776-A",
            "opening_qty": Decimal("340"),
            "used_qty": Decimal("160"),
        },
    ]
    payload = write_standard_co_stock(rows)
    parsed, errors = read_standard_co_stock(payload)
    assert errors == []
    assert len(parsed) == 2
    assert parsed[0]["registration_date"] == date(2025, 11, 12)
    assert parsed[0]["opening_qty"] == Decimal("500")
    assert parsed[0]["used_qty"] == Decimal("11.489")
    assert parsed[0]["source_co_no"] == "VNG26000001"
    assert parsed[1]["declaration_no"] == "107700957930"


def test_standard_template_accepts_vietnamese_header_labels():
    wb = Workbook()
    ws = wb.active
    ws.title = CO_STOCK_SHEET_NAME
    ws.append([CO_STOCK_HEADER_LABELS[col] for col in CO_STOCK_COLUMNS])
    ws.append([
        "107700688100", date(2025, 11, 12), "E11", "1", "0000096074",
        "39079940", "Sơn bột tĩnh", "CHINA", None, None,
        500, "KG", None, None, None, None, 11.489, "VNG26000001", "",
    ])
    payload = BytesIO()
    wb.save(payload)
    rows, errors = read_standard_co_stock(payload.getvalue())
    assert errors == []
    assert rows[0]["customs_code"] == "0000096074"
    assert rows[0]["opening_qty"] == Decimal("500")
    assert rows[0]["used_qty"] == Decimal("11.489")


def test_standard_template_rejects_missing_required_columns():
    wb = Workbook()
    ws = wb.active
    ws.title = CO_STOCK_SHEET_NAME
    ws.append(["declaration_no", "line_no"])  # missing customs_code
    ws.append(["107700688100", "1"])
    payload = BytesIO()
    wb.save(payload)
    with pytest.raises(CoStockTemplateError) as exc_info:
        read_standard_co_stock(payload.getvalue())
    assert "customs_code" in str(exc_info.value)


def test_standard_template_roundtrip_smoke():
    rows = [{"declaration_no": "D1", "line_no": "1", "customs_code": "M-A",
             "opening_qty": Decimal("100"), "used_qty": Decimal("5")}]
    xlsx = write_standard_co_stock(rows)
    parsed, errors = read_standard_co_stock(xlsx)
    assert not errors
    assert parsed[0]["declaration_no"] == "D1"
    assert parsed[0]["customs_code"] == "M-A"


def test_converter_save_sheet_consolidates_duplicate_keys_sums_used_qty(tmp_path: Path):
    """Save sheet = event log; same key across multiple events → sum used_qty."""
    input_path = tmp_path / "agency.xlsx"
    rows = [
        {"declaration_no": "D1", "registration_date": date(2025, 11, 1), "declaration_type": "E11",
         "line_no": "1", "customs_code": "M-A", "hs_code": "39079940", "goods_name": "A",
         "origin_country": "CHINA", "opening_qty": 500, "unit": "KG", "used_qty": 10,
         "source_co_no": "VNG-001"},
        {"declaration_no": "D1", "registration_date": date(2025, 11, 1), "declaration_type": "E11",
         "line_no": "1", "customs_code": "M-A", "hs_code": "39079940", "goods_name": "A",
         "origin_country": "CHINA", "opening_qty": 500, "unit": "KG", "used_qty": 25,
         "source_co_no": "VNG-002"},
        {"declaration_no": "D2", "registration_date": date(2025, 11, 5), "declaration_type": "E11",
         "line_no": "3", "customs_code": "M-B", "hs_code": "40069090", "goods_name": "B",
         "origin_country": "CHINA", "opening_qty": 200, "unit": "PCS", "used_qty": 80,
         "source_co_no": "VNG-001"},
        {"declaration_no": "", "line_no": "", "customs_code": ""},
    ]
    input_path.write_bytes(_make_agency_save_workbook(rows))
    output_path = tmp_path / "standard.xlsx"
    summary = convert_co_stock(input_path, output_path, "Save")
    assert summary["rows_in"] == 4
    assert summary["rows_unique"] == 2
    assert summary["duplicates_merged"] == 1
    assert summary["rows_dropped"] == 1
    parsed, errors = read_standard_co_stock(output_path.read_bytes())
    assert errors == []
    by_key = {(r["declaration_no"], r["line_no"], r["customs_code"]): r for r in parsed}
    d1 = by_key[("D1", "1", "M-A")]
    assert d1["used_qty"] == Decimal("35")  # 10 + 25
    assert "VNG-001" in d1["source_co_no"] and "VNG-002" in d1["source_co_no"]
    assert by_key[("D2", "3", "M-B")]["used_qty"] == Decimal("80")


def test_converter_nk2_sheet_keeps_one_row_per_lot_skips_zero_used(tmp_path: Path):
    """NK2 sheet = reconciled per-lot snapshot; rows with Đã_xuất=0 skipped."""
    input_path = tmp_path / "agency.xlsx"
    rows = [
        {"declaration_no": "D1", "registration_date": date(2025, 11, 1), "declaration_type": "E11",
         "line_no": "1", "customs_code": "M-A", "hs_code": "39079940", "goods_name": "A",
         "origin_country": "CHINA", "opening_qty": 500, "unit": "KG", "used_qty": 197.631},
        # Zero-used row → skipped by default.
        {"declaration_no": "D2", "registration_date": date(2025, 11, 5), "declaration_type": "E11",
         "line_no": "1", "customs_code": "M-Z", "opening_qty": 100, "unit": "PCS", "used_qty": 0},
        {"declaration_no": "D3", "registration_date": date(2025, 11, 8), "declaration_type": "E11",
         "line_no": "5", "customs_code": "M-C", "opening_qty": 2392, "unit": "PCS", "used_qty": 7.524},
        # Has some data (used_qty) but missing key fields → counted then dropped.
        {"declaration_no": "", "line_no": "", "customs_code": "", "used_qty": 99},
    ]
    input_path.write_bytes(_make_agency_nk2_workbook(rows))
    output_path = tmp_path / "out.xlsx"
    summary = convert_co_stock(input_path, output_path, "NK2")
    assert summary["sheet"] == "NK2"
    assert summary["rows_in"] == 4
    assert summary["rows_unique"] == 2
    assert summary["rows_skipped_zero_used"] == 1
    assert summary["rows_dropped"] == 1
    parsed, errors = read_standard_co_stock(output_path.read_bytes())
    assert errors == []
    by_key = {(r["declaration_no"], r["line_no"], r["customs_code"]): r for r in parsed}
    d1 = by_key[("D1", "1", "M-A")]
    assert d1["opening_qty"] == Decimal("500")
    assert d1["used_qty"] == Decimal("197.631")
    assert d1["used_qty"] < d1["opening_qty"]  # NK2 = reconciled = no overclaim
    assert ("D2", "1", "M-Z") not in by_key  # zero-used skipped


def test_converter_nk2_include_zero_used_flag_imports_all_rows(tmp_path: Path):
    input_path = tmp_path / "agency.xlsx"
    input_path.write_bytes(_make_agency_nk2_workbook([
        {"declaration_no": "D1", "line_no": "1", "customs_code": "M-A", "opening_qty": 100, "used_qty": 0},
        {"declaration_no": "D2", "line_no": "1", "customs_code": "M-B", "opening_qty": 200, "used_qty": 50},
    ]))
    output_path = tmp_path / "out.xlsx"
    summary = convert_co_stock(input_path, output_path, "NK2", include_zero_used=True)
    assert summary["rows_unique"] == 2
    assert summary["rows_skipped_zero_used"] == 0


def test_converter_writes_errors_log_when_rows_dropped(tmp_path: Path):
    input_path = tmp_path / "agency.xlsx"
    input_path.write_bytes(_make_agency_save_workbook([
        {"declaration_no": "", "line_no": "", "customs_code": ""},
        {"declaration_no": "D1", "line_no": "1", "customs_code": "M-A", "opening_qty": 100, "used_qty": 5},
    ]))
    output_path = tmp_path / "out.xlsx"
    summary = convert_co_stock(input_path, output_path, "Save")
    assert summary["rows_dropped"] == 1
    log_path = Path(summary["errors_log"])
    assert log_path.exists()
    assert "thiếu declaration_no" in log_path.read_text(encoding="utf-8")
