"""Slice 1 — NXT + inventory-snapshot tier foundation.

Covers the canonical system_template round-trip (render → parse), adapter
detect/fallback, and store create→read with runtime-derived closing_implied /
variance. See .ai/features/2026-06-14-nxt-inventory-tier/brief.md.
"""
from __future__ import annotations

from datetime import date

import pytest
from fastapi.testclient import TestClient

from app.auth.session import SESSION_COOKIE, create_session
from app.database import connect
from app.main import app
from app.parsers import inventory_adapters, nxt_adapters
from app.parsers.nxt_adapters.system_template import (
    SystemTemplateNxtAdapter, render_template_xlsx as render_nxt,
)
from app.parsers.inventory_adapters.system_template import (
    SystemTemplateInventoryAdapter, render_template_xlsx as render_inv,
)
from app.stores import inventory_snapshots as inv_store
from app.stores import nxt as nxt_store

CLIENT = "nxt_tier_test"


@pytest.fixture(autouse=True)
def _client():
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "insert into hub.clients (client_id, name) values (%s, %s) "
            "on conflict (client_id) do nothing",
            (CLIENT, "NXT tier test"),
        )
    yield


@pytest.fixture
def isolated_files_dir(monkeypatch, tmp_path):
    monkeypatch.setenv("DATA_HUB_FILES_ROOT", str(tmp_path / "files"))
    import app.storage as storage_mod
    storage_mod._BACKEND = None
    yield
    storage_mod._BACKEND = None


# ── NXT adapter round-trip ──────────────────────────────────────────────

def test_nxt_template_roundtrip():
    lines = SystemTemplateNxtAdapter().parse(render_nxt())
    assert len(lines) == 3  # one sample row per NVL/TP/BTP sheet
    by_role = {l["reported_role"]: l for l in lines}
    assert set(by_role) == {"nvl", "tp", "btp"}

    nvl = by_role["nvl"]
    assert nvl["internal_code"] == "NVL-001"
    assert nvl["customs_code"] == "A0123"
    assert nvl["uom"] == "KG"
    assert nvl["opening"] == 1000
    assert nvl["inbound_total"] == 5000
    assert nvl["out_xuat_sx"] == 4500
    assert nvl["closing_reported"] == 1500

    # BTP sample has an empty customs code → None, not "".
    assert by_role["btp"]["customs_code"] is None


def test_nxt_detect_and_fallback():
    blob = render_nxt()
    assert SystemTemplateNxtAdapter().detect(blob) == pytest.approx(0.95)
    result = nxt_adapters.parse_with_fallback(blob)
    assert result is not None
    lines, name = result
    assert name == "system_template"
    assert len(lines) == 3


def test_nxt_detect_abstains_on_foreign_workbook():
    # An inventory template has no NVL/TP/BTP sheets → NXT adapter abstains.
    assert SystemTemplateNxtAdapter().detect(render_inv()) is None


# ── NXT store round-trip + derived closing_implied ──────────────────────

def test_nxt_store_roundtrip_and_closing_implied():
    lines = SystemTemplateNxtAdapter().parse(render_nxt())
    aid = nxt_store.create_artifact(
        client_id=CLIENT, lines=lines,
        period_from=date(2025, 1, 1), period_to=date(2025, 12, 31),
        source_kind="system_template", adapter_name="system_template",
    )
    art = nxt_store.get_artifact(aid)
    assert art is not None
    assert art["client_id"] == CLIENT
    assert art["period_to"] == date(2025, 12, 31)
    assert len(art["lines"]) == 3

    nvl = next(l for l in art["lines"] if l["reported_role"] == "nvl")
    # closing_implied = opening + inbound - Σ out = 1000 + 5000 - 4500 = 1500
    assert nvl["closing_implied"] == pytest.approx(1500.0)
    assert nvl["closing_reported"] == pytest.approx(1500.0)

    listed = nxt_store.list_artifacts(CLIENT)
    assert any(a["id"] == aid and a["n_lines"] == 3 for a in listed)


# ── Inventory adapter + store round-trip + derived variance ─────────────

def test_inventory_template_roundtrip():
    lines = SystemTemplateInventoryAdapter().parse(render_inv())
    assert len(lines) == 2
    first = lines[0]
    assert first["code"] == "NVL-001"
    assert first["warehouse"] == "Kho NVL"
    assert first["qty_book"] == 1500
    assert first["qty_physical"] == 1495


def test_inventory_detect_and_fallback():
    blob = render_inv()
    assert SystemTemplateInventoryAdapter().detect(blob) == pytest.approx(0.9)
    result = inventory_adapters.parse_with_fallback(blob)
    assert result is not None
    lines, name = result
    assert name == "system_template"
    assert len(lines) == 2


def test_inventory_store_roundtrip_and_variance():
    lines = SystemTemplateInventoryAdapter().parse(render_inv())
    sid = inv_store.create_snapshot(
        client_id=CLIENT, lines=lines, snapshot_date=date(2025, 12, 31),
        source_kind="system_template", adapter_name="system_template",
    )
    snap = inv_store.get_snapshot(sid)
    assert snap is not None
    assert snap["snapshot_date"] == date(2025, 12, 31)
    assert len(snap["lines"]) == 2
    first = snap["lines"][0]
    # variance = physical - book = 1495 - 1500 = -5
    assert first["variance"] == pytest.approx(-5.0)

    listed = inv_store.list_snapshots(CLIENT)
    assert any(s["id"] == sid and s["n_lines"] == 2 for s in listed)


# ── Template download routes ────────────────────────────────────────────

def _dev_client() -> TestClient:
    with connect() as conn, conn.cursor() as cur:
        cur.execute("select user_id from hub.users where role='dev' "
                    "and status='active' limit 1")
        row = cur.fetchone()
    assert row, "no seeded dev user"
    c = TestClient(app)
    c.cookies.set(SESSION_COOKIE, create_session(row[0]))
    return c


_XLSX_MAGIC = b"PK"


@pytest.mark.parametrize("path, expect_name", [
    ("/templates/nxt.xlsx", "mau-nxt.xlsx"),
    ("/templates/inventory-snapshot.xlsx", "mau-ton-kho-cuoi-ky.xlsx"),
])
def test_template_download(path, expect_name):
    c = _dev_client()
    r = c.get(path)
    assert r.status_code == 200
    assert r.content.startswith(_XLSX_MAGIC)
    assert expect_name in r.headers["content-disposition"]


def test_template_download_requires_auth():
    r = TestClient(app).get("/templates/nxt.xlsx", follow_redirects=False)
    assert r.status_code in (302, 303, 401, 403)


# ── Upload → preview → confirm flow ─────────────────────────────────────

def test_nxt_upload_preview_confirm_flow(isolated_files_dir):
    c = _dev_client()
    files = {"file": ("mau-nxt.xlsx", render_nxt(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = c.post(f"/clients/{CLIENT}/nxt/upload",
               files=files,
               data={"period_from": "2025-01-01", "period_to": "2025-12-31"},
               follow_redirects=False)
    assert r.status_code == 303
    preview_url = r.headers["location"]
    assert "/nxt/preview/" in preview_url

    pv = c.get(preview_url)
    assert pv.status_code == 200
    assert "Tồn cuối khớp" in pv.text  # sample rows balance

    upload_id = preview_url.rstrip("/").split("/")[-1]
    cf = c.post(f"/clients/{CLIENT}/nxt/preview/{upload_id}/confirm",
                follow_redirects=False)
    assert cf.status_code == 303

    arts = nxt_store.list_artifacts(CLIENT)
    assert arts and arts[0]["n_lines"] == 3
    art = nxt_store.get_artifact(arts[0]["id"])
    assert art["period_to"] == date(2025, 12, 31)
    assert art["adapter_name"] == "system_template"


def test_inventory_upload_preview_confirm_flow(isolated_files_dir):
    c = _dev_client()
    files = {"file": ("mau-ton.xlsx", render_inv(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = c.post(f"/clients/{CLIENT}/inventory-snapshots/upload",
               files=files, data={"snapshot_date": "2025-12-31"},
               follow_redirects=False)
    assert r.status_code == 303
    preview_url = r.headers["location"]
    assert "/inventory-snapshots/preview/" in preview_url

    pv = c.get(preview_url)
    assert pv.status_code == 200
    assert "lệch sổ-thực" in pv.text  # sample has a 5-unit discrepancy

    upload_id = preview_url.rstrip("/").split("/")[-1]
    cf = c.post(f"/clients/{CLIENT}/inventory-snapshots/preview/{upload_id}/confirm",
                follow_redirects=False)
    assert cf.status_code == 303

    snaps = inv_store.list_snapshots(CLIENT)
    assert snaps and snaps[0]["n_lines"] == 2
    snap = inv_store.get_snapshot(snaps[0]["id"])
    assert snap["snapshot_date"] == date(2025, 12, 31)


def test_nxt_nav_tab_present():
    c = _dev_client()
    r = c.get(f"/clients/{CLIENT}/nxt")
    assert r.status_code == 200
    assert "Nhập-Xuất-Tồn" in r.text
    assert "Quyết toán" in r.text  # settlement nav group


# ── ezsoft_3tsoft adapter (bilingual, group-row roles, lumped Xuất) ──────

def _build_ezsoft_xlsx() -> bytes:
    """Synthetic EZSOFT/3TSoft NXT sheet: title block + bilingual header band
    (category row + sub-rows) + group rows that set role + data rows."""
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "EZSOFT - 3TSoft"
    ws["A1"] = "Công ty ABC"
    ws["A3"] = "TỔNG HỢP NHẬP - XUẤT - TỒN"
    ws["A4"] = "Từ ngày 01/01/2025"
    ws.cell(6, 1, "Vật tư - 物資")
    ws.cell(6, 5, "Tồn đầu - 期初庫存"); ws.cell(6, 6, "Nhập - 入庫")
    ws.cell(6, 7, "Xuất - 輸出"); ws.cell(6, 8, "Tồn cuối - 期末庫存")
    ws.cell(7, 1, "Mã 代碼"); ws.cell(7, 2, "Tên - 名稱"); ws.cell(7, 4, "Đvt 單位")
    for col in (5, 6, 7, 8):
        ws.cell(7, col, "Số lượng 數量")
    ws.cell(8, 2, "Tiếng Việt 越文"); ws.cell(8, 3, "Tiếng Hoa 中文")
    # group BTP (subtotal row — skipped, sets role)
    ws.cell(9, 1, "BTP"); ws.cell(9, 2, "Nhóm bán thành phẩm")
    ws.cell(9, 5, 150); ws.cell(9, 8, 150)
    ws.cell(10, 1, "100.001"); ws.cell(10, 2, "Bo mạch"); ws.cell(10, 4, "Cái")
    ws.cell(10, 5, 100); ws.cell(10, 6, 50); ws.cell(10, 7, 30); ws.cell(10, 8, 120)
    # group NVL
    ws.cell(11, 1, "NVL"); ws.cell(11, 2, "Nhóm nguyên vật liệu")
    ws.cell(12, 1, "001.002"); ws.cell(12, 2, "Nhựa"); ws.cell(12, 4, "KG")
    ws.cell(12, 5, 1000); ws.cell(12, 6, 500); ws.cell(12, 7, 300); ws.cell(12, 8, 1200)
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def test_ezsoft_adapter_parse():
    from app.parsers.nxt_adapters.ezsoft_3tsoft import Ezsoft3TSoftAdapter
    blob = _build_ezsoft_xlsx()
    a = Ezsoft3TSoftAdapter()
    assert a.detect(blob) == pytest.approx(0.97)
    lines = a.parse(blob)
    assert len(lines) == 2  # 2 data rows; the 2 group rows are skipped
    btp, nvl = lines
    assert btp["reported_role"] == "btp"
    assert btp["internal_code"] == "100.001"
    assert btp["uom"] == "Cái"
    assert btp["opening"] == 100 and btp["inbound_total"] == 50
    assert btp["outbound_total"] == 30  # lumped Xuất → outbound_total
    assert btp["out_xuat_sx"] is None   # not split into buckets
    assert btp["closing_reported"] == 120
    assert nvl["reported_role"] == "nvl"
    assert nvl["internal_code"] == "001.002"


def test_ezsoft_outranks_system_template_and_closing_implied():
    from app.parsers.nxt_adapters._common import closing_implied
    from app.parsers.nxt_adapters.system_template import SystemTemplateNxtAdapter
    blob = _build_ezsoft_xlsx()
    # system_template must abstain (no canonical NVL/TP/BTP sheets).
    assert SystemTemplateNxtAdapter().detect(blob) is None
    lines, name = nxt_adapters.parse_with_fallback(blob)
    assert name == "ezsoft_3tsoft"
    btp = lines[0]
    # closing_implied = opening + inbound − outbound_total = 100 + 50 − 30 = 120
    assert closing_implied(btp) == pytest.approx(120.0)


def test_ezsoft_store_roundtrip_preserves_outbound_total(isolated_files_dir):
    from app.parsers.nxt_adapters.ezsoft_3tsoft import Ezsoft3TSoftAdapter
    lines = Ezsoft3TSoftAdapter().parse(_build_ezsoft_xlsx())
    aid = nxt_store.create_artifact(
        client_id=CLIENT, lines=lines, source_kind="ezsoft_3tsoft",
        adapter_name="ezsoft_3tsoft")
    art = nxt_store.get_artifact(aid)
    btp = next(l for l in art["lines"] if l["reported_role"] == "btp")
    assert btp["outbound_total"] == pytest.approx(30.0)
    assert btp["closing_implied"] == pytest.approx(120.0)


# ── Per-client adapter binding + admin registry view ────────────────────

def test_settlement_adapter_binding_roundtrip():
    from app.stores import settlement_adapter_binding as b
    assert b.get_default_adapter(CLIENT, "nxt") == "auto"
    b.set_default_adapter(CLIENT, "nxt", "ezsoft_3tsoft")
    assert b.get_default_adapter(CLIENT, "nxt") == "ezsoft_3tsoft"
    # Unknown adapter rejected; unknown module rejected.
    with pytest.raises(ValueError):
        b.set_default_adapter(CLIENT, "nxt", "nope")
    with pytest.raises(ValueError):
        b.set_default_adapter(CLIENT, "bogus", "auto")
    rows = {r["client_id"]: r for r in b.list_bindings()}
    assert rows[CLIENT]["nxt"] == "ezsoft_3tsoft"
    assert rows[CLIENT]["inventory"] == "auto"
    b.set_default_adapter(CLIENT, "nxt", "auto")  # reset


def test_settlement_adapters_admin_view():
    c = _dev_client()
    r = c.get("/admin/settlement-adapters")
    assert r.status_code == 200
    assert "ezsoft_3tsoft" in r.text
    assert "system_template" in r.text
    assert "Adapter Quyết toán" in r.text  # nav link present


def test_nxt_upload_explicit_adapter_pick(isolated_files_dir):
    c = _dev_client()
    files = {"file": ("ez.xlsx", _build_ezsoft_xlsx(),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = c.post(f"/clients/{CLIENT}/nxt/upload",
               files=files, data={"adapter": "ezsoft_3tsoft"},
               follow_redirects=False)
    assert r.status_code == 303
    preview_url = r.headers["location"]
    upload_id = preview_url.rstrip("/").split("/")[-1]
    cf = c.post(f"/clients/{CLIENT}/nxt/preview/{upload_id}/confirm",
                follow_redirects=False)
    assert cf.status_code == 303
    adapters = {nxt_store.get_artifact(a["id"])["adapter_name"]
                for a in nxt_store.list_artifacts(CLIENT)}
    assert "ezsoft_3tsoft" in adapters


def test_nxt_set_default_adapter_route():
    c = _dev_client()
    r = c.post(f"/clients/{CLIENT}/nxt/default-adapter",
               data={"adapter": "system_template"}, follow_redirects=False)
    assert r.status_code == 303
    from app.stores import settlement_adapter_binding as b
    assert b.get_default_adapter(CLIENT, "nxt") == "system_template"
    b.set_default_adapter(CLIENT, "nxt", "auto")  # reset


# ── manual_generic adapter (alias auto-match + mapping_override) ─────────

def _build_generic_nxt(headers: list[str], row: list) -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    for ci, h in enumerate(headers, 1):
        ws.cell(1, ci, h)
    for ci, v in enumerate(row, 1):
        ws.cell(2, ci, v)
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def test_manual_generic_alias_automatch():
    from app.parsers.nxt_adapters.manual_generic import ManualGenericNxtAdapter
    blob = _build_generic_nxt(
        ["Mã nội bộ", "Tên", "ĐVT", "Tồn đầu kỳ", "Nhập trong kỳ", "Xuất", "Tồn cuối kỳ"],
        ["X1", "Vật tư X", "KG", 100, 50, 30, 120])
    a = ManualGenericNxtAdapter()
    assert a.detect(blob) is None  # abstains
    lines = a.parse(blob)
    assert len(lines) == 1
    ln = lines[0]
    assert ln["internal_code"] == "X1" and ln["uom"] == "KG"
    assert ln["opening"] == 100 and ln["inbound_total"] == 50
    assert ln["outbound_total"] == 30  # single "Xuất" → outbound_total
    assert ln["closing_reported"] == 120
    # fallback routes here when no high-precision adapter matches
    _lines, name = nxt_adapters.parse_with_fallback(blob)
    assert name == "manual_generic"


def test_manual_generic_mapping_override():
    from app.parsers.nxt_adapters import NxtParseError
    from app.parsers.nxt_adapters.manual_generic import ManualGenericNxtAdapter
    blob = _build_generic_nxt(
        ["Code", "Name", "Begin", "In", "Out", "End"],
        ["Y1", "Item Y", 10, 5, 2, 13])
    a = ManualGenericNxtAdapter()
    # Unknown headers → alias auto-match can't identify code/qty → raises.
    with pytest.raises(NxtParseError):
        a.parse(blob)
    # With a confirmed mapping override it parses.
    override = {"Code": "internal_code", "Name": "name", "Begin": "opening",
                "In": "inbound_total", "Out": "outbound_total", "End": "closing_reported"}
    lines = a.parse(blob, mapping_override=override)
    assert len(lines) == 1
    assert lines[0]["internal_code"] == "Y1"
    assert lines[0]["outbound_total"] == 2 and lines[0]["closing_reported"] == 13


def test_nxt_mapping_page_flow_and_cache(isolated_files_dir):
    c = _dev_client()
    # Hermetic: drop any cached mapping for this client from prior runs (the DB
    # persists across pytest invocations).
    with connect() as conn, conn.cursor() as cur:
        cur.execute("delete from hub.parser_mappings where client_id=%s "
                    "and module='nxt'", (CLIENT,))
    blob = _build_generic_nxt(
        ["Code", "Name", "Begin", "In", "Out", "End"],
        ["Z1", "Item Z", 10, 5, 2, 13])

    # Unknown headers → no adapter parses → routed to the mapping page.
    r = c.post(f"/clients/{CLIENT}/nxt/upload",
               files={"file": ("weird.xlsx", blob, "application/vnd.ms-excel")},
               data={"period_to": "2025-12-31"}, follow_redirects=False)
    assert r.status_code == 303
    assert "/nxt/upload/mapping/" in r.headers["location"]
    upload_id = r.headers["location"].rstrip("/").split("/")[-1]

    mp = c.get(f"/clients/{CLIENT}/nxt/upload/mapping/{upload_id}")
    assert mp.status_code == 200 and "Code" in mp.text

    form = {
        "col_0__header": "Code", "col_0__field": "internal_code",
        "col_1__header": "Name", "col_1__field": "name",
        "col_2__header": "Begin", "col_2__field": "opening",
        "col_3__header": "In", "col_3__field": "inbound_total",
        "col_4__header": "Out", "col_4__field": "outbound_total",
        "col_5__header": "End", "col_5__field": "closing_reported",
    }
    pr = c.post(f"/clients/{CLIENT}/nxt/upload/mapping/{upload_id}/parse",
                data=form, follow_redirects=False)
    assert pr.status_code == 303 and "/nxt/preview/" in pr.headers["location"]
    pid = pr.headers["location"].rstrip("/").split("/")[-1]
    cf = c.post(f"/clients/{CLIENT}/nxt/preview/{pid}/confirm", follow_redirects=False)
    assert cf.status_code == 303
    adapters = {nxt_store.get_artifact(a["id"])["adapter_name"]
                for a in nxt_store.list_artifacts(CLIENT)}
    assert "manual_generic" in adapters

    # Re-upload the same shape → cached mapping → straight to preview (no map page).
    r2 = c.post(f"/clients/{CLIENT}/nxt/upload",
                files={"file": ("weird2.xlsx", blob, "application/vnd.ms-excel")},
                follow_redirects=False)
    assert r2.status_code == 303
    assert "/nxt/preview/" in r2.headers["location"]


# ── sap_mb5b adapter (SAP MB5B per-material detail) ──────────────────────

def _build_mb5b_xlsx() -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws.append(["material", "material_description", "gl_account", "valuation_class",
               "base_unit", "opening_stock_qty", "total_receipt_qty",
               "total_issue_qty", "closing_stock_qty"])
    ws.append(["0000080692", "磁控連接線", "12150000", "3000", "EA", 0, 720, 0, 720])
    ws.append(["0000080058", "橡膠", "12150000", "3000", "EA", 100, 50, 30, 120])
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def test_sap_mb5b_adapter():
    from app.parsers.nxt_adapters._common import closing_implied
    from app.parsers.nxt_adapters.sap_mb5b import SapMb5bAdapter
    blob = _build_mb5b_xlsx()
    assert SapMb5bAdapter().detect(blob) == pytest.approx(0.93)
    lines, name = nxt_adapters.parse_with_fallback(blob)
    assert name == "sap_mb5b"
    assert len(lines) == 2
    a = lines[1]
    assert a["internal_code"] == "0000080058" and a["uom"] == "EA"
    assert a["opening"] == 100 and a["inbound_total"] == 50
    assert a["outbound_total"] == 30 and a["closing_reported"] == 120
    assert a["reported_role"] is None  # GL→role is per-client, not in adapter
    assert closing_implied(a) == pytest.approx(120.0)


# ── misa_can_doi_ton adapter (merged header + warehouse section breaks) ──

def _build_misa_xlsx() -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Page1"
    ws.cell(4, 1, "CÂN ĐỐI TỒN KHO")
    ws.cell(5, 1, "Từ ngày: 01/01/2024 - Đến ngày: 31/12/2024")
    # row 7 (code/name/uom) + row 8 (qty sub-cols) at scattered columns
    ws.cell(7, 1, "STT"); ws.cell(7, 2, "Mã Hàng"); ws.cell(7, 6, "Tên Hàng")
    ws.cell(7, 10, "ĐVT"); ws.cell(7, 11, "Số Lượng")
    ws.cell(8, 11, "Đầu Kỳ"); ws.cell(8, 13, "Nhập"); ws.cell(8, 17, "Xuất")
    ws.cell(8, 19, "Cuối Kỳ")
    ws.cell(9, 1, "Kho hàng: KHO-NK")  # section break — no code → skipped
    ws.cell(10, 1, 1); ws.cell(10, 2, "NK-DAYG"); ws.cell(10, 6, "Dây giầy")
    ws.cell(10, 10, "Đôi"); ws.cell(10, 11, 41913); ws.cell(10, 17, 1912)
    ws.cell(10, 19, 40001)
    ws.cell(11, 1, 2); ws.cell(11, 2, "NK-CK"); ws.cell(11, 6, "Chỉ khâu")
    ws.cell(11, 10, "Mét"); ws.cell(11, 11, 100); ws.cell(11, 13, 20)
    ws.cell(11, 17, 30); ws.cell(11, 19, 90)
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def test_v1_hub_read_api_and_period_end_link(monkeypatch):
    import secrets
    monkeypatch.setenv("DATA_HUB_API_AUTH_DISABLED", "1")
    # Fresh throwaway client — the DB persists across pytest runs, so a shared
    # client would accumulate artifacts and inflate the period-end aggregates.
    cid = "nxt-api-" + secrets.token_hex(4)
    with connect() as conn, conn.cursor() as cur:
        cur.execute("insert into hub.clients (client_id, name) values (%s, %s)",
                    (cid, "nxt api test"))
    d = date(2025, 12, 31)
    aid = nxt_store.create_artifact(
        client_id=cid, period_to=d, source_kind="system_template",
        adapter_name="system_template",
        lines=[{"internal_code": "MAT-LINK", "opening": 0, "inbound_total": 500,
                "out_xuat_sx": 0, "outbound_total": 0, "closing_reported": 500,
                "reported_role": "nvl"}])
    sid = inv_store.create_snapshot(
        client_id=cid, snapshot_date=d, source_kind="system_template",
        adapter_name="system_template",
        lines=[{"code": "MAT-LINK", "qty_book": 500, "qty_physical": 495}])
    c = TestClient(app)
    h = {"authorization": "Bearer test"}

    r = c.get(f"/v1/hub/dncxs/{cid}/nxt", headers=h)
    assert r.status_code == 200
    assert any(a["id"] == aid for a in r.json()["items"])

    r = c.get(f"/v1/hub/dncxs/{cid}/nxt/{aid}", headers=h)
    assert r.status_code == 200
    ln = r.json()["lines"][0]
    assert ln["internal_code"] == "MAT-LINK" and ln["closing_implied"] == 500

    r = c.get(f"/v1/hub/dncxs/{cid}/inventory-snapshots/{sid}", headers=h)
    assert r.status_code == 200
    assert r.json()["lines"][0]["variance"] == -5

    r = c.get(f"/v1/hub/dncxs/{cid}/period-end-link?date=2025-12-31", headers=h)
    assert r.status_code == 200
    link = {row["code"]: row for row in r.json()["items"]}
    assert link["MAT-LINK"]["nxt_closing"] == 500
    assert link["MAT-LINK"]["snapshot_book"] == 500
    assert link["MAT-LINK"]["snapshot_physical"] == 495

    # date is required
    assert c.get(f"/v1/hub/dncxs/{cid}/period-end-link", headers=h).status_code == 400


def _build_kiemke_xlsx() -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    hdr = ["序号", "物料编码", "物料名称", "批号", "库存主单位",
           "Số lượng tồn kho (theo đơn vị chính) .库存量(主单位)", "实盘数量"]
    data = {
        "Nguyên phụ liệu 主材": [
            [1, "03.03.11.090", "电子纸", "VRP055", "片", 22, 1],
            [2, "03.03.11.090", "电子纸", "VRP057", "片", 1505, 1500]],
        "Thành phẩm 成品仓": [
            [1, "0105A0101035", "成品A", "B001", "EA", 200, 200]],
    }
    for title, rows in data.items():
        ws = wb.create_sheet(title)
        ws.cell(1, 1, f"VN{title}仓")
        for ci, h in enumerate(hdr, 1):
            ws.cell(2, ci, h)
        for ri, row in enumerate(rows, start=3):
            for ci, v in enumerate(row, 1):
                ws.cell(ri, ci, v)
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def test_kiem_ke_multi_kho_adapter():
    from app.parsers.inventory_adapters._common import variance
    from app.parsers.inventory_adapters.kiem_ke_multi_kho import KiemKeMultiKhoAdapter
    blob = _build_kiemke_xlsx()
    assert KiemKeMultiKhoAdapter().detect(blob) == pytest.approx(0.93)
    lines, name = inventory_adapters.parse_with_fallback(blob)
    assert name == "kiem_ke_multi_kho"  # outranks system_template
    assert len(lines) == 3
    nvl = lines[0]
    assert nvl["code"] == "03.03.11.090" and nvl["uom"] == "片"
    assert nvl["warehouse"] == "Nguyên phụ liệu 主材"  # sheet title = kho
    assert nvl["batch"] == "VRP055"
    assert nvl["qty_book"] == 22 and nvl["qty_physical"] == 1
    assert variance(nvl) == pytest.approx(-21.0)


def test_misa_can_doi_ton_adapter():
    from app.parsers.nxt_adapters.misa_can_doi_ton import MisaCanDoiTonAdapter
    blob = _build_misa_xlsx()
    assert MisaCanDoiTonAdapter().detect(blob) == pytest.approx(0.9)
    lines, name = nxt_adapters.parse_with_fallback(blob)
    assert name == "misa_can_doi_ton"
    assert len(lines) == 2  # the "Kho hàng:" section break row is skipped
    dayg = lines[0]
    assert dayg["internal_code"] == "NK-DAYG" and dayg["uom"] == "Đôi"
    assert dayg["opening"] == 41913 and dayg["outbound_total"] == 1912
    assert dayg["closing_reported"] == 40001
    ck = lines[1]
    assert ck["inbound_total"] == 20 and ck["closing_reported"] == 90
