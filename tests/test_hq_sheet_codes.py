from app.workbook_io import hq_sheet_codes_for_product, write_hq_sheet_materials, _hq_layout_for


def _product(**fields):
    return fields


def test_hq_export_writes_customs_code_not_internal():
    """The HQ bảng kê must show the customs item code (mã HQ from the matched
    lot), not the internal allocation/BOM code (lookup-only). Growatt diverges:
    HQ 'DAYTINHIEU' vs internal '012.0002700'."""
    from openpyxl import Workbook
    mat = {
        "material_code": "012.0002700",           # internal — lookup only
        "customs_material_code": "DAYTINHIEU",     # mã HQ — matched lot's customs_item_code
        "internal_material_code": "012.0002700",
        "material_description": "Dây tín hiệu", "hs_code": "85339000", "uom": "ST",
        "bom_qty_per": "1", "consumed_qty": "368", "unit_value": "6708",
        "material_value": "2468835.82", "origin_status": "non_origin",
    }
    for sheet_code in ("LVC", "PL_VII"):
        col = _hq_layout_for(sheet_code)["cols"]["mat_code"]
        wb = Workbook(); ws = wb.active
        write_hq_sheet_materials(ws, {"materials": [mat]}, 5, sheet_code=sheet_code)
        assert ws.cell(row=5, column=col).value == "DAYTINHIEU", sheet_code


def test_hq_export_falls_back_to_material_code_when_no_customs_code():
    """Unmatched NVL has no matched-lot customs code → fall back to material_code."""
    from openpyxl import Workbook
    mat = {"material_code": "INT-1", "material_description": "x", "origin_status": "non_origin"}
    col = _hq_layout_for("LVC")["cols"]["mat_code"]
    wb = Workbook(); ws = wb.active
    write_hq_sheet_materials(ws, {"materials": [mat]}, 5, sheet_code="LVC")
    assert ws.cell(row=5, column=col).value == "INT-1"


def test_config_renderer_material_column_uses_customs_hq_code():
    """The config-driven renderer is the ACTIVE export path (write_hq_sheet_materials
    is only the legacy fallback). It too must put the customs item code (mã HQ) in
    the material column, not the internal allocation code."""
    from openpyxl import Workbook
    from app.bang_ke_renderer import load_form_config, render_into_sheet
    cfg = load_form_config("LVC")
    wb = Workbook(); ws = wb.active
    product = {"code": "P1", "materials": [{
        "material_code": "012.0002700", "customs_material_code": "DAYTINHIEU",
        "internal_material_code": "012.0002700", "material_description": "Dây tín hiệu",
        "origin_status": "non_origin", "hs_code": "85339000", "uom": "ST",
        "bom_qty_per": "1", "consumed_qty": "1", "unit_value": "1", "material_value": "1",
    }]}
    render_into_sheet(ws, cfg, case={"case_code": "C", "products": [product]}, product=product, sheet_title="P1")
    col, start = cfg["body"]["columns"]["material_code"], cfg["body"]["start_row"]
    assert ws[f"{col}{start}"].value == "DAYTINHIEU"


def test_eur1_with_lvc_override_picks_lvc_sheet():
    product = _product(
        origin_sheet_effective_form_code="EUR.1",
        origin_sheet_criteria_override="LVC 30%",
        origin_sheet_effective_criteria_text="LVC 30%",
        documented_result="Cần đối chiếu mô tả hàng hóa trước khi áp dụng: ...",
    )
    assert hq_sheet_codes_for_product(product) == {"LVC"}


def test_eur1_with_cth_override_picks_cth_sheet():
    product = _product(
        origin_sheet_effective_form_code="EUR.1",
        origin_sheet_criteria_override="CTH",
        documented_result="Cần đối chiếu mô tả hàng hóa trước khi áp dụng: ...",
    )
    assert hq_sheet_codes_for_product(product) == {"CTH"}


def test_eur1_with_no_criterion_match_falls_back_to_eur1():
    product = _product(
        origin_sheet_effective_form_code="EUR.1",
        origin_sheet_criteria_override="",
        documented_result="Cần đối chiếu mô tả hàng hóa trước khi áp dụng: ...",
    )
    assert hq_sheet_codes_for_product(product) == {"EUR1"}


def test_eur1_with_psr_documented_result_picks_psr():
    product = _product(
        origin_sheet_effective_form_code="EUR.1",
        documented_result="Tra PSR Form EUR.1 theo Phụ lục",
    )
    assert hq_sheet_codes_for_product(product) == {"PSR"}


def test_non_eur1_with_lvc_override_picks_lvc():
    product = _product(
        origin_sheet_effective_form_code="B",
        origin_sheet_criteria_override="LVC 30%",
    )
    assert hq_sheet_codes_for_product(product) == {"LVC"}


def test_non_eur1_with_no_criterion_defaults_to_lvc():
    product = _product(origin_sheet_effective_form_code="B")
    assert hq_sheet_codes_for_product(product) == {"LVC"}


def test_override_wins_over_documented_result():
    product = _product(
        origin_sheet_effective_form_code="AI",
        origin_sheet_criteria_override="CTH",
        documented_result="LVC 30%",
    )
    assert hq_sheet_codes_for_product(product) == {"CTH"}


# --- Compound criterion ("X hoặc Y") must follow the ORDER it is listed, not a
# fixed LVC>RVC>CTSH>CTH precedence. Client report 2026-06-18: a CPTPP dossier
# whose recommendation is "CTH hoặc RVC 30/40/50 tùy công thức" exported an RVC
# value-buildup sheet ("tiêu chí tổng") instead of CTH, because RVC was matched
# before CTH. The first-listed criterion is the primary one. ---

def test_cptpp_recommendation_cth_first_picks_cth():
    """The real CPTPP recommendation text (co_forms.py) lists CTH first; with no
    override the bảng kê must be CTH, not RVC."""
    product = _product(
        origin_sheet_effective_form_code="CPTPP",
        origin_sheet_effective_criteria_text="CTH hoặc RVC 30/40/50 tùy công thức",
    )
    assert hq_sheet_codes_for_product(product) == {"CTH"}


def test_cptpp_recommendation_ctsh_first_picks_ctsh():
    product = _product(
        origin_sheet_effective_form_code="CPTPP",
        origin_sheet_effective_criteria_text="CTSH hoặc RVC 30/40/50 tùy công thức",
    )
    assert hq_sheet_codes_for_product(product) == {"CTSH"}


def test_compound_rvc_first_still_picks_rvc():
    """First-listed wins both ways: RVC before CTH stays RVC."""
    product = _product(origin_sheet_criteria_override="RVC 40% hoặc CTH")
    assert hq_sheet_codes_for_product(product) == {"RVC"}


def test_compound_lvc_first_still_picks_lvc():
    """Regression: the documented 'LVC 30% hoặc CTH' → LVC must not change."""
    product = _product(origin_sheet_criteria_override="LVC 30% hoặc CTH")
    assert hq_sheet_codes_for_product(product) == {"LVC"}


def test_cth_before_ctsh_compound_picks_cth():
    product = _product(origin_sheet_criteria_override="CTH hoặc CTSH")
    assert hq_sheet_codes_for_product(product) == {"CTH"}


def test_ctsh_before_cth_compound_picks_ctsh():
    product = _product(origin_sheet_criteria_override="CTSH hoặc CTH")
    assert hq_sheet_codes_for_product(product) == {"CTSH"}


def test_override_cth_wins_over_cptpp_compound_documented():
    """Operator explicitly overrides a CPTPP product to plain CTH → CTH."""
    product = _product(
        origin_sheet_effective_form_code="CPTPP",
        origin_sheet_criteria_override="CTH",
        documented_result="CTH hoặc RVC 30/40/50 tùy công thức",
    )
    assert hq_sheet_codes_for_product(product) == {"CTH"}


def test_single_criteria_unchanged():
    for text, expected in (
        ("CTH", "CTH"), ("CTSH", "CTSH"), ("RVC 40%", "RVC"),
        ("LVC 30%", "LVC"), ("Tra PSR theo Phụ lục", "PSR"),
        ("RVC 40% MaxNOM", "RVC"),
    ):
        product = _product(origin_sheet_criteria_override=text)
        assert hq_sheet_codes_for_product(product) == {expected}, text
