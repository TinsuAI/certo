"""Generate diverse upload workbooks for training sessions.

Run:
    uv run python scripts/generate_training_input_scenarios.py

Output:
    docs/training/input-scenarios/
"""
from __future__ import annotations

import json
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


OUT = Path("docs/training/input-scenarios")


def style(ws) -> None:
    fill = PatternFill("solid", fgColor="EAF4EF")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 46)


def write_xlsx(name: str, sheet: str, headers: list[str], rows: list[list[Any]]) -> dict:
    path = OUT / name
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style(ws)
    wb.save(path)
    return {"file": name, "rows": len(rows), "columns": len(headers)}


def write_with_header_offset(name: str) -> dict:
    path = OUT / name
    wb = Workbook()
    ws = wb.active
    ws.title = "BCCT header row 5"
    ws.append(["CÔNG TY DEMO PRECISION MANUFACTURING VN"])
    ws.append(["Báo cáo chi tiết hàng hóa nhập khẩu"])
    ws.append(["Kỳ báo cáo", "2026"])
    ws.append([])
    headers = [
        "Số tờ khai", "STT hàng", "Mã loại hình", "Ngày đăng ký",
        "Mã HQ", "Tên hàng", "Mã HS", "Số lượng", "ĐVT", "Trị giá",
        "Đơn vị tiền tệ", "Số hóa đơn",
    ]
    ws.append(headers)
    rows = [
        ["TRN-HDR5-001", 1, "E11", date(2026, 4, 2), "NVL-001", "Demo raw material 001", "39069099", 125, "kg", 410, "USD", "HDR5-001"],
        ["TRN-HDR5-002", 1, "E31", date(2026, 4, 3), "NVL-002", "Demo raw material 002", "39069099", 90, "kg", 290, "USD", "HDR5-002"],
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[5]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EAF4EF")
    ws.freeze_panes = "A6"
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 46)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return {"file": name, "rows": len(rows), "columns": len(headers), "header_row": 5}


def catalog_rows() -> list[list[Any]]:
    return [
        ["NVL-101", "NVL-101", "Demo imported resin", "nvl", "kg", "39069099", "active"],
        ["NVL-102", "NVL-102", "Demo copper terminal", "nvl", "pcs", "85369099", "active"],
        ["BTP-901", "BTP-901", "Demo self-produced controller board", "btp_sx", "pcs", "85371019", "active"],
        ["TP-901", "TP-901", "Demo finished control unit", "tp", "pcs", "85044090", "active"],
        ["CCDC-901", "CCDC-901", "Demo assembly fixture", "ccdc", "pcs", "90318090", "active"],
    ]


def bcct_row(seq: int, decl: str, dtype: str, code: str, qty: float, unit: str, value: float, partner: str) -> list[Any]:
    reg = date(2026, 5, 1) + timedelta(days=seq)
    return [
        decl, 1, dtype, reg, code, f"Training item {code}", "85044090" if code.startswith("TP") else "39069099",
        qty, unit, round(value / qty, 4), value, "USD", "VN" if dtype in {"E62", "B11"} else "CN",
        f"TRN-INV-{seq:04d}", "Demo Precision Manufacturing VN", "0312345678", partner,
        "FOB" if dtype in {"E62", "B11"} else "CIF", round(qty * 1.1, 3), "KGM",
        max(1, int(qty // 10) + 1), "CT", reg, reg + timedelta(days=1), "VNCLI",
        "Demo bonded warehouse", "SEA", 24500,
    ]


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    manifest: list[dict] = []

    manifest.append(write_xlsx(
        "01_catalog_vi_hq_registered.xlsx",
        "Danh muc VI",
        ["Mã HQ", "Mã NB", "Tên hàng", "Loại", "ĐVT", "Mã HS", "Trạng thái"],
        catalog_rows(),
    ))
    manifest.append(write_xlsx(
        "02_catalog_english_headers.xlsx",
        "Materials EN",
        ["Customs Code", "Internal Code", "Description", "Category", "Unit", "HS Code", "Status"],
        [
            ["NVL-201", "NB-201", "English header raw material", "raw material", "kg", "39069099", "active"],
            ["TP-201", "FG-201", "English header finished good", "finished", "pcs", "85044090", "active"],
        ],
    ))
    manifest.append(write_xlsx(
        "03_catalog_with_skipped_rows.xlsx",
        "Danh muc skipped",
        ["Mã HQ", "Mã NB", "Tên hàng", "Loại", "ĐVT", "Mã HS", "Trạng thái"],
        [
            ["NVL-301", "NVL-301", "Valid raw material", "nvl", "kg", "39069099", "active"],
            ["", "", "Missing both identifiers", "nvl", "kg", "39069099", "active"],
            ["TP-301", "", "Valid TP with only HQ code", "tp", "pcs", "85044090", "active"],
        ],
    ))
    manifest.append(write_xlsx(
        "04_bqd_identity.xlsx",
        "BQD identity",
        ["Mã nội bộ", "Mã HQ", "Loại", "Ghi chú"],
        [[code, code, "tp" if code.startswith(("TP", "BTP")) else "nvl", "Identity training pair"] for code in ["NVL-101", "NVL-102", "BTP-901", "TP-901"]],
    ))
    manifest.append(write_xlsx(
        "05_bqd_dual_code_one_to_many.xlsx",
        "BQD dual",
        ["Mã nội bộ", "Mã HQ", "Loại", "Ghi chú"],
        [
            ["NB-COPPER-01", "HQ-COPPER-A", "nvl", "One internal code maps to first HQ bucket"],
            ["NB-COPPER-01", "HQ-COPPER-B", "nvl", "1-N case"],
            ["FG-CTRL-01", "HQ-FG-CTRL", "tp", "Finished good mapping"],
            ["BTP-BOARD-01", "HQ-BTP-BOARD", "tp", "BTP stored as tp for BQD legacy enum"],
        ],
    ))
    manifest.append(write_xlsx(
        "06_bcct_full_co_fields.xlsx",
        "BCCT full",
        [
            "Số tờ khai", "STT hàng", "Mã loại hình", "Ngày đăng ký",
            "Mã HQ", "Tên hàng", "Mã HS", "Số lượng", "ĐVT",
            "Đơn giá", "Trị giá", "Đơn vị tiền tệ", "Xuất xứ",
            "Số hóa đơn", "Tên doanh nghiệp", "Mã doanh nghiệp",
            "Tên đối tác", "Điều kiện giá hóa đơn", "Trọng lượng",
            "Mã ĐVT trọng lượng", "Số lượng kiện", "Mã ĐVT kiện",
            "Ngày hóa đơn", "Ngày khởi hành vận chuyển", "Mã địa điểm đích",
            "Tên địa điểm đích", "Mã hiệu PTVT", "Tỷ giá thanh toán",
        ],
        [
            bcct_row(1, "TRN-BCCT-IMPORT-001", "E11", "NVL-101", 120, "kg", 360, "Demo Supplier A"),
            bcct_row(2, "TRN-BCCT-IMPORT-002", "E31", "NVL-102", 300, "pcs", 900, "Demo Supplier B"),
            bcct_row(3, "TRN-BCCT-EXPORT-001", "E62", "TP-901", 25, "pcs", 4250, "Demo Buyer A"),
            bcct_row(4, "TRN-BCCT-EXPORT-002", "B11", "TP-901", 18, "pcs", 3060, "Demo Buyer B"),
        ],
    ))
    manifest.append(write_with_header_offset("07_bcct_header_row_5.xlsx"))
    manifest.append(write_xlsx(
        "08_bcct_update_baseline.xlsx",
        "BCCT baseline",
        ["Số tờ khai", "STT hàng", "Mã loại hình", "Ngày đăng ký", "Mã HQ", "Tên hàng", "Số lượng", "ĐVT", "Trị giá", "Đơn vị tiền tệ"],
        [
            ["TRN-UPD-A", 1, "E11", date(2026, 3, 1), "NVL-401", "Baseline A", 100, "kg", 250, "USD"],
            ["TRN-UPD-B", 1, "E11", date(2026, 3, 2), "NVL-402", "Baseline B", 200, "kg", 500, "USD"],
            ["TRN-UPD-C", 1, "E11", date(2026, 3, 3), "NVL-403", "Baseline C", 300, "kg", 750, "USD"],
        ],
    ))
    manifest.append(write_xlsx(
        "09_bcct_update_changed.xlsx",
        "BCCT changed",
        ["Số tờ khai", "STT hàng", "Mã loại hình", "Ngày đăng ký", "Mã HQ", "Tên hàng", "Số lượng", "ĐVT", "Trị giá", "Đơn vị tiền tệ"],
        [
            ["TRN-UPD-A", 1, "E11", date(2026, 3, 1), "NVL-401", "Baseline A", 88, "kg", 250, "USD"],
            ["TRN-UPD-B", 1, "E11", date(2026, 3, 2), "NVL-402", "Baseline B", 200, "kg", 500, "USD"],
            ["TRN-UPD-D", 1, "E11", date(2026, 3, 4), "NVL-404", "New D", 400, "kg", 1000, "USD"],
        ],
    ))
    manifest.append(write_xlsx(
        "10_wrong_file_bom_uploaded_as_bcct.xlsx",
        "Wrong module",
        ["Mã SP", "Mã NVL", "Định mức", "ĐVT"],
        [["TP-901", "NVL-101", 2.5, "kg"], ["TP-901", "NVL-102", 4, "pcs"]],
    ))
    manifest.append(write_xlsx(
        "11_bom_manual_flat.xlsx",
        "BOM manual flat",
        ["Mã SP", "Mã NVL", "Định mức", "ĐVT", "Mã BOM", "Phiên bản BOM"],
        [
            ["TP-901", "NVL-101", 2.5, "kg", "BOM-TP901", "A"],
            ["TP-901", "NVL-102", 4, "pcs", "BOM-TP901", "A"],
            ["TP-902", "NVL-101", 1.7, "kg", "BOM-TP902", "A"],
        ],
    ))
    manifest.append(write_xlsx(
        "12_bom_technical_multilevel.xlsx",
        "BOM technical",
        ["Mã SP", "Mã NVL", "Định mức", "ĐVT", "Mã BOM", "Phiên bản BOM"],
        [
            ["BTP-901", "NVL-101", 0.8, "kg", "BOM-BTP901", "A"],
            ["BTP-901", "NVL-102", 2, "pcs", "BOM-BTP901", "A"],
            ["TP-901", "BTP-901", 1, "pcs", "BOM-TP901", "A"],
            ["TP-901", "NVL-103", 0.35, "kg", "BOM-TP901", "A"],
            ["TP-902", "BTP-901", 0.5, "pcs", "BOM-TP902", "A"],
            ["TP-902", "NVL-104", 1.2, "m", "BOM-TP902", "A"],
        ],
    ))
    manifest.append(write_xlsx(
        "13_bom_sap_style_english.xlsx",
        "SAP exploded",
        ["Level", "Component number", "Object description", "Comp. Qty (CUn)", "Component unit"],
        [
            [0, "TP-991", "Demo SAP finished good", 1, "PCS"],
            [1, "BTP-991", "Demo SAP sub assembly", 1, "PCS"],
            [2, "NVL-991", "Demo SAP resin", 0.45, "KG"],
            [2, "NVL-992", "Demo SAP screw", 3, "PCS"],
            [1, "NVL-993", "Demo SAP wire", 1.2, "M"],
        ],
    ))

    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUT), "files": len(manifest)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
