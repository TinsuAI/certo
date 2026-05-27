"""Combine 4 FORM MAU template files into a single workbook.

Output: data/local/hq-templates/form-mau-combined.xlsx with sheets
LVC, RVC, CTH, CTSH (clone of CTH), PSR. Each sheet is a deep copy of the
corresponding source sheet, preserving styles, merges, column widths,
row heights, page setup, and print area so wb.copy_worksheet works as
expected at export time.
"""
from __future__ import annotations

import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.copier import WorksheetCopy

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "local" / "hq-templates" / "form-mau"
OUT = ROOT / "data" / "local" / "hq-templates" / "form-mau-combined.xlsx"

SOURCES = [
    ("LVC", "FORM LVC.xlsx"),
    ("RVC", "FORM RVC.xlsx"),
    ("CTH", "FORM CTH.xlsx"),
    ("PSR", "FORM PSR.xlsx"),
]


def pick_template_sheet(wb) -> str:
    for name in wb.sheetnames:
        if name.lower() == "sheet1" or name == "foxz":
            continue
        return name
    return wb.sheetnames[0]


def copy_sheet(src_ws, dst_ws) -> None:
    # Column dimensions.
    for col_letter, dim in src_ws.column_dimensions.items():
        dd = dst_ws.column_dimensions[col_letter]
        dd.width = dim.width
        dd.hidden = dim.hidden
        dd.outlineLevel = dim.outlineLevel
        dd.bestFit = dim.bestFit

    # Row dimensions.
    for row_idx, dim in src_ws.row_dimensions.items():
        dd = dst_ws.row_dimensions[row_idx]
        dd.height = dim.height
        dd.hidden = dim.hidden
        dd.outlineLevel = dim.outlineLevel

    # Cells (value + style).
    max_row = src_ws.max_row
    max_col = src_ws.max_column
    for row in src_ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for src_cell in row:
            if src_cell.value is None and not src_cell.has_style:
                continue
            dst_cell = dst_ws.cell(row=src_cell.row, column=src_cell.column, value=src_cell.value)
            if src_cell.has_style:
                dst_cell.font = copy.copy(src_cell.font)
                dst_cell.fill = copy.copy(src_cell.fill)
                dst_cell.border = copy.copy(src_cell.border)
                dst_cell.alignment = copy.copy(src_cell.alignment)
                dst_cell.number_format = src_cell.number_format
                dst_cell.protection = copy.copy(src_cell.protection)

    # Merged ranges.
    for rng in list(src_ws.merged_cells.ranges):
        dst_ws.merge_cells(str(rng))

    # Page setup, margins, print area, freeze panes.
    dst_ws.page_setup = copy.copy(src_ws.page_setup)
    dst_ws.page_margins = copy.copy(src_ws.page_margins)
    dst_ws.print_options = copy.copy(src_ws.print_options)
    dst_ws.print_title_rows = src_ws.print_title_rows
    dst_ws.print_title_cols = src_ws.print_title_cols
    if src_ws.print_area:
        dst_ws.print_area = src_ws.print_area
    dst_ws.freeze_panes = src_ws.freeze_panes
    dst_ws.sheet_view.showGridLines = src_ws.sheet_view.showGridLines


def main() -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for criterion, filename in SOURCES:
        src_path = SRC / filename
        if not src_path.exists():
            raise FileNotFoundError(src_path)
        src_wb = load_workbook(src_path, data_only=False)
        src_name = pick_template_sheet(src_wb)
        src_ws = src_wb[src_name]
        dst_ws = wb.create_sheet(criterion)
        copy_sheet(src_ws, dst_ws)
        print(f"Copied {filename}::{src_name!r} -> {criterion}")

    # CTSH shares the CTH structure; clone within the same workbook. openpyxl's
    # copy_worksheet drops print_area, so re-apply it under the new title.
    if "CTH" in wb.sheetnames:
        cth = wb["CTH"]
        ctsh = wb.copy_worksheet(cth)
        ctsh.title = "CTSH"
        ctsh.print_area = "'CTSH'!$A$1:$N$1623"
        print("Cloned CTH -> CTSH")

    # Reorder for stable picking.
    order = ["LVC", "RVC", "CTH", "CTSH", "PSR"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]

    wb.save(OUT)
    print(f"\nWrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
