"""Convert the agency `tru-lui-co-template.xlsm` into the standard CO stock
template xlsx that the import endpoint accepts.

Usage:
  uv run python scripts/convert_co_stock.py <input.xlsm> <output.xlsx>
  uv run python scripts/convert_co_stock.py --sheet NK2 input.xlsm out.xlsx
  uv run python scripts/convert_co_stock.py --sheet Save input.xlsm out.xlsx

Behaviour:
- Default source sheet: `NK2` (the agency's reconciled per-lot snapshot —
  one row per BCCT lot with `opening_qty` and `Đã_xuất` columns).
- Alternative source: `Save` (legacy event log per CO write-down) — use
  --sheet Save when reconstructing tồn from raw write-down events.
  WARNING: Save can have data quality issues (sum-of-events vs lot opening),
  flagged as overclaim. Prefer NK2 unless the agency hasn't reconciled.
- Rows with the same (declaration_no, line_no, customs_code) triplet are
  consolidated: used_qty is summed, static fields take the first non-empty.
  For NK2 each lot is one row so consolidation is a no-op.
- Rows missing declaration_no/line_no/customs_code are logged + dropped;
  errors written to a sibling `.errors.log` file.
- Rows with `Đã_xuất = 0` are skipped on NK2 import (they convey no
  adjustment — system already has BCCT opening from Data Hub). Override
  with --include-zero-used to import every lot.

Output: standard CO stock template xlsx (see app/co_stock_template.py).
"""
from __future__ import annotations

import argparse
import sys
import warnings
from collections import OrderedDict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

# Allow running as `uv run python scripts/convert_co_stock.py` without
# editable-install of the package.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.co_stock_template import (  # noqa: E402
    CO_STOCK_COLUMNS,
    write_standard_co_stock,
)

DEFAULT_SHEET = "NK2"

# Agency Save sheet column index (0-based) → standard template key.
# Save is an EVENT LOG: each row = one CO write-down for one export decl.
# Header row at index 0 (data starts row 2).
AGENCY_SAVE_COLUMN_MAP: dict[int, str] = {
    0: "declaration_no",      # A: Số TK
    1: "registration_date",   # B: Ngày ĐK
    2: "declaration_type",    # C: Mã loại hình
    3: "line_no",             # D: STT hàng
    4: "customs_code",        # E: Mã NPL/SP
    5: "hs_code",             # F: Mã HS
    6: "goods_name",          # G: Tên hàng
    7: "origin_country",      # H: Xuất xứ
    8: "unit_price",          # I: Đơn giá
    9: "taxable_unit_price",  # J: Đơn giá tính thuế
    10: "opening_qty",        # K: Tồn-before-this-event (first row's K = true opening per consolidation)
    11: "unit",               # L: Đơn vị tính
    12: "partner",            # M: Tên đối tác
    13: "invoice_no",         # N: Số hóa đơn
    14: "invoice_date",       # O: Ngày hóa đơn
    15: "exchange_rate",      # P: Tỷ giá thanh toán
    19: "used_qty",           # T: Số lượng XUẤT (this row's used qty)
    20: "source_co_no",       # U: Số CO (usually empty — agency tracks via TKX)
    22: "transaction_key",    # W: composite key
}

# Agency NK2 sheet column index (0-based) → standard template key.
# NK2 is the RECONCILED per-lot snapshot: one row per BCCT lot with the
# agency's official `Đã xuất` cumulative used. Header row at index 3
# (data starts row 5, see ws.iter_rows min_row=5 below).
AGENCY_NK2_COLUMN_MAP: dict[int, str] = {
    0: "declaration_no",      # A: Số TK
    1: "registration_date",   # B: Ngày ĐK
    2: "declaration_type",    # C: Mã loại hình
    3: "line_no",             # D: STT hàng
    4: "customs_code",        # E: Mã NPL/SP
    5: "hs_code",             # F: Mã HS
    6: "goods_name",          # G: Tên hàng
    7: "origin_country",      # H: Xuất xứ
    8: "unit_price",          # I: Đơn giá
    9: "taxable_unit_price",  # J: Đơn giá tính thuế
    10: "opening_qty",        # K: Tổng số lượng (= BCCT raw qty)
    11: "unit",               # L: Đơn vị tính
    12: "partner",            # M: Tên đối tác
    13: "invoice_no",         # N: Số hóa đơn
    14: "invoice_date",       # O: Ngày hóa đơn
    15: "exchange_rate",      # P: Tỷ giá thanh toán
    16: "used_qty",           # Q: Đã xuất (agency snapshot — source of truth)
    # 17: "Tồn" (R) = K - Q via formula; we derive remaining from K - used_qty.
    19: "transaction_key",    # T: TKX Đã cộng dồn (last consuming export decl)
}

# (sheet_name) -> (column_map, data_start_row)
SHEET_PROFILES: dict[str, tuple[dict[int, str], int]] = {
    "NK2": (AGENCY_NK2_COLUMN_MAP, 5),
    "Save": (AGENCY_SAVE_COLUMN_MAP, 2),
}

_DECIMAL_KEYS = {"opening_qty", "used_qty", "unit_price", "taxable_unit_price", "exchange_rate"}
_DATE_KEYS = {"registration_date", "invoice_date"}


def _coerce_value(key: str, raw):
    if raw in (None, ""):
        return None
    if key in _DECIMAL_KEYS:
        try:
            return Decimal(str(raw).replace(",", ""))
        except (InvalidOperation, ValueError):
            return None
    if key in _DATE_KEYS:
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        return None
    if isinstance(raw, str):
        return raw.strip()
    return raw


def _merge(existing: dict, incoming: dict, *, dup_source_cos: set[str]) -> dict:
    out = dict(existing)
    for key in CO_STOCK_COLUMNS:
        new_value = incoming.get(key)
        if new_value in (None, ""):
            continue
        if key == "used_qty":
            prev = out.get("used_qty") or Decimal("0")
            out["used_qty"] = (prev or Decimal("0")) + new_value
        elif key == "source_co_no":
            text = str(new_value).strip()
            if text and text not in dup_source_cos:
                dup_source_cos.add(text)
                if out.get("source_co_no"):
                    out["source_co_no"] = f"{out['source_co_no']}; {text}"
                else:
                    out["source_co_no"] = text
        elif not out.get(key):
            out[key] = new_value
    return out


def convert(input_path: Path, output_path: Path, sheet_name: str = DEFAULT_SHEET, *, include_zero_used: bool = False) -> dict:
    warnings.filterwarnings("ignore")
    wb = load_workbook(input_path, read_only=True, data_only=True, keep_vba=False)
    if sheet_name not in wb.sheetnames:
        raise SystemExit(
            f"Sheet '{sheet_name}' không có trong {input_path.name}. "
            f"Sheets có sẵn: {wb.sheetnames}"
        )
    if sheet_name not in SHEET_PROFILES:
        raise SystemExit(
            f"Sheet '{sheet_name}' chưa có profile. Hỗ trợ: {list(SHEET_PROFILES.keys())}"
        )
    column_map, data_start_row = SHEET_PROFILES[sheet_name]
    ws = wb[sheet_name]
    rows_in = 0
    rows_skipped_zero = 0
    rows_dropped: list[str] = []
    consolidated: OrderedDict[tuple, dict] = OrderedDict()
    dup_tracker: dict[tuple, set[str]] = {}
    # Skip rows before data_start_row (header + any preamble rows).
    iterator = ws.iter_rows(min_row=data_start_row, values_only=True)
    for line_no, raw_row in enumerate(iterator, start=data_start_row):
        if raw_row is None:
            continue
        if not any(cell not in (None, "") for cell in raw_row):
            continue
        record: dict = {}
        for col_idx, key in column_map.items():
            if col_idx >= len(raw_row):
                continue
            record[key] = _coerce_value(key, raw_row[col_idx])
        rows_in += 1
        # NK2 row filter: skip rows with no usage signal (system already has
        # opening from Data Hub BCCT; importing zero-used rows would just be
        # database churn). Save rows always carry a usage event so include all.
        if not include_zero_used and sheet_name == "NK2":
            used_val = record.get("used_qty")
            if used_val is None or used_val == Decimal("0"):
                rows_skipped_zero += 1
                continue
        declaration_no = (record.get("declaration_no") or "").strip() if isinstance(record.get("declaration_no"), str) else (record.get("declaration_no") or "")
        line_no_v = str(record.get("line_no") or "").strip()
        customs_code = (record.get("customs_code") or "").strip() if isinstance(record.get("customs_code"), str) else (record.get("customs_code") or "")
        # Coerce non-string identifiers.
        declaration_no = str(declaration_no).strip()
        customs_code = str(customs_code).strip()
        if not (declaration_no and line_no_v and customs_code):
            rows_dropped.append(
                f"Dòng {line_no}: thiếu declaration_no/line_no/customs_code "
                f"(decl={declaration_no!r}, line={line_no_v!r}, code={customs_code!r})"
            )
            continue
        # Normalize back into record so downstream sees stripped strings.
        record["declaration_no"] = declaration_no
        record["line_no"] = line_no_v
        record["customs_code"] = customs_code
        key = (declaration_no, line_no_v, customs_code)
        dup_tracker.setdefault(key, set())
        if key in consolidated:
            consolidated[key] = _merge(consolidated[key], record, dup_source_cos=dup_tracker[key])
        else:
            consolidated[key] = record
            if record.get("source_co_no"):
                dup_tracker[key].add(str(record["source_co_no"]).strip())
    wb.close()
    output_rows = list(consolidated.values())
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(write_standard_co_stock(output_rows))
    summary = {
        "input": str(input_path),
        "output": str(output_path),
        "sheet": sheet_name,
        "rows_in": rows_in,
        "rows_unique": len(output_rows),
        "rows_dropped": len(rows_dropped),
        "rows_skipped_zero_used": rows_skipped_zero,
        "duplicates_merged": rows_in - len(output_rows) - len(rows_dropped) - rows_skipped_zero,
    }
    if rows_dropped:
        log_path = output_path.with_suffix(output_path.suffix + ".errors.log")
        log_path.write_text("\n".join(rows_dropped), encoding="utf-8")
        summary["errors_log"] = str(log_path)
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Đường dẫn workbook gốc (e.g. tru-lui-co-template.xlsm)")
    parser.add_argument("output", type=Path, help="Đường dẫn xlsx output (standard CO stock template)")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, choices=list(SHEET_PROFILES.keys()), help=f"Sheet nguồn (mặc định: {DEFAULT_SHEET!r})")
    parser.add_argument("--include-zero-used", action="store_true", help="Import cả lot có Đã_xuất=0 (mặc định NK2 bỏ qua)")
    args = parser.parse_args(argv)
    if not args.input.exists():
        raise SystemExit(f"Không tìm thấy file: {args.input}")
    summary = convert(args.input, args.output, args.sheet, include_zero_used=args.include_zero_used)
    print("Converted CO stock workbook:")
    for key, value in summary.items():
        print(f"  {key}: {value}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
