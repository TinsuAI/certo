"""Generate and feed one demo company through the live web UI.

Output root:
    .ai/features/2026-05-04-demo-company-feed/

Run:
    uv run python scripts/feed_demo_company.py

Pre-reqs:
- Data Hub server live on http://127.0.0.1:8754.
- Login admin@data-hub.local / admin123.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import random
import re
import os
import sys
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.async_api import Page, async_playwright

sys.path.insert(0, os.fspath(Path(__file__).resolve().parents[1]))

from app.database import connect


BASE = "http://127.0.0.1:8754"
EMAIL = "admin@data-hub.local"
PASSWORD = "admin123"

ROOT = Path(".ai/features/2026-05-04-demo-company-feed")
INPUT = ROOT / "input"
REPORTS = ROOT / "reports"

FEATURE_DIRS = {
    "client_config": ROOT / "01-client-config",
    "catalog": ROOT / "02-catalog",
    "bqd": ROOT / "03-bqd",
    "bcct": ROOT / "04-bcct",
    "bom": ROOT / "05-bom",
    "readiness": ROOT / "06-readiness",
}

CLIENT_NAME = "Demo Precision Manufacturing VN"
TAX_CODE = "0312345678"
NOTES_MARKER = "demo_feed_20260504"


@dataclass(frozen=True)
class DemoData:
    catalog: list[dict[str, Any]]
    mappings: list[dict[str, Any]]
    bcct_rows: list[dict[str, Any]]
    bom_rows: list[dict[str, Any]]


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 42)


def _write_xlsx(path: Path, sheet_name: str, headers: list[str], rows: list[list[Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style_sheet(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_demo_data() -> DemoData:
    catalog: list[dict[str, Any]] = []

    def add(code: str, name: str, category: str, unit: str, hs: str) -> None:
        catalog.append({
            "customs_code": code,
            "internal_code": code,
            "name": name,
            "category": category,
            "unit": unit,
            "hs_code": hs,
            "status": "active",
        })

    for i in range(1, 35):
        unit = "kg" if i <= 26 else ("m" if i <= 30 else "pcs")
        hs = "39069099" if i <= 12 else ("85049090" if i <= 24 else "39269099")
        add(f"NVL-{i:03d}", f"Demo raw material {i:03d}", "nvl", unit, hs)
    for i in range(1, 6):
        add(f"BTP-{i:03d}", f"Demo sub-assembly level 1 {i:03d}", "btp_sx", "pcs", "85049090")
    for i in range(101, 103):
        add(f"BTP-{i}", f"Demo sub-assembly level 2 {i}", "btp_sx", "pcs", "85049090")
    for i in range(1, 9):
        add(f"TP-{i:03d}", f"Demo export finished product {i:03d}", "tp", "pcs", "85044090")
    add("CCDC-001", "Demo calibration fixture", "ccdc", "pcs", "90318090")

    assert len(catalog) == 50
    mappings = [{
        "internal_code": r["internal_code"],
        "customs_code": r["customs_code"],
        # hub.code_mappings has the legacy narrow category enum
        # nvl|tp|ccdc; BTP codes map under tp for BQD purposes.
        "category": "tp" if str(r["category"]).startswith("btp") else r["category"],
        "notes": "Identity HQ/NB mapping for demo",
    } for r in catalog]

    bom_rows = _build_bom_rows()
    bcct_rows = _build_bcct_rows(catalog)
    return DemoData(catalog=catalog, mappings=mappings, bcct_rows=bcct_rows, bom_rows=bom_rows)


def _build_bom_rows() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    def add(parent: str, child: str, qty: float, uom: str = "pcs") -> None:
        rows.append({
            "product_code": parent,
            "material_code": child,
            "qty_per_unit": qty,
            "uom": uom,
            "bom_code": "",
            "bom_variant_id": "default",
        })

    level2_defs = {
        "BTP-101": [("NVL-001", 0.35, "kg"), ("NVL-002", 0.18, "kg"), ("NVL-027", 0.8, "m")],
        "BTP-102": [("NVL-003", 0.22, "kg"), ("NVL-004", 0.16, "kg"), ("NVL-031", 2.0, "pcs")],
    }
    for parent, children in level2_defs.items():
        for child, qty, uom in children:
            add(parent, child, qty, uom)

    level1_defs = {
        "BTP-001": [("BTP-101", 1.0, "pcs"), ("NVL-005", 0.42, "kg"), ("NVL-006", 0.31, "kg")],
        "BTP-002": [("BTP-101", 0.5, "pcs"), ("NVL-007", 0.27, "kg"), ("NVL-028", 0.6, "m")],
        "BTP-003": [("BTP-102", 1.0, "pcs"), ("NVL-008", 0.55, "kg"), ("NVL-009", 0.21, "kg")],
        "BTP-004": [("NVL-010", 0.48, "kg"), ("NVL-011", 0.32, "kg"), ("NVL-032", 3.0, "pcs")],
        "BTP-005": [("BTP-102", 0.7, "pcs"), ("NVL-012", 0.25, "kg"), ("NVL-029", 0.45, "m")],
    }
    for parent, children in level1_defs.items():
        for child, qty, uom in children:
            add(parent, child, qty, uom)

    tp_components = {
        "TP-001": [("BTP-001", 1.0, "pcs"), ("BTP-004", 0.5, "pcs"), ("NVL-013", 0.4, "kg"), ("NVL-030", 1.2, "m")],
        "TP-002": [("BTP-002", 1.0, "pcs"), ("NVL-014", 0.35, "kg"), ("NVL-015", 0.11, "kg"), ("NVL-033", 2.0, "pcs")],
        "TP-003": [("BTP-003", 1.2, "pcs"), ("NVL-016", 0.5, "kg"), ("NVL-017", 0.2, "kg")],
        "TP-004": [("BTP-001", 0.8, "pcs"), ("BTP-005", 0.6, "pcs"), ("NVL-018", 0.42, "kg")],
        "TP-005": [("BTP-004", 1.1, "pcs"), ("NVL-019", 0.28, "kg"), ("NVL-020", 0.15, "kg")],
        "TP-006": [("BTP-002", 0.7, "pcs"), ("BTP-003", 0.4, "pcs"), ("NVL-021", 0.6, "kg")],
        "TP-007": [("BTP-005", 1.0, "pcs"), ("NVL-022", 0.33, "kg"), ("NVL-023", 0.22, "kg"), ("NVL-034", 1.0, "pcs")],
        "TP-008": [("BTP-003", 0.9, "pcs"), ("NVL-024", 0.38, "kg"), ("NVL-025", 0.18, "kg"), ("NVL-026", 0.09, "kg")],
    }
    for parent, children in tp_components.items():
        for child, qty, uom in children:
            add(parent, child, qty, uom)
    return rows


def _build_bcct_rows(catalog: list[dict[str, Any]]) -> list[dict[str, Any]]:
    random.seed(20260504)
    by_cat = {r["customs_code"]: r for r in catalog}
    nvl_codes = [f"NVL-{i:03d}" for i in range(1, 35)]
    tp_codes = [f"TP-{i:03d}" for i in range(1, 9)]
    rows: list[dict[str, Any]] = []
    start = date(2026, 1, 2)

    for i in range(650):
        code = nvl_codes[i % len(nvl_codes)]
        mat = by_cat[code]
        reg = start + timedelta(days=i % 120)
        qty = round(80 + (i % 17) * 9.5 + (i % 5) * 0.25, 3)
        unit_price = round(1.2 + (i % 11) * 0.37, 4)
        rows.append(_bcct_row(
            seq=i + 1,
            declaration_no=f"DEMOIM{202600000 + i:09d}",
            declaration_type="E11" if i % 3 else "E31",
            registration_date=reg,
            code=code,
            goods_name=f"{mat['name']} ({code})",
            hs=mat["hs_code"],
            qty=qty,
            unit=mat["unit"],
            unit_price=unit_price,
            direction_partner="Demo Supplier " + chr(65 + i % 6),
            invoice_prefix="DMI",
            origin=["CN", "JP", "KR", "VN"][i % 4],
        ))

    for j in range(350):
        code = tp_codes[j % len(tp_codes)]
        mat = by_cat[code]
        reg = start + timedelta(days=20 + (j % 140))
        qty = round(12 + (j % 13) * 1.5, 3)
        unit_price = round(85 + (j % 9) * 7.25, 4)
        rows.append(_bcct_row(
            seq=651 + j,
            declaration_no=f"DEMOEX{202600000 + j:09d}",
            declaration_type="E62" if j % 4 else "B11",
            registration_date=reg,
            code=code,
            goods_name=f"{mat['name']} ({code})",
            hs=mat["hs_code"],
            qty=qty,
            unit=mat["unit"],
            unit_price=unit_price,
            direction_partner="Demo Buyer " + chr(65 + j % 7),
            invoice_prefix="DMX",
            origin="VN",
        ))
    assert len(rows) == 1000
    return rows


def _bcct_row(
    *,
    seq: int,
    declaration_no: str,
    declaration_type: str,
    registration_date: date,
    code: str,
    goods_name: str,
    hs: str,
    qty: float,
    unit: str,
    unit_price: float,
    direction_partner: str,
    invoice_prefix: str,
    origin: str,
) -> dict[str, Any]:
    total = round(qty * unit_price, 2)
    return {
        "declaration_no": declaration_no,
        "line_no": "1",
        "declaration_type": declaration_type,
        "registration_date": registration_date,
        "customs_code": code,
        "goods_name": goods_name,
        "hs_code": hs,
        "quantity": qty,
        "unit": unit,
        "unit_price": unit_price,
        "total_value": total,
        "currency": "USD",
        "origin": origin,
        "invoice_ref": f"{invoice_prefix}-{seq:05d}",
        "exporter_name": CLIENT_NAME,
        "exporter_tax_code": TAX_CODE,
        "consignee_name": direction_partner,
        "incoterms": "FOB" if declaration_type in {"E62", "B11"} else "CIF",
        "weight": round(qty * 1.15, 3),
        "weight_unit": "KGM",
        "package_count": max(1, int(qty // 10) + 1),
        "package_unit": "CT",
        "invoice_date": registration_date - timedelta(days=2),
        "departure_date": registration_date + timedelta(days=1),
        "destination_code": "VNCLI",
        "destination_name": "Demo bonded warehouse",
        "transport_mode": "SEA",
        "exchange_rate": 24500,
    }


def generate_inputs(data: DemoData) -> None:
    rows = [[
        r["customs_code"], r["internal_code"], r["name"], r["category"],
        r["unit"], r["hs_code"], r["status"],
    ] for r in data.catalog]
    _write_xlsx(
        INPUT / "demo_catalog_50_codes.xlsx",
        "Danh muc demo",
        ["Mã HQ", "Mã NB", "Tên hàng", "Loại", "ĐVT", "Mã HS", "Trạng thái"],
        rows,
    )
    rows = [[r["internal_code"], r["customs_code"], r["category"], r["notes"]] for r in data.mappings]
    _write_xlsx(
        INPUT / "demo_bqd_identity_50_codes.xlsx",
        "BQD identity",
        ["Mã nội bộ", "Mã HQ", "Loại", "Ghi chú"],
        rows,
    )
    rows = [[
        r["declaration_no"], r["line_no"], r["declaration_type"],
        r["registration_date"], r["customs_code"], r["goods_name"],
        r["hs_code"], r["quantity"], r["unit"], r["unit_price"],
        r["total_value"], r["currency"], r["origin"], r["invoice_ref"],
        r["exporter_name"], r["exporter_tax_code"], r["consignee_name"],
        r["incoterms"], r["weight"], r["weight_unit"], r["package_count"],
        r["package_unit"], r["invoice_date"], r["departure_date"],
        r["destination_code"], r["destination_name"], r["transport_mode"],
        r["exchange_rate"],
    ] for r in data.bcct_rows]
    _write_xlsx(
        INPUT / "demo_bcct_1000_rows.xlsx",
        "BCCT demo",
        [
            "Số tờ khai", "STT hàng", "Mã loại hình", "Ngày đăng ký",
            "Mã HQ", "Tên hàng", "Mã HS", "Số lượng", "ĐVT",
            "Đơn giá", "Trị giá", "Đơn vị tiền tệ", "Xuất xứ",
            "Số hóa đơn", "Tên doanh nghiệp", "Mã doanh nghiệp",
            "Tên đối tác", "Điều kiện giá hóa đơn", "Trọng lượng",
            "Mã ĐVT trọng lượng", "Số lượng kiện", "Mã ĐVT kiện",
            "Ngày hóa đơn", "Ngày khởi hành vận chuyển", "Mã địa điểm đích",
            "Tên địa điểm đích", "Mã hiệu PTVT", "Tỷ giá thanh toán",
        ],
        rows,
    )
    rows = [[
        r["product_code"], r["material_code"], r["qty_per_unit"], r["uom"],
        r["bom_code"], r["bom_variant_id"],
    ] for r in data.bom_rows]
    _write_xlsx(
        INPUT / "demo_bom_multilevel_technical.xlsx",
        "BOM multi-level",
        ["Mã SP", "Mã NVL", "Định mức", "ĐVT", "Mã BOM", "Phiên bản BOM"],
        rows,
    )


def write_feature_docs(data: DemoData) -> None:
    for path in FEATURE_DIRS.values():
        (path / "tests").mkdir(parents=True, exist_ok=True)
        (path / "screenshots").mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(parents=True, exist_ok=True)

    _write_text(FEATURE_DIRS["client_config"] / "brief.md", f"""# Demo Company Setup

Creates `{CLIENT_NAME}` through the web UI and configures:
- code resolution mode: `identity` because HQ code and internal code are the same.
- BOM proposal mode: `auto`, tolerance `3.0`, approver tier `edit`.
- declaration config: import `E11,E31`, export `E62,B11`, fiscal year starts in January.
""")
    _write_text(FEATURE_DIRS["client_config"] / "tests" / "manual_cases.md", """# Manual Cases

1. Client exists and opens from `/clients`.
2. General config shows identity code mode and active status.
3. Declaration config API is versioned and includes `E11,E31,E62,B11`.
""")

    _write_text(FEATURE_DIRS["catalog"] / "brief.md", """# Demo Catalog

Input: `input/demo_catalog_50_codes.xlsx`.

The catalog has exactly 50 codes with `Mã HQ == Mã NB`: 34 NVL, 7 BTP, 8 TP, 1 CCDC. All rows are active and uploaded as HQ-registered catalog data.
""")
    _write_text(FEATURE_DIRS["catalog"] / "tests" / "expected_counts.json", json.dumps({
        "total": 50,
        "by_category": {"nvl": 34, "btp_sx": 7, "tp": 8, "ccdc": 1},
        "hq_equals_internal": True,
    }, indent=2))

    _write_text(FEATURE_DIRS["bqd"] / "brief.md", """# Demo BQD

Input: `input/demo_bqd_identity_50_codes.xlsx`.

Every row maps the same internal code to the same customs code. This keeps identity mode explicit while still exercising the BQD upload feature.
""")
    _write_text(FEATURE_DIRS["bqd"] / "tests" / "expected_counts.json", json.dumps({
        "total": 50,
        "identity_pairs": 50,
    }, indent=2))

    _write_text(FEATURE_DIRS["bcct"] / "brief.md", """# Demo BCCT

Input: `input/demo_bcct_1000_rows.xlsx`.

The workbook has 1,000 declaration lines: 650 import lines for NVL and 350 export lines for TP. It includes typed CO fields such as invoice, partner, incoterms, weight, packages, destination, transport mode, and exchange rate.
""")
    _write_text(FEATURE_DIRS["bcct"] / "tests" / "expected_counts.json", json.dumps({
        "total": 1000,
        "import_rows": 650,
        "export_rows": 350,
        "declaration_types": ["E11", "E31", "E62", "B11"],
    }, indent=2))

    _write_text(FEATURE_DIRS["bom"] / "brief.md", """# Demo BOM

Input: `input/demo_bom_multilevel_technical.xlsx`.

The workbook is a multi-level technical BOM with 8 export finished products, 5 level-1 BTPs, and 2 level-2 BTPs. It is uploaded through the `technical_flatten` profile so Data Hub materializes calculation-ready flattened BOM versions.
""")
    _write_text(FEATURE_DIRS["bom"] / "tests" / "expected_counts.json", json.dumps({
        "input_parent_products": 15,
        "export_products": 8,
        "btp_products": 7,
        "input_edges": len(data.bom_rows),
        "expected_flatten_status": "flattened",
    }, indent=2))

    _write_text(FEATURE_DIRS["readiness"] / "brief.md", """# Demo Readiness

This folder captures final UI screenshots and DB/API readiness checks after the demo company is fed through the web UI.
""")
    _write_text(FEATURE_DIRS["readiness"] / "tests" / "acceptance.md", """# Acceptance

- Client exists in Data Hub.
- Catalog count is 50.
- BQD count is 50.
- BCCT count is 1,000.
- BOM has flattened versions for 8 TP and 7 BTP products.
- `/v1/hub/dncxs/{client_id}/client-config` returns configured declaration filters.
""")


def _write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8")


def validate_input_files() -> dict[str, Any]:
    checks: dict[str, Any] = {}
    for name in [
        "demo_catalog_50_codes.xlsx",
        "demo_bqd_identity_50_codes.xlsx",
        "demo_bcct_1000_rows.xlsx",
        "demo_bom_multilevel_technical.xlsx",
    ]:
        wb = load_workbook(INPUT / name, read_only=True, data_only=True)
        ws = wb.active
        checks[name] = {"rows": max(ws.max_row - 1, 0), "columns": ws.max_column}
    return checks


async def login(page: Page) -> None:
    await page.goto(f"{BASE}/login")
    await page.fill('input[name="email"]', EMAIL)
    await page.fill('input[name="password"]', PASSWORD)
    await page.click('button[type="submit"]')
    await page.wait_for_url(f"{BASE}/clients", timeout=10_000)


async def capture(page: Page, feature: str, slug: str) -> Path:
    path = FEATURE_DIRS[feature] / "screenshots" / f"{slug}.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    await page.screenshot(path=str(path), full_page=True)
    return path


async def create_client(page: Page) -> str:
    await page.goto(f"{BASE}/clients/new")
    await page.fill('input[name="name"]', CLIENT_NAME)
    await page.fill('input[name="tax_code"]', TAX_CODE)
    await page.select_option('select[name="code_resolution_mode"]', "identity")
    await page.select_option('select[name="bom_proposal_mode"]', "auto")
    await page.fill('input[name="bom_proposal_qty_tolerance_pct"]', "3.0")
    await page.select_option('select[name="bom_approver_tier"]', "edit")
    await page.fill('textarea[name="notes"]', NOTES_MARKER)
    await capture(page, "client_config", "01_new_client_form")
    await page.locator("form.settings-form button.btn-primary").last.click()
    await page.wait_for_url(re.compile(f"{BASE}/clients/[^/]+$"), timeout=15_000)
    await page.wait_for_load_state("networkidle", timeout=15_000)
    client_id = page.url.rstrip("/").split("/")[-1]
    await capture(page, "client_config", "02_client_workspace_created")
    return client_id


async def configure_declaration_types(page: Page, client_id: str) -> None:
    await page.goto(f"{BASE}/clients/{client_id}/declaration-config")
    await page.wait_for_load_state("networkidle", timeout=15_000)
    await page.locator("input[name='eligible_import_declaration_types']").evaluate_all(
        "(els) => els.forEach((e) => e.checked = false)"
    )
    await page.locator("input[name='relevant_export_declaration_types']").evaluate_all(
        "(els) => els.forEach((e) => e.checked = false)"
    )
    for code in ["E11", "E31"]:
        await page.locator(f"input[name='eligible_import_declaration_types'][value='{code}']").set_checked(True)
    for code in ["E62", "B11"]:
        await page.locator(f"input[name='relevant_export_declaration_types'][value='{code}']").set_checked(True)
    await page.select_option("select[name='fiscal_year_start_month']", "1")
    await capture(page, "client_config", "03_declaration_config_form")
    await page.locator("form.settings-form button.btn-primary").click()
    await page.wait_for_url(re.compile(".*/declaration-config\\?saved=1"), timeout=15_000)
    await page.wait_for_load_state("networkidle", timeout=15_000)
    await capture(page, "client_config", "04_declaration_config_saved")


async def upload_mapping_preview_confirm(
    page: Page,
    *,
    client_id: str,
    feature: str,
    module: str,
    file_path: Path,
    profile: str | None = None,
    hq_registered: bool = False,
) -> None:
    upload_url = f"{BASE}/clients/{client_id}/{module}/upload"
    await page.goto(upload_url)
    await page.wait_for_load_state("networkidle", timeout=15_000)
    if profile is not None:
        await page.select_option('select[name="profile"]', profile)
    if hq_registered:
        await page.locator('input[name="is_hq_registered"]').set_checked(True)
    await page.set_input_files('input[type="file"]', str(file_path))
    await capture(page, feature, "01_upload_form_filled")
    await page.locator('form[enctype="multipart/form-data"] button[type="submit"]').first.click()
    await page.wait_for_load_state("networkidle", timeout=60_000)

    if "/upload/mapping/" in page.url:
        await capture(page, feature, "02_mapping_page")
        await page.locator('button[type="submit"]:has-text("Lưu mapping")').first.click()
        await page.wait_for_load_state("networkidle", timeout=60_000)

    if "/flatten-preview/" in page.url:
        await capture(page, feature, "03_flatten_preview")
        pending_decisions = await page.locator("input[name^='confirm_']").count()
        if pending_decisions:
            raise RuntimeError(f"Unexpected flatten decisions for demo BOM: {pending_decisions}")
        await page.locator("form button.btn-primary").first.click()
        await page.wait_for_url(re.compile(f"{BASE}/clients/{client_id}/{module}.*"), timeout=60_000)
        await page.wait_for_load_state("networkidle", timeout=60_000)
        await capture(page, feature, "04_after_confirm")
        return

    if "/preview/" in page.url or "/upload/preview/" in page.url:
        await capture(page, feature, "03_preview")
        if module == "bcct":
            await page.locator("form button.btn-primary").first.click()
        else:
            await page.locator("form#confirm-form button.btn-primary").first.click()
        await page.wait_for_url(re.compile(f"{BASE}/clients/{client_id}/{module}.*"), timeout=60_000)
        await page.wait_for_load_state("networkidle", timeout=60_000)
        await capture(page, feature, "04_after_confirm")
        return

    raise RuntimeError(f"Unexpected URL after {module} upload: {page.url}")


def summarize_db(client_id: str) -> dict[str, Any]:
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute("select name, tax_code, code_resolution_mode, bom_proposal_mode from hub.clients where client_id=%s", (client_id,))
            client_row = cur.fetchone()
            cur.execute("select category, count(*) from hub.materials where client_id=%s group by category order by category", (client_id,))
            catalog_by_category = dict(cur.fetchall())
            cur.execute("select count(*) from hub.materials where client_id=%s", (client_id,))
            catalog_count = cur.fetchone()[0]
            cur.execute("select count(*) from hub.code_mappings where client_id=%s", (client_id,))
            bqd_count = cur.fetchone()[0]
            cur.execute("select direction, count(*) from hub.bcct_rows where client_id=%s group by direction order by direction", (client_id,))
            bcct_by_direction = dict(cur.fetchall())
            cur.execute("select count(*) from hub.bcct_rows where client_id=%s", (client_id,))
            bcct_count = cur.fetchone()[0]
            cur.execute(
                """
                select coalesce(m.category, 'unknown') as category,
                       bv.flatten_status, count(distinct bv.product_code), sum(bv.row_count)
                from hub.bom_artifacts bv
                left join hub.materials m
                  on m.client_id=bv.client_id and m.customs_code=bv.product_code
                where bv.client_id=%s and bv.tombstoned_at is null
                group by coalesce(m.category, 'unknown'), bv.flatten_status
                order by 1,2
                """,
                (client_id,),
            )
            bom_summary = [
                {"category": r[0], "flatten_status": r[1], "products": r[2], "rows": int(r[3] or 0)}
                for r in cur.fetchall()
            ]
            cur.execute(
                """
                select count(*) from hub.bom_artifacts
                where client_id=%s and flatten_status='non_flattened'
                  and tombstoned_at is null
                """,
                (client_id,),
            )
            non_flattened = cur.fetchone()[0]
            cur.execute(
                """
                select count(*) from hub.bom_unresolved_nodes u
                join hub.bom_artifacts v on v.artifact_id=u.artifact_id
                where v.client_id=%s and v.tombstoned_at is null
                """,
                (client_id,),
            )
            unresolved = cur.fetchone()[0]
    return {
        "client_id": client_id,
        "client": {
            "name": client_row[0] if client_row else None,
            "tax_code": client_row[1] if client_row else None,
            "code_resolution_mode": client_row[2] if client_row else None,
            "bom_proposal_mode": client_row[3] if client_row else None,
        },
        "catalog_count": catalog_count,
        "catalog_by_category": catalog_by_category,
        "bqd_count": bqd_count,
        "bcct_count": bcct_count,
        "bcct_by_direction": bcct_by_direction,
        "bom_summary": bom_summary,
        "bom_non_flattened_versions": non_flattened,
        "bom_unresolved_nodes": unresolved,
    }


def pending_ids(client_id: str, module: str) -> list[str]:
    with connect() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                select pending_id from hub.upload_pending
                where client_id=%s and module=%s
                order by created_at
                """,
                (client_id, module),
            )
            return [r[0] for r in cur.fetchall()]


async def reject_pending_for_module(page: Page, *, client_id: str, module: str, feature: str) -> None:
    for idx, pending_id in enumerate(pending_ids(client_id, module), start=1):
        await page.goto(f"{BASE}/clients/{client_id}/{module}/preview/{pending_id}")
        await page.wait_for_load_state("networkidle", timeout=15_000)
        await capture(page, feature, f"00_reject_stale_pending_{idx}")
        page.once("dialog", lambda d: asyncio.create_task(d.accept()))
        await page.locator('form#confirm-form button.btn-secondary:has-text("Reject")').click()
        await page.wait_for_load_state("networkidle", timeout=15_000)


async def run_ui_feed(*, client_id: str | None = None, only: str = "all") -> dict[str, Any]:
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        ctx = await browser.new_context(viewport={"width": 1440, "height": 900})
        await ctx.add_cookies([{"name": "data_hub_lang", "value": "vi", "url": BASE}])
        page = await ctx.new_page()
        await login(page)
        if client_id is None:
            client_id = await create_client(page)
            await configure_declaration_types(page, client_id)
        if only in {"all", "catalog"}:
            await upload_mapping_preview_confirm(
                page, client_id=client_id, feature="catalog", module="catalog",
                file_path=INPUT / "demo_catalog_50_codes.xlsx", hq_registered=True,
            )
        if only in {"all", "bqd"}:
            await reject_pending_for_module(page, client_id=client_id, module="bqd", feature="bqd")
            await upload_mapping_preview_confirm(
                page, client_id=client_id, feature="bqd", module="bqd",
                file_path=INPUT / "demo_bqd_identity_50_codes.xlsx",
            )
        if only in {"all", "bcct"}:
            await upload_mapping_preview_confirm(
                page, client_id=client_id, feature="bcct", module="bcct",
                file_path=INPUT / "demo_bcct_1000_rows.xlsx",
            )
        if only in {"all", "bom"}:
            await upload_mapping_preview_confirm(
                page, client_id=client_id, feature="bom", module="bom",
                file_path=INPUT / "demo_bom_multilevel_technical.xlsx",
                profile="technical_flatten",
            )
        for module in ["catalog", "bqd", "bcct", "bom"]:
            await page.goto(f"{BASE}/clients/{client_id}/{module}")
            await page.wait_for_load_state("networkidle", timeout=15_000)
            await capture(page, "readiness", f"{module}_final_list")
        await page.goto(f"{BASE}/clients/{client_id}")
        await page.wait_for_load_state("networkidle", timeout=15_000)
        await capture(page, "readiness", "client_workspace_final")
        await browser.close()
    return summarize_db(client_id)


def assert_summary(summary: dict[str, Any]) -> None:
    expected = {
        "catalog_count": 50,
        "bqd_count": 50,
        "bcct_count": 1000,
        "bom_non_flattened_versions": 0,
        "bom_unresolved_nodes": 0,
    }
    for key, value in expected.items():
        if summary.get(key) != value:
            raise AssertionError(f"{key}: expected {value}, got {summary.get(key)}")
    if summary.get("bcct_by_direction", {}).get("import") != 650:
        raise AssertionError(f"bcct import count mismatch: {summary.get('bcct_by_direction')}")
    if summary.get("bcct_by_direction", {}).get("export") != 350:
        raise AssertionError(f"bcct export count mismatch: {summary.get('bcct_by_direction')}")
    bom_flat_products = sum(
        item["products"] for item in summary.get("bom_summary", [])
        if item["flatten_status"] == "flattened"
    )
    if bom_flat_products < 15:
        raise AssertionError(f"expected at least 15 flattened BOM products, got {bom_flat_products}")


def write_manifest(*, input_checks: dict[str, Any], summary: dict[str, Any] | None) -> None:
    manifest = {
        "base_url": BASE,
        "artifact_root": str(ROOT),
        "input_files": input_checks,
        "db_summary": summary,
    }
    _write_text(ROOT / "manifest.json", json.dumps(manifest, indent=2, ensure_ascii=False, default=str))
    if summary:
        _write_text(REPORTS / "final_db_summary.json", json.dumps(summary, indent=2, ensure_ascii=False, default=str))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--generate-only", action="store_true", help="Only generate input files and docs.")
    parser.add_argument("--client-id", help="Resume against an existing client instead of creating a new one.")
    parser.add_argument(
        "--only",
        choices=["all", "catalog", "bqd", "bcct", "bom"],
        default="all",
        help="When resuming, run only one upload module.",
    )
    args = parser.parse_args()

    data = build_demo_data()
    generate_inputs(data)
    write_feature_docs(data)
    input_checks = validate_input_files()
    summary = None
    if not args.generate_only:
        summary = asyncio.run(run_ui_feed(client_id=args.client_id, only=args.only))
        assert_summary(summary)
    write_manifest(input_checks=input_checks, summary=summary)
    print(json.dumps({"artifact_root": str(ROOT), "summary": summary or input_checks}, indent=2, ensure_ascii=False, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
