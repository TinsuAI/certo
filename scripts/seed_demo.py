"""Seed demo data: synthesize Excel files, upload via HTTP, verify rows persist."""
from __future__ import annotations

import io
import sys
import time
from pathlib import Path
from urllib.parse import urlencode

import httpx
from openpyxl import Workbook

BASE = "http://127.0.0.1:8754"
EMAIL = "admin@data-hub.local"
PASSWORD = "admin123"


def _login(client: httpx.Client) -> None:
    r = client.post(f"{BASE}/login", data={"email": EMAIL, "password": PASSWORD},
                    follow_redirects=False)
    if r.status_code not in (303, 302):
        raise SystemExit(f"login failed: {r.status_code} {r.text[:200]}")


def _make_materials_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DM NVL"
    ws.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    rows = [
        ("PE-001", "PE-001", "Polyethylene resin grade A", "nvl", "kg", "39011010"),
        ("PE-002", "PE-002", "Polyethylene resin grade B", "nvl", "kg", "39011010"),
        ("AL-100", "AL-100", "Aluminum sheet 1mm", "nvl", "kg", "76061110"),
        ("CU-WIRE-2", "CU-WIRE-2", "Copper wire 2.5mm", "nvl", "m", "85447000"),
        ("PCB-12", "PCB-12", "Printed circuit board 12-layer", "nvl", "pcs", "85340010"),
    ]
    for r in rows:
        ws.append(r)
    ws_tp = wb.create_sheet("DM TP")
    ws_tp.append(["Mã HQ", "Mã NB", "Tên", "Loại", "ĐVT", "HS"])
    ws_tp.append(("INV-3000", "INV-3000", "Solar inverter 3kW", "tp", "pcs", "85044090"))
    ws_tp.append(("INV-5000", "INV-5000", "Solar inverter 5kW", "tp", "pcs", "85044090"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_bqd_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BQD NVL"
    ws.append(["Mã nội bộ", "Mã hải quan"])
    pairs = [
        ("PE-001", "PE-001"),
        ("PE-002", "PE-002"),
        ("AL-100", "AL-100"),
        ("CU-WIRE-2", "CU-WIRE-2"),
        ("PCB-12", "PCB-12"),
        ("PE-001", "PE-ALT-FALLBACK"),  # 1:n example
    ]
    for ic, cc in pairs:
        ws.append((ic, cc))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_bcct_xlsx() -> bytes:
    """Embed internal codes Growatt-style: '<code>#&description'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BCCT"
    ws.append([
        "Số tờ khai", "Dòng", "Mã loại hình", "Ngày đăng ký", "Mã NPL/SP",
        "Tên hàng", "HS", "Tổng số lượng", "ĐVT", "Trị giá", "Nguyên tệ",
    ])
    rows = [
        ("104123456789", 1, "E11", "2025-03-15", "PE-001",
         "PE-001#&Polyethylene resin grade A premium", "39011010", 1000.0, "kg", 2500.0, "USD"),
        ("104123456789", 2, "E11", "2025-03-15", "PE-002",
         "PE-002#&Polyethylene resin grade B standard", "39011010", 500.0, "kg", 1100.0, "USD"),
        ("104123456790", 1, "E11", "2025-04-02", "AL-100",
         "AL-100#&Aluminum sheet 1mm rolled", "76061110", 200.0, "kg", 1600.0, "USD"),
        ("104123456791", 1, "E15", "2025-04-10", "CU-WIRE-2",
         "Copper wire 2.5mm 100m roll (CU-WIRE-2.A1)", "85447000", 100.0, "m", 350.0, "USD"),
        ("104123456800", 1, "E42", "2025-09-01", "INV-3000",
         "INV-3000#&Solar inverter 3kW model 2025", "85044090", 50.0, "pcs", 12500.0, "USD"),
    ]
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_bom_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(["Mã SP", "Mã NVL", "Định mức", "ĐVT"])
    rows = [
        ("INV-3000", "PE-001", 0.45, "kg"),
        ("INV-3000", "AL-100", 0.30, "kg"),
        ("INV-3000", "PCB-12", 1.0, "pcs"),
        ("INV-5000", "PE-001", 0.60, "kg"),
        ("INV-5000", "PE-002", 0.20, "kg"),
        ("INV-5000", "AL-100", 0.50, "kg"),
        ("INV-5000", "CU-WIRE-2", 1.20, "m"),
        ("INV-5000", "PCB-12", 1.0, "pcs"),
    ]
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> int:
    with httpx.Client(timeout=30.0, follow_redirects=False) as client:
        _login(client)

        # Find an existing batch_aggregate_resolution DNCX (Growatt) or create one.
        r = client.get(f"{BASE}/dncxs")
        if "Growatt VN" not in r.text:
            client.post(f"{BASE}/dncxs/new", data={
                "name": "Growatt VN", "tax_code": "0307123456",
                "code_resolution_mode": "batch_aggregate_resolution",
                "bom_proposal_mode": "auto",
                "bom_proposal_qty_tolerance_pct": "5.0",
            })
        # Locate the slug.
        r = client.get(f"{BASE}/dncxs")
        # crude slug extraction — pick first growatt-vn-* link
        import re
        m = re.search(r'href="/dncxs/(growatt-vn-[a-f0-9]+)"', r.text)
        if not m:
            print("Could not find Growatt DNCX in list")
            return 1
        dncx_id = m.group(1)
        print(f"Using DNCX: {dncx_id}")

        # Upload Materials
        r = client.post(f"{BASE}/materials/upload",
                        data={"dncx_id": dncx_id},
                        files={"file": ("materials.xlsx", _make_materials_xlsx(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        print(f"  Materials upload -> {r.status_code}")

        # Upload BQD
        r = client.post(f"{BASE}/code-mappings/upload",
                        data={"dncx_id": dncx_id},
                        files={"file": ("bqd.xlsx", _make_bqd_xlsx(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        print(f"  BQD upload -> {r.status_code}")

        # Upload BCCT
        r = client.post(f"{BASE}/bcct/upload",
                        data={"dncx_id": dncx_id, "year": "2025"},
                        files={"file": ("bcct.xlsx", _make_bcct_xlsx(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        print(f"  BCCT upload -> {r.status_code}")

        # Upload BOM (manual_flat profile)
        r = client.post(f"{BASE}/bom/upload",
                        data={"dncx_id": dncx_id, "profile": "manual_flat"},
                        files={"file": ("bom.xlsx", _make_bom_xlsx(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        print(f"  BOM upload -> {r.status_code}")

        # Verify by reading detail page.
        r = client.get(f"{BASE}/dncxs/{dncx_id}")
        print(f"  DNCX detail page -> {r.status_code}")
        for line in r.text.split("\n"):
            if "card-meta" in line and (
                "mã NVL/SP/BTP" in line or "mapping" in line
                or "tờ khai" in line or "version" in line
            ):
                print(f"    {line.strip()}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
