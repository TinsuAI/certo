"""Playwright shots of the BOM upload flow signals (destination banner +
honest post-ingest toast). Drives the live dev server on :8754.

Covers each case in the flow:
  - upload form (auto default)
  - preview: technical BOM  -> "tự động sinh BOM phẳng" banner
  - preview: flat BOM       -> "lưu nguyên trạng" banner
  - preview: multi-level file landing flat -> warning banner (synthetic)
  - list toast: technical OK (raw + flat shapes derived)
  - list toast: technical WARNING (raw + no flat shapes)
  - list toast: flat stored
"""
from __future__ import annotations

import asyncio
import io
import json
import secrets
from pathlib import Path

import openpyxl
from playwright.async_api import async_playwright

from app.database import connect
from app.auth.session import create_session, hash_password, SESSION_COOKIE

BASE = "http://127.0.0.1:8754"
OUT = Path(".ai/features/2026-06-05-bom-auto-tree-flat-fix/screenshots")
OUT.mkdir(parents=True, exist_ok=True)

CLIENT = "uxshot-vn"
UID = "u_uxshot"


def _tree_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in {
        "整机": [("工厂", "成品物料", "组件物料", "单位", "标准用量"),
                ("6180", "TP-TREE-01", "BTP-A", "ST", 2)],
        "B700": [("工厂", "成品物料", "组件物料", "单位", "标准用量"),
                ("6180", "BTP-A", "NVL-X", "KG", 3)],
    }.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _flat_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SP-FLAT-01"
    ws.append(("Mã NVL", "Định mức", "ĐVT"))
    ws.append(("NVL-A", 1, "kg"))
    ws.append(("NVL-B", 2, "pcs"))
    ws.append(("NVL-C", 0.5, "m"))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _db_setup():
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "insert into hub.clients (client_id, name, auto_derive_shallow_from_raw) "
            "values (%s,%s,'publish') on conflict (client_id) do update "
            "set auto_derive_shallow_from_raw='publish'",
            (CLIENT, "UX Shot Co."))
        cur.execute(
            "insert into hub.users (user_id,email,display_name,password_hash,role,status) "
            "values (%s,'uxshot@local','UX Shooter',%s,'admin','active') "
            "on conflict (user_id) do update set role='admin',status='active'",
            (UID, hash_password("x")))


def _stash_multilevel_pending() -> str:
    """Synthesise a flat pending whose rows carry a multi-level node_path,
    so the preview renders the 'multi-level landing flat' warning."""
    upload_id = "upl_" + secrets.token_urlsafe(10)
    pending_id = "pnd_" + secrets.token_urlsafe(10)
    products = {
        "SP-MULTI-01": [
            {"material_code": "NVL-DEEP-1", "qty_per_unit": 1.0, "uom": "kg",
             "bom_code": None, "bom_variant_id": None,
             "_node_path": "SP-MULTI-01 > BTP-SUB > NVL-DEEP-1"},
            {"material_code": "NVL-DEEP-2", "qty_per_unit": 0.4, "uom": "pcs",
             "bom_code": None, "bom_variant_id": None,
             "_node_path": "SP-MULTI-01 > BTP-SUB > CỤM-C > NVL-DEEP-2"},
        ]
    }
    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            "insert into hub.file_uploads "
            "(upload_id, client_id, module, original_filename, stored_path, "
            " storage_backend, content_sha256, size_bytes, parse_status, "
            " result, uploader_user_id) "
            "values (%s,%s,'bom','SP-MULTI-01.xlsx','x/y.xlsx','local',%s,123,"
            "'pending_preview','{}'::jsonb,%s)",
            (upload_id, CLIENT, secrets.token_hex(16), UID))
        cur.execute(
            "insert into hub.upload_pending "
            "(pending_id, client_id, module, upload_id, parsed_rows, "
            " diff_summary, created_by, expires_at) "
            "values (%s,%s,'bom',%s,%s::jsonb,%s::jsonb,%s, now() + interval '2 hours')",
            (pending_id, CLIENT, upload_id,
             json.dumps({"products": products}),
             json.dumps({"profile": "manual_flat",
                         "proposed_by": "parser_auto:manual_flat"}),
             UID))
    return pending_id


def _db_teardown():
    with connect() as conn, conn.cursor() as cur:
        cur.execute("delete from hub.bom_edges where artifact_id in "
                    "(select artifact_id from hub.bom_artifacts where client_id=%s)", (CLIENT,))
        cur.execute("delete from hub.bom_artifact_rows where artifact_id in "
                    "(select artifact_id from hub.bom_artifacts where client_id=%s)", (CLIENT,))
        cur.execute("delete from hub.bom_artifacts where client_id=%s", (CLIENT,))
        cur.execute("delete from hub.upload_pending where client_id=%s", (CLIENT,))
        cur.execute("delete from hub.file_uploads where uploader_user_id=%s", (UID,))
        cur.execute("delete from hub.sessions where user_id=%s", (UID,))
        cur.execute("delete from hub.users where user_id=%s", (UID,))
        cur.execute("delete from hub.clients where client_id=%s", (CLIENT,))


async def _upload_get_pending(page, blob: bytes, filename: str) -> str:
    resp = await page.request.post(
        f"{BASE}/clients/{CLIENT}/bom/upload",
        multipart={
            "profile": "auto",
            "file": {"name": filename, "mimeType":
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     "buffer": blob},
        },
        max_redirects=0,
    )
    loc = resp.headers["location"]
    return loc  # /clients/.../bom/preview/<pid>


async def shot(page, slug: str):
    await page.wait_for_timeout(350)
    out = OUT / f"{slug}.png"
    await page.screenshot(path=str(out), full_page=True)
    print(f"  saved {out}")


async def main():
    _db_setup()
    sid = create_session(UID)
    pending_ml = _stash_multilevel_pending()
    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            context = await browser.new_context(viewport={"width": 1280, "height": 900})
            await context.add_cookies([
                {"name": SESSION_COOKIE, "value": sid, "url": BASE},
                {"name": "data_hub_lang", "value": "vi", "url": BASE},
                {"name": "data_hub_theme", "value": "light", "url": BASE},
            ])
            page = await context.new_page()

            # 1. upload form
            await page.goto(f"{BASE}/clients/{CLIENT}/bom/upload")
            await shot(page, "01_upload_form")

            # 2. preview — technical BOM (auto -> raw, will flatten)
            loc = await _upload_get_pending(page, _tree_xlsx(), "TP-TREE-01.xlsx")
            await page.goto(BASE + loc)
            await shot(page, "02_preview_technical_will_flatten")

            # 3. preview — flat BOM (stored as-is)
            loc = await _upload_get_pending(page, _flat_xlsx(), "SP-FLAT-01.xlsx")
            await page.goto(BASE + loc)
            await shot(page, "03_preview_flat_as_is")

            # 4. preview — multi-level file landing flat (warning)
            await page.goto(f"{BASE}/clients/{CLIENT}/bom/preview/{pending_ml}")
            await shot(page, "04_preview_multilevel_warning")

            # 5-7. list toasts
            await page.goto(f"{BASE}/clients/{CLIENT}/bom?ingested=1&kind=raw&flat=1")
            await shot(page, "05_toast_technical_ok")
            await page.goto(f"{BASE}/clients/{CLIENT}/bom?ingested=1&kind=raw&flat=0")
            await shot(page, "06_toast_technical_warning")
            await page.goto(f"{BASE}/clients/{CLIENT}/bom?ingested=2&kind=flat&flat=0")
            await shot(page, "07_toast_flat_stored")

            await browser.close()
    finally:
        _db_teardown()
    print("done")


if __name__ == "__main__":
    asyncio.run(main())
